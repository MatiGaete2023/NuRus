"""Optional acceptance on the Windows workstation with desktop Office installed."""
import argparse,json,os,tempfile
from pathlib import Path


def verify(folder,outlook_draft=False):
    report={'plataforma':'Windows','Excel':'no disponible','Word':'no disponible','Outlook':'no comprobado'}
    if os.name!='nt':raise ValueError('Ejecuta esta comprobación en Windows con Office instalado.')
    from nurus.personal.reports import _book
    from nurus.personal.word_preview import export_pdf
    from docx import Document
    from win32com.client import DispatchEx
    import pythoncom
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    excel_file=folder/'Nomina_prueba.xlsx';_book([('Nómina',['RIT','NOMBRE','OBSERVACION'],[['001','Persona ficticia','=Texto literal']])],excel_file)
    pythoncom.CoInitialize();excel=None;book=None
    try:
        excel=DispatchEx('Excel.Application');excel.Visible=False;excel.DisplayAlerts=False;excel.AutomationSecurity=3
        book=excel.Workbooks.Open(str(excel_file.resolve()),UpdateLinks=0,ReadOnly=True)
        assert book.Worksheets(1).Cells(2,1).Value=='001'
        assert book.Worksheets(1).Cells(2,3).Value=='=Texto literal'
        report['Excel']='abierto, valores y literalidad comprobados'
    except Exception as exc:report['Excel']='no verificado: '+str(exc)
    finally:
        try:
            if book is not None:book.Close(False)
        finally:
            try:
                if excel is not None:excel.Quit()
            finally:pythoncom.CoUninitialize()
    doc=Document();doc.add_heading('Documento ficticio de comprobación',1);doc.add_paragraph('Persona ficticia. '+'Texto largo de revisión. '*30)
    source=folder/'Word_prueba.docx';doc.save(source)
    try:
        pdf=Path(export_pdf(source,folder/'Word_prueba.pdf'));assert pdf.read_bytes().startswith(b'%PDF-')
        report['Word']='abierto y convertido por Word a PDF'
    except Exception as exc:report['Word']='no verificado: '+str(exc)
    if outlook_draft:
        from nurus.personal.manual import ManualContext
        from nurus.personal.config import defaults
        from nurus.personal.outputs import Draft,create_draft
        # Creates a draft with a reserved test address; never sends it.
        context=ManualContext(defaults(),folder/'prueba_manual.json')
        draft=Draft('PRUEBA FICTICIA CSMP — revisar y eliminar','Buen día:\n\nPrueba de presentación y firma. No enviar.',to='prueba@example.test',attachments=[str(excel_file)])
        try:
            result=create_draft(context,draft,confirmed=True)
            report['Outlook']={'estado':'borrador guardado; revisar firma y adjunto en Outlook','entry_id':result.entry_id,'store_id':result.store_id}
        except Exception as exc:report['Outlook']='no verificado: '+str(exc)
    return report


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,default=Path('Aceptacion_Office.json'));p.add_argument('--outlook-draft',action='store_true',help='Crear un borrador ficticio para revisar firma y adjunto; no envía correos.')
    args=p.parse_args()
    with tempfile.TemporaryDirectory(prefix='CSMP-aceptacion-') as folder:
        result=verify(folder,args.outlook_draft)
        args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2))
