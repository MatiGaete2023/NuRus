from pathlib import Path


def replace_once(path, old, new):
    path=Path(path)
    text=path.read_text(encoding='utf-8')
    if text.count(old)!=1:
        raise SystemExit(f'{path}: expected one source block, found {text.count(old)}')
    path.write_text(text.replace(old,new,1),encoding='utf-8')


replace_once('src/nurus/personal/outputs.py', '''def _table(work,rows,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book=Workbook();sheet=book.active;sheet.title='Nómina'
    headers=['RIT','TRIBUNAL','RUT','NOMBRE','PROGRAMA','VENCIMIENTO / ESPERA','OBSERVACION']
    sheet.append(headers)
    for row in rows:
        sheet.append([value(work,row,k) for k in ('rit','tribunal','rut','nombre','programa')]+[value(work,row,'vencimiento') or value(work,row,'espera'),str(row.review.get('OBSERVACION',row.observation))])
        for cell in sheet[sheet.max_row]:cell.data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True)
    for col in ('A','B','C','D','E','F','G'):sheet.column_dimensions[col].width=25 if col!='G' else 65
    write_new_file(Path(path),book.save)
    return str(path)
''', '''def _table(work,rows,path,kind=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='Nómina'
    if kind=='programa_espera':
        headers=['RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA']
        keys=('rit','tribunal','nombre','programa','espera')
    elif kind in {'programa_vencido','programa_por_vencer'}:
        headers=['RIT','TRIBUNAL','NOMBRE','DERIVACION','F. VENCIMIENTO']
        keys=('rit','tribunal','nombre','programa','vencimiento')
    else:
        headers=['RIT','TRIBUNAL','RUT','NOMBRE','PROGRAMA','VENCIMIENTO / ESPERA','OBSERVACION']
        keys=None
    sheet.append(headers)
    for row in rows:
        if keys:
            sheet.append([value(work,row,key) for key in keys])
        else:
            sheet.append([value(work,row,k) for k in ('rit','tribunal','rut','nombre','programa')]+[value(work,row,'vencimiento') or value(work,row,'espera'),str(row.review.get('OBSERVACION',row.observation))])
        for cell in sheet[sheet.max_row]:cell.data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True)
    for index,header in enumerate(headers,1):
        sheet.column_dimensions[get_column_letter(index)].width=38 if header=='NOMBRE' else 24
    write_new_file(Path(path),book.save)
    return str(path)
''')

replace_once('src/nurus/personal/outputs.py', "draft.attachments.append(_table(work,subset,target/(program_filename(name)+'.xlsx')))", "draft.attachments.append(_table(work,subset,target/(program_filename(name)+'.xlsx'),kind))")

replace_once('src/nurus/personal/importing.py', "'RES':('res','resolucion generada')", "'RES':('res','resolucion generada','resolucion','generar resolucion','proyecto resolucion','proyecto de resolucion')")

replace_once('src/nurus/personal/resolutions.py', '''def resolution_marked(value):
    """Interpreta la columna humana RES sin convertir valores dudosos en proyectos."""
''', '''def resolution_kind(value):
    text=normalize(value).replace('_',' ')
    return {'pc ie':'PC_IE','pc info':'PC_INFO','nomencl':'NOMENCL','nomenclatura':'NOMENCL'}.get(text)

def _case_key(work,row):
    court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
    return normalize(court),normalize(value(work,row,'rit'))

def resolution_marked(value):
    """Interpreta la columna humana RES sin convertir valores dudosos en proyectos."""
    if resolution_kind(value):return True
''')

replace_once('src/nurus/personal/resolutions.py', '''def automatic_project_selections(work,fallback_kind='PC_IE'):
    """Proyectos automáticos respetando RES cuando la planilla revisada lo utiliza."""
    reviewed=reviewed_resolution_ids(work);result=[]
    for row in work.rows:
        if row.excluded:continue
        kinds=[kind for kind in row.actions if kind in KINDS]
        if reviewed is not None:
            if row.id not in reviewed:continue
            if not kinds:kinds=[fallback_kind]
        for kind in kinds:result.append((row.id,kind))
    return result
''', '''def automatic_project_selections(work,fallback_kind='PC_IE'):
    """Proyectos automáticos respetando RES y tratando cada tribunal/RIT como una causa."""
    reviewed=reviewed_resolution_ids(work)
    result=[]
    if reviewed is None:
        for row in work.rows:
            if row.excluded:continue
            for kind in (kind for kind in row.actions if kind in KINDS):result.append((row.id,kind))
        return result

    by_id={row.id:row for row in work.rows}
    case_kinds=OrderedDict()
    for rid in reviewed:
        row=by_id.get(rid)
        if not row or row.excluded:continue
        explicit=resolution_kind(row.review.get('RES'))
        kinds=[explicit] if explicit else [kind for kind in row.actions if kind in KINDS]
        if not kinds:kinds=[fallback_kind]
        bucket=case_kinds.setdefault(_case_key(work,row),[])
        for kind in kinds:
            if kind not in bucket:bucket.append(kind)
    for row in work.rows:
        if row.excluded:continue
        for kind in case_kinds.get(_case_key(work,row),[]):result.append((row.id,kind))
    return result
''')

replace_once('src/nurus/personal/resolutions.py', '''def prepare_projects(work,selections,template_dir):
    work.refresh()
    by_id={r.id:r for r in work.rows};groups=OrderedDict()
    for rid,kind in selections:
        if rid not in by_id or kind not in KINDS:continue
        row=by_id[rid]
        if row.excluded:continue
        court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
        rit=value(work,row,'rit').strip()
        key=(court,normalize(rit),kind)
        group=groups.setdefault(key,[])
        if row.id not in {r.id for r in group}:group.append(row)
''', '''def prepare_projects(work,selections,template_dir):
    work.refresh()
    by_id={r.id:r for r in work.rows};groups=OrderedDict();case_rows=OrderedDict()
    for row in work.rows:
        if row.excluded:continue
        court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
        rit=value(work,row,'rit').strip()
        case_rows.setdefault((court,normalize(rit)),[]).append(row)
    for rid,kind in selections:
        if rid not in by_id or kind not in KINDS:continue
        row=by_id[rid]
        if row.excluded:continue
        court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
        rit=value(work,row,'rit').strip()
        key=(court,normalize(rit),kind)
        group=groups.setdefault(key,[])
        existing={r.id for r in group}
        for candidate in case_rows.get((court,normalize(rit)),[]):
            if candidate.id not in existing:group.append(candidate);existing.add(candidate.id)
''')

test_path=Path('tests/test_personal_batch_revision.py')
test=test_path.read_text(encoding='utf-8')
anchor="""    all_drafts=prepare_drafts(work,'programa_espera')
    assert len(all_drafts)==2
"""
addition="""

def test_waiting_attachment_contains_only_operational_columns(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION'])
    sheet.append(['X-10-2026','MULCHEN','Persona Espera','11111111-1','AFT EJEMPLO',72,'Texto interno que no debe salir'])
    path=tmp_path/'espera_adj.xlsx';book.save(path)
    work=Work.external(path,defaults(),mode='ESPERA')
    drafts=prepare_drafts(work,'programa_espera',directory=tmp_path/'salida')
    assert len(drafts)==1 and len(drafts[0].attachments)==1
    out=load_workbook(drafts[0].attachments[0]).active
    assert [c.value for c in out[1]]==['RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA']
    assert out.max_column==5
    assert [out.cell(2,i).value for i in range(1,6)]==['X-10-2026','MULCHEN','Persona Espera','AFT EJEMPLO','72']


def test_due_report_attachment_contains_only_operational_columns(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Informes'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','FECHA VENCIMIENTO','OBSERVACION'])
    sheet.append(['X-20-2026','LAJA','Persona Informe','22222222-2','PRM EJEMPLO','30/09/2026','Texto interno que no debe salir'])
    path=tmp_path/'informes_adj.xlsx';book.save(path)
    work=Work.external(path,defaults(),mode='INFORMES')
    drafts=prepare_drafts(work,'programa_por_vencer',directory=tmp_path/'salida')
    assert len(drafts)==1 and len(drafts[0].attachments)==1
    out=load_workbook(drafts[0].attachments[0]).active
    assert [c.value for c in out[1]]==['RIT','TRIBUNAL','NOMBRE','DERIVACION','F. VENCIMIENTO']
    assert out.max_column==5
    assert [out.cell(2,i).value for i in range(1,6)]==['X-20-2026','LAJA','Persona Informe','PRM EJEMPLO','30/09/2026']
"""
if test.count(anchor)!=1:raise SystemExit('test mail anchor not unique')
test=test.replace(anchor,anchor+addition,1)
anchor2="""    assert automatic_project_selections(work,'PC_INFO')==[(work.rows[0].id,'PC_INFO')]
"""
addition2="""

def test_resolution_mark_on_one_row_recognizes_complete_rit(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Registros'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','OBSERVACION','GENERAR RESOLUCIÓN'])
    sheet.append(['X-149-2025','MULCHEN','SOFÍA IGNACIA DAROCH VERDUGO','24010461-K','AFT MULCHEN','Texto','1'])
    sheet.append(['X-149-2025','MULCHEN','MIA VALENTINA DAROCH VERDUGO','24624826-5','AFT MULCHEN','Texto',''])
    sheet.append(['X-150-2025','MULCHEN','Otra Persona','11111111-1','AFT MULCHEN','Texto',''])
    path=tmp_path/'res_por_causa.xlsx';book.save(path)
    work=Work.external(path,defaults())
    selected=automatic_project_selections(work,'PC_IE')
    assert {rid for rid,kind in selected}=={work.rows[0].id,work.rows[1].id}
    assert {kind for rid,kind in selected}=={'PC_IE'}
    projects,errors=prepare_projects(work,selected,BASE/'plantillas_word')
    assert not errors and len(projects)==1
    assert set(projects[0].record_ids)=={work.rows[0].id,work.rows[1].id}
    assert 'SOFÍA IGNACIA DAROCH VERDUGO' in projects[0].text
    assert 'MIA VALENTINA DAROCH VERDUGO' in projects[0].text
"""
if test.count(anchor2)!=1:raise SystemExit('test resolution anchor not unique')
test_path.write_text(test.replace(anchor2,anchor2+addition2,1),encoding='utf-8')
