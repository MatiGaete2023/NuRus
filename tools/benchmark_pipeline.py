"""Medición reproducible con datos sintéticos; no abre Excel ni Outlook.

Ejecutar desde la raíz: python tools/benchmark_pipeline.py --output resultado.json
Las duraciones son observaciones, nunca un requisito temporal de CI.
"""
from __future__ import annotations

import argparse
from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import platform
from statistics import median
import sys
import tempfile
from time import perf_counter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import pandas as pd
from openpyxl import Workbook, load_workbook
from nurus import __version__
from nurus.rus import Mode, evaluate_batch, read_workbook
from nurus.rus.reader import _records
from nurus.services.exports import export_proposal_workbook
from nurus.storage.database import Database


def measure(count, folder):
    headers = ["RIT", "TRIBUNAL", "NOMBRE", "DERIVACION", "T ESPERA", "RUT"]
    rows = [[f"X-{i+1}", "Juzgado de Laja", f"Persona {i+1}",
             "CESFAM" if i % 20 == 0 else "PRM Ejemplo", 70, f"{i+1}-X"]
            for i in range(count)]
    mapping = dict(zip(("rit", "tribunal", "nombre", "programa", "espera", "rut"), headers))
    frame = pd.DataFrame(rows, columns=headers, dtype=object)
    micro = []
    for _ in range(5):
        start = perf_counter()
        converted = _records(frame, "fixture.xlsx", "hash", "Espera", 1, mapping)
        micro.append(perf_counter() - start)
    assert len(converted) == count
    assert converted[-1].source.row_number == count + 1
    source = folder / "entrada.xlsx"
    book = Workbook()
    book.active.title = "Espera"
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    book.active.auto_filter.ref = book.active.dimensions
    book.create_sheet("OB")["A1"] = "=1+2"
    book.save(source)
    book.close()
    start = perf_counter()
    batch = read_workbook(source, Mode.ESPERA)
    read_seconds = perf_counter() - start
    assert len(batch.records) == count
    assert batch.workbook_sha256 == sha256(source.read_bytes()).hexdigest()
    start = perf_counter()
    result = evaluate_batch(batch, as_of=date(2026, 9, 8))
    evaluate_seconds = perf_counter() - start
    db = Database(folder / "work.sqlite3")
    start = perf_counter()
    db.save_evaluation_batch(result, source_bytes=batch.source_bytes)
    persist_seconds = perf_counter() - start
    start = perf_counter()
    exported = export_proposal_workbook(db, result.batch_id, folder / "propuesta.xlsx",
                                        backend="portable", allow_reduced_fidelity=True)
    export_seconds = perf_counter() - start
    generated = load_workbook(exported.path, read_only=True, data_only=False)
    assert generated["Espera"].max_row == count + 1
    assert generated["OB"]["A1"].value == "=1+2"
    generated.close()
    start = perf_counter()
    restored = Database(db.path)
    assert len(restored.list_review_records(result.batch_id)) == count
    assert restored.get_original_workbook(result.batch_id) == source.read_bytes()
    reopen_seconds = perf_counter() - start
    return dict(rows=count, conversion_median_seconds=median(micro), read_seconds=read_seconds,
                evaluate_seconds=evaluate_seconds, persist_seconds=persist_seconds,
                export_portable_seconds=export_seconds, reopen_seconds=reopen_seconds,
                total_pipeline_seconds=sum((read_seconds, evaluate_seconds, persist_seconds,
                                            export_seconds, reopen_seconds)), verified=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--sizes", nargs="+", type=int, default=[100, 1000, 10000])
    args = parser.parse_args()
    if any(size < 1 or size > 100000 for size in args.sizes):
        parser.error("sizes debe estar entre 1 y 100000")
    report = dict(version=__version__, python=platform.python_version(), platform=platform.platform(),
                  input="sintético", office_tested=False, results=[])
    for size in args.sizes:
        with tempfile.TemporaryDirectory(prefix="nurus-benchmark-") as root:
            result = measure(size, Path(root))
            report["results"].append(result)
            print(json.dumps(result), flush=True)
    with args.output.open("x", encoding="utf-8") as output:
        json.dump(report, output, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
