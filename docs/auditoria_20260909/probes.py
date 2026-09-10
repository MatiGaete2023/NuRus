"""Sondas de auditoría; no modifica originales ni crea mensajes de Outlook.

Desde la raíz del repositorio:
PYTHONPATH=src:tests python docs/auditoria_20260909/probes.py --sources /ruta/upload
Los defectos se registran, no se ocultan como pruebas de aceptación aprobadas.
"""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import date
import hashlib
import json
from pathlib import Path
import tempfile
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from nurus.rus import Mode, read_workbook, evaluate_batch
from nurus.services.workflow import WorkController
from nurus.services.exports import export_preserved_workbook
from nurus.storage.database import Database
from test_preserved_workbook_export import _approved_excluded

TODAY = date(2026, 9, 9)
HEADERS = ['DERIVACION', 'TRIBUNAL', 'NOMBRE', 'RIT', 'DIAS DE CUMPLIMIENTO',
           'DIAS PARA EGRESAR', 'FEC.INGRESO EFECTIVO', 'FEC.EGRESO PROYECTADO']

def compliance(path, *, program='PRM Norte', days=100, exit_days=100,
               exit_date=date(2026, 12, 18), broken_cross=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'CUMPLIMIENTO'
    sheet.append(HEADERS)
    sheet.append([program, 'Juzgado de Laja', 'Persona ficticia', 'C-SINTETICO',
                  days, exit_days, date(2025, 1, 1), exit_date])
    if broken_cross:
        cross = book.create_sheet('Hoja2')
        cross.append(['nota'])
        cross.append(['sin columnas de cruce'])
    book.save(path)
    book.close()

def run(sources):
    result = {'audited_commit': 'afd82b80ae69375e1d2f6fc03f78c081b0bab45c',
              'as_of': TODAY.isoformat(), 'synthetic': {}, 'provided_files': []}
    with tempfile.TemporaryDirectory(prefix='nurus-audit-') as folder:
        root = Path(folder)
        path = root / 'cross.xlsx'
        compliance(path, broken_cross=True)
        db = Database(root / 'cross.sqlite3')
        controller = WorkController(db)
        batch = controller.analyze(path, Mode.CUMPLIMIENTO, as_of=TODAY)
        try:
            controller.approve_batch(batch.batch_id)
            approved, error = True, None
        except ValueError as exc:
            approved, error = False, str(exc)
        result['synthetic']['invalid_cross'] = {
            'warnings': batch.warnings, 'approved_without_exception': approved,
            'exception_recorded': db.get_batch_exception(batch.batch_id) is not None,
            'error': error}
        for name, options in (
            ('missing_projected_exit', {'exit_days': 10, 'exit_date': None}),
            ('future_marked_expired', {'days': -1}),
            ('absent_fae_column', {'program': 'FAE Norte'}),
        ):
            path = root / (name + '.xlsx')
            compliance(path, **options)
            batch = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=TODAY)
            item = batch.evaluations[0]
            result['synthetic'][name] = {
                'status': item.status.value, 'rule_ids': item.rule_ids,
                'issues': [i.code for i in item.issues], 'observation': item.observation}

        db = Database(root / 'export.sqlite3')
        batch, _ = _approved_excluded(db, root / 'source.xlsx')
        output = root / 'export.xlsx'
        export_preserved_workbook(db, batch.batch_id, output,
                                 backend='portable', allow_reduced_fidelity=True)
        book = load_workbook(output, data_only=False)
        trace = book['NURUS_TRAZABILIDAD']
        result['synthetic']['trace_formula'] = {
            'new_formula_cells': [cell.coordinate for row in trace for cell in row
                                  if cell.data_type == 'f'],
            'visible_observation': book['Espera']['F2'].value}
        book.close()

        partial = root / 'partial.xlsx'
        def fail_copy(source, output, *args, **kwargs):
            output.write(b'PK-partial')
            raise OSError('synthetic disk failure')
        try:
            # Aislar la copia final: shutil también puede usarse al serializar XLSX.
            with patch('nurus.services.exports._portable_preserved',
                       side_effect=lambda content, target, snapshot: target.write_bytes(output.read_bytes())), \
                 patch('nurus.services.exports.shutil.copyfileobj', fail_copy):
                export_preserved_workbook(db, batch.batch_id, partial,
                                         backend='portable', allow_reduced_fidelity=True)
        except Exception as exc:
            result['synthetic']['partial_export'] = {
                'error_type': type(exc).__name__, 'destination_remains': partial.exists(),
                'bytes': partial.stat().st_size if partial.exists() else 0}

        path = root / 'displaced.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'ESPERA'
        for _ in range(13):
            sheet.append(['Portada'])
        sheet.append(['DERIVACION', 'TRIBUNAL', 'NOMBRE', 'RIT', 'T ESPERA'])
        for index in range(60):
            sheet.append(['PRM Norte', 'Laja', 'Persona ficticia', str(index), 40])
        book.save(path)
        book.close()
        with patch('pandas.read_excel', wraps=pd.read_excel) as spy:
            batch = read_workbook(path, Mode.ESPERA)
            result['synthetic']['header_reads'] = {
                'read_excel_calls': spy.call_count, 'header_row': batch.header_row,
                'first_source_row': batch.records[0].source.row_number,
                'records': len(batch.records)}

    if sources:
        for filename, modes in (
            ('AGOSTO.xlsx', list(Mode)),
            ('RUS_CUMPLIMIENTO_20260904_114132_856292.xlsx', [Mode.CUMPLIMIENTO]),
        ):
            path = sources / filename
            if not path.exists():
                result['provided_files'].append({'file': filename, 'available': False})
                continue
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            for mode in modes:
                entry = {'file': filename, 'sha256': digest, 'mode': mode.value}
                try:
                    read = read_workbook(path, mode)
                    evaluated = evaluate_batch(read, as_of=TODAY)
                    entry.update(header_row=read.header_row, sheet=read.primary_sheet,
                                 rows=len(read.records), warnings=evaluated.warnings,
                                 status_counts=dict(Counter(x.status.value for x in evaluated.evaluations)),
                                 issue_counts=dict(Counter(i.code for x in evaluated.evaluations for i in x.issues)))
                except Exception as exc:
                    entry.update(error_type=type(exc).__name__, error=str(exc))
                result['provided_files'].append(entry)
    return result

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--sources', type=Path)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    encoded = json.dumps(run(args.sources), ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(encoded + '\n', encoding='utf-8')
    print(encoded)
