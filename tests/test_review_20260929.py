from copy import deepcopy
from pathlib import Path
from queue import Queue
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import Workbook, load_workbook

from nurus.personal.app_base import App as BaseApp
from nurus.personal.config import defaults
from nurus.personal.mail_controls import empty_preparation_message
from nurus.personal.outputs import Draft, create_draft, prepare_drafts, word_values
from nurus.personal.personalization import duplicate_template, save_contact
from nurus.personal.product_state import stale
from nurus.personal.session import merge_draft_edits
from nurus.personal.work import Row, Work
from test_audit_integrity_20260928 import source, edit
from test_personal_usability_20260921 import Var


def test_duplicate_program_template_preserves_destination_rules_and_attachment_columns(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    work.rows[1].actions=[]
    work.config,key=duplicate_template(work.config,'programa_espera','Mi espera')
    work.config,key=duplicate_template(work.config,key,'Copia de copia')
    original=prepare_drafts(work,'programa_espera')[0]
    duplicated=prepare_drafts(work,key)[0]
    assert duplicated.kind==key
    assert duplicated.recipient_type==original.recipient_type=='programas'
    assert duplicated.record_ids==original.record_ids==[work.rows[0].id]
    book=load_workbook(duplicated.attachments[0])
    assert list(book.active.values)[0]==('RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA')
    work.rows[0].decisions['mail']['programa_espera']='omit'
    assert prepare_drafts(work,key)==[]


def test_duplicate_informative_keeps_court_destination(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    work.config,key=duplicate_template(work.config,'espera','Mi revisión')
    draft=prepare_drafts(work,key,modalities='todas')[0]
    assert draft.recipient_type=='tribunales'
    assert draft.to=='; '.join(work.config['correos']['tribunales']['MULCHEN']['para'])


def test_corrected_measure_due_and_resolution_date_are_used_in_products(tmp_path):
    book=Workbook();sheet=book.active
    sheet.append(['RIT','TRIBUNAL','NOMBRE','PROGRAMA','FEC. EGRESO PROYECTADO','FEC. RESOLUCION','OBSERVACION'])
    sheet.append(['X-1','MULCHEN','Persona','RTA - Prueba','10/10/2026','01/09/2026','Revisado'])
    path=tmp_path/'cumplimiento.xlsx';book.save(path)
    work=Work.external(path,defaults(),'CUMPLIMIENTO')
    work.rows[0].overrides={'egreso_proy':'15/10/2026','resolucion':'02/09/2026'}
    draft=prepare_drafts(work,'medidas',modalities='Residencial',manual_selection=True)[0]
    assert draft.due=='2026-10-15'
    attachment=load_workbook(draft.attachments[0])
    assert attachment.active.cell(2,6).value=='15/10/2026'
    assert word_values(work,work.rows[0])['FECHA_RESOLUCION']=='dos de septiembre de dos mil veintiséis'


def test_mail_dependencies_cover_whole_canonical_court_group(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    draft=prepare_drafts(work,'espera',modalities='todas')[0]
    work.rows.append(Row('other',50,{'TRIBUNAL':'LAJA','RIT':'X-99'},'',[],[],[]))
    assert not stale(work,draft)
    work.rows.append(Row('same',51,{'TRIBUNAL':'Jgdo. L. y G. de Mulchén','RIT':'X-100'},'',[],[],[]))
    assert stale(work,draft)


def test_partial_attachment_removal_and_manual_addition_survive_regeneration():
    original=['old/a.xlsx','old/b.xlsx','old/c.xlsx']
    old=Draft('Asunto','Texto',kind='medidas',court='Mulchén',
              attachments=['old/a.xlsx','old/c.xlsx','manual.pdf'],original={'attachments':original})
    new=Draft('Asunto','Texto',kind='medidas',court='Mulchén',
              attachments=['new/a.xlsx','new/b.xlsx','new/c.xlsx'])
    assert merge_draft_edits([old],[new])[0].attachments==['new/a.xlsx','new/c.xlsx','manual.pdf']


def test_sync_recalculates_against_new_valid_header_mapping(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    assert 'programa_espera' in work.rows[0].actions
    def change(book):
        sheet=book['Espera'];sheet['F1']='DIAS_ESPERA';sheet['F2']=5
    edit(work,change)
    assert work.refresh()
    assert work.mapping['espera']=='DIAS_ESPERA'
    assert 'programa_espera' not in work.rows[0].actions
    assert 'programa_espera' in work.rows[1].actions


def test_contact_name_cannot_hide_another_contact_alias():
    config=defaults();config['contactos']={'Primero':'uno@example.cl'};config['aliases']={'Alias':'Primero'}
    before=deepcopy(config)
    with pytest.raises(ValueError,match='alias'):
        save_contact(config,'ÁLIAS','dos@example.cl')
    assert config==before
    updated=save_contact(config,'Alias','uno@example.cl',['Alias'],previous='Primero')
    assert updated['aliases']=={'Alias':'Alias'}


def test_receipt_metadata_and_activity_date_survive_session_reload(tmp_path,monkeypatch):
    work=source(tmp_path/'origen.xlsx')
    work.save(tmp_path/'session')
    monkeypatch.setattr('nurus.personal.outputs.save_draft',lambda *a,**k:SimpleNamespace(entry_id='e',store_id='s'))
    draft=prepare_drafts(work,'espera',modalities='todas')[0]
    create_draft(work,draft,confirmed=True)
    receipt=work.receipts[draft.key]
    assert receipt['record_ids']==draft.record_ids and receipt['subject']==draft.subject
    assert receipt['created_at']
    recovered=Work.load(tmp_path/'session')
    assert recovered.receipts==work.receipts


def test_uncertain_outlook_save_preserves_activity_metadata(tmp_path,monkeypatch):
    from nurus.adapters.outlook import DraftSaveUncertain
    work=source(tmp_path/'origen.xlsx')
    draft=prepare_drafts(work,'espera',modalities='todas')[0]
    def uncertain(*a,**k):raise DraftSaveUncertain('incierto')
    monkeypatch.setattr('nurus.personal.outputs.save_draft',uncertain)
    with pytest.raises(DraftSaveUncertain):create_draft(work,draft,confirmed=True)
    item=work.receipts[draft.key]
    assert item['state']=='uncertain' and item['created_at'] and item['subject']==draft.subject


def test_close_keeps_window_open_when_recovery_cannot_be_written(monkeypatch):
    def denied():raise PermissionError('Sin permiso')
    app=SimpleNamespace(busy=False,_save_text=Mock(),_save_tpl=Mock(),_save_session=denied,destroy=Mock())
    errors=[];monkeypatch.setattr('tkinter.messagebox.showerror',lambda *a:errors.append(a))
    BaseApp._close(app)
    app.destroy.assert_not_called()
    assert errors and 'Sin permiso' in errors[0][1]


def test_session_write_error_is_propagated_with_recovery_data_kept_in_memory(tmp_path):
    work=SimpleNamespace(save=Mock(side_effect=PermissionError('Sin permiso')))
    app=SimpleNamespace(work=work,busy=False,_capture_observation=Mock(),_capture_mail=Mock(),
                        _capture_project=Mock(),drafts=[],projects=[],last_word='',
                        cfg=SimpleNamespace(directory=tmp_path),_update_local_activity=Mock(),status=Var(''))
    with pytest.raises(PermissionError):BaseApp._save_session(app)
    assert work.session['version']==2
    app._update_local_activity.assert_not_called()


def test_poll_is_rescheduled_even_if_recovery_fails(monkeypatch):
    events=Queue();events.put((False,ValueError('Error del trabajo'),None))
    def denied():raise PermissionError('Sin permiso')
    app=SimpleNamespace(events=events,progress=Mock(),status=Var(''),_save_session=denied,
                        _poll=Mock(),after=Mock(),_locked_inputs=[])
    monkeypatch.setattr('tkinter.messagebox.showerror',lambda *a:None)
    BaseApp._poll(app)
    app.after.assert_called_once_with(100,app._poll)
    assert 'Sin permiso' in app.status.get()


def test_empty_preparation_explains_modality_and_exclusion(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    ids=[row.id for row in work.rows]
    assert 'modalidades' in empty_preparation_message(work,ids,['FAE'],'espera')
    for row in work.rows:row.excluded=True
    assert 'excluidas' in empty_preparation_message(work,ids,['FAE'],'espera')


@pytest.mark.parametrize('invalid', ['address','attachment'])
def test_invalid_editable_draft_is_rejected_before_outlook(tmp_path,monkeypatch,invalid):
    work=source(tmp_path/'origen.xlsx')
    draft=prepare_drafts(work,'espera',modalities='todas')[0]
    if invalid=='address':draft.to='sin arroba'
    else:draft.attachments=[str(tmp_path/'ausente.xlsx')]
    save=Mock();monkeypatch.setattr('nurus.personal.outputs.save_draft',save)
    with pytest.raises(ValueError):create_draft(work,draft,confirmed=True)
    save.assert_not_called()
    assert not work.receipts
