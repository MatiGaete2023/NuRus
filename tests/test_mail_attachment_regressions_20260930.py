from copy import deepcopy
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from nurus.personal.config import defaults
from nurus.personal.outputs import prepare_drafts
from nurus.personal.work import Work


def source(tmp_path,mode='INFORMES',due_header='FECHA VENCIMIENTO',due=datetime(2026,10,15),waiting=0):
    book=Workbook();sheet=book.active
    sheet.append(['RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA',due_header,'OBSERVACION'])
    sheet.append(['X-1','MULCHËN','Persona Prueba','PRM PRUEBA',waiting,due,'Revisado'])
    path=tmp_path/'origen.xlsx';book.save(path)
    return Work.external(path,defaults(),mode)


@pytest.mark.parametrize('header',['F. VENCIMIENTO','F.VENCIMIENTO','FVENCIMIENTO','FECHA DE VENCIMIENTO'])
def test_due_aliases_are_read_in_import_and_generated_attachment(tmp_path,header):
    work=source(tmp_path,due_header=header)
    assert work.mapping.get('vencimiento')==header
    draft=prepare_drafts(work,'programa_por_vencer')[0]
    assert load_workbook(draft.attachments[0]).active['E2'].value=='15/10/2026'
    assert draft.due=='2026-10-15'


@pytest.mark.parametrize('mode,kind,expected',[('ESPERA','programa_por_vencer','15/10/2026'),
                                              ('INFORMES','programa_espera','0')])
def test_other_mail_type_reads_its_metric_from_same_sheet_without_changing_mode(tmp_path,mode,kind,expected):
    work=source(tmp_path,mode)
    draft=prepare_drafts(work,kind,manual_selection=True)[0]
    assert load_workbook(draft.attachments[0]).active['E2'].value==expected
    assert work.mode==mode


@pytest.mark.parametrize('kind,header,expected',[('informes','F. VENCIMIENTO','15/10/2026'),
                                               ('espera','T ESPERA','0')])
def test_informative_attachments_label_and_use_only_the_correct_metric(tmp_path,kind,header,expected):
    work=source(tmp_path,'ESPERA')
    work.config['correos']['plantillas'][kind]['adjunto']='obligatorio'
    draft=prepare_drafts(work,kind,modalities='Todas')[0]
    out=load_workbook(draft.attachments[0]).active
    assert out['F1'].value==header and out['F2'].value==expected
    for row in out.iter_rows():
        for cell in row:
            for side in ('left','right','top','bottom'):
                border=getattr(cell.border,side)
                assert border.style=='thin' and border.color.rgb=='FF000000'


def test_missing_due_is_explicit_and_never_replaced_by_waiting_time(tmp_path):
    work=source(tmp_path,due=None,waiting=45)
    work.config['correos']['plantillas']['informes']['adjunto']='obligatorio'
    draft=prepare_drafts(work,'informes',modalities='Todas')[0]
    assert load_workbook(draft.attachments[0]).active['F2'].value=='Sin dato'
    assert draft.due==''


def test_metric_uses_correction_including_explicit_empty_and_zero(tmp_path):
    work=source(tmp_path)
    work.rows[0].overrides={'vencimiento':'20/10/2026','espera':0}
    draft=prepare_drafts(work,'programa_por_vencer')[0]
    assert load_workbook(draft.attachments[0]).active['E2'].value=='20/10/2026'
    work.rows[0].overrides['vencimiento']=''
    empty=prepare_drafts(work,'programa_por_vencer')[0]
    assert load_workbook(empty.attachments[0]).active['E2'].value=='Sin dato'
    assert empty.due==''


def test_legacy_mapping_with_two_due_candidates_is_rejected_without_mutation(tmp_path):
    from nurus.rus.columns import ColumnMappingError
    work=source(tmp_path,'ESPERA')
    work.rows[0].values['FEC. VENCIMIENTO']='20/10/2026'
    original=deepcopy(work.rows[0].values)
    with pytest.raises(ColumnMappingError):prepare_drafts(work,'programa_por_vencer')
    assert work.rows[0].values==original


def test_due_outside_active_mode_is_tracked_for_staleness(tmp_path):
    from nurus.personal.product_state import stale
    work=source(tmp_path,'ESPERA')
    draft=prepare_drafts(work,'programa_por_vencer')[0]
    assert not stale(work,draft)
    work.rows[0].values['FECHA VENCIMIENTO']='20/10/2026'
    assert stale(work,draft)


@pytest.mark.parametrize('mode,kind',[('ESPERA','programa_espera'),('INFORMES','programa_por_vencer')])
def test_generated_nomina_can_be_reimported_without_losing_its_metric(tmp_path,mode,kind):
    work=source(tmp_path,mode)
    first=prepare_drafts(work,kind)[0]
    reloaded=Work.external(first.attachments[0],defaults(),mode)
    second=prepare_drafts(reloaded,kind,directory=tmp_path/'nuevo')[0]
    before=load_workbook(first.attachments[0]).active
    after=load_workbook(second.attachments[0]).active
    assert list(before.values)==list(after.values)
    for row in after.iter_rows():
        for cell in row:
            assert all(getattr(cell.border,side).style=='thin' and getattr(cell.border,side).color.rgb=='FF000000'
                       for side in ('left','right','top','bottom'))


def test_empty_whitespace_metric_is_explicit(tmp_path):
    work=source(tmp_path,due='   ')
    draft=prepare_drafts(work,'programa_por_vencer')[0]
    assert load_workbook(draft.attachments[0]).active['E2'].value=='Sin dato'
