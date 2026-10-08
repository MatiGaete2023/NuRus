import base64
from datetime import date
import html
import json
import pytest
from nurus.personal.bitacora_har import read_popup, import_har, export_har_audit, HEADERS
from nurus.personal.bitacoras import analyze, earliest, cc_for_type

PARAMS = {'tipo_popUp':'12','COD_Tribunal':'11','CRR_IdCausa':'22','ID_Ingreso':'33',
          'COD_Etapa':'3','TIP_Consulta':'2'}


def popup(text='Texto completo que la tabla oculta.',response='---',kind='Al tribunal'):
    fields={**PARAMS,'tipo_popUp':'','RIT_Causa':'P-12-2026','GLS_Nombre':'Persona de prueba','GLS_Centro':'Centro de prueba'}
    inputs=''.join(f'<input type="hidden" name="{k}" value="{html.escape(v,quote=True)}">' for k,v in fields.items())
    hidden={'HOBS_Centro_':'','HOBS_Tribunal_':response,'HCOD_Estado_':'6','HFUN_Tribunal_':'USR','HFLG_Envio_':'1'}
    hidden['HOBS_Centro_']=text
    hidden=''.join(f'<input type="hidden" name="{k}44" value="{html.escape(v,quote=True)}">' for k,v in hidden.items())
    headers=''.join(f'<td>{h}</td>' for h in HEADERS)
    values=[kind,'Cumplimiento','12/08/2026 10:13','Centro','Texto recortado','---','---','Recortado']
    row=''.join(f'<td>{v}</td>' for v in values)+'<td><img id="44" onclick="MuestraDiv(this.id,3)"><img id="44"></td>'
    return f'<html><body><form name="InformesPpalForm" action="/SITFAWEB/InformesDAction.do">{inputs}<table id="TablaInforme"><tr>{headers}</tr><tr>{row}</tr>{hidden}</table></form></body></html><script>/* RUS añade un script al final */</script>'


def har(path, texts, *, encoded=False):
    from urllib.parse import urlencode
    entries=[]
    for text in texts:
        content={'text':text}
        if encoded:content={'text':base64.b64encode(text.encode()).decode(),'encoding':'base64'}
        entries.append({'startedDateTime':'2026-10-07T10:00:00-03:00',
            'request':{'method':'GET','url':'https://familia.pjud.cl/SITFAWEB/IrPopUpInformesAccion.do?'+urlencode(PARAMS)},
            'response':{'status':200,'content':content}})
    path.write_text(json.dumps({'log':{'entries':entries}}),encoding='utf-8')
    return path


def test_full_text_response_and_actual_identity():
    text='Línea uno & comillas "exactas"\nLínea dos '*12
    query=read_popup(popup(text),PARAMS)
    entry=query['entries'][0]
    assert entry.texto==text and entry.entry_id=='44'
    assert entry.respuesta=='' and entry.respuesta_comprobada
    assert cc_for_type(entry.tipo)==1
    assert query['coverage']=='PARCIAL'
    assert entry.fecha=='2026-08-12T10:13:00'


def test_identical_captures_not_triple_counted(tmp_path):
    source=har(tmp_path/'capture.har',[popup()]*3,encoded=True)
    captured=import_har([source]);query=captured['queries'][0]
    assert query['captures']==3 and len(query['entries'])==1
    assert captured['sources'][0]['openings']==3


def test_conflicting_snapshots_rejected(tmp_path):
    source=har(tmp_path/'capture.har',[popup('Antes'),popup('Después')])
    with pytest.raises(ValueError,match='difieren'):import_har([source])


@pytest.mark.parametrize('change',[
    lambda t:t.replace('name="ID_Ingreso" value="33"','name="ID_Ingreso" value="99"'),
    lambda t:t.replace('HOBS_Centro_44','HOBS_Centro_55'),
    lambda t:t.replace('Fecha Centro','Fecha desconocida'),
    lambda t:t.replace('</html>',''),
    lambda t:t.replace('<img id="44"></td>','<img id="45"></td>'),
])
def test_unverifiable_content_rejected(change):
    with pytest.raises(ValueError):read_popup(change(popup()),PARAMS)


def test_four_calendar_months_and_full_historical_copy(tmp_path):
    body=popup('=Literal, no fórmula',kind='Administrativa').replace('12/08/2026','04/06/2026')
    captured=import_har([har(tmp_path/'capture.har',[body])])
    today=date(2026,10,7)
    assert earliest(today)==date(2026,6,7)
    result=analyze(captured['queries'][0]['entries'],earliest(today),today,today=today)
    assert not result['entries']
    destination=tmp_path/'output.xlsx'
    export_har_audit(captured,destination,earliest(today),today,today=today)
    from openpyxl import load_workbook
    book=load_workbook(destination)
    assert book['Bitácoras'].max_row==1
    assert book['Copia íntegra']['M2'].value=='=Literal, no fórmula'
    assert book['Copia íntegra']['M2'].data_type=='s'
    assert book['Copia íntegra']['K2'].value==0
    assert book['Copia íntegra']['M2'].border.left.style=='thin'
    book.close()


def test_empty_har_not_success(tmp_path):
    with pytest.raises(ValueError,match='No hay respuestas'):import_har([har(tmp_path/'empty.har',[])])


def test_foreign_origin_rejected(tmp_path):
    source=har(tmp_path/'capture.har',[popup()])
    source.write_text(source.read_text().replace('familia.pjud.cl','example.org'))
    with pytest.raises(ValueError,match='Origen'):import_har([source])
