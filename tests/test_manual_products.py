from copy import deepcopy
from dataclasses import asdict
from pathlib import Path
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
from docx import Document
from nurus.personal.config import defaults
from nurus.personal.manual_products import ManualStore,update_table,generate_table,free_resolution,manual_project
from nurus.personal.outputs import prepare_drafts,create_draft
from nurus.personal.session import merge_draft_edits
from test_mail_attachment_regressions_20260930 import source


def test_deferred_attachment_edits_are_local_and_formula_safe(tmp_path):
    work=source(tmp_path);before=deepcopy(work.rows[0].values);original=Path(work.path).read_bytes()
    draft=prepare_drafts(work,'programa_por_vencer',defer_attachments=True)[0]
    assert draft.attachments==[] and not list(tmp_path.rglob('adjuntos'))
    table=deepcopy(draft.options['table']);table['records'][0]['values'][2]='=Persona revisada'
    table['records'][0]['values'][4]='20/10/2026';update_table(draft,table)
    files=generate_table(draft,tmp_path);book=load_workbook(files[0]);sheet=book.active
    assert sheet['C2'].value=='=Persona revisada' and sheet['C2'].data_type=='s'
    assert sheet['E2'].value=='20/10/2026'
    for cells in sheet:
        for cell in cells:
            assert all(getattr(cell.border,side).style=='thin' for side in ('left','right','top','bottom'))
    book.close();assert work.rows[0].values==before and Path(work.path).read_bytes()==original
    assert not draft.options['pending_table']


def test_pending_review_blocks_outlook_and_empty_nomina(tmp_path,monkeypatch):
    work=source(tmp_path);draft=prepare_drafts(work,'programa_por_vencer',defer_attachments=True)[0]
    monkeypatch.setattr('nurus.personal.outputs.save_draft',lambda **kw:pytest.fail('Outlook no debe invocarse'))
    with pytest.raises(ValueError,match='Revisa los registros'):create_draft(work,draft,confirmed=True)
    with pytest.raises(ValueError,match='Revisa los registros'):generate_table(draft,tmp_path)
    table=deepcopy(draft.options['table']);table['records'][0]['include']=False;update_table(draft,table)
    with pytest.raises(ValueError,match='no tiene registros'):generate_table(draft,tmp_path)


def test_regeneration_preserves_corrections_but_requires_new_review(tmp_path):
    work=source(tmp_path);old=prepare_drafts(work,'programa_por_vencer',defer_attachments=True)[0]
    table=deepcopy(old.options['table']);table['records'][0]['values'][2]='Nombre corregido';update_table(old,table)
    previous=generate_table(old,tmp_path)[0]
    new=prepare_drafts(work,'programa_por_vencer',defer_attachments=True)[0]
    merge_draft_edits([old],[new])
    assert new.options['table']['records'][0]['values'][2]=='Nombre corregido'
    assert previous not in new.attachments and new.options['pending_table']
    assert Path(previous).exists()


def test_manual_mail_recovery_and_receipts_without_excel(tmp_path,monkeypatch):
    store=ManualStore(tmp_path/'manuales',defaults());draft=store.new_mail()
    draft.to='persona@example.test';draft.subject='Consulta manual';draft.body='Texto revisado'
    captured=[]
    def save(product,**kwargs):
        captured.append(product);return SimpleNamespace(entry_id='ficticio',store_id='local')
    monkeypatch.setattr('nurus.personal.outputs.save_draft',save)
    create_draft(store,draft,confirmed=True)
    reloaded=ManualStore(tmp_path/'manuales',defaults())
    assert reloaded.drafts[0].body=='Texto revisado' and len(reloaded.receipts)==1
    assert not list(tmp_path.rglob('*.xlsx')) and len(captured)==1
    with pytest.raises(ValueError,match='ya se creó'):create_draft(reloaded,reloaded.drafts[0],confirmed=True)


def test_free_manual_resolution_and_template_identity(tmp_path):
    path=tmp_path/'libre.docx';free_resolution(path,'Tribunal de prueba','X-1-2026','Texto revisado\nSegunda línea')
    assert [p.text for p in Document(path).paragraphs]==['Tribunal de prueba','RIT: X-1-2026','Texto revisado','Segunda línea']
    with pytest.raises(ValueError,match='ya existe'):free_resolution(path,'Tribunal','X-2','No reemplazar')
    matrix=Path(__file__).parents[1]/'src/nurus/personal/plantillas_word/MULCHEN/PC_IE.docx'
    from nurus.personal.outputs import template_variables
    values={key:'Dato ficticio' for key in template_variables(matrix)}
    project=manual_project(matrix,'MULCHEN','X-1-2026',values)
    assert project.rit=='X-1-2026' and '{{' not in project.text
    from nurus.personal.resolutions import generate_projects
    store=ManualStore(tmp_path/'manuales',defaults())
    generated=generate_projects(store,[project],tmp_path/'matriz.docx')
    assert 'X-1-2026' in '\n'.join(p.text for p in Document(generated).paragraphs)
    with pytest.raises(ValueError,match='no corresponde'):manual_project(matrix,'LAJA','X-1-2026',values)
    with pytest.raises(ValueError,match='Completa los campos'):manual_project(matrix,'MULCHEN','X-1-2026',{})


def test_manual_gui_and_suppressed_workflow_recovery(tmp_path):
    import os,subprocess,sys
    env={**os.environ,'PYTHONPATH':os.pathsep.join(sys.path)}
    code='from pathlib import Path; from test_manual_products import gui_probe; import sys; gui_probe(Path(sys.argv[1]))'
    result=subprocess.run([sys.executable,'-c',code,str(tmp_path)],env=env,capture_output=True,text=True,timeout=45)
    assert result.returncode==0,result.stdout+'\n'+result.stderr


def gui_probe(root):
    from nurus.personal.app import App
    from nurus.personal.config import Configuration
    from nurus.personal.work import Work
    app=App(Configuration(root/'perfil'));app.withdraw()
    try:
        app._new_manual_mail();assert app.work is None and len(app.drafts)==1
        app.subject.set('Correo independiente');app.body.insert('1.0','Texto independiente');app._save_session()
        assert ManualStore(root/'perfil/manuales',defaults()).drafts[0].subject=='Correo independiente'
        app._delete_mail();assert app.drafts==[]
        work=source(root);app.work=work;app._show_work()
        drafts=prepare_drafts(work,'programa_por_vencer',defer_attachments=True)
        app._display_prepared_drafts(drafts,'{count}');app._delete_mail()
        assert not app.drafts
        app._display_prepared_drafts(prepare_drafts(work,'programa_por_vencer',defer_attachments=True),'{count}')
        assert not app.drafts
        saved=Work.load(root/'perfil/sesion');assert saved.dismissed_mail==work.dismissed_mail
        app._restore_mail();app._display_prepared_drafts(prepare_drafts(work,'programa_por_vencer',defer_attachments=True),'{count}')
        assert len(app.drafts)==1
    finally:app.destroy()
