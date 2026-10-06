from pathlib import Path
from dataclasses import asdict
from unittest.mock import Mock
from datetime import date
import json
import zipfile
import pytest
from openpyxl import load_workbook
from docx import Document
from nurus.personal.config import defaults, Configuration
from nurus.personal.work import Work
from nurus.personal.outputs import Draft, create_draft, CC
from nurus.personal.roster import snapshot, generate
from nurus.personal.manual import ManualContext, generate_word
from nurus.personal.current_tools import summary, selected_zip, pending
from nurus.personal.work_tools import quality, deadlines, matching_contacts
from current_fixture import source


def test_save_skips_unchanged_state_but_saves_receipts_and_products(tmp_path,monkeypatch):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA');folder=tmp_path/'session';work.save(folder)
    before=(folder/'trabajo.json').read_bytes()
    import nurus.personal.work as module
    actual=module.atomic_json;calls=[]
    def write(path,data):calls.append(Path(path).name);actual(path,data)
    monkeypatch.setattr(module,'atomic_json',write)
    work.save(folder);assert calls==[]
    work.receipts['critical']={'state':'saving'};work.save(folder);assert 'trabajo.json' in calls
    calls.clear();work.session={'drafts':[asdict(Draft('edit','body'))]};work.save(folder);assert calls
    restored=Work.load(folder);assert restored.receipts['critical']['state']=='saving';assert restored.session['drafts'][0]['subject']=='edit'
    assert before!=(folder/'trabajo.json').read_bytes()


def test_roster_edits_are_in_export_before_mail_and_never_touch_origin(tmp_path):
    path=source(tmp_path);origin=path.read_bytes();work=Work(defaults()).analyze(path,'ESPERA')
    draft=Draft('Roster','Reviewed',to='a@example.test',required=True,roster_reviewed=False)
    draft.roster=[snapshot(work,work.rows,'Programa','espera')]
    with pytest.raises(ValueError,match='nómina'):create_draft(work,draft,confirmed=True)
    rows=draft.roster[0]['rows'];rows[0]['cells'][3]='=Texto editado';rows[1]['include']=False
    files=generate(draft,tmp_path);book=load_workbook(files[0]);assert book.active.max_row==3
    assert book.active['D2'].value=='=Texto editado';assert book.active['D2'].data_type=='s';book.close()
    assert path.read_bytes()==origin
    assert '=Texto editado' in summary([draft]);assert CC in summary([draft])
    json.dumps(asdict(draft))


def test_manual_mail_persists_current_edits_and_critical_receipt_without_work(tmp_path,monkeypatch):
    ctx=ManualContext(defaults(),tmp_path/'manual_actual.json');ctx.current['mail']={'subject':'Actual','body':'Editado'};ctx.save()
    restored=ManualContext(defaults(),ctx.path);assert restored.current['mail']['body']=='Editado'
    receipt=Mock(entry_id='entry',store_id='store')
    def save(*args,**kwargs):
        durable=json.loads(ctx.path.read_text());assert list(durable['receipts'].values())[0]['state']=='saving'
        return receipt
    monkeypatch.setattr('nurus.personal.outputs.save_draft',save)
    draft=Draft('Libre','Texto',to='a@example.test');create_draft(restored,draft,confirmed=True)
    assert list(json.loads(ctx.path.read_text())['receipts'].values())[0]['state']=='created'
    with pytest.raises(ValueError,match='ya se creó'):create_draft(restored,draft,confirmed=True)


def test_manual_word_free_and_matrix_validate_before_writing(tmp_path):
    free=tmp_path/'free.docx';generate_word(free,'Título','Párrafo uno\nPárrafo dos');assert Document(free).paragraphs[-1].text=='Párrafo dos'
    matrix=tmp_path/'matrix.docx';doc=Document();doc.add_paragraph('{{NOMBRE}}').runs[0].bold=True;doc.save(matrix)
    target=tmp_path/'filled.docx'
    with pytest.raises(ValueError,match='NOMBRE'):generate_word(target,'','',str(matrix),{})
    assert not target.exists()
    generate_word(target,'','',str(matrix),{'NOMBRE':'Persona'});doc=Document(target);assert doc.paragraphs[0].text=='Persona';assert doc.paragraphs[0].runs[0].bold
    with pytest.raises(ValueError,match="ya existe"):generate_word(target,'','Otro')


def test_selected_delivery_includes_only_chosen_files_with_duplicate_names(tmp_path):
    a=tmp_path/'one';b=tmp_path/'two';a.mkdir();b.mkdir();(a/'file.txt').write_text('one');(b/'file.txt').write_text('two')
    (a/'excluded.txt').write_text('exclude');target=tmp_path/'task.zip';selected_zip([a/'file.txt',b/'file.txt'],target)
    with zipfile.ZipFile(target) as z:assert len(z.namelist())==2;assert {z.read(n) for n in z.namelist()}=={b'one',b'two'}


def test_quality_and_dates_keep_duplicates_and_date_types_distinct(tmp_path):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA');work.mapping.update(vencimiento='DUE',egreso_proy='EXIT')
    row=work.rows[0];row.values['DUE']='07/10/2026';row.values['EXIT']='10/11/2026';row.values[work.mapping['espera']]='texto'
    assert any(i[2]=='espera' for i in quality(work))
    dates=deadlines(work);assert {i[1] for i in dates}=={'Vencimiento de informe','Egreso proyectado'}


def test_contacts_show_ambiguous_normalized_names_for_explicit_choice():
    cfg=defaults();cfg['contactos']={'PRM uno':'a@example.test','prm UNO':'b@example.test','Malo':'invalid'}
    choices=matching_contacts(cfg,'prm uno');assert len(choices)==2;assert matching_contacts(cfg,'invalid')==[]


def test_favorite_block_never_inserts_missing_or_unknown_variables():
    from nurus.personal.manual import render_block,block_variables
    with pytest.raises(ValueError):render_block('Causa {RIT} de {NOMBRE}',{'RIT':'X-1'})
    with pytest.raises(ValueError):block_variables('{DESCONOCIDA}')
    assert render_block('{RIT}: {NOMBRE}',{'RIT':'X-1','NOMBRE':'Persona'})=='X-1: Persona'


def test_preloaded_workbook_is_analyzed_without_another_read(tmp_path,monkeypatch):
    from nurus.rus.reader import read_workbook
    path=source(tmp_path);batch=read_workbook(path,'ESPERA')
    monkeypatch.setattr('nurus.personal.work.read_workbook',lambda *a,**k:pytest.fail('Lectura repetida'))
    work=Work(defaults()).analyze(path,'ESPERA',batch=batch)
    assert len(work.rows)==3 and work.source_hash==batch.workbook_sha256


def test_reprepared_roster_keeps_explicit_edits_but_requires_review(tmp_path):
    from nurus.personal.session import merge_draft_edits
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA')
    old=Draft('Actual','Texto',kind='espera',court='Laja');old.roster=[snapshot(work,work.rows,'Programa','espera')]
    old.roster[0]['rows'][0]['cells'][3]='Persona corregida';old.roster[0]['rows'][1]['include']=False
    new=Draft('Nuevo','Texto',kind='espera',court='Laja',roster_reviewed=False);new.roster=[snapshot(work,work.rows,'Programa','espera')]
    merged=merge_draft_edits([old],[new])[0]
    assert merged.roster[0]['rows'][0]['cells'][3]=='Persona corregida'
    assert not merged.roster[0]['rows'][1]['include'] and not merged.roster_reviewed
