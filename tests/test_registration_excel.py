from datetime import date, datetime
from hashlib import sha256

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side
import pytest

from nurus.personal.config import defaults
from nurus.personal.registration import Journal, RemoteEntry, attach_receipts, export_journal, intent_from_excel, submit
from nurus.personal.registration_excel import reconcile_excel
from nurus.personal.work import Work
from test_registration import Adapter, Clock


def source(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Registros'
    sheet.append(['RIT', 'TRIBUNAL', 'NOMBRE', 'DERIVACION', 'OBSERVACIONES', 'Fecha observación',
                  'CC', 'TT', 'RES', 'Fórmula', 'SITFA_TRIBUNAL_CODIGO', 'SITFA_CAUSA_RUS',
                  'SITFA_INGRESO_RUS', 'SITFA_PERSONA_RUS', 'SITFA_CENTRO_RUS', 'SITFA_VINCULO_ESTADO',
                  'SITFA_ETAPA_RUS', 'SITFA_ANTIGUO_RUS', 'SITFA_MODALIDAD'])
    for n in range(2):
        sheet.append(['X-1-2026', 'Tribunal distinto de los tres históricos', f'Persona ficticia {n}',
                      'Programa ficticio', f'  Observación {n} del Excel.\nTexto sin modificar.  ', None,
                      None, 0, 'Manual', '=1+2', '666', '10', str(20+n), str(30+n), '40',
                      'Vinculado desde respuesta actual', '3', '1', '1'])
    edge = Side(style='thin', color='000000')
    for row in sheet:
        for cell in row:
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
    sheet['E2'].fill = PatternFill('solid', fgColor='ABCDEF')
    sheet['F2'].number_format = 'dd/mm/yyyy'
    book.create_sheet('Otra')['A1'] = '=SUM(Registros!H2:H3)'
    path = tmp_path / 'registro.xlsx'
    book.save(path)
    book.close()
    return path


def registered(tmp_path, path=None):
    path = path or source(tmp_path)
    work = Work.external(path, defaults(), sheet='Registros')
    clock = Clock()
    data = intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                             send_to_tribunal=True, author_id='usuario-ficticio')
    journal = Journal(tmp_path / 'registro.sqlite', clock=clock)
    operation = journal.prepare(data)['operation_id']
    adapter = Adapter(data, clock)
    submit(journal, operation, adapter)
    return path, work, journal, operation, adapter


def portable(journal, source, target, **kw):
    return reconcile_excel(journal, source, target, sheet='Registros', backend='portable',
                           reduced_fidelity=True, **kw)


def test_excel_return_follows_real_ingreso_after_reorder_and_preserves_other_fields(tmp_path):
    path, work, journal, operation, adapter = registered(tmp_path)
    book = load_workbook(path)
    sheet = book['Registros']
    one, two = [cell.value for cell in sheet[2]], [cell.value for cell in sheet[3]]
    for index, value in enumerate(two, 1):
        sheet.cell(2, index).value = value
    for index, value in enumerate(one, 1):
        sheet.cell(3, index).value = value
    book.save(path)
    book.close()
    before = path.read_bytes()
    target = tmp_path / 'devuelto.xlsx'
    portable(journal, path, target)
    assert path.read_bytes() == before and adapter.saves == 1
    result = load_workbook(target)
    try:
        sheet = result['Registros']
        assert sheet['F2'].value is None and sheet['G2'].value is None
        assert sheet['F3'].value.date() == date(2026, 10, 5) and sheet['G3'].value == 1
        assert sheet['E3'].value == work.rows[0].review['OBSERVACION']
        assert sheet['H3'].value == 0 and sheet['I3'].value == 'Manual'
        assert sheet['J3'].value == '=1+2'
        assert result['Otra']['A1'].value == '=SUM(Registros!H2:H3)'
        assert sheet['E2'].fill.fgColor.rgb == '00ABCDEF'
        assert sheet['F3'].border.bottom.style == 'thin'
    finally:
        result.close()
    receipt = journal.get(operation)
    assert receipt['state'] == 'COMPROBADA'
    assert receipt['excel_sha256'] == sha256(target.read_bytes()).hexdigest()


@pytest.mark.parametrize('cell,value', [('E2', 'Nuevo texto'), ('F2', datetime(2026, 10, 1)), ('G2', 0)])
def test_manual_conflicts_leave_receipt_pending_and_source_untouched(tmp_path, cell, value):
    path, work, journal, operation, adapter = registered(tmp_path)
    book = load_workbook(path)
    book['Registros'][cell] = value
    book.save(path)
    book.close()
    before = path.read_bytes()
    with pytest.raises(ValueError):
        portable(journal, path, tmp_path / 'salida.xlsx')
    assert not (tmp_path / 'salida.xlsx').exists()
    assert path.read_bytes() == before
    assert journal.get(operation)['state'] == 'PENDIENTE_EXCEL' and adapter.saves == 1


def test_locked_excel_retries_only_return_and_not_remote_save(tmp_path, monkeypatch):
    path, work, journal, operation, adapter = registered(tmp_path)
    from nurus.personal import registration_excel as module
    original = module.write_new_file
    def locked(*args):
        raise PermissionError('Excel bloqueado')
    monkeypatch.setattr(module, 'write_new_file', locked)
    with pytest.raises(PermissionError):
        portable(journal, path, tmp_path / 'salida.xlsx')
    assert journal.get(operation)['state'] == 'PENDIENTE_EXCEL'
    monkeypatch.setattr(module, 'write_new_file', original)
    portable(journal, path, tmp_path / 'salida.xlsx')
    assert journal.get(operation)['state'] == 'COMPROBADA' and adapter.saves == 1


def test_crash_after_excel_creation_keeps_receipt_and_can_make_a_new_copy(tmp_path, monkeypatch):
    path, work, journal, operation, adapter = registered(tmp_path)
    original = journal.complete_excel_many
    def interrupted(*args):
        raise OSError('Corte al confirmar el retorno')
    monkeypatch.setattr(journal, 'complete_excel_many', interrupted)
    with pytest.raises(OSError):
        portable(journal, path, tmp_path / 'primera.xlsx')
    assert (tmp_path / 'primera.xlsx').exists()
    assert journal.get(operation)['state'] == 'PENDIENTE_EXCEL'
    monkeypatch.setattr(journal, 'complete_excel_many', original)
    portable(journal, path, tmp_path / 'segunda.xlsx')
    assert journal.get(operation)['state'] == 'COMPROBADA' and adapter.saves == 1


def test_unverified_operations_do_not_fill_date_or_charge(tmp_path):
    path = source(tmp_path)
    work = Work.external(path, defaults(), sheet='Registros')
    data = intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                             send_to_tribunal=True, author_id='usuario-ficticio')
    journal = Journal(tmp_path / 'registro.sqlite', clock=Clock())
    operation = journal.prepare(data)['operation_id']
    with pytest.raises(ValueError, match='carece de registro'):
        portable(journal, path, tmp_path / 'salida.xlsx', operation_ids=[operation])
    assert not (tmp_path / 'salida.xlsx').exists()


def test_motor_proposal_is_not_accepted_as_effective_excel_text(tmp_path):
    path = source(tmp_path)
    work = Work.external(path, defaults(), sheet='Registros')
    work.external_input = False
    with pytest.raises(ValueError, match='Excel de registro'):
        intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                          send_to_tribunal=True, author_id='usuario-ficticio')


def test_unexported_edit_in_csmp_is_not_substituted_for_excel_text(tmp_path):
    work = Work.external(source(tmp_path), defaults(), sheet='Registros')
    work.rows[0].review['OBSERVACION'] = 'Cambio local todavía no guardado en Excel'
    with pytest.raises(ValueError, match='guarda y vuelve'):
        intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                          send_to_tribunal=True, author_id='usuario-ficticio')


def test_receipts_are_attached_only_to_their_remote_ingreso(tmp_path):
    path, work, journal, operation, adapter = registered(tmp_path)
    attach_receipts(work, journal)
    assert work.receipts['rus:' + operation]['text'] == work.rows[0].review['OBSERVACION']
    work.rows = work.rows[1:]
    work.receipts = {}
    attach_receipts(work, journal)
    assert work.receipts == {}


def test_multiple_receipts_complete_atomically_or_none_do(tmp_path):
    path, work, journal, operation, adapter = registered(tmp_path)
    data = intent_from_excel(work, work.rows[1], type='Administrativa', state='Realizada',
                             send_to_tribunal=False, author_id='usuario-ficticio')
    unverified = journal.prepare(data)['operation_id']
    proof = tmp_path / 'prueba.xlsx'
    proof.write_bytes(path.read_bytes())
    with pytest.raises(ValueError):
        journal.complete_excel_many([operation, unverified], proof, sha256(proof.read_bytes()).hexdigest())
    assert journal.get(operation)['state'] == 'PENDIENTE_EXCEL'


def test_formula_in_management_field_is_not_replaced(tmp_path):
    path, work, journal, operation, adapter = registered(tmp_path)
    book = load_workbook(path)
    book['Registros']['F2'] = '=IF(1=1,"",TODAY())'
    book.save(path)
    book.close()
    before = path.read_bytes()
    with pytest.raises(ValueError, match='fórmula'):
        portable(journal, path, tmp_path / 'salida.xlsx')
    assert path.read_bytes() == before and journal.get(operation)['state'] == 'PENDIENTE_EXCEL'
    assert not (tmp_path / 'salida.xlsx').exists()


def test_source_modified_after_import_must_be_loaded_again(tmp_path):
    path = source(tmp_path)
    work = Work.external(path, defaults(), sheet='Registros')
    book = load_workbook(path)
    book['Registros']['E2'] = 'Texto corregido en el archivo'
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match='cambió desde la lectura'):
        intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                          send_to_tribunal=True, author_id='usuario-ficticio')


def test_registration_report_separates_intentions_and_confirmed_entries(tmp_path):
    path, work, journal, operation, adapter = registered(tmp_path)
    other = intent_from_excel(work, work.rows[1], type='Administrativa', state='Realizada',
                              send_to_tribunal=False, author_id='usuario-ficticio')
    journal.prepare(other)
    target = tmp_path / 'operaciones.xlsx'
    export_journal(journal, target)
    book = load_workbook(target)
    try:
        counts = dict(book['Resumen'].values)
        assert counts['Registros RUS comprobados'] == 1
        assert counts['Operaciones sin comprobación'] == 1
        entries = list(book['Operaciones'].values)[1:]
        assert [row[5] for row in entries] == [True, False]
        assert entries[0][8] == work.rows[0].review['OBSERVACION']
        assert entries[1][10] is None
    finally:
        book.close()


def test_zero_antiguo_flag_is_not_lost_when_reading_numeric_cell(tmp_path):
    path = source(tmp_path)
    book = load_workbook(path)
    book['Registros']['R2'] = 0
    book.save(path)
    book.close()
    work = Work.external(path, defaults(), sheet='Registros')
    data = intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                             send_to_tribunal=True, author_id='usuario-ficticio')
    assert data.antiguo == '0'


def test_existing_observation_can_return_its_date_without_becoming_new_management(tmp_path):
    from nurus.personal.reports import export_management
    path = source(tmp_path)
    work = Work.external(path, defaults(), sheet='Registros')
    clock = Clock()
    data = intent_from_excel(work, work.rows[0], type='Al Tribunal', state='Realizada',
                             send_to_tribunal=True, author_id='usuario-ficticio')
    adapter = Adapter(data, clock)
    adapter.entries = [RemoteEntry('previa', clock().isoformat(), data.author_id,
                                   data.text, data.type, data.state, data.send_to_tribunal)]
    journal = Journal(tmp_path / 'registro.sqlite', clock=clock)
    operation = journal.prepare(data)['operation_id']
    submit(journal, operation, adapter)
    portable(journal, path, tmp_path / 'fechas.xlsx')
    attach_receipts(work, journal)
    target = tmp_path / 'gestion.xlsx'
    export_management([work], target, date(2026, 10, 1), date(2026, 10, 5))
    book = load_workbook(target)
    try:
        counts = dict(book['Resumen'].values)
        assert counts['Observaciones nuevas comprobadas'] == 0
        assert counts['Entradas RUS anteriores o sin atribución de gestión nueva'] == 1
        assert book['Entradas ya existentes'].max_row == 2
    finally:
        book.close()
    assert adapter.saves == 0


def test_macro_enabled_format_without_a_vba_project_is_preserved(tmp_path):
    from io import BytesIO
    from zipfile import ZipFile
    original = source(tmp_path)
    content = BytesIO()
    with ZipFile(original) as before, ZipFile(content, 'w') as after:
        for item in before.infolist():
            data = before.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
                                    b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
            after.writestr(item, data)
    macro = tmp_path / 'registro.xlsm'
    macro.write_bytes(content.getvalue())
    path, work, journal, operation, adapter = registered(tmp_path, path=macro)
    target = tmp_path / 'devuelto.xlsm'
    portable(journal, path, target)
    with ZipFile(target) as book:
        assert b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' in book.read('[Content_Types].xml')
    assert journal.get(operation)['state'] == 'COMPROBADA' and adapter.saves == 1
