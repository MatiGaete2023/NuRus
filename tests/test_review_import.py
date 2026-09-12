from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from nurus.rus import Mode
from nurus.services.exports import export_proposal_workbook
from nurus.services.review_import import ReviewImportError, import_reviewed_workbook, preview_reviewed_workbook
from nurus.services.workflow import WorkController
from nurus.storage.database import Database


def _proposal(tmp_path):
    source = tmp_path / "origen.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Espera"
    sheet.append(["RIT", "TRIBUNAL", "DERIVACION", "T ESPERA"])
    sheet.append(["C-1", "Juzgado de Laja", "PRM Norte", 45])
    sheet.append(["C-2", "Juzgado de Tomé", "PRM Sur", 10])
    book.save(source)
    db = Database(tmp_path / "nurus.sqlite3")
    controller = WorkController(db)
    batch = controller.analyze(source, Mode.ESPERA, as_of=date(2026, 9, 8))
    proposal = tmp_path / "propuestas.xlsx"
    export_proposal_workbook(
        db, batch.batch_id, proposal, backend="portable", allow_reduced_fidelity=True
    )
    return db, controller, batch.batch_id, proposal


def _complete_review(path):
    book = load_workbook(path)
    sheet = book["Espera"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in (2, 3):
        sheet.cell(row, headers["FECHA_OBS"]).value = "08/09/2026"
    sheet.cell(2, headers["OBSERVACION"]).value += " Antecedente adicional constatado."
    sheet.cell(2, headers["TT"]).value = "Sí"
    sheet.cell(2, headers["CC"]).value = "Con carga"
    sheet.cell(2, headers["RES"]).value = "No"
    book.save(path)


def test_reviewed_excel_is_imported_and_required_before_freeze(tmp_path):
    db, controller, batch_id, proposal = _proposal(tmp_path)
    with pytest.raises(ValueError, match="Primero carga"):
        controller.approve_batch(batch_id, require_review_import=True)
    _complete_review(proposal)
    preview = preview_reviewed_workbook(db, batch_id, proposal)
    assert preview.valid
    assert preview.reviewed_count == 2
    assert preview.changed_observations == 1
    imported = import_reviewed_workbook(
        db, batch_id, proposal, responsible="Revisora CSMP", confirmed_in_rus=True
    )
    assert imported.sha256 == db.get_batch(batch_id)["review_import_hash"]
    rows = db.list_review_records(batch_id)
    assert all(row["decision"] == "approved" and row["rus_recorded"] == 1 for row in rows)
    assert rows[0]["review_date"] == "2026-09-08"
    assert rows[0]["workload_value"] == "Con carga"
    snapshot = controller.approve_batch(batch_id, require_review_import=True)
    assert len(snapshot) == 64
    assert db.get_snapshot(batch_id)["records"][0]["review_date"] == "2026-09-08"


def test_working_copy_survives_database_reopen_without_approving(tmp_path):
    db, _, batch_id, proposal = _proposal(tmp_path)
    reopened = Database(db.path)
    assert reopened.get_working_workbook(batch_id) == proposal.resolve()
    assert reopened.get_batch(batch_id)["status"] == "review"
    assert reopened.get_working_workbook("another-batch") is None


def test_confirmation_rejects_a_changed_workbook(tmp_path):
    db, _, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    preview = preview_reviewed_workbook(db, batch_id, proposal)
    book = load_workbook(proposal)
    book["Espera"].cell(2, 2, "Otro tribunal")
    book.save(proposal)
    with pytest.raises(ReviewImportError, match="cambió durante"):
        import_reviewed_workbook(db, batch_id, proposal, responsible="Prueba",
                                confirmed_in_rus=True, expected_sha256=preview.sha256)
    assert not db.get_batch(batch_id)["review_import_hash"]


def test_import_commits_the_bytes_actually_parsed(tmp_path, monkeypatch):
    db, _, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    expected = proposal.read_bytes()
    original_preview = preview_reviewed_workbook

    def preview_then_change(*args, **kwargs):
        result = original_preview(*args, **kwargs)
        proposal.write_bytes(b"changed after parsing")
        return result

    monkeypatch.setattr("nurus.services.review_import.preview_reviewed_workbook", preview_then_change)
    result = import_reviewed_workbook(db, batch_id, proposal, responsible="Prueba", confirmed_in_rus=True)
    from hashlib import sha256
    assert result.sha256 == sha256(expected).hexdigest()
    assert db.get_batch(batch_id)["review_import_hash"] == result.sha256


def test_ids_cannot_be_swapped_between_cases(tmp_path):
    db, _, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    book = load_workbook(proposal)
    sheet = book["Espera"]
    column = next(cell.column for cell in sheet[1] if cell.value == "NURUS_ID_REGISTRO")
    sheet.cell(2, column).value, sheet.cell(3, column).value = sheet.cell(3, column).value, sheet.cell(2, column).value
    book.save(proposal)
    preview = preview_reviewed_workbook(db, batch_id, proposal)
    assert not preview.valid
    assert any("no coincide con el ID" in issue for issue in preview.errors)


def test_partial_import_cannot_bypass_completeness_by_default_argument(tmp_path):
    db, controller, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    book = load_workbook(proposal)
    book["Espera"].delete_rows(3)
    book.save(proposal)
    import_reviewed_workbook(db, batch_id, proposal, responsible="Prueba", confirmed_in_rus=True)
    with pytest.raises(ValueError, match="sin constancia"):
        controller.approve_batch(batch_id)


@pytest.mark.parametrize("operation", ["edit", "restore"])
def test_edit_or_restore_invalidates_record_evidence(tmp_path, operation):
    db, controller, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    import_reviewed_workbook(db, batch_id, proposal, responsible="Prueba", confirmed_in_rus=True)
    rows = db.list_review_records(batch_id)
    if operation == "edit":
        db.set_record_decision(batch_id, rows[0]["record_id"], "approved",
                               observation="Nueva revisión", reason="Cambio")
    else:
        db.restore_record(batch_id, rows[0]["record_id"])
    changed = db.list_review_records(batch_id)[0]
    assert not changed["rus_recorded"]
    assert not changed["review_import_hash"]
    assert not changed["review_date"]
    # Otra devolución no puede reciclar la confirmación antigua de esa fila.
    book = load_workbook(proposal)
    book["Espera"].delete_rows(2)
    book.save(proposal)
    import_reviewed_workbook(db, batch_id, proposal, responsible="Prueba", confirmed_in_rus=True)
    with pytest.raises(ValueError, match="sin constancia"):
        controller.approve_batch(batch_id)


def test_final_export_contains_the_reviewed_fields(tmp_path):
    from nurus.services.exports import export_preserved_workbook
    db, controller, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    import_reviewed_workbook(db, batch_id, proposal, responsible="Prueba", confirmed_in_rus=True)
    controller.approve_batch(batch_id, require_review_import=True)
    row = db.list_review_records(batch_id)[0]
    result = export_preserved_workbook(db, batch_id, tmp_path / "final.xlsx",
                                       backend="portable", allow_reduced_fidelity=True)
    book = load_workbook(result.path)
    sheet = book["Espera"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for name, expected in {"OBSERVACION": row["edited_observation"], "FECHA_OBS": "2026-09-08",
                            "TT": "Sí", "CC": "Con carga", "RES": "No"}.items():
        assert sheet.cell(2, columns[name]).value == expected
    book.close()


def test_import_requires_explicit_rus_confirmation(tmp_path):
    db, _, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    with pytest.raises(ReviewImportError, match="confirmar"):
        import_reviewed_workbook(
            db, batch_id, proposal, responsible="Revisora", confirmed_in_rus=False
        )
    assert not db.get_batch(batch_id)["review_import_hash"]


def test_import_rejects_file_from_another_evaluation(tmp_path):
    db, _, batch_id, proposal = _proposal(tmp_path)
    book = load_workbook(proposal)
    trace = book["NURUS_TRAZABILIDAD"]
    trace["B3"] = "otra-evaluacion"
    book.save(proposal)
    preview = preview_reviewed_workbook(db, batch_id, proposal)
    assert not preview.valid
    assert any("otra evaluación" in error for error in preview.errors)


def test_partial_review_keeps_absent_rows_pending_and_traces_each_import(tmp_path):
    db, controller, batch_id, proposal = _proposal(tmp_path)
    _complete_review(proposal)
    book = load_workbook(proposal)
    sheet = book["Espera"]
    sheet.delete_rows(3)
    book.save(proposal)

    preview = preview_reviewed_workbook(db, batch_id, proposal)
    assert preview.valid
    assert preview.reviewed_count == 1
    assert any("permanecerán pendientes" in warning for warning in preview.warnings)
    imported = import_reviewed_workbook(
        db, batch_id, proposal, responsible="Revisora CSMP", confirmed_in_rus=True
    )
    rows = db.list_review_records(batch_id)
    assert rows[0]["decision"] == "approved"
    assert rows[0]["review_import_hash"] == imported.sha256
    assert rows[0]["review_import_name"] == proposal.name
    assert rows[1]["decision"] == "pending"
    assert rows[1]["rus_recorded"] == 0
    assert not rows[1]["review_import_hash"]
    with pytest.raises(ValueError, match="1 filas sin constancia"):
        controller.approve_batch(batch_id, require_review_import=True)
