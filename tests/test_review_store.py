from datetime import date
from openpyxl import load_workbook
import pytest

from nurus.personal.config import defaults
from nurus.personal.work import Work
from nurus.personal.review_store import bind
from nurus.personal.rus_activity import mark_reviewed
from test_rus_activity import source


def remote_source(folder, reverse=False):
    folder.mkdir()
    path=source(folder)
    book=load_workbook(path);sheet=book['ESPERA']
    for title in ('SITFA_CAUSA_RUS','SITFA_INGRESO_RUS','SITFA_PERSONA_RUS',
                  'SITFA_CENTRO_RUS','SITFA_VINCULO_ESTADO'):
        sheet.cell(1,sheet.max_column+1,title)
    for i in range(2,5):
        for j,value in enumerate(('101',str(200+i),str(300+i),'401','Vinculado desde respuesta actual'),8):
            sheet.cell(i,j,value)
    if reverse:
        rows=list(sheet.values);sheet.delete_rows(1,sheet.max_row)
        for row in [rows[0],*reversed(rows[1:])]:sheet.append(row)
    book.save(path);book.close();return path


def test_verified_ingreso_keeps_review_between_downloads_and_order_changes(tmp_path):
    path=tmp_path/'revisiones.json'
    first=Work(defaults()).analyze(remote_source(tmp_path/'uno'),'ESPERA',as_of=date(2026,10,2))
    bind(first,path);row=first.rows[0];mark_reviewed(first,row.id,'firma-1')
    second=Work(defaults()).analyze(remote_source(tmp_path/'dos',True),'ESPERA',as_of=date(2026,10,2))
    assert first.source_hash!=second.source_hash
    bind(second,path)
    target=next(r for r in second.rows if r.values['SITFA_INGRESO_RUS']=='202')
    assert second.signed_activity[target.id]['pendientes']==1
    other=next(r for r in second.rows if r.values['SITFA_INGRESO_RUS']=='203')
    assert second.signed_activity[other.id]['pendientes']==2
    target.overrides['NOMBRE']='Corrección local'
    mark_reviewed(second,target.id,'firma-2')
    bind(first,path)
    assert first.signed_activity[row.id]['pendientes']==0
    assert all(not r.review for r in second.rows)


def test_unverified_ids_cannot_transfer_review_between_downloads(tmp_path):
    first=Work(defaults()).analyze(remote_source(tmp_path/'uno'),'ESPERA')
    first.rows[0].values['SITFA_VINCULO_ESTADO']='Identidad ambigua'
    bind(first,tmp_path/'revisiones.json');mark_reviewed(first,first.rows[0].id,'firma-1')
    second=Work(defaults()).analyze(remote_source(tmp_path/'dos',True),'ESPERA')
    bind(second,tmp_path/'revisiones.json')
    assert second.signed_activity[second.rows[-1].id]['pendientes']==2


def test_damaged_store_does_not_claim_review_success(tmp_path):
    work=Work(defaults()).analyze(remote_source(tmp_path/'uno'),'ESPERA')
    path=tmp_path/'revisiones.json';bind(work,path)
    path.write_text('{"version":999,"ingresos":{}}')
    with pytest.raises(ValueError):mark_reviewed(work,work.rows[0].id,'firma-1')
    assert work.signed_activity[work.rows[0].id]['pendientes']==2
