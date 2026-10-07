from datetime import date
from openpyxl import load_workbook
import pytest
from nurus.personal.config import defaults
from nurus.personal.work import Work
from nurus.personal.reports import export_management,export_activity
from nurus.personal.rus_activity import mark_reviewed
from test_rus_activity import source


def test_historical_excel_is_never_counted_as_new_registration(tmp_path):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    work.rows[0].review={'OBSERVACION':'Constancia histórica','FECHA_OBS':'2026-10-01','TT':1,'CC':1}
    work.receipts={'draft':dict(kind='draft',state='created',created_at='2026-10-02T09:00:00'),
        'uncertain':dict(kind='rus_observation',state='INCIERTA',verified=False,operation_id='u'),
        'verified':dict(kind='rus_observation',state='PENDIENTE_EXCEL',verified=True,new_registration=True,operation_id='op-1',
            registered_at='2026-10-02T10:00:00',remote_entry_id='12',type='Al Tribunal',cc=1,text='Texto guardado')}
    path=export_management([work,work],tmp_path/'gestion.xlsx',date(2026,10,1),date(2026,10,5))
    book=load_workbook(path);counts=dict(book['Resumen'].values)
    assert counts['Observaciones nuevas comprobadas']==1
    assert counts['Constancias históricas del Excel']==1
    assert counts['Nuevas con carga']==1
    assert book['Nuevas comprobadas'].max_row==2
    book.close()
    work.receipts['verified']['cc']=0
    with pytest.raises(ValueError):export_management([work],tmp_path/'wrong.xlsx',date(2026,10,1),date(2026,10,5))


def test_resolution_report_keeps_per_ingreso_and_partial_coverage(tmp_path):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    mark_reviewed(work,work.rows[0].id,'firma-1')
    path=export_activity([work,work],tmp_path/'firmas.xlsx',date(2026,10,1),date(2026,10,3));book=load_workbook(path)
    assert book['Firmas por ingreso'].max_row==5
    counts=dict(book['Resumen'].values)
    assert counts['Revisadas para el ingreso']==1 and counts['Por revisar']==3
    assert all(r[4]=='PARCIAL' for r in list(book['Cobertura'].values)[1:])
    book.close()


def test_activity_report_filters_signatures_inclusive_and_labels_coverage(tmp_path):
    from nurus.personal.rus_activity import attach
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    attach(work)
    work.content=b''  # Keep the following source-backed signatures as the test fixture.
    work.rows=work.rows[:1]
    work.signed_activity[work.rows[0].id]['firmas']=[
        {'firma':'2026-09-30','hora':'12:00','tramite':'Anterior','huella':'before','identidad':'b'},
        {'firma':'2026-10-01','hora':'09:00','tramite':'Inicio','huella':'start','identidad':'s'},
        {'firma':'2026-10-02','hora':'10:00','tramite':'Fin','huella':'end','identidad':'e'},
        {'firma':'2026-10-03','hora':'11:00','tramite':'Posterior','huella':'after','identidad':'a'},
    ]
    coverage={'tribunal_codigo':'777','desde':'2026-10-01','hasta':'2026-10-03',
              'estado':'COMPLETA','criterio_temporal':'firmada','sin_coincidencias':False}
    work.activity_sources={'cobertura':{'777':coverage}}
    path=export_activity([work],tmp_path/'periodo.xlsx',date(2026,9,30),date(2026,10,2))
    book=load_workbook(path)
    try:
        assert book['Firmas por ingreso'].max_row==4
        assert [row[4] for row in list(book['Firmas por ingreso'].values)[1:]]==[
            '2026-09-30','2026-10-01','2026-10-02']
        summary=dict(book['Resumen'].values)
        assert summary['Período solicitado desde']=='2026-09-30'
        assert summary['Período solicitado hasta']=='2026-10-02'
        scope=list(book['Cobertura'].values)[1]
        assert scope[4]=='PARCIAL'
        assert scope[10]=='La consulta de origen no cubre todo el período solicitado.'
    finally:book.close()


def test_activity_report_validates_dates_and_preserves_failed_or_unqueried_coverage(tmp_path):
    from datetime import datetime
    from nurus.personal.rus_activity import attach
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    attach(work)
    work.content=b''
    work.signed_activity[work.rows[0].id]['firmas']=[]
    work.activity_sources={'cobertura':{
        '777':{'tribunal_codigo':'777','desde':'2026-10-01','hasta':'2026-10-03','estado':'FALLIDA','criterio_temporal':'firmada','sin_coincidencias':False},
        '778':{'tribunal_codigo':'778','desde':'2026-09-01','hasta':'2026-09-30','estado':'COMPLETA','criterio_temporal':'firmada','sin_coincidencias':True},
        '779':{'tribunal_codigo':'779','desde':'2026-10-01','hasta':'2026-10-03','estado':'NO_CONSULTADA','criterio_temporal':'firmada','sin_coincidencias':False},
    }}
    with pytest.raises(ValueError):export_activity([work],tmp_path/'invalido.xlsx',date(2026,10,3),date(2026,10,2))
    with pytest.raises(ValueError):export_activity([work],tmp_path/'datetime.xlsx',datetime(2026,10,1),date(2026,10,2))
    path=export_activity([work],tmp_path/'cobertura.xlsx',date(2026,10,1),date(2026,10,2))
    book=load_workbook(path)
    try:
        states={row[1]:row[4] for row in list(book['Cobertura'].values)[1:]}
        assert states=={'777':'FALLIDA','778':'NO_CONSULTADA','779':'NO_CONSULTADA'}
    finally:book.close()


@pytest.mark.parametrize('start,end,expected', [
    (date(2026,9,30),date(2026,10,2),'PARCIAL'),
    (date(2026,9,1),date(2026,9,2),'NO_CONSULTADA'),
])
def test_verified_empty_source_does_not_claim_empty_outside_its_period(tmp_path,start,end,expected):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    work.content=b''
    for row in work.rows:
        work.signed_activity[row.id]['firmas']=[]
    work.activity_sources={'cobertura':{
        '777':{'tribunal_codigo':'777','desde':'2026-10-01','hasta':'2026-10-03',
               'estado':'VACIA_COMPROBADA','criterio_temporal':'firmada','sin_coincidencias':True}}}
    path=export_activity([work],tmp_path/'empty.xlsx',start,end)
    book=load_workbook(path)
    try:
        scope=list(book['Cobertura'].values)[1]
        assert scope[4]==expected
        assert scope[10]
    finally:book.close()
