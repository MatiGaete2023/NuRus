from __future__ import annotations

from datetime import date
from io import BytesIO
import platform

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from nurus.rus import Mode
from nurus.services.exports import ExportError, export_preserved_workbook, export_proposal_workbook
from nurus.services.workflow import WorkController
from nurus.storage.database import Database


def _source(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Espera"
    sheet.append(["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "T ESPERA"])
    sheet.append(["PRM Norte", "Juzgado de Laja", "Ana", "C-1", 45])
    sheet.append(["PRM Norte", "Juzgado de Laja", "Beto", "C-2", 45])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E3"
    sheet.column_dimensions["B"].width = 31
    sheet["A1"].font = Font(bold=True)
    sheet["C2"].fill = PatternFill(fill_type="solid", fgColor="92D050")
    validation = DataValidation(type="list", formula1='"uno,dos"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add("E2:E3")
    ob = book.create_sheet("OB")
    ob["A1"] = "=1+2"
    ob.column_dimensions["A"].width = 27
    ob.sheet_state = "hidden"
    expired = book.create_sheet("Medidas vencidas")
    expired["A1"] = "No procesar"
    other = book.create_sheet("Otra")
    other["A1"] = "Conservar"
    book.save(path)
    return path.read_bytes()


def _approved_excluded(db, path):
    original = _source(path)
    controller = WorkController(db)
    batch = controller.analyze(path, Mode.ESPERA, as_of=date(2026, 9, 8))
    first, second = controller.rows_from_batch(batch)
    controller.approve_record(
        batch.batch_id, first.record_id, observation="=texto de prueba", reason="Validación"
    )
    controller.exclude_record(batch.batch_id, second.record_id, reason="Duplicado confirmado")
    controller.approve_batch(batch.batch_id)
    return batch, original


def test_preserved_export_uses_archived_bytes_and_marks_excluded(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    source = tmp_path / "entrada.xlsx"
    batch, original = _approved_excluded(db, source)
    source.write_bytes(b"cambio externo")
    source.unlink()

    target = tmp_path / "salida.xlsx"
    result = export_preserved_workbook(
        db, batch.batch_id, target, backend="portable", allow_reduced_fidelity=True
    )

    assert result.path == target.resolve()
    assert result.row_count == 2
    expected = load_workbook(BytesIO(original), data_only=False)
    generated = load_workbook(target, data_only=False)
    try:
        assert generated.sheetnames[:4] == expected.sheetnames
        assert generated["OB"]["A1"].value == "=1+2"
        assert generated["OB"].sheet_state == "hidden"
        assert generated["OB"].column_dimensions["A"].width == 27
        assert generated["Medidas vencidas"]["A1"].value == "No procesar"
        sheet = generated["Espera"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:E3"
        assert sheet.column_dimensions["B"].width == 31
        assert sheet["A1"].font.bold is True
        assert sheet["C2"].fill.fgColor.rgb == "0092D050"
        assert len(sheet.data_validations.dataValidation) == 1
        headers = [cell.value for cell in sheet[1]]
        observation = headers.index("NURUS_OBSERVACION_FINAL") + 1
        state = headers.index("NURUS_ESTADO_REVISION") + 1
        assert sheet.cell(2, observation).value == "'=texto de prueba"
        assert sheet.cell(2, state).value == "REVISADO"
        assert sheet.cell(3, state).value == "EXCLUIDO"
        rules = list(sheet.conditional_formatting)
        assert rules
        assert "EXCLUIDO" in str(sheet.conditional_formatting[rules[0]][0].formula)
        trace = generated["NURUS_TRAZABILIDAD"]
        assert trace.sheet_state == "hidden"
        assert trace["B2"].value == "CONSTANCIA_REVISADA"
        assert trace["B4"].value == db.get_batch(batch.batch_id)["snapshot_hash"]
    finally:
        expected.close()
        generated.close()


def test_proposal_export_does_not_require_human_approval(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    source = tmp_path / "entrada.xlsx"
    original = _source(source)
    controller = WorkController(db)
    batch = controller.analyze(source, Mode.ESPERA, as_of=date(2026, 9, 8))
    assert db.get_batch(batch.batch_id)["status"] == "review"

    source.unlink()
    target = tmp_path / "propuestas.xlsx"
    result = export_proposal_workbook(
        db, batch.batch_id, target, backend="portable", allow_reduced_fidelity=True
    )

    expected = load_workbook(BytesIO(original), data_only=False)
    generated = load_workbook(target, data_only=False)
    try:
        assert generated.sheetnames[:4] == expected.sheetnames
        assert generated["OB"]["A1"].value == "=1+2"
        sheet = generated["Espera"]
        headers = [cell.value for cell in sheet[1]]
        proposal = headers.index("NURUS_PROPUESTA") + 1
        state = headers.index("NURUS_ESTADO_REVISION") + 1
        assert sheet.cell(2, proposal).value
        assert sheet.cell(2, state).value == "PROPUESTA"
        trace = generated["NURUS_TRAZABILIDAD"]
        assert trace["B2"].value == "PROPUESTA_NO_REVISADA"
        assert trace["B3"].value == db.get_batch(batch.batch_id)["evaluation_hash"]
        assert result.row_count == 2
    finally:
        expected.close()
        generated.close()


def test_preserved_export_requires_explicit_portable_confirmation(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    batch, _ = _approved_excluded(db, tmp_path / "entrada.xlsx")
    with pytest.raises(ExportError, match="fidelidad reducida"):
        export_preserved_workbook(db, batch.batch_id, tmp_path / "salida.xlsx", backend="portable")


def test_preserved_export_rejects_source_and_existing_destination(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    source = tmp_path / "entrada.xlsx"
    batch, _ = _approved_excluded(db, source)
    with pytest.raises(ExportError, match="archivo de origen"):
        export_preserved_workbook(
            db, batch.batch_id, source, backend="portable", allow_reduced_fidelity=True
        )
    existing = tmp_path / "existente.xlsx"
    existing.write_bytes(b"keep")
    with pytest.raises(ExportError, match="ya existe"):
        export_preserved_workbook(
            db, batch.batch_id, existing, backend="portable", allow_reduced_fidelity=True
        )
    assert existing.read_bytes() == b"keep"


@pytest.mark.skipif(platform.system() == "Windows", reason="requiere Excel de escritorio real")
def test_native_backend_requires_windows_excel_environment(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    batch, _ = _approved_excluded(db, tmp_path / "entrada.xlsx")
    with pytest.raises(ExportError, match="Windows"):
        export_preserved_workbook(db, batch.batch_id, tmp_path / "salida.xlsx")


@pytest.mark.parametrize("suffix", [".xls", ".xlsx", ".xlsm"])
@pytest.mark.parametrize("failure", [False, True])
def test_native_savecopyas_separates_input_output_and_closes_excel(tmp_path, monkeypatch, suffix, failure):
    """Contrato COM simulado; no sustituye la aceptación en Excel real."""
    import sys
    from pathlib import Path
    from unittest.mock import MagicMock
    from nurus.services.exports import _native_preserved

    content = b"legacy xls fixture"
    if suffix != ".xls":
        buffer = BytesIO()
        source_book = Workbook()
        source_book.save(buffer)
        content = buffer.getvalue()
    target = tmp_path / ("output" + suffix)
    target.touch()
    sheet = MagicMock()
    sheet.Name = "Espera"
    sheet.UsedRange.Column = sheet.UsedRange.Row = 1
    sheet.UsedRange.Columns.Count = 1
    sheet.UsedRange.Rows.Count = 2
    cells = {}

    def cell(row, column):
        if (row, column) not in cells:
            cells[row, column] = MagicMock(Value2="RIT" if (row, column) == (1, 1) else None)
        return cells[row, column]

    sheet.Cells.side_effect = cell
    book = MagicMock()
    book.Worksheets.return_value = sheet
    book.Worksheets.Count = 1
    placeholder = MagicMock()
    app = MagicMock()
    app.Workbooks.Add.return_value = placeholder
    opened = []

    def open_book(path, *args):
        opened.append(Path(path))
        assert Path(path).read_bytes() == content
        assert Path(path) != target
        assert args == (0, False)
        return book

    def save_copy_as(path):
        assert path == str(target)
        assert not target.exists()
        if failure:
            raise RuntimeError("SaveCopyAs refused")
        target.write_bytes(b"output-created-by-fake-excel")

    app.Workbooks.Open.side_effect = open_book
    book.SaveCopyAs.side_effect = save_copy_as
    pythoncom = MagicMock()
    win32com = MagicMock()
    win32com.client.DispatchEx.return_value = app
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", win32com.client)
    monkeypatch.setattr("platform.system", lambda: "Windows")
    snapshot = {
        "batch": {"primary_sheet": "Espera", "header_row": 1, "source_hash": "hash",
                  "mode": "ESPERA", "export_stage": "proposal"},
        "records": [{"record_id": "id", "source_sheet": "Espera", "source_row": 2,
                     "decision": "excluded", "evaluation_status": "excluded",
                     "edited_observation": "Propuesta", "edit_reason": "", "rule_ids_json": "[]",
                     "source_hash": "hash"}],
    }
    if failure:
        with pytest.raises(ExportError, match="guardar la copia"):
            _native_preserved(content, target, snapshot)
    else:
        _native_preserved(content, target, snapshot)
        assert target.read_bytes() == b"output-created-by-fake-excel"
    book.Save.assert_not_called()
    book.SaveAs.assert_not_called()
    book.SaveCopyAs.assert_called_once_with(str(target))
    placeholder.Close.assert_called_once_with(False)
    book.Close.assert_called_once_with(False)
    app.Quit.assert_called_once()
    pythoncom.CoUninitialize.assert_called_once()
    assert sheet.Range.return_value.Interior.Color == 255 + 242 * 256 + 204 * 65536
    assert opened and not opened[0].exists()
