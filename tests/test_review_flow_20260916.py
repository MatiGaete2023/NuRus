from pathlib import Path

from openpyxl import Workbook, load_workbook

from nurus.personal.config import defaults
from nurus.personal.outputs import Draft, draft_fingerprint
from nurus.personal.resolutions import automatic_project_selections
from nurus.personal.work import Work, Row


def _espera_book(path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION'])
    sheet.append(['X-1','Jgdo. L. y G. de Mulchén','NNA UNO','11111111-1','AFT PRUEBA',45,'texto previo'])
    book.save(path)


def test_edit_in_work_is_exported_to_visible_observation(tmp_path):
    source = tmp_path/'entrada.xlsx'
    _espera_book(source)
    work = Work(defaults()).analyze(source,'ESPERA',as_of=None)
    row = work.rows[0]
    row.review['OBSERVACION'] = 'observación humana corregida'
    row.review['RES'] = 'PC_IE'
    target = tmp_path/'salida.xlsx'
    work.export(target,backend='portable',reduced_fidelity=True)
    book = load_workbook(target,data_only=False)
    sheet = book['Espera']
    headers = {str(sheet.cell(1,c).value):c for c in range(1,sheet.max_column+1)}
    assert sheet.cell(2,headers['OBSERVACION']).value == 'observación humana corregida'
    assert sheet.cell(2,headers['RES']).value == 'PC_IE'
    assert sheet.cell(2,headers['NURUS_ESTADO_REVISION']).value == 'REVISADO'
    book.close()


def test_unedited_rows_remain_pending_when_another_row_is_reviewed(tmp_path):
    source = tmp_path/'entrada.xlsx'
    _espera_book(source)
    book = load_workbook(source)
    sheet = book['Espera']
    sheet.append(['X-2','Jgdo. L. y G. de Mulchén','NNA DOS','22222222-2','AFT PRUEBA',45,'original dos'])
    book.save(source);book.close()
    work = Work(defaults()).analyze(source,'ESPERA')
    work.rows[0].review['OBSERVACION'] = 'editada uno'
    target = tmp_path/'salida.xlsx'
    work.export(target,backend='portable',reduced_fidelity=True)
    book = load_workbook(target)
    sheet = book['Espera']
    headers = {str(sheet.cell(1,c).value):c for c in range(1,sheet.max_column+1)}
    assert sheet.cell(2,headers['NURUS_ESTADO_REVISION']).value == 'REVISADO'
    assert sheet.cell(3,headers['NURUS_ESTADO_REVISION']).value == 'PENDIENTE'
    assert sheet.cell(3,headers['OBSERVACION']).value == 'original dos'
    book.close()


def test_reviewed_export_preserves_unedited_existing_review_fields(tmp_path):
    source = tmp_path/'entrada.xlsx'
    book = Workbook();sheet=book.active;sheet.title='Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION','TT','CC'])
    sheet.append(['X-1','Jgdo. L. y G. de Mulchén','NNA UNO','11111111-1','AFT PRUEBA',45,'texto previo',1,0])
    book.save(source);book.close()
    work=Work(defaults()).analyze(source,'ESPERA')
    work.rows[0].review['OBSERVACION']='corregida'
    target=tmp_path/'salida.xlsx'
    work.export(target,backend='portable',reduced_fidelity=True)
    book=load_workbook(target);sheet=book['Espera']
    headers={str(sheet.cell(1,c).value):c for c in range(1,sheet.max_column+1)}
    assert sheet.cell(2,headers['TT']).value==1
    assert sheet.cell(2,headers['CC']).value==0
    book.close()


def test_draft_identity_changes_with_real_edit_but_not_temp_directory(tmp_path):
    left=tmp_path/'a';right=tmp_path/'b';left.mkdir();right.mkdir()
    first_path=left/'nomina.xlsx';second_path=right/'nomina.xlsx'
    first_path.write_bytes(b'mismos bytes');second_path.write_bytes(b'mismos bytes')
    first=Draft('Asunto','Cuerpo','a@example.cl','ucc_concepcion@pjud.cl',[str(first_path)])
    same=Draft('Asunto','Cuerpo','a@example.cl','ucc_concepcion@pjud.cl',[str(second_path)])
    edited=Draft('Asunto','Cuerpo corregido','a@example.cl','ucc_concepcion@pjud.cl',[str(second_path)])
    assert draft_fingerprint(first)==draft_fingerprint(same)
    assert draft_fingerprint(first)!=draft_fingerprint(edited)
    second_path.write_bytes(b'otros bytes')
    assert draft_fingerprint(first)!=draft_fingerprint(same)


def _row(rid, observation, actions, res=''):
    return Row(
        rid,2,
        {'RIT':'X-1','TRIBUNAL':'Jgdo. L. y G. de Mulchén','NOMBRE':'NNA','RUT':'11111111-1','DERIVACION':'AFT TEST'},
        observation,[],actions,[],False,
        {'RES':res} if res != '' else {},
    )


def _fake_work(rows):
    work = type('W',(),{})()
    work.rows = rows
    work.mapping = {'rit':'RIT','tribunal':'TRIBUNAL','nombre':'NOMBRE','rut':'RUT','programa':'DERIVACION'}
    return work


def test_resolution_type_comes_from_reviewed_observation_when_res_only_marks_case():
    row = _row('r1','Se remite proyecto de resolución pidiendo cuenta respecto del informe de avance.',['PC_IE'],'X')
    assert automatic_project_selections(_fake_work([row]),'PC_IE') == [('r1','PC_INFO')]


def test_resolution_type_explicit_in_res_wins_over_observation():
    row = _row('r1','Se remite proyecto de resolución pidiendo cuenta respecto del ingreso efectivo.',['PC_IE'],'NOMENCL')
    assert automatic_project_selections(_fake_work([row]),'PC_IE') == [('r1','NOMENCL')]


def test_nomenclatura_can_be_inferred_from_observation():
    row = _row('r1','Corresponde proyecto de nomenclatura para regularizar la causa.',['PC_IE'],'X')
    assert automatic_project_selections(_fake_work([row]),'PC_IE') == [('r1','NOMENCL')]


def test_observation_without_project_action_does_not_create_resolution():
    row=_row('r1','Se revisa informe de avance. No corresponde proyecto.',[], '')
    assert automatic_project_selections(_fake_work([row]),'PC_IE') == []


class _Words:
    def __init__(self):
        self.data = {
            'a|PC_IE': ['X-1','MULCHEN','PC_IE','auto'],
            'b|PC_IE': ['X-2','MULCHEN','PC_IE','auto'],
        }
        self.selected = ('a|PC_IE',)
    def selection(self): return self.selected
    def item(self,iid,what): return {'values':self.data[iid]}[what]
    def delete(self,iid): self.data.pop(iid)
    def exists(self,iid): return iid in self.data
    def insert(self,parent,where,iid,values,tags=()): self.data[iid]=list(values)
    def selection_set(self,items): self.selected=tuple(items)
    def get_children(self): return tuple(self.data)


class _List:
    def delete(self,*args): pass


class _Var:
    def get(self): return 'PC_INFO'


class _Status:
    def set(self,value): self.value=value


def test_individual_resolution_type_edit_changes_only_selected_row():
    from nurus.personal.app import App
    dummy = type('Dummy',(),{})()
    dummy.words = _Words();dummy.manual_word = _Var();dummy.projects=['old'];dummy.project_index=0;dummy.project_list=_List();dummy.status=_Status()
    dummy.project_editor=_List()
    dummy._clear_projects=lambda: App._clear_projects(dummy)
    dummy._visible_project_key=App._visible_project_key
    dummy._existing_project_iid=lambda key,exclude=None: App._existing_project_iid(dummy,key,exclude)
    App._assign_word_type(dummy)
    assert 'a|PC_INFO' in dummy.words.data
    assert 'b|PC_IE' in dummy.words.data
    assert 'b|PC_INFO' not in dummy.words.data


def test_res_categorical_values_and_legacy_marks_are_distinguished():
    from nurus.personal.resolutions import resolution_kind,resolution_review_issue,resolution_selection_source
    assert resolution_kind('PC_IE')=='PC_IE'
    assert resolution_kind('PC_INFO')=='PC_INFO'
    assert resolution_kind('NOMENCL')=='NOMENCL'
    assert resolution_review_issue('X')==''
    assert resolution_review_issue(1)==''
    assert 'RES no reconocido' in resolution_review_issue('PC_INOF')
    explicit=_row('typed','Texto cualquiera',[],'PC_INFO')
    work=_fake_work([explicit])
    assert resolution_selection_source(work,explicit,'PC_INFO')=='Definido en RES'
    legacy=_row('legacy','Se remite proyecto de resolución pidiendo cuenta respecto del informe.',[],'X')
    work=_fake_work([legacy])
    assert resolution_selection_source(work,legacy,'PC_INFO')=='RES antiguo · tipo inferido'


def test_blank_res_column_explicitly_means_no_project():
    from nurus.personal.resolutions import reviewed_resolution_ids
    row=_row('r1','Se remite proyecto de resolución pidiendo cuenta respecto del ingreso efectivo.',['PC_IE'])
    row.review={'RES':''}
    work=_fake_work([row])
    assert reviewed_resolution_ids(work)==set()
    assert automatic_project_selections(work,'PC_IE')==[]
