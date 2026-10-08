from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from test_bitacora_har import popup,PARAMS
from nurus.personal.bitacora_lote import import_lote,export_lote
from nurus.personal.bitacora_link import verify_reply

def manifest(tmp_path):
    body=popup().encode();(tmp_path/'captura.html').write_bytes(body)
    record={'tribunal_codigo':'11','causa_id':'22','ingreso_id':'33','rit':'P-12-2026',
        'nombre':'Persona de prueba','programa':'Centro de prueba','estado':'LEIDA',
        'params':PARAMS,'archivo':'captura.html','sha256':sha256(body).hexdigest(),'capturada':'2026-10-08T12:00:00Z',
        'modalidad_nombre':'FAE','pestana':'Cumplimiento','pagina':1,'fila':2}
    data={'version':1,'tipo':'BITACORAS_LECTURA','estado':'INCOMPLETA','fecha':'2026-10-08',
        'seleccion':{'tribunals':['11'],'modalities':['1','2','3','4'],'tabs':['Cumplimiento']},
        'consultas':[{'tribunal':'11','modality':'3','tab':'Cumplimiento','estado':'ENUMERADA','registros':3,'paginas':1}],
        'registros':[record,{**record,'ingreso_id':'34','estado':'FALLIDA','error':'Sesión vencida'},
            {**record,'ingreso_id':'','causa_id':'','estado':'SIN_VINCULO','error':'Sin enlace comprobado'}]}
    path=tmp_path/'bitacoras.json';path.write_text(json.dumps(data),encoding='utf-8');return path

def test_excel_copies_full_text_and_reports_every_row_and_unattempted_scope(tmp_path):
    source=manifest(tmp_path);out=tmp_path/'lote.xlsx'
    export_lote(source,out,date(2026,6,8),date(2026,10,8),today=date(2026,10,8))
    book=load_workbook(out)
    assert book['Resumen'].max_row==3 and book['Lecturas'].max_row==4
    assert book['Consultas'].max_row==5
    assert book['Copia íntegra']['M2'].value=='Texto completo que la tabla oculta.'
    assert book['Copia íntegra']['K2'].value==1
    assert book['Copia íntegra']['M2'].border.left.style=='thin'
    assert book['Lecturas']['K3'].value=='FALLIDA'
    assert book['Lecturas']['K4'].value=='SIN_VINCULO'
    assert book['Resumen']['E3'].value=='FALLIDA'
    book.close()

def test_changed_html_is_explicit_failed_read(tmp_path):
    source=manifest(tmp_path);(tmp_path/'captura.html').write_text('alterada')
    capture=import_lote(source)
    assert capture['queries'][0]['coverage']=='FALLIDA'
    assert 'cambió' in capture['queries'][0]['error']

def test_remote_identity_mismatch_not_exported(tmp_path):
    source=manifest(tmp_path);data=json.loads(source.read_text())
    data['registros'][0]['ingreso_id']='99';source.write_text(json.dumps(data))
    assert import_lote(source)['queries'][0]['coverage']=='FALLIDA'

def test_reply_id_and_scope_and_hash_checked(tmp_path):
    folder=tmp_path/'solicitud';lot=folder/'lotes'/'lote';lot.mkdir(parents=True)
    source=manifest(lot)
    request={'folder':str(folder),'id':'prueba'}
    reply=folder/'respuesta.json';data={'id':'prueba','manifest':str(source),'sha256':sha256(source.read_bytes()).hexdigest()}
    reply.write_text(json.dumps(data));assert verify_reply(request)==source.resolve()
    data['id']='otra';reply.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='otra solicitud'):verify_reply(request)
    data['id']='prueba';data['manifest']=str(manifest(tmp_path));reply.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='fuera'):verify_reply(request)

def test_date_window_never_extends_beyond_four_calendar_months(tmp_path):
    source=manifest(tmp_path)
    with pytest.raises(ValueError):export_lote(source,tmp_path/'bad.xlsx',date(2026,6,7),date(2026,10,8),today=date(2026,10,8))
