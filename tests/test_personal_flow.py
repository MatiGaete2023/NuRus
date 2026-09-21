from datetime import date
from pathlib import Path
from copy import deepcopy
import json
import pytest
from openpyxl import Workbook, load_workbook
from docx import Document
from nurus.personal.config import defaults, Configuration, CC
from nurus.personal.work import Work
from nurus.personal.outputs import prepare_drafts, prepare_required_drafts, date_in_words, fill_docx, resolve_contact, create_draft, Draft

def source(tmp_path,mode='ESPERA',program='AFT EJEMPLO',days=30):
    b=Workbook();s=b.active;s.title='informe_1'
    s.append(['Reporte RUS']);s.append([])
    headers=['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION']
    values=['X-1-2026','Jgdo. L. y G. de Laja','Persona Ejemplo','12345678-9',program]
    if mode=='ESPERA':headers+=['T ESPERA'];values+=[days]
    elif mode=='INFORMES':headers+=['FECHA VENCIMIENTO'];values+=[date(2026,9,12)]
    else:headers+=['DIAS DE CUMPLIMIENTO','DIAS PARA EGRESAR','FEC.EGRESO PROYECTADO'];values+=[50,10,date(2026,9,23)]
    s.append(headers);s.append(values);s.column_dimensions['A'].width=27;s.auto_filter.ref=f'A3:F4'
    other=b.create_sheet('Original');other['B2']='=1+2'
    path=tmp_path/'origen.xlsx';b.save(path);return path

def exported(tmp_path,mode='ESPERA'):
    work=Work(defaults()).analyze(source(tmp_path,mode),mode,as_of=date(2026,9,13))
    if work.needs_cross:work.document_exception('Revisión parcial: descarga sin cruce.')
    work.export(tmp_path/'salida.xlsx',backend='portable',reduced_fidelity=True)
    return work

@pytest.mark.parametrize('mode',['ESPERA','CUMPLIMIENTO','INFORMES'])
def test_vertical_preserves_source_and_makes_result_available(tmp_path,mode):
    w=exported(tmp_path,mode)
    assert w.rows and w.rows[0].observation
    b=load_workbook(w.output);assert b['Original']['B2'].value=='=1+2'
    assert b['informe_1'].column_dimensions['A'].width==27
    assert 'NURUS_ID_REGISTRO' in [c.value for c in b['informe_1'][3]]
    assert w.refresh() is False
    assert not w.rows[0].review
    assert 'Se sugiere remitir' not in w.rows[0].observation
    assert 'Se sugiere preparar' not in w.rows[0].observation

def test_no_cross_is_warning_and_export_continues(tmp_path):
    w=Work(defaults()).analyze(source(tmp_path,'CUMPLIMIENTO'),'CUMPLIMIENTO')
    assert w.rows and w.cross_missing and not w.needs_cross
    assert any('Sin cruce utilizable' in warning for warning in w.warnings)
    out=w.export(tmp_path/'out.xlsx',backend='portable',reduced_fidelity=True)
    assert Path(out).is_file()

def test_text_edit_cannot_change_actions(tmp_path):
    path=source(tmp_path);a=Work(defaults()).analyze(path,'ESPERA')
    cfg=defaults();cfg['textos']['ESPERA']['E05_PROYECTO_Y_CORREO']['texto']='Redacción completamente diferente.'
    b=Work(cfg).analyze(path,'ESPERA')
    assert a.rows[0].actions==b.rows[0].actions==['programa_espera','PC_IE']

def test_refresh_and_recovery_without_reselection(tmp_path):
    w=exported(tmp_path);b=load_workbook(w.output);s=b[w.sheet]
    headers={c.value:c.column for c in s[3]}
    s.cell(4,headers['OBSERVACION']).value='Constancia editada';s.cell(4,headers['CC']).value=1;b.save(w.output)
    assert w.refresh();assert w.rows[0].review['OBSERVACION']=='Constancia editada'
    w.save(tmp_path/'sesion');r=Work.load(tmp_path/'sesion')
    assert r.output==w.output and r.rows[0].review['CC']==1

def test_unrecognized_contact_still_prepares_draft_with_subset(tmp_path):
    w=exported(tmp_path);drafts=prepare_drafts(w,'programa_espera')
    assert len(drafts)==1 and drafts[0].to=='' and CC in drafts[0].cc
    assert len(drafts[0].attachments)==1
    assert len(load_workbook(drafts[0].attachments[0]).sheetnames)==1

def test_general_mail_does_not_require_an_approval_step(tmp_path):
    w=exported(tmp_path)
    ds=prepare_drafts(w,'espera',modalities='Ambulatoria')
    assert 'Jgdo. L. y G. de Laja' in ds[0].subject

def test_contact_ambiguity_and_config_recovery(tmp_path):
    cfg=defaults();cfg['contactos']={'AFT X':'a@example.org','aft x':'b@example.org'}
    assert resolve_contact(cfg,'AFT X')==''
    conf=Configuration(tmp_path/'config');conf.save(cfg);assert conf.path.with_suffix('.bak').exists()
    conf.path.write_text('invalid')
    with pytest.raises(ValueError,match='copia'):Configuration(tmp_path/'config')

def test_split_word_tokens_keep_styles_and_fail_missing_variables(tmp_path):
    doc=Document();p=doc.add_paragraph();p.add_run('{{NO').bold=True;p.add_run('MBRE}}')
    doc.sections[0].header.paragraphs[0].text='{{RIT}}'
    source=tmp_path/'template.docx';doc.save(source)
    out=tmp_path/'out.docx';fill_docx(source,out,{'NOMBRE':'Persona Ejemplo','RIT':'X-1'})
    got=Document(out);assert got.paragraphs[0].text=='Persona Ejemplo' and got.paragraphs[0].runs[0].bold
    assert got.sections[0].header.paragraphs[0].text=='X-1'
    with pytest.raises(ValueError):fill_docx(source,tmp_path/'bad.docx',{'NOMBRE':'Persona'})
    assert not (tmp_path/'bad.docx').exists()

def test_empty_to_is_saved_with_mandatory_cc_and_repeats_blocked(tmp_path,monkeypatch):
    from types import SimpleNamespace
    w=exported(tmp_path);d=Draft('Asunto','Cuerpo',cc='',key='uno');seen=[]
    monkeypatch.setattr('nurus.personal.outputs.save_draft',lambda p,**kw:seen.append(p) or SimpleNamespace(entry_id='e',store_id='s'))
    create_draft(w,d,confirmed=True)
    assert seen[0].recipient=='' and seen[0].cc==CC
    with pytest.raises(ValueError,match='ya se creó'):create_draft(w,d,confirmed=True)

@pytest.mark.parametrize('court,program,days,expected',[
    ('LAJA','DCE PRUEBA',29,[]),('LAJA','DCE PRUEBA',30,['programa_espera']),
    ('LAJA','AFT PRUEBA',29,[]),('LAJA','AFT PRUEBA',30,['programa_espera','PC_IE']),
    ('MULCHEN','AFT PRUEBA',30,['programa_espera','PC_IE']),
    ('TOME','AFT PRUEBA',30,['programa_espera']),('TOME','AFT PRUEBA',59,['programa_espera']),
    ('TOME','AFT PRUEBA',60,['programa_espera','PC_IE'])])
def test_assistant_waiting_profile(tmp_path,court,program,days,expected):
    path=source(tmp_path,program=program,days=days);book=load_workbook(path);book.active['B4']=court;book.save(path)
    work=Work(defaults()).analyze(path,'ESPERA',as_of=date(2026,9,13))
    assert work.rows[0].actions==expected

def test_excluded_rows_remain_colored_and_not_in_mails(tmp_path):
    w=Work(defaults()).analyze(source(tmp_path,program='CESFAM PRUEBA'),'ESPERA')
    w.export(tmp_path/'out.xlsx',backend='portable',reduced_fidelity=True)
    book=load_workbook(w.output);assert book[w.sheet].max_row==4
    assert len(book[w.sheet].conditional_formatting)>0
    assert prepare_drafts(w,'programa_espera')==[]

def test_config_threshold_changes_only_next_analysis(tmp_path):
    path=source(tmp_path,program='DCE PRUEBA',days=40);cfg=defaults();work=Work(cfg).analyze(path,'ESPERA')
    cfg['umbrales']['espera_dce']=45
    assert work.rows[0].actions==['programa_espera']
    assert Work(cfg).analyze(path,'ESPERA').rows[0].actions==[]

def test_identity_change_cannot_attach_another_person_to_review(tmp_path):
    w=exported(tmp_path);b=load_workbook(w.output);b[w.sheet]['C4']='Otra persona';b.save(w.output)
    with pytest.raises(ValueError,match='identidad'):w.refresh()
    assert not w.rows[0].review

def test_invalid_future_exit_is_an_integrity_warning(tmp_path):
    path=source(tmp_path,'CUMPLIMIENTO');b=load_workbook(path);b.active['G4']=-1;b.save(path)
    w=Work(defaults()).analyze(path,'CUMPLIMIENTO',as_of=date(2026,9,13))
    assert w.rows[0].warnings and 'se visualiza vencida' not in w.rows[0].observation

def test_external_records_do_not_require_motor_roundtrip(tmp_path):
    b=Workbook();b.active.append(['RIT','TRIBUNAL','NOMBRE','PROGRAMA']);b.active.append(['X-1','LAJA','Persona','AFT EJEMPLO']);path=tmp_path/'externos.xlsx';b.save(path)
    w=Work.external(path,defaults());assert not w.rows[0].rules
    ds=prepare_drafts(w,'programa_espera',selected=[w.rows[0].id],manual_selection=True)
    assert len(ds)==1 and ds[0].to==''

def test_word_template_inventory_and_real_generation(tmp_path):
    from nurus.personal.config import BASE
    from nurus.personal.resolutions import prepare_projects,generate_projects
    from nurus.personal.outputs import template_variables
    folder=BASE/'plantillas_word';assert len(list(folder.rglob('*.docx')))==6
    w=exported(tmp_path);out=tmp_path/'proyecto.docx'
    projects,errors=prepare_projects(w,[(w.rows[0].id,'PC_IE')],folder)
    assert not errors and len(projects)==1
    generate_projects(w,projects,out)
    assert not template_variables(out)
    text='\n'.join(p.text for p in Document(out).paragraphs)
    assert 'Persona Ejemplo' in text and 'X-1-2026' in text

def test_statistics_do_not_count_proposals_as_completed_reviews(tmp_path):
    from nurus.personal.statistics import summarize
    w=exported(tmp_path);assert summarize(w)['constancias']==0
    w.rows[0].review={'OBSERVACION':'Revisado','FECHA_OBS':'2026-09-13','TT':1,'CC':0}
    totals=summarize(w);assert totals['constancias']==1 and totals['sin_carga']==1


def test_resolution_dates_are_fully_written_in_words():
    assert date_in_words(date(2026,9,14))=='catorce de septiembre de dos mil veintiséis'
    assert date_in_words(date(2024,1,1))=='uno de enero de dos mil veinticuatro'
    assert date_in_words(date(2000,12,31))=='treinta y uno de diciembre de dos mil'

@pytest.mark.parametrize('mode,expected_specific',[('ESPERA','Lista de espera'),('CUMPLIMIENTO','Medidas sin vigencia'),('INFORMES','Informes pendientes')])
def test_prepare_required_drafts_includes_general_and_specific_mail(tmp_path,mode,expected_specific):
    w=exported(tmp_path,mode)
    drafts=prepare_required_drafts(w,modalities='Intervención ambulatoria de reparación',modality_keys=['AMB'])
    subjects=[d.subject for d in drafts]
    assert any('Revisión pestaña' in subject for subject in subjects)
    assert any(expected_specific in subject for subject in subjects)
    assert len(drafts)>=2

def test_exported_review_recovers_rule_actions_when_reloaded(tmp_path):
    w=exported(tmp_path,'ESPERA')
    book=load_workbook(w.output)
    headers={cell.value:cell.column for cell in book[w.sheet][w.header]}
    assert 'NURUS_REGLAS' in headers
    assert book[w.sheet].column_dimensions[book[w.sheet].cell(w.header,headers['NURUS_REGLAS']).column_letter].hidden
    reloaded=Work.external(w.output,defaults(),mode='ESPERA')
    assert reloaded.rows[0].rules==w.rows[0].rules
    assert reloaded.rows[0].actions==w.rows[0].actions
    drafts=prepare_required_drafts(reloaded,modalities='Intervención ambulatoria de reparación',modality_keys=['AMB'])
    assert len(drafts)>=2

def test_previous_export_without_hidden_rules_uses_trace_fallback(tmp_path):
    w=exported(tmp_path,'ESPERA')
    book=load_workbook(w.output)
    sheet=book[w.sheet]
    headers={cell.value:cell.column for cell in sheet[w.header]}
    sheet.delete_cols(headers['NURUS_REGLAS'])
    book.save(w.output)
    reloaded=Work.external(w.output,defaults(),mode='ESPERA')
    assert reloaded.rows[0].actions==w.rows[0].actions


def test_prepare_required_drafts_creates_each_program_mail(tmp_path):
    path=source(tmp_path,'ESPERA',program='AFT UNO',days=30)
    book=load_workbook(path);sheet=book.active
    values=[cell.value for cell in sheet[4]]
    values[0]='X-2-2026';values[2]='Persona Dos';values[3]='22222222-2';values[4]='AFT DOS'
    sheet.append(values);book.save(path)
    work=Work(defaults()).analyze(path,'ESPERA',as_of=date(2026,9,13))
    work.export(tmp_path/'salida_multi.xlsx',backend='portable',reduced_fidelity=True)
    drafts=prepare_required_drafts(work,modalities='Intervención ambulatoria de reparación',modality_keys=['AMB'])
    program_subjects=[d.subject for d in drafts if d.required]
    assert len(drafts)==3
    assert any('AFT UNO' in subject for subject in program_subjects)
    assert any('AFT DOS' in subject for subject in program_subjects)
