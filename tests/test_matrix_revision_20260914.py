from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document

from nurus.personal.config import BASE
from nurus.personal.template_package import BUNDLED_REVISION, install_bundled_templates


def document_hash(path):
    with ZipFile(path) as archive:
        return hashlib.sha256(archive.read('word/document.xml')).hexdigest()


def test_user_matrix_revision_patches_all_five_bodies_with_backup(tmp_path):
    source=BASE/'plantillas_word';target=tmp_path/'plantillas_word'
    shutil.copytree(source,target)
    custom=target/'LAJA/PC_IE.docx'
    doc=Document(custom);doc.add_paragraph('EDICION PREVIA DEL USUARIO');doc.save(custom)

    result=install_bundled_templates(source,target)
    data=json.loads((BASE/'matrix_document_patches.json').read_text(encoding='utf-8'))
    assert data['revision']==BUNDLED_REVISION
    assert len(data['templates'])==5
    assert len(result['patched'])==5
    for relative,item in data['templates'].items():
        path=target/relative
        assert path.is_file()
        assert document_hash(path)==item['sha256']
        Document(path)

    backups=[Path(p) for p in result['backups'] if 'PC_IE' in p and 'LAJA' in p]
    assert backups
    assert 'EDICION PREVIA DEL USUARIO' in '\n'.join(p.text for p in Document(backups[0]).paragraphs)


def test_same_revision_preserves_later_manual_edit(tmp_path):
    source=BASE/'plantillas_word';target=tmp_path/'plantillas_word'
    install_bundled_templates(source,target)
    path=target/'MULCHEN/PC_INFO.docx'
    doc=Document(path);doc.add_paragraph('EDICION POSTERIOR DEL USUARIO');doc.save(path)
    before=path.read_bytes()

    result=install_bundled_templates(source,target)
    assert path.read_bytes()==before
    assert not result['patched']
    assert 'EDICION POSTERIOR DEL USUARIO' in '\n'.join(p.text for p in Document(path).paragraphs)


def test_same_revision_restores_deleted_matrix_with_current_body(tmp_path):
    source=BASE/'plantillas_word';target=tmp_path/'plantillas_word'
    install_bundled_templates(source,target)
    data=json.loads((BASE/'matrix_document_patches.json').read_text(encoding='utf-8'))
    relative='LAJA/NOMENCL.docx';path=target/relative
    expected=data['templates'][relative]['sha256']
    assert document_hash(path)==expected

    path.unlink()
    result=install_bundled_templates(source,target)

    assert path.is_file()
    assert str(path) in result['installed']
    assert str(path) in result['patched']
    assert document_hash(path)==expected
    Document(path)
