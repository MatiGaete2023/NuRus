from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from nurus.rus import Mode
from nurus.services.exports import ExportError, export_review_snapshot
from nurus.services.workflow import WorkController
from nurus.storage.database import Database


def _approved_batch(db, path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Espera"
    sheet.append(["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "T ESPERA"])
    sheet.append(["PRM Norte", "Juzgado de Laja", "Ana", "C-1", 45])
    book.save(path)
    controller = WorkController(db)
    batch = controller.analyze(path, Mode.ESPERA, as_of=date(2026, 9, 8))
    row = controller.rows_from_batch(batch)[0]
    controller.approve_record(
        batch.batch_id,
        row.record_id,
        observation="=2+2",
        reason="Prueba de texto potencialmente interpretable como fórmula",
    )
    controller.approve_batch(batch.batch_id)
    return batch


def test_xlsx_export_uses_approved_snapshot_and_neutralizes_formula_text(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    batch = _approved_batch(db, tmp_path / "source.xlsx")
    target = tmp_path / "salida.xlsx"

    result = export_review_snapshot(db, batch.batch_id, target)

    assert result.path == target.resolve()
    assert len(result.sha256) == 64
    assert result.row_count == 1
    book = load_workbook(target, data_only=False)
    sheet = book["Revision"]
    headers = [cell.value for cell in sheet[1]]
    observation_col = headers.index("NURUS_OBSERVACION") + 1
    assert sheet.cell(2, observation_col).value == "'=2+2"
    trace = book["Trazabilidad"]
    trace_values = {trace.cell(row, 1).value: trace.cell(row, 2).value for row in range(2, trace.max_row + 1)}
    assert trace_values["snapshot_hash"] == db.get_batch(batch.batch_id)["snapshot_hash"]


def test_csv_export_uses_semicolon_utf8_and_formula_is_text(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    batch = _approved_batch(db, tmp_path / "source.xlsx")
    target = tmp_path / "salida.csv"

    export_review_snapshot(db, batch.batch_id, target)
    content = target.read_text(encoding="utf-8-sig")
    assert ";" in content.splitlines()[0]
    assert "'=2+2" in content


def test_existing_destination_is_not_overwritten_without_explicit_permission(tmp_path):
    db = Database(tmp_path / "nurus.sqlite3")
    batch = _approved_batch(db, tmp_path / "source.xlsx")
    target = tmp_path / "salida.csv"
    target.write_text("existente", encoding="utf-8")

    with pytest.raises(ExportError, match="ya existe"):
        export_review_snapshot(db, batch.batch_id, target)
    assert target.read_text(encoding="utf-8") == "existente"
