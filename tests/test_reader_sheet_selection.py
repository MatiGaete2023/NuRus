from hashlib import sha256
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from nurus.rus.models import Mode
from nurus.rus.reader import read_workbook, WorkbookReadError


def make_book(path, names):
    book = Workbook()
    book.remove(book.active)
    for name in names:
        sheet = book.create_sheet(name)
        sheet.append(["RIT", "TRIBUNAL", "DERIVACION"])
        sheet.append(["X-1", "Juzgado de Laja", "PRM Prueba"])
    book.save(path)


@pytest.mark.parametrize("mode,title", [
    (Mode.ESPERA, "Espera"),
    (Mode.CUMPLIMIENTO, "Cumplimiento"),
    (Mode.INFORMES, "Informes"),
])
def test_mode_selects_its_sheet_and_preserves_source(tmp_path, mode, title):
    path = tmp_path / "multi.xlsx"
    make_book(path, ["OB", "Espera", "Cumplimiento", "Informes", "Medidas vencidas"])
    original = path.read_bytes()
    result = read_workbook(path, mode)
    assert result.primary_sheet == title
    assert result.records[0].source.sheet_name == title
    assert result.records[0].source.row_number == 2
    assert result.workbook_sha256 == sha256(original).hexdigest()
    assert path.read_bytes() == original


def test_unnamed_multisheet_requires_explicit_selection(tmp_path):
    path = tmp_path / "unknown.xlsx"
    make_book(path, ["Primera", "Segunda"])
    with pytest.raises(WorkbookReadError, match="explícitamente"):
        read_workbook(path, Mode.ESPERA)
    assert read_workbook(path, Mode.ESPERA, sheet_name="Segunda").primary_sheet == "Segunda"


@pytest.mark.parametrize("name", ["OB", "Medidas vencidas"])
def test_excluded_sheet_cannot_be_explicit_primary(tmp_path, name):
    path = tmp_path / "excluded.xlsx"
    make_book(path, [name])
    with pytest.raises(WorkbookReadError, match="no se procesan"):
        read_workbook(path, Mode.ESPERA, sheet_name=name)


def test_single_unnamed_sheet_remains_supported(tmp_path):
    path = tmp_path / "single.xlsx"
    make_book(path, ["Hoja1"])
    assert read_workbook(path, Mode.ESPERA).primary_sheet == "Hoja1"


def test_autodetects_displaced_header_and_preserves_physical_rows(tmp_path):
    path = tmp_path / "displaced.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "informe_1"
    for _ in range(13):
        sheet.append(["REPORTE DE CUMPLIMIENTO"])
    sheet.append([
        "RIT", "TRIBUNAL", "RUT", "NOMBRE", "DERIVACIÓN",
        "DÍAS DE CUMPLIMIENTO", "DÍAS PARA EGRESAR",
    ])
    sheet.append([
        "X-1-2026", "Jgdo. L. y G. de Laja", "11.111.111-1", "NNA Prueba",
        "PRM Prueba", 120, 60,
    ])
    book.save(path)

    result = read_workbook(path, Mode.CUMPLIMIENTO)

    assert result.header_row == 14
    assert result.records[0].source.sheet_name == "informe_1"
    assert result.records[0].source.row_number == 15
    assert result.column_mapping["programa"] == "DERIVACIÓN"
    assert result.column_mapping["dias_cumpl"] == "DÍAS DE CUMPLIMIENTO"
    assert any(item.startswith("HEADER_ROW_AUTODETECTED") for item in result.warnings)


def test_displaced_header_scans_loaded_rows_without_reopening_each_prefix(tmp_path):
    path = tmp_path / "displaced.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "informe_1"
    for _ in range(13):
        sheet.append(["REPORTE DE CUMPLIMIENTO"])
    sheet.append([
        "RIT", "TRIBUNAL", "NOMBRE", "DERIVACION",
        "DIAS DE CUMPLIMIENTO", "DIAS PARA EGRESAR",
    ])
    sheet.append(["X-1", "Laja", "Persona", "PRM", 50, 50])
    book.save(path)

    import pandas as pd
    with patch("pandas.read_excel", wraps=pd.read_excel) as reader:
        result = read_workbook(path, Mode.CUMPLIMIENTO)

    assert result.header_row == 14
    assert reader.call_count == 3


def test_structural_formula_footer_is_not_a_review_record(tmp_path):
    path = tmp_path / "totals.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "ESPERA"
    sheet.append(["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "T ESPERA", "TT"])
    sheet.append(["PRM Norte", "Laja", "Persona", "X-1", 40, 1])
    sheet.append([None, None, None, None, None, "=SUM(F2:F2)"])
    book.save(path)

    result = read_workbook(path, Mode.ESPERA)

    assert len(result.records) == 1
    assert result.records[0].source.row_number == 2


def test_reports_preserves_both_ingress_dates_without_ambiguity(tmp_path):
    path = tmp_path / "informes.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "INFORMES"
    sheet.append([
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "FECHA VENCIMIENTO",
        "FECHA INGRESO", "FEC. INGRESO EFECTIVO",
    ])
    sheet.append(["PRM Norte", "Laja", "Persona", "X-1", "20/09/2026", "01/01/2025", "02/02/2025"])
    book.save(path)

    result = read_workbook(path, Mode.INFORMES)

    assert result.column_mapping["ingreso"] == "FEC. INGRESO EFECTIVO"
    assert result.records[0].values["FECHA INGRESO"] == "01/01/2025"
    assert result.records[0].values["FEC. INGRESO EFECTIVO"] == "02/02/2025"


def generic_cumplimiento(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Hoja1"
    sheet.append(["RIT", "TRIBUNAL", "DERIVACIÓN", "DÍAS DE CUMPLIMIENTO", "DÍAS PARA EGRESAR"])
    sheet.append(["X-1", "Laja", "PRM Norte", 50, 60])
    cross = book.create_sheet("Hoja2")
    cross.append(["RIT", "TRIBUNAL", "RUT MENOR", "NOMBRE MENOR", "NOMBRE CENTRO", "FECHA VENCIMIENTO"])
    cross.append(["X-1", "Laja", "111-1", "Prueba", "PRM Norte", "2026-09-30"])
    book.save(path)


def test_detects_primary_by_columns_and_keeps_cross(tmp_path):
    path = tmp_path / "generic.xlsx"
    generic_cumplimiento(path)
    result = read_workbook(path, Mode.CUMPLIMIENTO)
    assert result.primary_sheet == "Hoja1"
    assert result.cross_sheet == "Hoja2"
    assert result.records[0].source.row_number == 2
    assert len(result.cross_records) == 1


def test_wrong_mode_is_explained_not_silently_changed(tmp_path):
    path = tmp_path / "generic.xlsx"
    generic_cumplimiento(path)
    with pytest.raises(WorkbookReadError, match="No hay una tabla compatible con ESPERA.*CUMPLIMIENTO"):
        read_workbook(path, Mode.ESPERA)


def test_does_not_choose_richer_sheet_over_another_valid_table(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    book = Workbook()
    book.active.title = "Primera"
    for sheet in (book.active, book.create_sheet("Segunda")):
        sheet.append(["RIT", "TRIBUNAL", "DERIVACION", "T ESPERA"])
        sheet.append(["X-1", "Laja", "PRM", 45])
    book["Primera"].cell(1, 5, "NOMBRE")
    book["Primera"].cell(2, 5, "Persona")
    book.save(path)
    with pytest.raises(WorkbookReadError, match="más de una tabla"):
        read_workbook(path, Mode.ESPERA)


def test_generic_espera_with_notes_detects_header_once_per_sheet(tmp_path):
    path = tmp_path / "espera.xlsx"
    book = Workbook()
    book.active.title = "informe_1"
    book.active.append(["Reporte"])
    book.active.append(["RIT", "TRIBUNAL", "DERIVACION", "T ESPERA"])
    book.active.append(["X-1", "Laja", "PRM", 45])
    book.create_sheet("Notas").append(["No contiene casos"])
    book.save(path)
    import pandas as pd
    with patch("pandas.read_excel", wraps=pd.read_excel) as reader:
        result = read_workbook(path, Mode.ESPERA)
    assert result.primary_sheet == "informe_1"
    assert result.records[0].source.row_number == 3
    assert reader.call_count == 3  # dos cabeceras y una tabla completa


def test_sheet_inspection_lists_exact_names_without_analysis(tmp_path):
    from nurus.rus.reader import list_workbook_sheets
    path = tmp_path / "source.xlsx"
    make_book(path, ["OB", "Primera", "Segunda"])
    original = path.read_bytes()
    assert list_workbook_sheets(path) == ("OB", "Primera", "Segunda")
    assert path.read_bytes() == original
