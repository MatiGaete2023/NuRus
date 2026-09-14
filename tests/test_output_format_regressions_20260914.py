from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from nurus.personal.config import defaults
from nurus.personal.outputs import prepare_drafts
from nurus.personal.work import Work


def test_native_excel_date_is_rendered_without_time_in_program_attachment(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Informes'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','FECHA VENCIMIENTO','OBSERVACION'])
    sheet.append(['X-20-2026','LAJA','Persona Informe','22222222-2','PRM EJEMPLO',datetime(2026,9,30),'Texto'])
    path=tmp_path/'informes_fecha.xlsx';book.save(path)

    work=Work.external(path,defaults(),mode='INFORMES')
    drafts=prepare_drafts(work,'programa_por_vencer',directory=tmp_path/'salida')

    assert len(drafts)==1
    out=load_workbook(drafts[0].attachments[0]).active
    assert out['E2'].value=='30/09/2026'
    assert out['E2'].data_type=='s'


def test_zero_waiting_days_is_not_lost_in_program_attachment(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION'])
    sheet.append(['X-21-2026','LAJA','Persona Espera','11111111-1','PRM EJEMPLO',0,'Texto'])
    path=tmp_path/'espera_cero.xlsx';book.save(path)

    work=Work.external(path,defaults(),mode='ESPERA')
    drafts=prepare_drafts(work,'programa_espera',directory=tmp_path/'salida')

    assert len(drafts)==1
    out=load_workbook(drafts[0].attachments[0]).active
    assert out['E2'].value=='0'
    assert out['E2'].data_type=='s'
