from datetime import date
import json
import sqlite3

import pytest
from openpyxl import Workbook

from nurus.rus import Mode
from nurus.services.workflow import WorkController
from nurus.services.products import prepare_from_snapshot, approve_product, persist_approved_product
from nurus.services.exports import export_review_snapshot, ExportError
from nurus.storage.database import Database


def fixture(tmp_path):
    source = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Espera"
    sheet.append(["RIT", "TRIBUNAL", "DERIVACION", "T ESPERA"])
    sheet.append(["X-1", "Juzgado de Laja", "PRM Prueba", 45])
    book.save(source)
    db = Database(tmp_path / "test.sqlite3")
    controller = WorkController(db)
    batch = controller.analyze(source, Mode.ESPERA, as_of=date(2026, 9, 8))
    return db, controller, batch.batch_id, source


@pytest.mark.parametrize("failure", ["snapshot", "batch"])
def test_failure_rolls_back_rows_and_snapshot(tmp_path, failure):
    db, controller, batch, _ = fixture(tmp_path)
    with db.connect() as conn:
        if failure == "snapshot":
            conn.execute("CREATE TRIGGER fail_snapshot BEFORE INSERT ON approved_snapshots BEGIN SELECT RAISE(ABORT, 'injected'); END")
        else:
            conn.execute("CREATE TRIGGER fail_batch BEFORE UPDATE ON batches WHEN NEW.status='approved' BEGIN SELECT RAISE(ABORT, 'injected'); END")
    with pytest.raises(sqlite3.IntegrityError, match="injected"):
        controller.approve_batch(batch)
    assert db.list_review_records(batch)[0]["decision"] == "pending"
    assert db.get_batch(batch)["status"] == "review"
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM approved_snapshots").fetchone()[0] == 0


def test_historical_snapshot_survives_reapproval_and_source_removal(tmp_path):
    db, controller, batch, source = fixture(tmp_path)
    first = controller.approve_batch(batch)
    original = db.get_snapshot(batch, first)
    record = original["records"][0]
    source.unlink()
    controller.approve_record(batch, record["record_id"], observation="Nueva revisión", reason="Verificado")
    with pytest.raises(ValueError, match="aprobado"):
        db.get_snapshot(batch)
    assert db.get_snapshot(batch, first) == original
    second = controller.approve_batch(batch)
    assert second != first
    assert db.get_snapshot(batch)["records"][0]["edited_observation"] == "Nueva revisión"
    assert db.get_snapshot(batch, first) == original
    assert original["records"][0]["source_row"] == 2
    assert original["records"][0]["source_sheet"] == "Espera"


def test_product_and_export_read_frozen_records_not_working_table(tmp_path):
    db, controller, batch, _ = fixture(tmp_path)
    controller.approve_batch(batch)
    before = prepare_from_snapshot(db, batch, "email-ingreso")
    # Simula una modificación fuera del servicio: el snapshot sigue siendo la fuente.
    with db.connect() as conn:
        conn.execute("UPDATE review_records SET edited_observation='ALTERADO',values_json='{}' WHERE batch_id=?", (batch,))
        conn.execute("UPDATE batches SET column_mapping='{}' WHERE id=?", (batch,))
    after = prepare_from_snapshot(db, batch, "email-ingreso")
    assert before.rendered_body == after.rendered_body
    target = tmp_path / "frozen.csv"
    export_review_snapshot(db, batch, target)
    assert "ALTERADO" not in target.read_text(encoding="utf-8-sig")
    assert "X-1" in target.read_text(encoding="utf-8-sig")


def test_old_preview_cannot_be_saved_as_new_snapshot(tmp_path):
    db, controller, batch, _ = fixture(tmp_path)
    first = controller.approve_batch(batch)
    old = prepare_from_snapshot(db, batch, "email-ingreso")
    assert old.source_snapshot_hash == first
    row = db.list_review_records(batch)[0]
    controller.approve_record(batch, row["record_id"], observation="Nueva", reason="Corrección")
    controller.approve_batch(batch)
    approve_product(old)
    with pytest.raises(ValueError, match="vista previa"):
        persist_approved_product(db, old)
    assert not db.list_products()


def test_repeated_approval_is_idempotent(tmp_path):
    db, controller, batch, _ = fixture(tmp_path)
    first = controller.approve_batch(batch)
    approved_at = db.get_batch(batch)["approved_at"]
    assert controller.approve_batch(batch) == first
    assert db.get_batch(batch)["approved_at"] == approved_at
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM approved_snapshots").fetchone()[0] == 1


@pytest.mark.parametrize("statement", [
    "UPDATE approved_snapshots SET payload='{}'",
    "DELETE FROM approved_snapshots",
])
def test_snapshot_cannot_be_updated_or_deleted(tmp_path, statement):
    db, controller, batch, _ = fixture(tmp_path)
    controller.approve_batch(batch)
    with pytest.raises(sqlite3.IntegrityError, match="inmutable"):
        with db.connect() as conn:
            conn.execute(statement)


def test_legacy_hash_is_not_reconstructed_silently(tmp_path):
    db, controller, batch, _ = fixture(tmp_path)
    with db.connect() as conn:
        conn.execute("UPDATE review_records SET decision='approved' WHERE batch_id=?", (batch,))
        conn.execute("UPDATE batches SET status='approved',snapshot_hash='legacy-hash' WHERE id=?", (batch,))
    with pytest.raises(ValueError, match="histórico"):
        db.get_snapshot(batch)
    with pytest.raises(ExportError, match="histórico"):
        export_review_snapshot(db, batch, tmp_path / "old.csv")
    new_hash = controller.approve_batch(batch)
    assert db.get_snapshot(batch)["batch"]["snapshot_hash"] == new_hash


def test_snapshot_cannot_be_read_under_another_batch(tmp_path):
    db, controller, batch, _ = fixture(tmp_path)
    digest = controller.approve_batch(batch)
    with pytest.raises(ValueError, match="histórico"):
        db.get_snapshot("another-batch", digest)


def test_migration_backs_up_existing_database_and_rejects_downgrade(tmp_path):
    path = tmp_path / "old.sqlite3"
    with sqlite3.connect(path) as conn:
        conn.execute("CREATE TABLE preserved(value TEXT)")
        conn.execute("INSERT INTO preserved VALUES('original')")
        conn.execute("PRAGMA user_version=2")
    db = Database(path)
    assert db.migration_backup.is_file()
    with sqlite3.connect(db.migration_backup) as backup:
        assert backup.execute("PRAGMA user_version").fetchone()[0] == 2
        assert backup.execute("SELECT value FROM preserved").fetchone()[0] == "original"
    assert Database(path).migration_backup is None
    with db.connect() as conn:
        conn.execute("PRAGMA user_version=99")
    with pytest.raises(ValueError, match="degradarla"):
        Database(path)


@pytest.mark.parametrize("alias", ["direct", "relative", "hardlink"])
def test_export_never_overwrites_source_even_with_confirmation(tmp_path, monkeypatch, alias):
    db, controller, batch, source = fixture(tmp_path)
    controller.approve_batch(batch)
    original = source.read_bytes()
    destination = source
    if alias == "relative":
        monkeypatch.chdir(tmp_path)
        destination = "input.xlsx"
    elif alias == "hardlink":
        import os
        destination = tmp_path / "alias.xlsx"
        os.link(source, destination)
    with pytest.raises(ExportError, match="archivo de origen"):
        export_review_snapshot(db, batch, destination, overwrite=True)
    assert source.read_bytes() == original
    assert not list(tmp_path.glob("*.tmp.xlsx"))


def test_export_to_another_file_preserves_source(tmp_path):
    db, controller, batch, source = fixture(tmp_path)
    controller.approve_batch(batch)
    original = source.read_bytes()
    destination = tmp_path / "review.xlsx"
    destination.write_bytes(b"previous export")
    result = export_review_snapshot(db, batch, destination, overwrite=True)
    assert result.path == destination
    assert result.row_count == 1
    assert source.read_bytes() == original
    from openpyxl import load_workbook
    book = load_workbook(destination)
    try:
        assert book["Revision"].max_row == 2
        assert "Trazabilidad" in book.sheetnames
    finally:
        book.close()
