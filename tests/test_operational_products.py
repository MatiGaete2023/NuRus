from datetime import date, datetime
from hashlib import sha256
import json

import pytest
from openpyxl import Workbook, load_workbook

from nurus.adapters.outlook import DraftReceipt, DraftSaveUncertain
from nurus.adapters.sent_mail import SentMailReport, export_sent_report, scan_sent_items
from nurus.domain.models import ProductStatus
from nurus.rus import Mode
from nurus.services.communications import attach_snapshot_table, prepare_communications
from nurus.services.delivery import save_approved_draft
from nurus.services.products import approve_product, persist_approved_product
from nurus.services.resolutions import export_resolution, prepare_resolution
from nurus.services.workflow import WorkController
from nurus.storage.database import Database


def _reviewed(tmp_path, court="Juzgado de Laja"):
    source = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "ESPERA"
    sheet.append(["RIT", "TRIBUNAL", "NOMBRE", "DERIVACION", "T ESPERA", "RUT", "DURACION", "FEC. RESOLUCIÓN"])
    sheet.append(["X-1", court, "Persona Uno", "PRM Uno", 70, "111-1", "6 meses", "2026-01-01"])
    book.save(source)
    db = Database(tmp_path / "nurus.sqlite3")
    controller = WorkController(db)
    batch = controller.analyze(source, Mode.ESPERA, as_of=date(2026, 9, 8))
    records = db.list_review_records(batch.batch_id)
    db.apply_review_import(
        batch.batch_id,
        source_name="revisado.xlsx",
        source_bytes=b"constancia-revisada",
        responsible="Revisora CSMP",
        updates=[{
            "record_id": records[0]["record_id"],
            "observation": records[0]["original_observation"],
            "review_date": "2026-09-08",
            "tt": "", "workload": "", "resolution": "",
        }],
    )
    controller.approve_batch(batch.batch_id, require_review_import=True)
    return db, batch, records[0]


def test_grouped_email_uses_constancy_mandatory_cc_and_limited_attachment(tmp_path):
    db, batch, record = _reviewed(tmp_path)
    products = prepare_communications(
        db, batch.batch_id, "programa-espera", modalities="PRM"
    )
    assert len(products) == 1
    product = products[0]
    assert product.status is ProductStatus.READY
    assert product.cc == "ucc_concepcion@pjud.cl"
    assert product.record_ids == (record["record_id"],)
    path = attach_snapshot_table(db, product, tmp_path / "adjuntos")
    assert sha256(path.read_bytes()).hexdigest() == product.attachments[0][1]
    assert load_workbook(path).active.max_row == 2


def test_explicit_empty_selection_does_not_prepare_entire_batch(tmp_path):
    db, batch, _ = _reviewed(tmp_path)
    assert prepare_communications(db, batch.batch_id, "programa-espera", modalities="PRM",
                                  selected_record_ids=[]) == []
    assert len(prepare_communications(db, batch.batch_id, "programa-espera", modalities="PRM",
                                     selected_record_ids=None)) == 1


def test_attachment_rejects_record_without_rus_constancy(tmp_path):
    from types import SimpleNamespace
    db = SimpleNamespace(get_snapshot=lambda *_: {"records": [
        {"record_id": "id", "decision": "approved", "rus_recorded": False}]})
    product = SimpleNamespace(batch_id="batch", source_snapshot_hash="hash", record_ids=("id",))
    with pytest.raises(ValueError, match="Selección"):
        attach_snapshot_table(db, product, tmp_path / "adjuntos")
    assert not (tmp_path / "adjuntos").exists()


def test_resolution_word_is_bound_to_approved_product_and_snapshot(tmp_path):
    db, batch, record = _reviewed(tmp_path)
    product = prepare_resolution(
        db, batch.batch_id, record["record_id"], "historica-laja-pc-ie",
        confirmed_review=True,
    )
    assert product.status is ProductStatus.READY
    approve_product(product)
    persist_approved_product(db, product)
    target = export_resolution(db, product, tmp_path / "proyecto.docx")
    assert target.is_file()
    with pytest.raises(ValueError, match="nuevo"):
        export_resolution(db, product, target)


@pytest.mark.parametrize("uncertain", [False, True])
def test_draft_receipt_or_uncertainty_prevents_automatic_duplicate(tmp_path, monkeypatch, uncertain):
    db, batch, _ = _reviewed(tmp_path)
    product = prepare_communications(db, batch.batch_id, "correo-espera", modalities="PRM")[0]
    approve_product(product)
    persist_approved_product(db, product)

    def fake_save(*_args, **_kwargs):
        if uncertain:
            raise DraftSaveUncertain("Guardado incierto")
        return DraftReceipt("entry", "store", "cuenta", "Borradores")

    monkeypatch.setattr("nurus.services.delivery.save_draft", fake_save)
    if uncertain:
        with pytest.raises(DraftSaveUncertain):
            save_approved_draft(db, product, confirmed=True)
    else:
        assert save_approved_draft(db, product, confirmed=True).entry_id == "entry"
    with pytest.raises(ValueError):
        save_approved_draft(db, product, confirmed=True)


def test_sent_mail_range_and_report_expose_limits(tmp_path):
    class Items(list):
        def Sort(self, *_args):
            self.sort(key=lambda item: getattr(item, "SentOn", datetime.min), reverse=True)

    class Mail:
        Class = 43
        To = "destino@example.test"
        Subject = "Asunto"
        EntryID = "id"
        def __init__(self, day):
            self.SentOn = datetime.combine(day, datetime.min.time())

    items = Items([Mail(date(2026, 9, 8)), Mail(date(2026, 9, 7))])
    rows, skipped, errors, truncated = scan_sent_items(
        items, date(2026, 9, 8), date(2026, 9, 8), limit=10
    )
    assert len(rows) == 1 and not skipped and not errors and not truncated
    report = SentMailReport(rows, 0, 1, True, "cuenta", "Enviados")
    path = export_sent_report(report, tmp_path / "enviados.xlsx")
    control = load_workbook(path)["Control"]
    assert dict(control.values)["Consulta limitada"] is True


def test_changed_constancy_blocks_materializing_old_email_and_word(tmp_path, monkeypatch):
    db, batch, record = _reviewed(tmp_path)
    email = prepare_communications(db, batch.batch_id, "correo-espera", modalities="PRM")[0]
    resolution = prepare_resolution(
        db, batch.batch_id, record["record_id"], "historica-laja-pc-ie",
        confirmed_review=True,
    )
    for product in (email, resolution):
        approve_product(product)
        persist_approved_product(db, product)

    db.set_record_decision(
        batch.batch_id,
        record["record_id"],
        "approved",
        observation=record["edited_observation"] + " Actualizada.",
        reason="Nueva constancia",
    )
    monkeypatch.setattr(
        "nurus.services.delivery.save_draft",
        lambda *_args, **_kwargs: pytest.fail("Outlook no debe invocarse"),
    )
    with pytest.raises(ValueError, match="constancia cambió"):
        save_approved_draft(db, email, confirmed=True)
    with pytest.raises(ValueError, match="constancia cambió"):
        export_resolution(db, resolution, tmp_path / "obsoleto.docx")
    assert not (tmp_path / "obsoleto.docx").exists()
