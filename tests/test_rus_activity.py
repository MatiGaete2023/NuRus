from datetime import date
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Border,Side,PatternFill
import pytest
from nurus.personal.config import defaults
from nurus.personal.work import Work
from nurus.personal.rus_activity import mark_reviewed,TITLES


def source(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='ESPERA'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','NOMBRE CENTRO','T ESPERA','SITFA_TRIBUNAL_CODIGO'])
    for code,person,program in [('777','Persona uno','PRM Centro uno'),('777','Persona dos','PRM Centro dos'),('888','Persona tres','PRM Centro uno')]:
        sheet.append(['X-1-2026','Tribunal ficticio '+code,person,'11111111-1',program,40,code])
        for cell in sheet[sheet.max_row]:cell.border=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    sheet.cell(2,6).fill=PatternFill(fill_type='solid',fgColor='ABCDEF')
    details=book.create_sheet('Resoluciones firmadas');details.append(['tribunal_codigo','tribunal','rit','tramite','firma','hora','documento','identidad','huella'])
    details.append(['777','Tribunal ficticio 777','X-1-2026','Resolución ficticia','2026-10-01','09:30','SI','huella_fila_sin_id_remoto','firma-1'])
    details.append(['777','Tribunal ficticio 777','X-1-2026','Otra resolución ficticia','2026-10-02','10:00','SI','huella_fila_sin_id_remoto','firma-2'])
    coverage=book.create_sheet('SITFA_COBERTURA');coverage.append(['tribunal_codigo','desde','hasta','estado','criterio_temporal','sin_coincidencias'])
    for code in ('777','888'):coverage.append([code,'2026-08-04','2026-10-02','PARCIAL','filtro_carga_no_confirmado','Sin coincidencias en los informes consultados'])
    path=tmp_path/'fuentes.xlsx';book.save(path);book.close();return path


def test_case_scope_and_nonhistorical_court_has_no_proposals(tmp_path):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    assert len(work.rows)==3
    first,second,other=work.rows
    assert work.signed_activity[first.id]['pendientes']==2
    assert work.signed_activity[second.id]['pendientes']==2
    assert work.signed_activity[other.id]['pendientes']==0
    assert 'informes consultados' in work.signed_activity[other.id]['valores']['ACTIVIDAD RECIENTE']
    assert all(not r.observation and not r.actions and not r.rules and not r.review for r in work.rows)


def test_review_persists_per_ingreso_never_per_case(tmp_path):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA',as_of=date(2026,10,2))
    first,second,other=work.rows;mark_reviewed(work,first.id,'firma-1');mark_reviewed(work,first.id,'firma-2')
    assert work.signed_activity[first.id]['pendientes']==0
    assert work.signed_activity[second.id]['pendientes']==2
    with pytest.raises(ValueError):mark_reviewed(work,other.id,'firma-1')
    work.save(tmp_path/'sesion');restored=Work.load(tmp_path/'sesion')
    assert set(restored.activity_reviewed[first.id])=={'firma-1','firma-2'}
    assert restored.signed_activity[second.id]['pendientes']==2
    assert all(not r.review for r in restored.rows)


def test_export_alerts_preserves_sources_styles_and_gestion_fields(tmp_path):
    path=source(tmp_path);original=path.read_bytes();work=Work(defaults()).analyze(path,'ESPERA',as_of=date(2026,10,2))
    output=Path(work.export(tmp_path/'final.xlsx',backend='portable',reduced_fidelity=True))
    assert path.read_bytes()==original
    book=load_workbook(output);sheet=book['ESPERA'];columns={c.value:c.column for c in sheet[1]}
    assert all(k in columns for k in TITLES)
    assert sheet.max_row==4
    assert sheet.cell(2,columns['ACTIVIDAD RECIENTE']).fill.fgColor.rgb.endswith('DDEBF7')
    assert sheet.cell(4,columns['ACTIVIDAD RECIENTE']).fill.fgColor.rgb.endswith('FFF2CC')
    assert sheet.cell(2,6).fill.fgColor.rgb.endswith('ABCDEF')
    assert sheet.cell(2,1).border.left.style=='thin'
    assert sheet.cell(2,columns['RESOLUCIONES EN EL PERÍODO']).value==2
    for key in ('FECHA_OBS','TT','CC'):assert sheet.cell(2,columns[key]).value in (None,'')
    assert {'Resoluciones firmadas','SITFA_COBERTURA'}.issubset(book.sheetnames)
    book.close()


def test_invalid_signature_or_coverage_does_not_turn_into_zero(tmp_path):
    path=source(tmp_path);book=load_workbook(path);book['Resoluciones firmadas'].cell(2,5,'2026-07-31');book.save(path);book.close()
    with pytest.raises(ValueError,match='fuera'):Work(defaults()).analyze(path,'ESPERA')
