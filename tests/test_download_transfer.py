from datetime import date
from hashlib import sha256
import json
from pathlib import Path
from types import SimpleNamespace
import pytest

from nurus.personal import download_link as link
from nurus.personal.download_transfer import Transfer, recovery_work, archive_current
from nurus.personal.config import defaults
from nurus.personal.work import Work
from current_fixture import source


class Variable:
    def __init__(self, value=''):
        self.value = value
    def get(self):
        return self.value
    def set(self, value):
        self.value = value


class App:
    def __init__(self, root, work=None):
        self.cfg = SimpleNamespace(directory=root/'config', data=defaults(), save=lambda data: None)
        self.work = work
        self.busy = False
        self.folder = Variable(str(root/'salida'))
        self.sheet = Variable()
        self.mode = Variable('ESPERA')
        self.status = Variable()
        self.exports = 0
        self.analyses = 0
        self.callbacks = []
    def _run(self, text, action, done):
        self.analyses += 1
        done(action())
    def _save_session(self):
        if self.work:
            self.work.save(self.cfg.directory/'sesion')
            archive_current(self.work,self.cfg.directory)
    def _show_work(self):
        pass
    def _clear_drafts(self):
        pass
    def _export_current(self):
        self.exports += 1
    def _continue_resume(self):
        self.work = self._resume_work
    def after(self, delay, callback):
        self.callbacks.append(callback)
    def winfo_exists(self):
        return True
    def _guard(self, callback):
        callback()


def ready(app, path, mode='ESPERA'):
    request = Transfer.create(app.cfg.directory/'intercambio', 'ESPERA', app.folder.get())
    request.reply.write_text(json.dumps(dict(archivo=str(path), modo=mode,
                                           sha256=sha256(path.read_bytes()).hexdigest())), encoding='utf-8')
    return request


def test_frozen_script_is_reselected_and_never_runs_csmp_as_python(tmp_path, monkeypatch):
    app = App(tmp_path)
    script = tmp_path/'descargador.py';script.write_text('raise RuntimeError')
    exe = tmp_path/'SITFA_Descargador.exe';exe.write_bytes(b'fake executable; never launched')
    app.cfg.data['descargador_integral'] = str(script)
    monkeypatch.setattr(link.sys, 'frozen', True, raising=False)
    choices = []
    def choose(**kwargs):
        choices.append(kwargs);return str(exe)
    monkeypatch.setattr(link.filedialog, 'askopenfilename', choose)
    calls = []
    monkeypatch.setattr(link.subprocess, 'Popen', lambda command, **kwargs:
                        (calls.append(command) or SimpleNamespace(poll=lambda: None)))
    link.start(app)
    assert choices[0]['filetypes'] == [('Descargador Windows', '*.exe')]
    assert calls[0][0] == str(exe.resolve()) and str(script) not in calls[0]
    ticket = next((app.cfg.directory/'intercambio').glob('*.solicitud.json'))
    request = Transfer.load(ticket.parent, ticket)
    assert calls[0][calls[0].index('--csmp-reply')+1] == str(request.reply)


def test_source_python_entry_point_remains_available(tmp_path, monkeypatch):
    script = tmp_path/'descargador.py';script.write_text('# not executed')
    monkeypatch.setattr(link.sys, 'frozen', False, raising=False)
    assert link.launch_command(script) == [link.sys.executable, str(script.resolve())]
    with pytest.raises(ValueError):link.launch_command(Path(link.sys.executable))


def test_failure_to_persist_request_prevents_launch(tmp_path, monkeypatch):
    app = App(tmp_path);exe=tmp_path/'download.exe';exe.write_bytes(b'fixture')
    app.cfg.data['descargador_integral'] = str(exe)
    import nurus.personal.download_transfer as storage
    monkeypatch.setattr(storage, 'atomic_json', lambda *args: (_ for _ in ()).throw(OSError('disk full')))
    calls=[];monkeypatch.setattr(link.subprocess, 'Popen', lambda *args, **kwargs: calls.append(args))
    with pytest.raises(OSError, match='disk full'):link.start(app)
    assert calls == []


def test_launch_failure_keeps_recoverable_request(tmp_path, monkeypatch):
    app=App(tmp_path);exe=tmp_path/'download.exe';exe.write_bytes(b'fixture')
    app.cfg.data['descargador_integral']=str(exe)
    monkeypatch.setattr(link.subprocess, 'Popen', lambda *args, **kwargs: (_ for _ in ()).throw(OSError('blocked')))
    with pytest.raises(OSError):link.start(app)
    path=next((app.cfg.directory/'intercambio').glob('*.solicitud.json'))
    assert Transfer.load(path.parent,path).data['state']=='ERROR_INICIO'


def test_second_launch_does_not_duplicate_live_download(tmp_path, monkeypatch):
    app=App(tmp_path);app._download_process=SimpleNamespace(poll=lambda:None)
    calls=[];monkeypatch.setattr(link.subprocess,'Popen',lambda *args,**kwargs:calls.append(args))
    with pytest.raises(ValueError,match='ya está abierto'):link.start(app)
    assert not calls and not (app.cfg.directory/'intercambio').exists()


@pytest.mark.parametrize('problem', ['mode','hash','columns'])
def test_invalid_return_preserves_current_work_and_does_not_export(tmp_path, problem):
    path=source(tmp_path);work=Work(defaults()).analyze(path,'ESPERA',as_of=date(2026,10,2))
    work.rows[0].review={'OBSERVACION':'Edición que se conserva','TT':0}
    app=App(tmp_path,work);request=ready(app,path,'CUMPLIMIENTO' if problem=='mode' else 'ESPERA')
    if problem=='hash':path.write_bytes(b'changed')
    if problem=='columns':
        from openpyxl import load_workbook
        book=load_workbook(path);book['ESPERA'].delete_cols(5);book.save(path);book.close()
        data=json.loads(request.reply.read_text());data['sha256']=sha256(path.read_bytes()).hexdigest()
        request.reply.write_text(json.dumps(data))
    with pytest.raises(ValueError):link.incorporate(app,request)
    assert app.work is work and work.rows[0].review['TT']==0 and app.exports==0
    assert Transfer.load(request.folder,request.path).data['state']=='PENDIENTE'


def test_result_after_csmp_restart_is_selected_without_starting_downloader(tmp_path, monkeypatch):
    path=source(tmp_path);app=App(tmp_path);request=ready(app,path)
    new_app=App(tmp_path)
    calls=[];monkeypatch.setattr(link.subprocess,'Popen',lambda *args,**kwargs:calls.append(args))
    link.resume(new_app)
    assert not calls and new_app.analyses==1 and new_app.exports==1
    assert new_app.work.source_hash==sha256(path.read_bytes()).hexdigest()
    assert Transfer.load(request.folder,request.path).data['state']=='INCORPORADA'
    assert (request.recovery_folder/'incorporada/trabajo.json').is_file()


def test_excel_failure_then_recovery_keeps_manual_edits_without_reanalysis(tmp_path):
    path=source(tmp_path);app=App(tmp_path);request=ready(app,path)
    app._export_current=lambda: (_ for _ in ()).throw(OSError('Excel unavailable'))
    with pytest.raises(OSError,match='Excel unavailable'):link.incorporate(app,request)
    app.work.rows[0].review={'OBSERVACION':'Edición posterior','CC':1,'TT':0}
    app._save_session()
    recovered=App(tmp_path)
    link.resume(recovered)
    assert recovered.analyses==0 and recovered.exports==0
    assert recovered.work.rows[0].review==app.work.rows[0].review


def test_cut_after_session_save_before_receipt_does_not_repeat_analysis(tmp_path, monkeypatch):
    path=source(tmp_path);app=App(tmp_path);request=ready(app,path)
    original_update=request.update
    def fail_finish(state, **facts):
        if state=='INCORPORADA':raise OSError('receipt failed')
        original_update(state,**facts)
    monkeypatch.setattr(request,'update',fail_finish)
    with pytest.raises(OSError,match='receipt failed'):link.incorporate(app,request)
    app.work.rows[0].review={'OBSERVACION':'Se conserva después del corte'};app._save_session()
    restarted=App(tmp_path);link.resume(restarted)
    assert restarted.analyses==0 and restarted.work.rows[0].review==app.work.rows[0].review
    assert Transfer.load(request.folder,request.path).data['state']=='INCORPORADA'


def test_incorporation_archives_previous_work_before_replacing_it(tmp_path):
    path=source(tmp_path);old=Work(defaults()).analyze(path,'ESPERA',as_of=date(2026,10,2))
    old.rows[0].review={'OBSERVACION':'Texto anterior'}
    app=App(tmp_path,old);request=ready(app,path);link.incorporate(app,request)
    snapshot=next((request.recovery_folder/'anteriores').iterdir())
    assert Work.load(snapshot).rows[0].review==old.rows[0].review


def test_request_path_escape_or_unknown_state_cannot_select_other_reply(tmp_path):
    request=Transfer.create(tmp_path/'a','ESPERA',tmp_path/'salida')
    with pytest.raises(ValueError):Transfer.load(tmp_path/'b',request.path)
    data=dict(request.data,state='pretend success');request.path.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='estado'):Transfer.load(request.folder,request.path)


def test_incorporated_session_mismatch_is_not_replaced_by_rit_match(tmp_path):
    path=source(tmp_path);app=App(tmp_path);request=ready(app,path);link.incorporate(app,request)
    data=dict(request.data,source_hash='a'*64);request.path.write_text(json.dumps(data))
    loaded=Transfer.load(request.folder,request.path)
    with pytest.raises(ValueError,match='corresponde'):recovery_work(loaded,app.work)


def test_starting_another_work_does_not_erase_edits_to_returned_download(tmp_path):
    path=source(tmp_path);app=App(tmp_path);request=ready(app,path);link.incorporate(app,request)
    app.work.rows[0].review={'OBSERVACION':'Texto editado después de descargar','CC':1}
    app._save_session()
    from openpyxl import load_workbook
    second=tmp_path/'otra.xlsx';book=load_workbook(path);book['ESPERA'].cell(2,3,'Otra persona');book.save(second);book.close()
    app.work=Work(defaults()).analyze(second,'ESPERA');app._save_session()
    restarted=App(tmp_path);link.resume(restarted)
    assert restarted.work.source_hash==request.data['source_hash']
    assert restarted.work.rows[0].review['OBSERVACION']=='Texto editado después de descargar'
    assert restarted.analyses==0


def test_other_transfer_cannot_receive_saved_edits_by_shared_rit(tmp_path):
    path=source(tmp_path);app=App(tmp_path);request=ready(app,path);link.incorporate(app,request)
    other=Transfer.create(request.folder,'CUMPLIMIENTO',tmp_path/'salida')
    app.work.download_transfer_id=other.identifier
    with pytest.raises(ValueError,match='corresponde'):archive_current(app.work,app.cfg.directory)
    assert not (other.recovery_folder/'incorporada/trabajo.json').exists()
