from copy import deepcopy
from pathlib import Path
import time

import pytest
from openpyxl import Workbook, load_workbook

from nurus.personal.config import defaults, validate, BASE
from nurus.personal.work import Work
from nurus.personal.sync import SyncConflict
from nurus.personal.record_edits import apply, undo
from nurus.personal.outputs import prepare_drafts, create_draft, draft_fingerprint
from nurus.personal.resolutions import automatic_project_selections, prepare_projects, resolution_review_issue


def source(path, *, external=False):
    book=Workbook(); sheet=book.active; sheet.title='Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION','RES','TT','CC'])
    sheet.append(['X-1','MULCHEN','Persona Uno','11111111-1','AFT PRUEBA',45,'Obs Uno','PC_IE',0,1])
    sheet.append(['X-2','MULCHEN','Persona Dos','22222222-2','AFT PRUEBA',45,'Obs Dos','PC_IE',1,0])
    book.create_sheet('Notas')['A1']='sin cambio'
    book.save(path)
    if external:
        return Work.external(path,defaults(),'ESPERA')
    work=Work(defaults()).analyze(path,'ESPERA')
    work.export(path.with_name('copia.xlsx'),backend='portable',reduced_fidelity=True)
    return work


def edit(work, change):
    book=load_workbook(work.output)
    change(book)
    book.save(work.output);book.close()


def test_refresh_keeps_local_edit_and_commits_snapshot_after_sort_and_header_move(tmp_path):
    work=source(tmp_path/'origen.xlsx',external=True)
    original=work.content
    first,second=work.rows
    first.review['OBSERVACION']='Edición local uno'
    def reorder(book):
        sheet=book['Espera']
        a=[cell.value for cell in sheet[2]];b=[cell.value for cell in sheet[3]]
        for index,value in enumerate(b,1):sheet.cell(2,index,value)
        for index,value in enumerate(a,1):sheet.cell(3,index,value)
        sheet.insert_rows(1,2)
    edit(work,reorder)
    assert work.refresh()
    assert work.rows[0].review['OBSERVACION']=='Edición local uno'
    assert work.rows[0].source_row==5
    assert work.header==3
    assert work.original_content==original
    destination=tmp_path/'resultado.xlsx'
    work.export(destination,backend='portable',reduced_fidelity=True)
    book=load_workbook(destination)
    data={row[0]:row for row in book['Espera'].iter_rows(min_row=4,values_only=True)}
    assert data['X-1'][6]=='Edición local uno'
    assert data['X-2'][6]=='Obs Dos'
    assert data['X-1'][8:10]==(0,1)
    assert data['X-2'][8:10]==(1,0)
    work.save(tmp_path/'session')
    recovered=Work.load(tmp_path/'session')
    assert recovered.original_content==original
    assert recovered.content==work.content


def test_unrelated_save_never_loses_local_edit_even_after_restart(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    work.rows[0].review['OBSERVACION']='Edición local'
    work.save(tmp_path/'session');work=Work.load(tmp_path/'session')
    edit(work,lambda book:setattr(book['Notas']['A1'],'value','cambio ajeno'))
    assert work.refresh()
    assert work.rows[0].review['OBSERVACION']=='Edición local'
    edit(work,lambda book:setattr(book['Notas']['A1'],'value','segundo cambio ajeno'))
    work.refresh()
    assert work.rows[0].review['OBSERVACION']=='Edición local'


def test_conflicts_are_transactional_and_resolution_is_explicit(tmp_path):
    work=source(tmp_path/'origen.xlsx',external=True)
    row=work.rows[0];row.review['OBSERVACION']='Local'
    before=deepcopy(work.__dict__)
    edit(work,lambda book:setattr(book['Espera']['G2'],'value','Excel'))
    with pytest.raises(SyncConflict) as caught:work.refresh()
    assert work.__dict__==before
    assert caught.value.conflicts[0].field=='OBSERVACION'
    work.refresh({(row.id,'OBSERVACION'):'Combinado'})
    assert work.rows[0].review['OBSERVACION']=='Combinado'
    edit(work,lambda book:setattr(book['Notas']['A1'],'value','otro cambio'))
    work.refresh()
    assert work.rows[0].review['OBSERVACION']=='Combinado'


def test_changed_identity_rejects_whole_transaction(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    before=deepcopy(work.__dict__)
    edit(work,lambda book:setattr(book['Espera']['C3'],'value','Otra Persona'))
    with pytest.raises(ValueError,match='identidad'):work.refresh()
    assert work.__dict__==before


def test_generators_require_explicit_sync_and_do_not_use_old_res_selection(tmp_path):
    work=source(tmp_path/'origen.xlsx',external=True)
    selected=automatic_project_selections(work)
    def remove(book):
        book['Espera']['H2']=None;book['Espera']['H3']=None
    edit(work,remove)
    before=deepcopy(work.__dict__)
    with pytest.raises(ValueError,match='Actualizar desde Excel'):
        prepare_projects(work,selected,BASE/'plantillas_word')
    assert work.__dict__==before
    work.refresh()
    assert automatic_project_selections(work)==[]


def test_generated_attachment_bytes_and_receipt_identity_survive_regeneration(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    first=prepare_drafts(work,'programa_espera',directory=tmp_path/'one')[0]
    time.sleep(2.1)
    second=prepare_drafts(work,'programa_espera',directory=tmp_path/'two')[0]
    assert Path(first.attachments[0]).read_bytes()==Path(second.attachments[0]).read_bytes()
    assert draft_fingerprint(first)==draft_fingerprint(second)
    book=load_workbook(second.attachments[0]);book.active['C2']='Editado';book.save(second.attachments[0])
    assert draft_fingerprint(first)!=draft_fingerprint(second)


def test_decision_survives_export_and_text_does_not_control_eligibility(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    for row in work.rows:row.review['OBSERVACION']='Texto personalizado con email'
    before=prepare_drafts(work,'programa_espera')
    work.export(tmp_path/'editado.xlsx',backend='portable',reduced_fidelity=True)
    imported=Work.external(work.output,work.config,'ESPERA')
    after=prepare_drafts(imported,'programa_espera')
    assert len(before)==len(after)==1
    assert before[0].record_ids==after[0].record_ids
    apply(imported,[row.id for row in imported.rows],decisions={'mail':{'programa_espera':'omit'}})
    imported.export(tmp_path/'omitido.xlsx',backend='portable',reduced_fidelity=True)
    restored=Work.external(imported.output,imported.config,'ESPERA')
    assert prepare_drafts(restored,'programa_espera')==[]


def test_prepared_product_cannot_be_saved_after_decision_change(tmp_path,monkeypatch):
    work=source(tmp_path/'origen.xlsx')
    draft=prepare_drafts(work,'programa_espera')[0]
    calls=[]
    monkeypatch.setattr('nurus.personal.outputs.save_draft',lambda *a,**k:calls.append(True))
    apply(work,[work.rows[0].id],decisions={'mail':{'programa_espera':'omit'}})
    with pytest.raises(ValueError,match='actualizarse'):create_draft(work,draft,confirmed=True)
    assert calls==[]


@pytest.mark.parametrize('value',[0,False,'0','0.0','False'])
def test_negative_res_is_valid(value):
    assert resolution_review_issue(value)==''


@pytest.mark.parametrize('text',['{TRIBUNAL:{NO_EXISTE}}','{TRIBUNAL!r}','{TRIBUNAL.foo}','{TRIBUNAL[0]}','{TRIBUNAL','{}'])
def test_template_rejects_unsupported_syntax(text):
    cfg=defaults();cfg['correos']['plantillas']['espera']['asunto']=text
    with pytest.raises(ValueError):validate(cfg)


def test_bulk_edit_and_undo_preserve_zero_and_empty_date(tmp_path):
    work=source(tmp_path/'origen.xlsx')
    snapshot=deepcopy(work.rows)
    previous=apply(work,[row.id for row in work.rows],review={'FECHA_OBS':'','TT':'0','CC':'1'},decisions={'resolution':'NOMENCL'})
    assert all(row.review['TT']==0 and row.review['FECHA_OBS']=='' for row in work.rows)
    assert all(kind=='NOMENCL' for _,kind in automatic_project_selections(work))
    undo(work,previous)
    assert work.rows==snapshot
