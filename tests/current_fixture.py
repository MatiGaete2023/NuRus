from openpyxl import Workbook
from openpyxl.styles import Border,Side,PatternFill

def source(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='ESPERA'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','NOMBRE CENTRO','T ESPERA','SITFA_TRIBUNAL_CODIGO'])
    for code,person,program in [('777','Persona uno','PRM Centro uno'),('777','Persona dos','PRM Centro dos'),('888','Persona tres','PRM Centro uno')]:
        sheet.append(['X-1-2026','Tribunal ficticio '+code,person,'11111111-1',program,40,code])
        for cell in sheet[sheet.max_row]:cell.border=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    sheet.cell(2,6).fill=PatternFill(fill_type='solid',fgColor='ABCDEF')
    details=book.create_sheet('Resoluciones firmadas');details.append(['tribunal_codigo','tribunal','rit','tramite','firma','hora','documento','identidad','huella'])
    details.append(['777','Tribunal ficticio 777','X-1-2026','Resolución ficticia','2026-10-01','09:30','SI','huella_fila_sin_id_remoto','firma-1'])
    details.append(['777','Tribunal ficticio 777','X-1-2026','Otra resolución ficticia','2026-10-02','10:00','SI','huella_fila_sin_id_remoto','firma-2'])
    coverage=book.create_sheet('SITFA_COBERTURA');coverage.append(['tribunal_codigo','desde','hasta','estado','criterio_temporal','sin_coincidencias'])
    for code in ('777','888'):coverage.append([code,'2026-08-04','2026-10-02','PARCIAL','filtro_carga_no_confirmado','Sin coincidencias en los informes consultados'])
    path=tmp_path/'fuentes.xlsx';book.save(path);book.close();return path

