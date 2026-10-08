from dataclasses import replace
from datetime import date
from hashlib import sha256
import json
from openpyxl import load_workbook
import pytest

from nurus.personal.bitacoras import Observation,earliest,analyze,cc_for_type,export_audit


TODAY=date(2026,10,5);START=date(2026,6,5)
def entry(id='1',fecha='2026-10-01T09:00:00',**extra):
    values=dict(tribunal_codigo='777',causa_id='100',ingreso_id='200',entry_id=id,fecha=fecha,
      autor='Autor ficticio',tipo='Al Tribunal',etapa='Cumplimiento',texto='Consulta ficticia de un informe pendiente.',
      origen='centro',fuente='captura-1',fila_fuente=1)
    return Observation(**{**values,**extra})


def test_four_calendar_months_and_no_older_fallback():
    assert earliest(date(2026,6,30))==date(2026,2,28)
    with pytest.raises(ValueError):analyze([],date(2026,6,4),TODAY,today=TODAY)
    old=entry(fecha='2026-06-04T15:00:00')
    result=analyze([old],START,TODAY,today=TODAY,coverage='COMPLETA')
    assert result['entries']==[] and result['latest_center']==[]
    at_limit=entry(fecha='2026-06-05T00:00:00')
    assert analyze([at_limit],START,TODAY,today=TODAY)['latest_center']==[at_limit]


def test_last_center_any_author_and_date_ties_kept():
    one=entry();two=entry('2',autor='Otro autor')
    tribunal=entry('3',fecha='2026-10-04T15:00:00',origen='tribunal')
    result=analyze([one,two,tribunal],START,TODAY,today=TODAY)
    assert result['latest_center']==[one,two]
    assert analyze([one,two],START,TODAY,today=TODAY,author='Otro autor')['latest_center']==[two]


def test_empty_confirmed_is_distinct_from_unqueried_and_cannot_hide_entries():
    result=analyze([],START,TODAY,today=TODAY,coverage='VACIA_COMPROBADA')
    assert result['coverage']=='VACIA_COMPROBADA'
    with pytest.raises(ValueError,match='no puede tener cobertura vacía'):
        analyze([entry()],START,TODAY,today=TODAY,coverage='VACIA_COMPROBADA')


def test_type_charge_and_response_uncertainty_are_separate(tmp_path):
    assert cc_for_type('Al Tribunal')==1
    assert cc_for_type('Administrativa')==0
    assert cc_for_type('Comentario con carga') is None
    from nurus.personal.bitacoras import response_state
    assert 'no comprobable' in response_state(entry())
    assert response_state(entry(respuesta_comprobada=True))=='Sin respuesta registrada'
    assert response_state(entry(respuesta_comprobada=True,respuesta='Respuesta ficticia'))=='Respuesta registrada'
    assert 'no exige' in response_state(entry(tipo='Administrativa'))


def test_same_remote_entry_from_two_tabs_counted_once_real_repetitions_kept():
    one=entry();copy=replace(one,fuente='otra-pestaña',fila_fuente=99);repeat=entry('2',fecha='2026-10-02T10:00:00')
    result=analyze([one,copy,repeat],START,TODAY,today=TODAY)
    assert len(result['entries'])==2 and len(result['repeated'])==2
    with pytest.raises(ValueError):analyze([one,replace(copy,texto='Contenido diferente')],START,TODAY,today=TODAY)
    local=replace(one,entry_id='')
    assert len(analyze([local,replace(local,fila_fuente=2)],START,TODAY,today=TODAY)['entries'])==2
    with pytest.raises(ValueError):analyze([one,replace(repeat,ingreso_id='300')],START,TODAY,today=TODAY)


def test_possible_repetition_is_found_past_five_lexical_neighbors():
    from hashlib import sha512
    from difflib import SequenceMatcher

    base='observacion seguimiento resolucion firmada del ingreso con antecedente judicial estable '
    similar_a,similar_b=base+'a',base+'z'
    texts=[similar_a,similar_b]
    texts.extend(base+letter+' '+sha512(letter.encode()).hexdigest()*12 for letter in 'mnopq')
    texts.sort()
    assert abs(texts.index(similar_a)-texts.index(similar_b))==6
    assert SequenceMatcher(None,similar_a,similar_b,autojunk=False).ratio()>=.92
    entries=[entry(str(i),texto=text) for i,text in enumerate(texts)]

    result=analyze(entries,START,TODAY,today=TODAY,coverage='COMPLETA')
    by_text={observation.texto:observation for observation in entries}
    pair={by_text[similar_a].identity(),by_text[similar_b].identity()}
    assert pair <= result['possible_repeated']
    assert len(result['entries'])==len(entries)  # una sugerencia no fusiona entradas
    assert result['duplicate_scan_complete'] is True
    assert result['coverage']=='COMPLETA'


def test_possible_repetition_below_threshold_is_not_marked():
    from difflib import SequenceMatcher

    left='a'*45+'x'*45
    right='b'*45+'y'*45
    assert SequenceMatcher(None,left,right,autojunk=False).ratio()<.92
    first=entry('1',texto=left);second=entry('2',texto=right)
    result=analyze([first,second],START,TODAY,today=TODAY)
    assert result['possible_repeated']==set()
    assert len(result['entries'])==2


def test_zero_one_and_high_pair_volume_are_reported_truthfully(monkeypatch):
    import nurus.personal.bitacoras as bitacoras

    empty=analyze([],START,TODAY,today=TODAY)
    single=analyze([entry(texto='Una observación única de longitud suficiente.')],START,TODAY,today=TODAY)
    assert empty['duplicate_scan_complete'] and empty['duplicate_pair_checks']==0
    assert single['duplicate_scan_complete'] and single['duplicate_pair_checks']==0

    monkeypatch.setattr(bitacoras,'MAX_DUPLICATE_PAIR_CHECKS',4)
    many=[entry(str(i),texto=f'observacion larga para comparar ingreso numero {i:03d} con texto estable')
          for i in range(4)]
    result=analyze(many,START,TODAY,today=TODAY,coverage='COMPLETA')
    assert result['duplicate_pair_checks']==4
    assert result['duplicate_scan_complete'] is False
    assert result['coverage']=='PARCIAL'
    assert result['duplicate_scan_reason']
    assert len(result['entries'])==len(many)


def test_audit_keeps_failed_queries_full_text_and_date_incidence(tmp_path):
    known=entry(texto='=Texto externo\n'+('detalle completo '*150))
    invalid=entry('2',fecha='fecha ilegible')
    queries=[dict(tribunal_codigo='777',causa_id='100',ingreso_id='200',rit='X-1-2026',
        entries=[known,invalid],coverage='COMPLETA'),
        dict(tribunal_codigo='888',causa_id='300',ingreso_id='400',entries=[],coverage='FALLIDA',error='Sesión vencida')]
    path=export_audit(queries,tmp_path/'auditoria.xlsx',START,TODAY,today=TODAY)
    book=load_workbook(path)
    assert book['Resumen'].max_row==3
    assert book['Resumen']['E2'].value=='PARCIAL'
    assert book['Resumen']['E3'].value=='FALLIDA'
    assert book['Bitácoras']['L2'].value==known.texto
    assert book['Bitácoras']['L2'].data_type=='s'
    assert book['Bitácoras']['I2'].value==1
    assert book['Bitácoras']['A2'].border.left.style=='thin'
    assert book['Incidencias'].max_row==4
    book.close()


def test_long_audit_text_is_split_and_recoverable_from_utf8_sidecar(tmp_path):
    text=('Resolución firmada y gestión verificable con caracteres multibyte 😀 á. '*700)
    known=entry(texto=text)
    tie=entry('2',texto='Otra observación del centro '*1100)
    query=dict(tribunal_codigo='777',causa_id='100',ingreso_id='200',rit='X-1-2026',
               entries=[known,tie],coverage='COMPLETA')
    target=tmp_path/'larga.xlsx'
    path=export_audit([query],target,START,TODAY,today=TODAY)
    sidecar=target.with_suffix('.bitacora-textos.json')
    payload=json.loads(sidecar.read_text(encoding='utf-8'))
    assert payload['workbook']=='larga.xlsx'
    items={item['id']:item for item in payload['items']}
    assert any(item['text']==text and item['sha256']==sha256(text.encode('utf-8')).hexdigest()
               for item in items.values())
    book=load_workbook(path)
    try:
        history=book['Bitácoras'];parts=book['Textos largos']
        assert history['L2'].value.startswith('Texto extenso en hoja')
        assert history['S2'].value==sidecar.name
        assert parts.max_row>2
        for row in parts.iter_rows(min_row=2,values_only=True):
            assert len(row[7])<=32767
            assert sum(2 if ord(char)>0xFFFF else 1 for char in row[7])<=32767
        for item in items.values():
            matching=[row for row in parts.iter_rows(min_row=2,values_only=True) if row[9]==item['id']]
            assert len(matching)==item['parts']
            assert ''.join(row[7] for row in matching)==item['text']
            assert sha256(''.join(row[7] for row in matching).encode('utf-8')).hexdigest()==item['sha256']
        assert book['Resumen']['P2'].value==sidecar.name
    finally:book.close()


def test_short_audit_does_not_create_an_unneeded_sidecar(tmp_path):
    target=tmp_path/'corta.xlsx'
    export_audit([dict(tribunal_codigo='777',causa_id='100',ingreso_id='200',entries=[entry()],coverage='COMPLETA')],
                 target,START,TODAY,today=TODAY)
    assert target.is_file()
    assert not target.with_suffix('.bitacora-textos.json').exists()
