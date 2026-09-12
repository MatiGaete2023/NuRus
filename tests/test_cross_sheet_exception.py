from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook

from nurus.rus import Mode
from nurus.services.workflow import WorkController
from nurus.storage.database import Database


def _write_compliance(path, *, extra_sheets=()):
    book = Workbook()
    sheet = book.active
    sheet.title = "Cumplimiento"
    sheet.append([
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "RUT",
        "DIAS DE CUMPLIMIENTO", "DIAS PARA EGRESAR",
    ])
    sheet.append(["PRM Centro", "Juzgado de Laja", "Ana", "C-1", "1-9", 100, 100])
    for name in extra_sheets:
        cross = book.create_sheet(name)
        cross.append([
            "NOMBRE CENTRO", "TRIBUNAL", "NOMBRE MENOR", "RIT", "RUT MENOR",
            "FECHA VENCIMIENTO",
        ])
        cross.append(["PRM Centro", "Juzgado de Laja", "Ana", "C-1", "1-9", "20/09/2026"])
    book.save(path)


def test_missing_cross_sheet_requires_documented_exception_and_freezes_it(tmp_path):
    path = tmp_path / "cumplimiento.xlsx"
    _write_compliance(path)
    db = Database(tmp_path / "nurus.sqlite3")
    controller = WorkController(db)
    batch = controller.analyze(path, Mode.CUMPLIMIENTO, as_of=date(2026, 9, 8))

    assert any(item.startswith("CROSS_SHEET_NOT_SELECTED") for item in batch.warnings)
    with pytest.raises(ValueError, match="excepción documentada"):
        controller.approve_batch(batch.batch_id)
    with pytest.raises(ValueError, match="responsable y motivo"):
        controller.document_missing_cross_sheet_exception(
            batch.batch_id, responsible="", reason=""
        )

    controller.document_missing_cross_sheet_exception(
        batch.batch_id,
        responsible="Matías Gaete",
        reason="La planilla recibida no contiene hoja de cruce; revisión manual autorizada.",
    )
    exception = db.get_batch_exception(batch.batch_id)
    assert exception["code"] == "CROSS_SHEET_NOT_SELECTED"
    assert exception["responsible"] == "Matías Gaete"

    first = controller.approve_batch(batch.batch_id)
    snapshot = db.get_snapshot(batch.batch_id, first)
    assert snapshot["schema_version"] == 2
    assert snapshot["exceptions"] == [{
        "code": "CROSS_SHEET_NOT_SELECTED",
        "responsible": "Matías Gaete",
        "reason": "La planilla recibida no contiene hoja de cruce; revisión manual autorizada.",
        "created_at": exception["created_at"],
    }]

    controller.document_missing_cross_sheet_exception(
        batch.batch_id, responsible="Otra revisión", reason="Motivo actualizado"
    )
    assert db.get_batch(batch.batch_id)["status"] == "review"
    second = controller.approve_batch(batch.batch_id)
    assert second != first
    assert db.get_snapshot(batch.batch_id, first) == snapshot


def test_ambiguous_cross_sheet_cannot_be_approved_by_exception(tmp_path):
    path = tmp_path / "cumplimiento-ambiguo.xlsx"
    _write_compliance(path, extra_sheets=("Cruce A", "Cruce B"))
    db = Database(tmp_path / "nurus.sqlite3")
    batch = WorkController(db).analyze(path, Mode.CUMPLIMIENTO, as_of=date(2026, 9, 8))

    with pytest.raises(ValueError, match="ambigua"):
        db.document_missing_cross_sheet_exception(
            batch.batch_id, responsible="Matías", reason="No corresponde"
        )


def test_hoja2_name_does_not_override_two_valid_cross_sheets(tmp_path):
    from nurus.rus.reader import read_workbook
    path = tmp_path / "ambiguous.xlsx"
    _write_compliance(path, extra_sheets=("Hoja2", "Otro cruce"))
    batch = read_workbook(path, Mode.CUMPLIMIENTO)
    assert not batch.cross_sheet
    assert any(item.startswith("CROSS_SHEET_AMBIGUOUS") for item in batch.warnings)
    explicit = read_workbook(path, Mode.CUMPLIMIENTO, cross_sheet_name="Otro cruce")
    assert explicit.cross_sheet == "Otro cruce"


def test_cross_and_primary_headers_can_be_on_different_rows(tmp_path):
    from openpyxl import load_workbook
    from nurus.rus.reader import read_workbook
    path = tmp_path / "displaced.xlsx"
    _write_compliance(path, extra_sheets=("Hoja2",))
    book = load_workbook(path)
    book["Cumplimiento"].insert_rows(1, 4)
    book["Hoja2"].insert_rows(1, 2)
    book.save(path)
    result = read_workbook(path, Mode.CUMPLIMIENTO, header_row=5)
    assert result.records[0].source.row_number == 6
    assert result.cross_records[0].source.row_number == 4


def test_invalid_hoja2_does_not_hide_a_valid_cross(tmp_path):
    from openpyxl import load_workbook
    from nurus.rus.reader import read_workbook
    path = tmp_path / "cross.xlsx"
    _write_compliance(path, extra_sheets=("Cruce",))
    book = load_workbook(path)
    book.create_sheet("Hoja2").append(["Solo notas"])
    book.save(path)
    assert read_workbook(path, Mode.CUMPLIMIENTO).cross_sheet == "Cruce"


@pytest.mark.parametrize("dates", [("20/09/2026", "30/09/2026"), ("30/09/2026", "20/09/2026")])
def test_conflicting_cross_dates_never_depend_on_row_order(tmp_path, dates):
    from openpyxl import load_workbook
    from nurus.rus import read_workbook, evaluate_batch
    path = tmp_path / "conflicting.xlsx"
    _write_compliance(path, extra_sheets=("Hoja2",))
    book = load_workbook(path)
    sheet = book["Hoja2"]
    sheet.cell(2, 6, dates[0])
    sheet.append(["PRM Centro", "Juzgado de Laja", "Ana", "C-1", "1-9", dates[1]])
    book.save(path)
    result = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8))
    row = result.evaluations[0]
    assert row.status.value == "blocked"
    assert "C-10" not in row.rule_ids
    assert any(item.code == "CROSS_RECORD_CONFLICT" for item in row.issues)
    assert [item.row_number for item in row.related_sources] == [2, 3]


def test_identical_cross_dates_do_not_create_false_conflict(tmp_path):
    from openpyxl import load_workbook
    from nurus.rus import read_workbook, evaluate_batch
    path = tmp_path / "duplicate.xlsx"
    _write_compliance(path, extra_sheets=("Hoja2",))
    book = load_workbook(path)
    sheet = book["Hoja2"]
    sheet.append([cell.value for cell in sheet[2]])
    book.save(path)
    row = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8)).evaluations[0]
    assert "C-10" in row.rule_ids
    assert not any(item.code == "CROSS_RECORD_CONFLICT" for item in row.issues)
