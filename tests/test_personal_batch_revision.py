from pathlib import Path
from copy import deepcopy
from types import SimpleNamespace
from zipfile import ZipFile
import json
import pytest
from openpyxl import Workbook,load_workbook
from docx import Document
from nurus.personal.config import defaults,Configuration,BASE,CC,LEGACY_CORREO_BODIES,CORREO_REVISION
from nurus.personal.work import Work,Row
from nurus.personal.importing import SheetChoice
from nurus.personal.outputs import prepare_drafts,create_drafts,Draft
from nurus.personal.resolutions import prepare_projects,generate_projects,replace_paragraph,automatic_project_selections,reviewed_resolution_ids
from nurus.personal.template_package import import_templates,install_bundled_templates

def external(tmp_path,rows=None):
    book=Workbook();book.active.title='Portada';book.active['A1']='REPORTE'
    sheet=book.create_sheet('Registros');sheet.append(['Reporte mensual']);sheet.append([])
    sheet.append(['RIT','TRIBUNAL','NOMBRE MENOR','RUT MENOR','NOMBRE CENTRO','Observaciones','FECHA_OBS','TT','CC'])
    for row in rows or [
        ['X-1-2026','LAJA','Persona Uno','11111111-1','AFT EJEMPLO','Se remite proyecto.','2026-09-14',1,0],
        ['X-1-2026','LAJA','Persona Dos','22222222-2','AFT EJEMPLO','Texto humano.','2026-09-14',1,1],
        ['X-2-2026','LAJA','Persona Tres','33333333-3','FAE EJEMPLO','Constancia.','2026-09-14',1,0],
    ]:sheet.append(row)
    path=tmp_path/'revisado.xlsx';book.save(path)
    return Work.external(path,defaults())

def test_headers_offset_other_sheet_aliases_and_human_review(tmp_path):
    work=external(tmp_path)
    assert work.sheet=='Registros' and work.header==3 and len(work.rows)==3
    assert work.rows[1].review['OBSERVACION']=='Texto humano.'
    assert work.rows[1].review['CC']==1
    book=load_workbook(work.output);book['Registros']['F5']='Edición posterior';book.save(work.output)
    assert work.refresh()
    assert work.rows[1].review['OBSERVACION']=='Edición posterior'

def test_multiple_tables_request_only_actual_sheet_choice(tmp_path):
    work=external(tmp_path);book=load_workbook(work.output);book.copy_worksheet(book['Registros']).title='Otra';book.save(work.output)
    with pytest.raises(SheetChoice):Work.external(work.output,defaults())
    assert Work.external(work.output,defaults(),sheet='Otra').sheet=='Otra'

def test_modalities_filter_and_program_attachment_names(tmp_path):
    work=external(tmp_path)
    drafts=prepare_drafts(work,'programa_espera',modality_keys=['FAE'])
    assert len(drafts)==1
    assert Path(drafts[0].attachments[0]).name=='FAE EJEMPLO.xlsx'
    sheet=load_workbook(drafts[0].attachments[0]).active
    assert sheet.max_row==2 and sheet['C2'].value=='Persona Tres'
    all_drafts=prepare_drafts(work,'programa_espera')
    assert len(all_drafts)==2


def test_waiting_attachment_contains_only_operational_columns(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION'])
    sheet.append(['X-10-2026','MULCHEN','Persona Espera','11111111-1','AFT EJEMPLO',72,'Texto interno que no debe salir'])
    path=tmp_path/'espera_adj.xlsx';book.save(path)
    work=Work.external(path,defaults(),mode='ESPERA')
    drafts=prepare_drafts(work,'programa_espera',directory=tmp_path/'salida')
    assert len(drafts)==1 and len(drafts[0].attachments)==1
    out=load_workbook(drafts[0].attachments[0]).active
    assert [c.value for c in out[1]]==['RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA']
    assert out.max_column==5
    assert [out.cell(2,i).value for i in range(1,6)]==['X-10-2026','MULCHEN','Persona Espera','AFT EJEMPLO','72']
    assert all(cell.border.left.style=='thin' and cell.border.right.style=='thin' and cell.border.top.style=='thin' and cell.border.bottom.style=='thin' for row in out.iter_rows() for cell in row)



def test_waiting_program_mail_respects_rule_per_row_after_external_reload(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Espera'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','T ESPERA','OBSERVACION','NURUS_REGLAS'])
    sheet.append(['X-1-2026','MULCHEN','Persona Gestion','11111111-1','AFT EJEMPLO',30,
                  'Medida revisada, a la espera de ingreso efectivo. Se remite correo al programa consultando fecha estimada de ingreso.',
                  json.dumps(['ESPERA.E05_PROYECTO_Y_CORREO'])])
    sheet.append(['X-2-2026','MULCHEN','Persona Sin Gestion','22222222-2','AFT EJEMPLO',12,
                  'Medida revisada, a la espera de ingreso efectivo al programa AFT Ejemplo.',
                  json.dumps(['ESPERA.E06_SIN_RESOLUCION'])])
    path=tmp_path/'espera_revisada.xlsx';book.save(path)
    work=Work.external(path,defaults(),mode='ESPERA')
    drafts=prepare_drafts(work,'programa_espera',directory=tmp_path/'salida')
    assert len(drafts)==1
    assert drafts[0].record_ids==[work.rows[0].id]
    attached=load_workbook(drafts[0].attachments[0]).active
    assert attached.max_row==2 and attached['A2'].value=='X-1-2026'


def test_human_observation_can_remove_automatic_program_mail(tmp_path):
    cfg=defaults()
    original='Medida revisada, a la espera de ingreso efectivo. Se remite correo al programa consultando fecha estimada de ingreso.'
    row=Row('r1',2,{'RIT':'X-1-2026','TRIBUNAL':'MULCHEN','NOMBRE':'Persona','DERIVACION':'AFT EJEMPLO','T ESPERA':45},
            original,['ESPERA.E05_SOLO_CORREO'],['programa_espera'],[],False,
            {'OBSERVACION':'Medida revisada, a la espera de ingreso efectivo al programa AFT Ejemplo.'})
    work=SimpleNamespace(config=cfg,rows=[row],mapping={'rit':'RIT','tribunal':'TRIBUNAL','nombre':'NOMBRE','programa':'DERIVACION','espera':'T ESPERA'},
                         output=str(tmp_path/'revisable.xlsx'),external_input=False,refresh=lambda:False)
    assert prepare_drafts(work,'programa_espera',directory=tmp_path/'salida')==[]
    row.review['OBSERVACION']='Medida revisada, a la espera de ingreso efectivo. Se remite correo al programa consultando fecha estimada de ingreso.'
    assert len(prepare_drafts(work,'programa_espera',directory=tmp_path/'salida2'))==1

def test_due_report_attachment_contains_only_operational_columns(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Informes'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','FECHA VENCIMIENTO','OBSERVACION'])
    sheet.append(['X-20-2026','LAJA','Persona Informe','22222222-2','PRM EJEMPLO','30/09/2026','Texto interno que no debe salir'])
    path=tmp_path/'informes_adj.xlsx';book.save(path)
    work=Work.external(path,defaults(),mode='INFORMES')
    drafts=prepare_drafts(work,'programa_por_vencer',directory=tmp_path/'salida')
    assert len(drafts)==1 and len(drafts[0].attachments)==1
    out=load_workbook(drafts[0].attachments[0]).active
    assert [c.value for c in out[1]]==['RIT','TRIBUNAL','NOMBRE','DERIVACION','F. VENCIMIENTO']
    assert out.max_column==5
    assert [out.cell(2,i).value for i in range(1,6)]==['X-20-2026','LAJA','Persona Informe','PRM EJEMPLO','30/09/2026']


def test_manual_email_bodies_match_csmp_2025_contract():
    templates=defaults()['correos']['plantillas']
    assert templates['espera']['cuerpo']==('Buen día:\n\nJunto con saludar, se informa que pestaña espera del módulo RUS en SITFA, '
                                             'se revisaron todos los registros pertenecientes a {ALCANCE_MODALIDADES} y se registraron observaciones en bitácora.\n\nAtte. a Ud.,')
    assert templates['cumplimiento']['cuerpo']==('Buen día:\n\nJunto con saludar, se informa que pestaña cumplimiento del modulo RUS en SITFA, '
                                                   'se revisaron todos los registros pertenecientes a {ALCANCE_MODALIDADES} y se registraron observaciones en bitácora.\n\nAtte. a Ud.,')
    assert templates['informes']['cuerpo']==('Buen día:\n\nJunto con saludar, se informa que pestaña informes del módulo RUS en SITFA, '
                                               'se revisaron todos los registros pertenecientes a {ALCANCE_MODALIDADES} y se registraron observaciones en bitácora.\n\nAtte. a Ud.,')
    assert templates['medidas']['cuerpo']==('Buen día:\n\nJunto con saludar, se remite archivo Excel adjunto con nomina de medidas sin vigencia y por vencer '
                                              'al día {FECHA} en {ALCANCE_MODALIDADES}.\n\nAtte. a Ud.,')
    assert templates['programa_por_vencer']['cuerpo']==('Buen día:\n\nJunto con saludar, se envía planilla con RIT de causas en las cuales, informes de avances '
                                                           'se encuentran pronto a vencer. Se ruega acusar recibo de la información.\n\nAtte. a Ud.,')
    assert templates['programa_por_vencer']['adjunto']=='obligatorio'


def test_existing_default_mail_bodies_migrate_but_custom_body_is_preserved(tmp_path):
    conf=Configuration(tmp_path/'config');cfg=deepcopy(conf.data);cfg.pop('revision_correos',None)
    cfg['correos']['plantillas']['espera']['cuerpo']=LEGACY_CORREO_BODIES['espera']
    cfg['correos']['plantillas']['cumplimiento']['cuerpo']='Mi formato personalizado.'
    cfg['correos']['plantillas']['programa_por_vencer']['cuerpo']=LEGACY_CORREO_BODIES['programa_por_vencer']
    conf.save(cfg);updated=Configuration(conf.directory)
    canonical=defaults()['correos']['plantillas']
    assert updated.data['revision_correos']==CORREO_REVISION
    assert updated.data['correos']['plantillas']['espera']['cuerpo']==canonical['espera']['cuerpo']
    assert updated.data['correos']['plantillas']['programa_por_vencer']['cuerpo']==canonical['programa_por_vencer']['cuerpo']
    assert updated.data['correos']['plantillas']['cumplimiento']['cuerpo']=='Mi formato personalizado.'
    assert updated.path.with_suffix('.bak').exists()


def test_batch_save_empty_to_and_partial_failure_do_not_duplicate(tmp_path,monkeypatch):
    work=external(tmp_path);seen=[]
    def save(product,**kwargs):
        seen.append(product)
        if product.rendered_subject=='Falla':raise ValueError('Error de prueba')
        return SimpleNamespace(entry_id='e',store_id='s')
    monkeypatch.setattr('nurus.personal.outputs.save_draft',save)
    drafts=[Draft('Uno','Texto',key='a'),Draft('Falla','Texto',key='b'),Draft('Tres','Texto',key='c')]
    result=create_drafts(work,drafts)
    assert result['created']==2 and len(result['errors'])==1 and all(CC in p.cc for p in seen)
    assert seen[0].recipient==''
    again=create_drafts(work,drafts)
    assert again['created']==0 and again['skipped']==2 and len(seen)==4

def test_migration_restores_old_defaults_and_preserves_customizations(tmp_path):
    conf=Configuration(tmp_path/'config');cfg=deepcopy(conf.data);cfg.pop('revision_textos',None)
    original=cfg['textos']['ESPERA']['E05_PROYECTO_Y_CORREO']['texto']
    cfg['textos']['ESPERA']['E05_PROYECTO_Y_CORREO']['texto']=original.replace('Se remite proyecto de resolución','Se sugiere preparar proyecto de resolución').replace('se remite correo electrónico','se sugiere remitir correo electrónico').replace('Se remite correo electrónico','Se sugiere remitir correo electrónico')
    cfg['textos']['ESPERA']['E05_SOLO_CORREO']['texto']='Mi texto personalizado.'
    conf.save(cfg);updated=Configuration(conf.directory)
    assert updated.data['textos']['ESPERA']['E05_PROYECTO_Y_CORREO']['texto']==original
    assert updated.data['textos']['ESPERA']['E05_SOLO_CORREO']['texto']=='Mi texto personalizado.'
    assert updated.path.with_suffix('.bak').exists()

def test_bundled_template_revision_replaces_once_with_backup_then_preserves_user_edit(tmp_path):
    source=tmp_path/'source';target=tmp_path/'target';(source/'LAJA').mkdir(parents=True);(target/'LAJA').mkdir(parents=True)
    bundled=source/'LAJA/PC_IE.docx';old=target/'LAJA/PC_IE.docx'
    Document().save(bundled);old.write_bytes(b'previous-template')
    result=install_bundled_templates(source,target,revision='r1')
    assert old.read_bytes()==bundled.read_bytes() and result['backups']
    backup=Path(result['backups'][0]);assert backup.read_bytes()==b'previous-template'
    old.write_bytes(b'user-edited-after-migration')
    result=install_bundled_templates(source,target,revision='r1')
    assert old.read_bytes()==b'user-edited-after-migration' and not result['installed']

def test_single_word_groups_people_and_starts_next_project_on_new_page(tmp_path):
    work=external(tmp_path)
    projects,errors=prepare_projects(work,[(r.id,'PC_IE') for r in work.rows],BASE/'plantillas_word')
    assert not errors and len(projects)==2
    assert len(projects[0].record_ids)==2
    assert 'Persona Uno' in projects[0].text and 'Persona Dos' in projects[0].text
    assert '11111111-1' in projects[0].text and '22222222-2' in projects[0].text
    projects[0].text+='\nPárrafo incorporado por el usuario.'
    path=tmp_path/'Resoluciones.docx';generate_projects(work,projects,path)
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
    assert 'Párrafo incorporado por el usuario.' in text
    assert text.count('RIT: X-1-2026')==1
    assert text.count('RIT: X-2-2026')==1
    assert 'w:type="page"' in doc.element.xml
    assert len(list(tmp_path.glob('*.docx')))==1

def test_manual_type_and_missing_data_produce_editable_project(tmp_path):
    work=external(tmp_path)
    projects,errors=prepare_projects(work,[(work.rows[0].id,'NOMENCL')],BASE/'plantillas_word')
    assert not errors and projects[0].kind=='NOMENCL'
    assert '[COMPLETAR NOMENCLATURA]' in projects[0].text

def test_same_rit_different_court_does_not_merge(tmp_path):
    work=external(tmp_path);work.rows[1].values[work.mapping['tribunal']]='MULCHEN'
    projects,errors=prepare_projects(work,[(r.id,'PC_IE') for r in work.rows[:2]],BASE/'plantillas_word')
    assert not errors and len(projects)==2
    by_court={p.court:p for p in projects}
    assert set(by_court)=={'LAJA','MULCHEN'}
    assert Path(by_court['LAJA'].template).parent.name=='LAJA'
    assert Path(by_court['MULCHEN'].template).parent.name=='MULCHEN'
    assert by_court['LAJA'].text.startswith('Laja,')
    assert by_court['MULCHEN'].text.startswith('Mulchén,')
    out=tmp_path/'mixto.docx';generate_projects(work,projects,out)
    combined='\n'.join(p.text for p in Document(out).paragraphs)
    assert 'Laja,' in combined and 'Mulchén,' in combined


def test_combined_word_preserves_each_matrix_style(tmp_path):
    rows=[
        ['X-10-2026','LAJA','Persona Laja','11111111-1','AFT EJEMPLO','Texto','2026-09-14',1,0],
        ['X-20-2026','MULCHEN','Persona Mulchen','22222222-2','AFT EJEMPLO','Texto','2026-09-14',1,0],
    ]
    work=external(tmp_path,rows)
    templates=tmp_path/'matrices'
    for court,font in [('LAJA','Arial'),('MULCHEN','Times New Roman')]:
        folder=templates/court;folder.mkdir(parents=True)
        doc=Document();doc.styles['Normal'].font.name=font
        doc.add_paragraph(court.title()+' {{RIT}}')
        doc.save(folder/'PC_IE.docx')
    projects,errors=prepare_projects(work,[(r.id,'PC_IE') for r in work.rows],templates)
    assert not errors and {p.court for p in projects}=={'LAJA','MULCHEN'}
    out=tmp_path/'estilos.docx';generate_projects(work,projects,out)
    doc=Document(out)
    found={p.text.split()[0].upper():p.style.font.name for p in doc.paragraphs if p.text.startswith(('Laja','Mulchen'))}
    assert found['LAJA']=='Arial'
    assert found['MULCHEN']=='Times New Roman'

def test_edit_preserves_unchanged_bold_runs():
    doc=Document();p=doc.add_paragraph();p.add_run('RIT: ').bold=True;p.add_run('X-1')
    replace_paragraph(p,'RIT: X-2')
    assert p.text=='RIT: X-2' and p.runs[0].bold

def test_zip_import_known_names_preserves_backup_and_rejects_ambiguous(tmp_path):
    package=tmp_path/'matrices.zip';source=BASE/'plantillas_word/LAJA/PC_IE.docx'
    with ZipFile(package,'w') as z:
        z.write(source,'plantillas_word/LAJA/PC_IE.docx')
        z.write(source,'sin_identificar.docx')
    directory=tmp_path/'matrices';result=import_templates(package,directory)
    assert len(result['imported'])==1 and result['unmatched']==['sin_identificar.docx']
    import_templates(package,directory)
    assert (directory/'LAJA/PC_IE.bak.docx').exists()

def test_summary_counts_one_document_and_two_projects(tmp_path):
    from nurus.personal.statistics import summarize
    work=external(tmp_path);projects,_=prepare_projects(work,[(r.id,'PC_IE') for r in work.rows],BASE/'plantillas_word')
    generate_projects(work,projects,tmp_path/'lote.docx')
    counts=summarize(work)
    assert counts['word_generados']==1 and counts['proyectos_generados']==2


def test_reviewed_res_column_is_authoritative_for_automatic_projects(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Registros'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','OBSERVACION','RES'])
    sheet.append(['X-1-2026','LAJA','Persona Uno','11111111-1','AFT EJEMPLO','Texto','1.0'])
    sheet.append(['X-2-2026','LAJA','Persona Dos','22222222-2','AFT EJEMPLO','Texto','No'])
    sheet.append(['X-3-2026','LAJA','Persona Tres','33333333-3','AFT EJEMPLO','Texto',''])
    path=tmp_path/'seleccion_res.xlsx';book.save(path)
    work=Work.external(path,defaults())
    assert reviewed_resolution_ids(work)=={work.rows[0].id}
    assert automatic_project_selections(work,'PC_INFO')==[(work.rows[0].id,'PC_INFO')]


def test_resolution_mark_on_one_row_recognizes_complete_rit(tmp_path):
    book=Workbook();sheet=book.active;sheet.title='Registros'
    sheet.append(['RIT','TRIBUNAL','NOMBRE','RUT','DERIVACION','OBSERVACION','GENERAR RESOLUCIÓN'])
    sheet.append(['X-149-2025','MULCHEN','SOFÍA IGNACIA DAROCH VERDUGO','24010461-K','AFT MULCHEN','Texto','1'])
    sheet.append(['X-149-2025','MULCHEN','MIA VALENTINA DAROCH VERDUGO','24624826-5','AFT MULCHEN','Texto',''])
    sheet.append(['X-150-2025','MULCHEN','Otra Persona','11111111-1','AFT MULCHEN','Texto',''])
    path=tmp_path/'res_por_causa.xlsx';book.save(path)
    work=Work.external(path,defaults())
    selected=automatic_project_selections(work,'PC_IE')
    assert {rid for rid,kind in selected}=={work.rows[0].id,work.rows[1].id}
    assert {kind for rid,kind in selected}=={'PC_IE'}
    projects,errors=prepare_projects(work,selected,BASE/'plantillas_word')
    assert not errors and len(projects)==1
    assert set(projects[0].record_ids)=={work.rows[0].id,work.rows[1].id}
    assert 'SOFÍA IGNACIA DAROCH VERDUGO' in projects[0].text
    assert 'MIA VALENTINA DAROCH VERDUGO' in projects[0].text


def test_grouped_people_keep_each_identity_number_next_to_its_name(tmp_path):
    rows=[
        ['X-149-2025','MULCHEN','SOFÍA IGNACIA DAROCH VERDUGO','24010461-K','AFT  MULCHEN','Texto','2026-09-14',1,0],
        ['X-149-2025','MULCHEN','MIA VALENTINA DAROCH VERDUGO','24624826-5','AFT  MULCHEN','Texto','2026-09-14',1,0],
    ]
    work=external(tmp_path,rows)
    projects,errors=prepare_projects(work,[(r.id,'PC_IE') for r in work.rows],BASE/'plantillas_word')
    assert not errors and len(projects)==1
    text=projects[0].text
    expected=('SOFÍA IGNACIA DAROCH VERDUGO, cédula de identidad N° 24010461-K y '
              'MIA VALENTINA DAROCH VERDUGO, cédula de identidad N° 24624826-5')
    assert expected in text
    assert 'SOFÍA IGNACIA DAROCH VERDUGO y MIA VALENTINA DAROCH VERDUGO, cédula' not in text
    assert 'cédulas de identidad' not in text