from __future__ import annotations

import hashlib
from datetime import date

from openpyxl import Workbook

from nurus.rus import Mode, evaluate_batch, read_workbook


def write_sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def test_waiting_reader_preserves_source_and_applies_rule(tmp_path):
    path = tmp_path / "espera.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Espera", ["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "T ESPERA"], [["FAE Norte", "Juzgado de Laja", "Ana Pérez", "C-1", 45]])
    book.save(path)

    batch = read_workbook(path, Mode.ESPERA)
    result = evaluate_batch(batch, as_of=date(2026, 9, 8))

    reference = result.evaluations[0].source
    assert reference.workbook_name == "espera.xlsx"
    assert reference.workbook_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert (reference.sheet_name, reference.row_number) == ("Espera", 2)
    assert "E-05" in result.evaluations[0].rule_ids
    assert "proyecto de resolución" in result.evaluations[0].observation


def test_compliance_cross_keeps_secondary_row_provenance(tmp_path):
    path = tmp_path / "cumplimiento.xlsx"
    book = Workbook(); book.remove(book.active)
    headers = ["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "RUT", "DIAS DE CUMPLIMIENTO", "DIAS PARA EGRESAR"]
    write_sheet(book, "Cumplimiento", headers, [["PRM Centro", "Juzgado de Tomé", "Juan Pérez", "C-2", "12.345.678-9", 100, 100]])
    write_sheet(book, "Hoja2", ["NOMBRE CENTRO", "TRIBUNAL", "NOMBRE MENOR", "RIT", "RUT MENOR", "FECHA VENCIMIENTO"], [["PRM Centro", "Juzgado de Tomé", "Juan Pérez", "C-2", "12.345.678-9", "20/09/2026"]])
    book.save(path)

    result = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8))

    evaluation = result.evaluations[0]
    assert "C-10" in evaluation.rule_ids
    assert "20 de septiembre de 2026" in evaluation.observation
    assert len(evaluation.related_sources) == 1
    assert (evaluation.related_sources[0].sheet_name, evaluation.related_sources[0].row_number) == ("Hoja2", 2)
    assert evaluation.related_sources[0].workbook_sha256 == evaluation.source.workbook_sha256


def test_compliance_does_not_use_an_unrelated_second_sheet(tmp_path):
    path = tmp_path / "sin_cruce.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Cumplimiento", ["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "RUT"], [["PRM Centro", "Juzgado de Tomé", "Juan", "C-3", "1-9"]])
    write_sheet(book, "VALIDACION", ["FILA_EXCEL", "MOTIVO"], [[2, "dato"]])
    book.save(path)

    result = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8))

    assert not result.evaluations[0].related_sources
    assert any(warning.startswith("CROSS_SHEET_NOT_SELECTED") for warning in result.warnings)


def test_named_cross_sheet_with_incomplete_mapping_requires_missing_cross_exception(tmp_path):
    path = tmp_path / "cruce_incompleto.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Cumplimiento", [
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "RUT",
        "DIAS DE CUMPLIMIENTO", "DIAS PARA EGRESAR",
    ], [["PRM Centro", "Juzgado de Tomé", "Juan Pérez", "C-2", "12.345.678-9", 100, 100]])
    write_sheet(book, "Hoja2", ["NOTA"], [["sin columnas de cruce"]])
    book.save(path)

    result = read_workbook(path, Mode.CUMPLIMIENTO)

    assert result.cross_sheet == ""
    assert not result.cross_records
    assert any(warning.startswith("CROSS_MAPPING_MISSING") for warning in result.warnings)
    assert any(warning.startswith("CROSS_SHEET_NOT_SELECTED") for warning in result.warnings)


def test_reports_missing_due_date_is_an_incident_not_an_exception(tmp_path):
    path = tmp_path / "informes.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Informes", ["DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "FECHA VENCIMIENTO"], [["DCE Centro", "Juzgado de Laja", "Luisa", "C-4", ""]])
    book.save(path)

    result = evaluate_batch(read_workbook(path, Mode.INFORMES), as_of=date(2026, 9, 8))

    assert result.evaluations[0].observation == ""
    assert result.evaluations[0].issues[0].code == "I-01/I-02"
    assert result.evaluations[0].issues[0].source.row_number == 2


def test_compliance_missing_projected_exit_does_not_report_no_observations(tmp_path):
    path = tmp_path / "sin_egreso.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Cumplimiento", [
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "DIAS DE CUMPLIMIENTO",
        "DIAS PARA EGRESAR", "FEC.INGRESO EFECTIVO", "FEC.EGRESO PROYECTADO",
    ], [["PRM Centro", "Laja", "Juan", "C-1", 100, 10, "01/01/2026", ""]])
    book.save(path)

    result = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8))
    evaluation = result.evaluations[0]

    assert evaluation.status.value == "blocked"
    assert any(issue.code == "C-05" for issue in evaluation.issues)
    assert "sin observaciones" not in evaluation.observation


def test_compliance_negative_days_with_future_exit_is_a_conflict(tmp_path):
    path = tmp_path / "egreso_incoherente.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Cumplimiento", [
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "DIAS DE CUMPLIMIENTO",
        "DIAS PARA EGRESAR", "FEC.EGRESO PROYECTADO",
    ], [["PRM Centro", "Laja", "Juan", "C-1", -1, 100, "18/12/2026"]])
    book.save(path)

    result = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8))
    evaluation = result.evaluations[0]

    assert evaluation.status.value == "blocked"
    assert any("contradictorios" in issue.message for issue in evaluation.issues)
    assert "vencida" not in evaluation.observation


def test_fae_rule_requires_the_fae_source_column(tmp_path):
    path = tmp_path / "fae_sin_columna.xlsx"
    book = Workbook(); book.remove(book.active)
    write_sheet(book, "Cumplimiento", [
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "DIAS DE CUMPLIMIENTO",
        "DIAS PARA EGRESAR", "FEC.INGRESO EFECTIVO", "FEC.EGRESO PROYECTADO",
    ], [["FAE Centro", "Laja", "Juan", "C-1", 200, 100, "01/01/2026", "18/12/2026"]])
    book.save(path)

    result = evaluate_batch(read_workbook(path, Mode.CUMPLIMIENTO), as_of=date(2026, 9, 8))
    evaluation = result.evaluations[0]

    assert evaluation.status.value == "blocked"
    assert any(issue.code == "C-08" for issue in evaluation.issues)
    assert "no tiene ficha FAE" not in evaluation.observation
