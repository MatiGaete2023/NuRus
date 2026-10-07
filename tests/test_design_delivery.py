from copy import deepcopy
from pathlib import Path
from hashlib import sha256
from threading import Event
from queue import Queue
import json
from zipfile import ZipFile
import pytest
from docx import Document
from openpyxl import load_workbook
from nurus.personal.config import defaults,Configuration
from nurus.personal.work import Work
from nurus.personal.outputs import Draft
from nurus.personal.product_checks import check_products
from nurus.personal.product_state import fingerprint,dependency_scope,stamp,stale
from nurus.personal.delivery import selected_zip,validate_product
from nurus.personal.grid_edit import paste_plan,apply_plan
from nurus.personal.mail_presentation import body_html,insert_before_signature
from nurus.personal.operations import Operation,operation_scope,checkpoint,progress,OperationCancelled
from nurus.personal.reports import _book
from nurus.personal.word_quality import require_complete,pending_fields
from current_fixture import source


def test_semantic_excel_layout_preserves_literal_text_and_prints_headers(tmp_path):
    path=tmp_path/'nomina.xlsx'
    _book([('Nómina',['RIT','NOMBRE','OBSERVACION'],[['001','Persona muy larga','=Texto literal']])],path)
    book=load_workbook(path);s=book.active
    assert s['A2'].value=='001' and s['C2'].value=='=Texto literal' and s['C2'].data_type=='s'
    assert s.column_dimensions['C'].width>s.column_dimensions['A'].width
    assert s.freeze_panes=='A2' and s.print_title_rows=='$1:$1'
    assert s.page_setup.fitToWidth==1 and s.page_setup.fitToHeight==0
    assert s.row_dimensions[2].height>=22
    book.close()


def test_delivery_index_manifest_and_collision_safety(tmp_path):
    a=tmp_path/'a';b=tmp_path/'b';a.mkdir();b.mkdir()
    for folder,text in [(a,'one'),(b,'two')]: (folder/'same.txt').write_text(text)
    path=tmp_path/'delivery.zip';selected_zip([a/'same.txt',b/'same.txt'],path)
    with ZipFile(path) as archive:
        data=json.loads(archive.read('MANIFIESTO.json'))
        assert len(data['archivos'])==2 and len({p['archivo'] for p in data['archivos']})==2
        for item in data['archivos']:assert sha256(archive.read(item['archivo'])).hexdigest()==item['sha256']
        assert 'Outlook' in archive.read('INDICE.txt').decode()
    before=path.read_bytes()
    with pytest.raises(ValueError,match='ya existe'):selected_zip([a/'same.txt'],path)
    assert path.read_bytes()==before


def test_delivery_rejects_review_markers_and_invalid_documents(tmp_path):
    path=tmp_path/'pending.docx';doc=Document();doc.add_paragraph('Persona [COMPLETAR RUT]');doc.save(path)
    require_complete(path,allow_review_markers=True)
    with pytest.raises(ValueError,match='pendientes'):selected_zip([path],tmp_path/'bad.zip')
    assert not (tmp_path/'bad.zip').exists()
    doc.paragraphs[0].text='Persona revisada';doc.save(path);validate_product(path)
    broken=tmp_path/'invalid.xlsx';broken.write_text('Not a workbook')
    with pytest.raises(Exception):validate_product(broken)


def test_header_fields_are_checked_across_runs(tmp_path):
    doc=Document();p=doc.sections[0].header.paragraphs[0];p.add_run('{{NO');p.add_run('MBRE}}');path=tmp_path/'header.docx';doc.save(path)
    assert '{{NOMBRE}}' in pending_fields(path)
    with pytest.raises(ValueError):require_complete(path,allow_review_markers=True)


def test_mail_html_escapes_literal_content_and_preserves_signature_once():
    text='Buen día <persona>\n\nTexto & más\nSegunda línea'
    signature='<html><body><p>Firma institucional</p><img src="cid:signature"></body></html>'
    result=insert_before_signature(text,signature)
    assert '&lt;persona&gt;' in result and '&amp;' in result
    assert result.count('Firma institucional')==1 and 'cid:signature' in result
    assert result.index('Buen día')<result.index('Firma institucional')
    assert 'script' not in body_html('<script>') or '&lt;script&gt;' in body_html('<script>')


def test_clipboard_plan_uses_visible_ids_and_applies_atomically():
    group={'headers':['Nombre','Observación'],'rows':[{'id':'hidden','cells':['H','Old']},{'id':'b','cells':['B','Old']},{'id':'a','cells':['A','Old']}]}
    before=deepcopy(group);plan=paste_plan(group,['a','b'],'a',0,'Nombre nuevo\t=literal\nOtro\tTexto')
    assert group==before;apply_plan(group,plan)
    assert group['rows'][0]['cells']==['H','Old']
    assert group['rows'][2]['cells']==['Nombre nuevo','=literal']
    before=deepcopy(group)
    with pytest.raises(ValueError):apply_plan(group,[('a',0,'change'),('missing',0,'bad')])
    assert group==before
    with pytest.raises(ValueError):paste_plan(group,['a'],'a',1,'A\tB')


def test_product_ids_distinguish_equal_subjects_and_persist(tmp_path):
    cfg=Configuration(tmp_path/'config');work=Work(cfg.data).analyze(source(tmp_path),'ESPERA')
    drafts=[Draft('Same','Text',to='a@example.test'),Draft('Same','Text',to='b@example.test')]
    checks=check_products(work,drafts,[])
    assert checks[0].id!=checks[1].id and checks[0].index==0 and checks[1].index==1
    from dataclasses import asdict
    restored=Draft(**asdict(drafts[0]));assert restored.product_id==drafts[0].product_id
    drafts[1].to='';checks=check_products(work,drafts,[])
    assert not checks[0].issues and 'Completar destinatario' in checks[1].issues


def test_dependency_batch_matches_uncached_and_expires_after_edits(tmp_path):
    work=Work(defaults()).analyze(source(tmp_path),'ESPERA')
    ids=[work.rows[0].id];before=fingerprint(work,ids,'espera','tribunales')
    with dependency_scope(work):
        assert fingerprint(work,ids,'espera','tribunales')==before
        assert fingerprint(work,ids,'espera','tribunales')==before
    work.rows[0].observation+=' Changed'
    assert fingerprint(work,ids,'espera','tribunales')!=before


def test_cancel_only_stops_at_checkpoint_and_keeps_confirmed_units():
    events=Queue();op=Operation(events.put);completed=[]
    with operation_scope(op):
        checkpoint();completed.append('confirmed');op.cancelled.set();progress('Current unit finishes',1,2)
        with pytest.raises(OperationCancelled):checkpoint()
    assert completed==['confirmed'];assert events.get()==('progress','Current unit finishes',1,2)
