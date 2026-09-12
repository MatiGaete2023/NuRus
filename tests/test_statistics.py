from datetime import date

from openpyxl import load_workbook

from nurus.services.statistics import ReviewStatistics, export_review_statistics


def test_statistics_export_keeps_raw_administrative_values(tmp_path):
    stats = ReviewStatistics(
        start=date(2026, 9, 1), end=date(2026, 9, 30),
        total_records=8, reviewed_in_period=5, excluded=1,
        outside_period=1, missing_or_invalid_date=1, observations_changed=2,
        tt_values=(("Sí", 3), ("(vacío)", 2)),
        workload_values=(("Con carga", 4), ("Sin carga", 1)),
        resolution_values=(("No", 5),),
    )
    path = export_review_statistics(stats, tmp_path / "estadisticas.xlsx")
    book = load_workbook(path)
    assert dict(book["Resumen"].values)["Revisados en período"] == 5
    assert dict(book["CC"].values)["Con carga"] == 4
    assert dict(book["TT"].values)["Sí"] == 3


def test_statistics_exports_formula_like_codes_as_text(tmp_path):
    stats = ReviewStatistics(
        start=date(2026, 9, 1), end=date(2026, 9, 30), total_records=1,
        reviewed_in_period=1, excluded=0, outside_period=0, missing_or_invalid_date=0,
        observations_changed=0, tt_values=(("=2+2", 1),),
        workload_values=(), resolution_values=(),
    )
    path = export_review_statistics(stats, tmp_path / "statistics.xlsx")
    book = load_workbook(path)
    assert book["TT"]["A2"].data_type == "s"
    assert book["TT"]["A2"].value == "'=2+2"
    book.close()


def test_statistics_separates_unconfirmed_from_bad_dates():
    from types import SimpleNamespace
    from nurus.services.statistics import build_review_statistics
    rows = [
        {"decision": "approved", "rus_recorded": False, "review_date": "2026-09-01"},
        {"decision": "approved", "rus_recorded": True, "review_date": "invalid"},
        {"decision": "excluded"},
    ]
    db = SimpleNamespace(get_snapshot=lambda _: {"batch": {"review_import_hash": "x"}, "records": rows})
    result = build_review_statistics(db, "batch", date(2026, 9, 1), date(2026, 9, 30))
    assert result.unconfirmed == result.missing_or_invalid_date == result.excluded == 1
    assert result.total_records == sum((result.reviewed_in_period, result.outside_period,
                                       result.excluded, result.missing_or_invalid_date, result.unconfirmed))
