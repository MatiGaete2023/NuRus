"""Consulta de Enviados. No modifica elementos ni carpetas de Outlook."""
from dataclasses import dataclass
from datetime import date
import platform

from nurus.adapters.outlook import OutlookUnavailable, _find_account
from nurus.services.file_output import excel_text, write_new_file


@dataclass(frozen=True)
class SentMailReport:
    rows: tuple[dict, ...]
    skipped: int
    errors: int
    truncated: bool
    account: str
    folder: str


def scan_sent_items(items, start: date, end: date, *, limit=50000):
    if end < start or limit < 1:
        raise ValueError("Período o límite inválido.")
    items.Sort("[SentOn]", True)
    rows, skipped, errors = [], 0, 0
    for index, item in enumerate(items):
        if index >= limit:
            return tuple(rows), skipped, errors, True
        try:
            if getattr(item, "Class", None) != 43:
                skipped += 1
                continue
            sent = item.SentOn
            if sent is None:
                errors += 1
                continue
            day = sent.date()
            if day < start:
                break
            if day > end:
                continue
            rows.append({
                "Destinatario": str(item.To or ""),
                "Fecha de envío": day.isoformat(),
                "Asunto": str(item.Subject or ""),
                "EntryID": str(item.EntryID or ""),
            })
        except Exception:
            errors += 1
    return tuple(rows), skipped, errors, False


def count_sent_mail(start: date, end: date, *, account_key=None, limit=50000):
    if end < start:
        raise ValueError("La fecha final debe ser igual o posterior a la inicial.")
    if platform.system() != "Windows":
        raise OutlookUnavailable("El contador requiere Windows y Outlook clásico.")
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise OutlookUnavailable("Falta pywin32 en el entorno NuRus.") from exc
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        session = outlook.Session
        account = _find_account(session, account_key) if account_key else None
        folder = account.DeliveryStore.GetDefaultFolder(5) if account else session.GetDefaultFolder(5)
        rows, skipped, errors, truncated = scan_sent_items(folder.Items, start, end, limit=limit)
        return SentMailReport(
            rows, skipped, errors, truncated, account_key or "predeterminada", str(folder.FolderPath)
        )
    finally:
        pythoncom.CoUninitialize()


def export_sent_report(report: SentMailReport, destination):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from pathlib import Path

    path = Path(destination).resolve()
    if path.suffix.lower() != ".xlsx" or path.exists():
        raise ValueError("Selecciona un archivo .xlsx nuevo.")
    book = Workbook()
    sheet = book.active
    sheet.title = "Correos enviados"
    headers = ("Destinatario", "Fecha de envío", "Asunto", "EntryID")
    sheet.append(headers)
    for row in report.rows:
        sheet.append([row[name] for name in headers])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in (("A", 45), ("B", 18), ("C", 80), ("D", 24)):
        sheet.column_dimensions[column].width = width
    control = book.create_sheet("Control")
    for row in (
        ("Cuenta", report.account), ("Carpeta", report.folder),
        ("Correos", len(report.rows)), ("Omitidos no correo", report.skipped),
        ("Errores", report.errors), ("Consulta limitada", report.truncated),
    ):
        control.append([excel_text(value) for value in row])
    try:
        write_new_file(path, book.save)
    finally:
        book.close()
    return path
