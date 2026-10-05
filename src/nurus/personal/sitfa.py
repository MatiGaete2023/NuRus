"""Contrato opcional de entrada/procedencia SITFA. No cambia reglas ni revisiones."""
from datetime import date,datetime
import hashlib
import json
from pathlib import Path

from nurus.rus.reader import read_workbook

REQUIRED={'ESPERA':('rit','tribunal','nombre','programa','espera'),
          'CUMPLIMIENTO':('rit','tribunal','nombre','programa','dias_cumpl','dias_egresar'),
          'INFORMES':('rit','tribunal','nombre','programa','vencimiento')}


def inspect_batch(batch,today=None):
    today=today or date.today();mapping=batch.column_mapping
    missing=[key for key in REQUIRED[str(batch.mode)] if key not in mapping]
    incomplete=sum(any(row.values.get(mapping.get(key,'')) in ('',None) for key in REQUIRED[str(batch.mode)] if key in mapping) for row in batch.records)
    dates=set()
    for row in batch.records:
        raw=str(row.values.get('SITFA_FECHA_DESCARGA',''))[:10]
        if raw:
            try:dates.add(date.fromisoformat(raw))
            except ValueError:pass
    ages=sorted((today-day).days for day in dates)
    state='Sin datos suficientes' if missing else 'Compatible con limitaciones' if incomplete else 'Compatible'
    messages=[]
    if missing:messages.append('Columnas funcionales ausentes: '+', '.join(missing)+'. Algunas reglas no pueden evaluarse.')
    if incomplete:messages.append(f'{incomplete} filas tienen campos obligatorios vacíos; revisar antes de usar los productos.')
    if ages and max(ages)>0:messages.append(f'La descarga tiene hasta {max(ages)} días. El análisis usa {today.isoformat()}; los días de espera/cumplimiento no se recalculan.')
    if ages and min(ages)<0:messages.append('La procedencia indica una fecha de descarga futura: revisar el origen.')
    return {'estado':state,'faltan':missing,'filas_incompletas':incomplete,'mensajes':messages}


def inspect_input(path,mode,sheet=None):
    return inspect_batch(read_workbook(path,mode,sheet_name=sheet))


def source_index(content):
    """Leer mapa de archivos conservado en el libro, sin confiar en rutas ejecutables."""
    import io
    import openpyxl
    if not content.startswith(b'PK\x03\x04'):return {}
    book=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    try:
        if 'SITFA_ARCHIVOS' not in book.sheetnames:return {}
        entries={}
        rows=book['SITFA_ARCHIVOS'].iter_rows(values_only=True)
        header=next(rows,())
        if tuple(header[:3])!=('Carpeta origen','Archivo origen','SHA256'):return {}
        for index,row in enumerate(rows):
            if index>20_000:raise ValueError('Demasiados archivos de procedencia.')
            if len(row)<3:continue
            folder,name,digest=map(lambda v:str(v or ''),row[:3])
            if name and digest:entries[json.dumps([Path(folder).name,name,digest],ensure_ascii=False)]=folder
        return entries
    finally:book.close()


def verified_source(work,row):
    name=str(row.values.get('SITFA_ARCHIVO',''));digest=str(row.values.get('SITFA_SHA256',''))
    key=json.dumps([str(row.values.get('SITFA_LOTE','')),name,digest],ensure_ascii=False)
    folder=getattr(work,'sitfa_sources',{}).get(key)
    if not folder:raise ValueError('No hay procedencia SITFA verificable para esta fila.')
    root=Path(folder).resolve()
    if Path(name).name!=name or '\\' in name:raise ValueError('Ruta de origen inválida.')
    path=(root/name).resolve()
    if path.parent!=root or path.suffix.lower() not in ('.xls','.xlsx') or not path.is_file():
        raise ValueError('El original fue movido o no está disponible. Conserva la carpeta del lote.')
    calculated=hashlib.sha256()
    with path.open('rb') as source:
        for block in iter(lambda:source.read(1024*1024),b''):calculated.update(block)
    if calculated.hexdigest()!=digest:raise ValueError('El original cambió; no coincide con el hash de la descarga.')
    return path


def startup_arguments(argv=None):
    import argparse
    parser=argparse.ArgumentParser(description='CSMP Assistant experimental: precarga sin análisis automático')
    parser.add_argument('--archivo',type=Path)
    parser.add_argument('--modo',choices=list(REQUIRED))
    parser.add_argument('--hoja')
    parser.add_argument('--verificar-paquete',type=Path,help=argparse.SUPPRESS)
    args=parser.parse_args(argv)
    if args.archivo:
        args.archivo=args.archivo.expanduser().resolve()
        if not args.archivo.is_file() or args.archivo.suffix.lower() not in ('.xls','.xlsx','.xlsm'):
            parser.error('El archivo Excel no existe o su extensión no está admitida.')
    return args
