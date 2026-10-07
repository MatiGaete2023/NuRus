"""Productos manuales y nóminas revisadas, separados del análisis de RUS."""
from copy import deepcopy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from uuid import uuid4
import json

from .config import atomic_json, CC
from .outputs import Draft, value, mail_metric_key, mail_metric_value, due_value, program_filename


def logical_mail_key(draft):
    return json.dumps([draft.kind,draft.recipient_type,draft.court,draft.program],ensure_ascii=False)


class ManualStore:
    def __init__(self,folder,config):
        self.folder=Path(folder);self.config=deepcopy(config)
        self.path=self.folder/'productos_manuales.json'
        self.storage_directory=str(self.folder)
        self.output='';self.rows=[];self.mapping={};self.receipts={};self.drafts=[];self.resolution={}
        if self.path.is_file():
            data=json.loads(self.path.read_text(encoding='utf-8'))
            if data.get('version')!=1:raise ValueError('Versión de productos manuales no compatible.')
            self.drafts=[Draft(**d) for d in data.get('drafts',[])]
            self.receipts=data.get('receipts',{});self.resolution=data.get('resolution',{})

    def save(self,directory=None):
        atomic_json(self.path,{'version':1,'drafts':[asdict(d) for d in self.drafts],
                              'receipts':self.receipts,'resolution':self.resolution})

    def new_mail(self):
        draft=Draft('','',cc='; '.join(filter(None,(CC,self.config['correos'].get('cc_adicional','')))),
                    key=uuid4().hex,kind='manual',recipient_type='manual',options={'manual':True})
        draft.original={key:deepcopy(getattr(draft,key)) for key in ('to','cc','subject','body','attachments')}
        self.drafts.append(draft);self.save();return draft


def table_snapshot(work,rows,kind):
    if kind=='programa_espera':
        headers=['RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA'];keys=('rit','tribunal','nombre','programa','espera')
    elif kind in {'programa_vencido','programa_por_vencer'}:
        headers=['RIT','TRIBUNAL','NOMBRE','DERIVACION','F. VENCIMIENTO'];keys=('rit','tribunal','nombre','programa','vencimiento')
    else:
        metric=mail_metric_key(work,kind)
        headers=['RIT','TRIBUNAL','RUT','NOMBRE','PROGRAMA',{'espera':'T ESPERA','vencimiento':'F. VENCIMIENTO','egreso_proy':'F. EGRESO PROYECTADO'}[metric],'OBSERVACION']
        keys=('rit','tribunal','rut','nombre','programa',metric,'observation')
    records=[]
    for row in rows:
        values=[]
        for key in keys:
            v=row.review.get('OBSERVACION',row.observation) if key=='observation' else mail_metric_value(work,row,key) if key in ('espera','vencimiento','egreso_proy') else value(work,row,key)
            values.append(str(v if v not in ('',None) else 'Sin dato' if key in ('espera','vencimiento','egreso_proy') else ''))
        records.append({'id':row.id,'values':values,'base_values':list(values),'include':True,'group':value(work,row,'programa')})
    return {'headers':headers,'records':records,'reviewed':False,'dirty':True,'generated':[]}


def empty_table():
    return {'headers':['RIT','TRIBUNAL','RUT','NOMBRE','PROGRAMA','F. VENCIMIENTO','T ESPERA','OBSERVACION'],
            'records':[],'reviewed':False,'dirty':True,'generated':[]}


def update_table(draft,table):
    old=draft.options.get('table',{})
    table=deepcopy(table)
    if not table.get('headers') or len(set(table['headers']))!=len(table['headers']):raise ValueError('Cabecera de nómina inválida.')
    for row in table.get('records',[]):
        if len(row['values'])!=len(table['headers']):raise ValueError('Una fila no coincide con las columnas de la nómina.')
        program=next((i for i,h in enumerate(table['headers']) if h in ('PROGRAMA','DERIVACION')),None)
        if program is not None:row['group']=row['values'][program] or 'Nomina revisada'
    draft.attachments=[p for p in draft.attachments if p not in old.get('generated',[])]
    table.update(reviewed=True,dirty=True,generated=[])
    draft.options['table']=table;draft.options['pending_table']=True


def generate_table(draft,directory):
    table=draft.options.get('table')
    if not table or not table.get('reviewed'):raise ValueError('Revisa los registros de la nómina antes de generar el adjunto.')
    included=[r for r in table['records'] if r.get('include',True)]
    if not included:raise ValueError('La nómina no tiene registros incluidos; no se genera un Excel vacío.')
    groups={}
    for row in included:groups.setdefault(row.get('group') or 'Nomina revisada',[]).append(row)
    folder=Path(directory)/'adjuntos'/uuid4().hex;folder.mkdir(parents=True)
    from openpyxl import Workbook
    from openpyxl.styles import Border,Side,Font,Alignment
    from openpyxl.utils import get_column_letter
    from .attachment_store import workbook_bytes
    from nurus.services.file_output import write_new_file
    generated=[]
    for name,rows in groups.items():
        book=Workbook();sheet=book.active;sheet.title='Nómina';sheet.append(table['headers'])
        for row in rows:
            sheet.append(row['values'])
            for cell in sheet[sheet.max_row]:cell.data_type='s'
        edge=Side(style='thin',color='FF000000');border=Border(left=edge,right=edge,top=edge,bottom=edge)
        for cells in sheet:
            for cell in cells:cell.border=border;cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in sheet[1]:cell.font=Font(bold=True);cell.data_type='s'
        for i,header in enumerate(table['headers'],1):sheet.column_dimensions[get_column_letter(i)].width=42 if 'NOMBRE' in header or 'OBSERV' in header else 25
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        content=workbook_bytes(book);book.close();path=folder/(program_filename(name)+'.xlsx')
        write_new_file(path,lambda target:target.write_bytes(content));generated.append(str(path))
    old=table.get('generated',[])
    draft.attachments=[p for p in draft.attachments if p not in old]+generated
    table.update(dirty=False,generated=generated);draft.options['pending_table']=False
    return generated


def merge_table_edits(old,new):
    before=old.options.get('table');after=new.options.get('table')
    if not before or not after or before['headers']!=after['headers']:return
    index={r['id']:r for r in before['records']}
    for row in after['records']:
        previous=index.get(row['id'])
        if previous:
            row['include']=previous.get('include',True)
            for i,(edited,base) in enumerate(zip(previous['values'],previous.get('base_values',previous['values']))):
                if edited!=base:row['values'][i]=edited
    present={r['id'] for r in after['records']}
    after['records'] += [deepcopy(r) for r in before['records'] if r.get('manual_added') and r['id'] not in present]
    # Un nuevo preparado necesita revisión; no reutiliza una nómina obsoleta.
    new.attachments=[p for p in new.attachments if p not in before.get('generated',[])]
    new.options['pending_table']=True


def free_resolution(destination,court,rit,text):
    if not court.strip() or not rit.strip() or not text.strip():raise ValueError('Completa tribunal, RIT y texto del proyecto.')
    from docx import Document
    from docx.shared import Pt
    from nurus.services.file_output import write_new_file
    doc=Document();doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(12)
    doc.add_paragraph(court.strip());doc.add_paragraph('RIT: '+rit.strip())
    for line in text.splitlines():doc.add_paragraph(line)
    write_new_file(Path(destination),doc.save);return str(destination)


def manual_project(template,court,rit,values):
    from .resolutions import Project,_validate_project_template,render_project,paragraphs
    from .outputs import template_variables
    from tempfile import TemporaryDirectory
    values={**values,'RIT':rit}
    missing=[key for key in template_variables(template) if not str(values.get(key,'')).strip()]
    if missing:raise ValueError('Completa los campos de la matriz: '+', '.join(sorted(missing)))
    project=Project(uuid4().hex,court,rit,Path(template).stem,[uuid4().hex],str(template),values,'','')
    _validate_project_template(project)
    with TemporaryDirectory() as tmp:
        doc=render_project(project,Path(tmp)/'vista.docx');project.text='\n'.join(p.text for p in paragraphs(doc))
    project.original_text=project.text;project.original_values=deepcopy(values)
    from hashlib import sha256
    project.template_hash=sha256(Path(template).read_bytes()).hexdigest()
    return project
