"""Cruce factual de firmas por causa. Nunca modifica propuestas, TT o CC."""
from collections import defaultdict
from datetime import date,datetime
from io import BytesIO
import json
import re

from nurus.rus.columns import normalize
from nurus.rus.rules import tribunal

TITLES=('ACTIVIDAD RECIENTE','ÚLTIMA FIRMA','RESOLUCIONES EN EL PERÍODO','DETALLE')


def _headers(sheet):
    rows=sheet.iter_rows(values_only=True);header=next(rows,())
    names=[str(value or '') for value in header]
    if len(set(names))!=len(names):raise ValueError('Columnas ambiguas en la fuente de actividad.')
    return names,rows


def _day(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    try:return date.fromisoformat(str(value)[:10])
    except ValueError:return None


def read_sources(content):
    if not content.startswith(b'PK\x03\x04'):return None
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        if 'Resoluciones firmadas' not in book.sheetnames or 'SITFA_COBERTURA' not in book.sheetnames:return None
        names,rows=_headers(book['SITFA_COBERTURA'])
        required=('tribunal_codigo','desde','hasta','estado','criterio_temporal','sin_coincidencias')
        if any(key not in names for key in required):raise ValueError('Falta el contrato de cobertura de actividad.')
        coverage={}
        for row in rows:
            item=dict(zip(names,row));code=str(item['tribunal_codigo'] or '')
            start,end=_day(item['desde']),_day(item['hasta'])
            if not code or code in coverage or not start or not end or start>end:raise ValueError('Cobertura inválida o duplicada.')
            if item['estado'] not in ('COMPLETA','VACIA_COMPROBADA','PARCIAL','FALLIDA','NO_CONSULTADA'):
                raise ValueError('Estado de cobertura no reconocido.')
            coverage[code]=item
        names,rows=_headers(book['Resoluciones firmadas']);required=('tribunal_codigo','tribunal','rit','tramite','firma','hora','identidad','huella')
        if any(key not in names for key in required):raise ValueError('Faltan columnas del detalle de firmas.')
        result=[];signature_counts=defaultdict(int)
        for row in rows:
            item=dict(zip(names,row));code=str(item['tribunal_codigo'] or '');day=_day(item['firma'])
            scope=coverage.get(code)
            if not scope or not day or not _day(scope['desde'])<=day<=_day(scope['hasta']):raise ValueError('Firma fuera de la cobertura declarada.')
            item['tribunal_codigo']=code;item['firma']=day.isoformat();result.append(item);signature_counts[code]+=1
        if any(item['estado']=='VACIA_COMPROBADA' and signature_counts[code]
               for code,item in coverage.items()):
            raise ValueError('Una fuente con firmas no puede declarar cobertura vacía comprobada.')
        return {'firmas':result,'cobertura':coverage}
    finally:book.close()


def court_key(value):
    return tribunal(value) or normalize(value)


def _legacy_alias_matches(signatures):
    alias_counts=defaultdict(int);fingerprint_counts=defaultdict(int)
    aliases={}
    for signature in signatures:
        fingerprint=str(signature.get('huella') or '')
        fingerprint_counts[fingerprint]+=1
        identity=str(signature.get('identidad') or '')
        match=re.fullmatch(r'huella_compuesta_local_sin_id_remoto;huella_legacy=([0-9a-f]{64}:\d+)',identity)
        if match:
            alias=match.group(1);aliases[fingerprint]=alias;alias_counts[alias]+=1
    return {fingerprint:alias for fingerprint,alias in aliases.items()
        if fingerprint_counts[fingerprint]==1 and alias_counts[alias]==1}


def _reviews_with_legacy_aliases(signatures, reviewed):
    """Dual-read only exact, one-to-one legacy aliases; keep every old flag."""
    result=dict(reviewed)
    aliases=_legacy_alias_matches(signatures)
    for fingerprint,alias in aliases.items():
        if alias in result:
            result.setdefault(fingerprint,result[alias])
    return result


def attach(work):
    sources=read_sources(work.content)
    if sources is None:return
    index=defaultdict(list);labels=defaultdict(set)
    for signature in sources['firmas']:
        index[(signature['tribunal_codigo'],normalize(signature['rit']))].append(signature)
        labels[court_key(signature['tribunal'])].add(signature['tribunal_codigo'])
    work.activity_sources=sources;work.signed_activity={}
    for row in work.rows:
        code=str(row.values.get('SITFA_TRIBUNAL_CODIGO','') or '')
        if not code:
            matches=labels.get(court_key(row.values.get(work.mapping.get('tribunal',''),'')),set())
            if len(matches)==1:code=next(iter(matches))
        scope=sources['cobertura'].get(code)
        found=index.get((code,normalize(row.values.get(work.mapping.get('rit',''),''))),[])
        prior_reviews=work.activity_reviewed.get(row.id,{})
        reviewed=_reviews_with_legacy_aliases(found,prior_reviews)
        work.activity_reviewed.setdefault(row.id,{}).update(reviewed)
        ambiguous=[s for s in found if str(s.get('identidad') or '').startswith('huella_compuesta_ambigua_')]
        pending=[s for s in found if str(s.get('identidad') or '').startswith('huella_compuesta_ambigua_') or s['huella'] not in reviewed]
        if found:
            state='Firma por revisar en la causa' if pending else 'Firmas revisadas de este ingreso'
            latest=max(s['firma'] for s in found)
            details='Resolución firmada en la causa; revisar pertinencia. '+str(len(found))+' filas firmadas; identidad remota no disponible.'
            color='azul' if pending else ''
        else:
            complete=scope and scope['estado'] in ('COMPLETA','VACIA_COMPROBADA')
            state='Sin coincidencias en los informes consultados' if scope and not complete else 'Sin firmas en el período' if scope else 'Actividad no consultada'
            latest='';details='';color='ambar' if not complete else ''
        if scope:
            details+=((' ' if details else '')+f"Período {scope['desde']} a {scope['hasta']}; cobertura {scope['estado'].lower()}.")
        if ambiguous:
            details+=((' ' if details else '')+f'Firmas con identidad ambigua: {len(ambiguous)}; sin ID remoto no se pueden marcar individualmente como revisadas con seguridad.')
        resolved_legacy=set(_legacy_alias_matches(found).values())
        unresolved_legacy=sum(1 for key in prior_reviews if re.fullmatch(r'[0-9a-f]{64}:\d+',str(key)) and key not in resolved_legacy)
        if unresolved_legacy:
            details+=((' ' if details else '')+f'Revisiones heredadas sin vínculo seguro: {unresolved_legacy}; historial conservado para revisión manual.')
        work.signed_activity[row.id]={'firmas':found,'pendientes':len(pending),'color':color,
            'valores':dict(zip(TITLES,(state,latest,len(found) if scope else '',details)))}


def mark_reviewed(work,row_id,fingerprint,reviewed_at=None):
    entry=work.signed_activity.get(row_id)
    if not entry or fingerprint not in {s['huella'] for s in entry['firmas']}:raise ValueError('La firma no corresponde al ingreso seleccionado.')
    signature=next(s for s in entry['firmas'] if s['huella']==fingerprint)
    if str(signature.get('identidad') or '').startswith('huella_compuesta_ambigua_'):
        raise ValueError('Esta firma no tiene identidad única. Revisa el detalle, pero no se puede marcar individualmente con seguridad.')
    timestamp=(reviewed_at or datetime.now().astimezone()).isoformat()
    if getattr(work,'review_store_path',''):
        from .review_store import ReviewStore
        row=next(r for r in work.rows if r.id==row_id)
        ReviewStore(work.review_store_path).remember(work,row,fingerprint,timestamp)
    work.activity_reviewed.setdefault(row_id,{})[fingerprint]=timestamp
    work.revision+=1;attach(work)
