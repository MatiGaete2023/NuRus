"""Conteos de constancias; propuestas y borradores no equivalen a gestiones."""
from collections import Counter
from pathlib import Path
from nurus.rus.rules import as_date, as_int
from nurus.services.file_output import write_new_file

def summarize(work):
    counts=Counter(propuestas=len(work.rows),excluidos=sum(r.excluded for r in work.rows),constancias=0,con_carga=0,sin_carga=0,datos_incompletos=0)
    for row in work.rows:
        if row.excluded:continue
        review=row.review
        if not review:continue
        day=as_date(review.get('FECHA_OBS'))
        total=as_int(review.get('TT'));charge=as_int(review.get('CC'))
        if day and review.get('OBSERVACION') and total==1 and charge in (0,1):
            counts['constancias']+=1;counts['con_carga']+=charge;counts['sin_carga']+=1-charge
        elif any(review.get(k) not in ('',None) for k in ('FECHA_OBS','TT','CC','RES')):counts['datos_incompletos']+=1
    counts['word_generados']=len({r.get('path') for r in work.receipts.values() if r.get('kind')=='word'})
    counts['proyectos_generados']=sum(r.get('kind')=='word' for r in work.receipts.values())
    counts['borradores_creados']=sum(r.get('kind')=='draft' and r.get('state')=='created' for r in work.receipts.values())
    return dict(counts)

def export_summary(work,path):
    from openpyxl import Workbook
    from .outputs import value
    work.refresh();book=Workbook();sheet=book.active;sheet.title='Resumen'
    sheet.append(['Categoría','Cantidad'])
    for name,n in summarize(work).items():sheet.append([name.replace('_',' '),n])
    rows=book.create_sheet('Constancias')
    rows.append(['Tribunal','RIT','Fecha','TT','CC','Sin carga','RES'])
    for row in work.rows:
        r=row.review;tt=as_int(r.get('TT'));cc=as_int(r.get('CC'))
        if row.excluded or not r.get('OBSERVACION') or not as_date(r.get('FECHA_OBS')) or tt!=1 or cc not in (0,1):continue
        rows.append([value(work,row,'tribunal'),value(work,row,'rit'),str(r['FECHA_OBS']),tt,cc,tt-cc,str(r.get('RES',''))])
        for c in (rows.cell(rows.max_row,1),rows.cell(rows.max_row,2),rows.cell(rows.max_row,7)):c.data_type='s'
    for ws in book:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in 'ABCDEFG':ws.column_dimensions[col].width=25
    write_new_file(Path(path),book.save)
    return str(path)
