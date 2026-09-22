"""Proyectos agrupados y editables; un documento por lote."""
from dataclasses import dataclass, field
from collections import OrderedDict
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4
from difflib import SequenceMatcher

from docx import Document
from docxcompose.composer import Composer

from nurus.rus.rules import tribunal
from nurus.rus.columns import normalize
from nurus.services.file_output import write_new_file
from .outputs import value,word_values,template_variables,fill_docx


KINDS=('PC_IE','PC_INFO','NOMENCL')


@dataclass
class Project:
    key: str
    court: str
    rit: str
    kind: str
    record_ids: list
    template: str
    values: dict
    text: str
    original_text: str
    warnings: list = field(default_factory=list)


def join_names(values):
    values=list(dict.fromkeys(str(v).strip() for v in values if str(v).strip()))
    return values[0] if len(values)==1 else ', '.join(values[:-1])+' y '+values[-1] if values else ''


def _has_explicit_value(value):
    if value is None:return False
    if isinstance(value,str):return bool(value.strip())
    return True


def resolution_kind(value):
    text=normalize(value).replace('_',' ')
    return {'pc ie':'PC_IE','pc info':'PC_INFO','nomencl':'NOMENCL','nomenclatura':'NOMENCL'}.get(text)


def _case_key(work,row):
    court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
    return normalize(court),normalize(value(work,row,'rit'))


def resolution_marked(value):
    """Compatibilidad: reconoce tipos nuevos y marcas usadas en planillas antiguas."""
    if resolution_kind(value):return True
    if isinstance(value,bool):return value
    if isinstance(value,(int,float)):
        try:return float(value)!=0
        except (TypeError,ValueError):return False
    raw=str(value or '').strip().lower().replace(',','.')
    if raw in {'1','1.0'}:return True
    if raw in {'0','0.0'}:return False
    text=normalize(value)
    if text in {'si','s','x','true','verdadero','res','resolucion','con resolucion'}:return True
    return False


def resolution_review_issue(value):
    """Advierte un RES escrito pero desconocido; nunca intenta corregirlo por aproximación."""
    if not _has_explicit_value(value):return ''
    if resolution_kind(value) or resolution_marked(value):return ''
    shown=str(value).strip()
    return f"RES no reconocido: {shown!r}. Usa PC_IE, PC_INFO o NOMENCL, o deja la celda vacía."


def resolution_selection_source(work,row,kind):
    """Explica si el tipo procede de RES, de compatibilidad antigua o del motor."""
    case=_case_key(work,row)
    peers=[item for item in work.rows if not item.excluded and _case_key(work,item)==case]
    if any(resolution_kind((item.review or {}).get('RES'))==kind for item in peers):return 'Definido en RES'
    if any(_has_explicit_value((item.review or {}).get('RES')) and resolution_marked((item.review or {}).get('RES')) for item in peers):return 'RES antiguo · tipo inferido'
    return 'Sugerencia automática revisable'

def reviewed_resolution_ids(work):
    """None si la planilla no tenía RES; si existe, incluso vacío, manda la selección humana."""
    explicit=any('RES' in (row.review or {}) for row in work.rows)
    if not explicit:return None
    selected=set()
    for row in work.rows:
        if 'RES' not in (row.review or {}):continue
        value=row.review.get('RES')
        if _has_explicit_value(value) and resolution_marked(value):selected.add(row.id)
    return selected


def observation_resolution_kind(row):
    """Obtiene el tipo de matriz desde la observación humana cuando RES no lo especifica."""
    text=normalize((row.review or {}).get('OBSERVACION',row.observation) or '')
    if not text:return None
    if 'nomencl' in text:return 'NOMENCL'
    if 'pc info' in text:return 'PC_INFO'
    if 'pc ie' in text:return 'PC_IE'
    if any(part in text for part in ('informe','diagnostico','diagnostico clinico')):return 'PC_INFO'
    if any(part in text for part in ('ingreso efectivo','fecha estimada de ingreso','ingreso al programa')):return 'PC_IE'
    return None


def automatic_project_selections(work,fallback_kind='PC_IE'):
    """Selecciona proyectos con precedencia RES explícito > observación > regla.

    Si RES no fue usado, una observación por sí sola no crea un proyecto: debe
    existir una acción de resolución del motor. Si RES sí fue usado como marca, la
    observación puede discriminar la matriz y el fallback se usa solo como último
    recurso para un caso expresamente marcado.
    """
    reviewed=reviewed_resolution_ids(work)
    result=[]
    if reviewed is None:
        for row in work.rows:
            if row.excluded:continue
            suggested=[kind for kind in row.actions if kind in KINDS]
            if not suggested:continue
            observed=observation_resolution_kind(row)
            kinds=[observed] if observed else suggested
            for kind in kinds:
                if kind in KINDS:result.append((row.id,kind))
        return result

    case_kinds=OrderedDict()
    for row in work.rows:
        if row.id not in reviewed or row.excluded:continue
        explicit=resolution_kind(row.review.get('RES'))
        observed=observation_resolution_kind(row)
        actions=[kind for kind in row.actions if kind in KINDS]
        kinds=[explicit] if explicit else ([observed] if observed else actions)
        if not kinds:kinds=[fallback_kind]
        bucket=case_kinds.setdefault(_case_key(work,row),[])
        for kind in kinds:
            if kind and kind not in bucket:bucket.append(kind)
    for row in work.rows:
        if row.excluded:continue
        for kind in case_kinds.get(_case_key(work,row),[]):result.append((row.id,kind))
    return result


def unique_case_selections(work,selections):
    """Una fila visible por tribunal + RIT + tipo, preservando el orden de origen."""
    by_id={row.id:row for row in work.rows}
    seen=set();result=[]
    for rid,kind in selections:
        row=by_id.get(rid)
        if not row or row.excluded or kind not in KINDS:continue
        key=(*_case_key(work,row),kind)
        if key in seen:continue
        seen.add(key);result.append((rid,kind))
    return result


def replace_paragraph(paragraph,text):
    """Modifica solo tramos cambiados conservando los runs del resto de la matriz."""
    original=paragraph.text
    for tag,i,j,a,b in reversed(SequenceMatcher(None,original,text,autojunk=False).get_opcodes()):
        if tag=='equal':continue
        runs=list(paragraph.runs)
        if not runs:paragraph.add_run(text);return
        positions=[];pos=0
        for run in runs:
            positions.append((run,pos,pos+len(run.text)));pos+=len(run.text)
        touched=[t for t in positions if t[1]<j and t[2]>i]
        if not touched:
            run=next((r for r,start,end in positions if start<=i<=end),runs[-1])
            start=next(start for r,start,end in positions if r is run)
            run.text=run.text[:i-start]+text[a:b]+run.text[i-start:]
            continue
        first,start,end=touched[0];last,lstart,lend=touched[-1]
        prefix=first.text[:i-start];suffix=last.text[j-lstart:]
        first.text=prefix+text[a:b]+(suffix if first is last else '')
        for run,_,_ in touched[1:]:run.text=suffix if run is last else ''


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def grouped_values(work,rows):
    vals=word_values(work,rows[0])
    for key in ('NOMBRE','RUT','PROGRAMA','FECHA_RESOLUCION','DURACION'):
        vals[key]=join_names(word_values(work,r).get(key,'') for r in rows)
    persons=OrderedDict()
    for r in rows:
        name=value(work,r,'nombre');rut=value(work,r,'rut');program=value(work,r,'programa')
        persons.setdefault((normalize(name),normalize(rut)),[name,rut,[]])[2].append(program)
    vals['PERSONAS']=join_names(name+', cédula de identidad N° '+(rut or '[COMPLETAR RUT]') for name,rut,_ in persons.values())
    vals['DETALLE_PERSONAS']='; '.join(name+', cédula de identidad N° '+(rut or '[COMPLETAR RUT]')+', programa '+join_names(programs) for name,rut,programs in persons.values())
    vals['_PLURAL']=len(persons)>1
    vals['_VARIOS_PROGRAMAS']=len({normalize(value(work,r,'programa')) for r in rows})>1
    vals['_VARIAS_FECHAS']=len({word_values(work,r).get('FECHA_RESOLUCION') for r in rows if word_values(work,r).get('FECHA_RESOLUCION')})>1
    return vals


def render_project(project,target):
    fill_docx(project.template,target,project.values)
    doc=Document(target)
    if project.values.get('_PLURAL'):
        replacements=[
            ('la persona aludida','las personas individualizadas'),
            ('del niño sujeto de protección','de los niños, niñas o adolescentes individualizados'),
            ('del niño, niña o adolescente','de los niños, niñas o adolescentes'),
        ]
        if project.values.get('_VARIOS_PROGRAMAS'):
            replacements.extend([('al organismo interventor','a los organismos interventores'),('al programa interventor','a los programas intervinientes'),('a la institución','a las instituciones correspondientes'),('de dicho programa','de dichos programas')])
        if project.values.get('_VARIAS_FECHAS'):replacements.append(('con fecha ','con fechas '))
        for p in paragraphs(doc):
            text=p.text
            for marker in ('cédula de identidad N°','cédula de identidad N.º'):
                old_pair=project.values['NOMBRE']+', '+marker+' '+project.values['RUT']
                if old_pair in text:
                    people=project.values['PERSONAS'].replace('cédula de identidad N°',marker)
                    text=text.replace(old_pair,people)
            for old,new in replacements:text=text.replace(old,new)
            replace_paragraph(p,text)
        if project.values.get('_VARIOS_PROGRAMAS'):
            detail='Personas comprendidas: '+project.values['DETALLE_PERSONAS']+'.'
            anchor=next((p for p in doc.paragraphs if '{{' not in p.text and p.text.strip().startswith('RIT')),None)
            if anchor:anchor.insert_paragraph_before(detail)
            else:doc.add_paragraph(detail)
    doc.save(target)
    return doc


def prepare_projects(work,selections,template_dir):
    work.refresh()
    by_id={r.id:r for r in work.rows};groups=OrderedDict();case_rows=OrderedDict()
    for row in work.rows:
        if row.excluded:continue
        court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
        rit=value(work,row,'rit').strip()
        case_rows.setdefault((court,normalize(rit)),[]).append(row)
    for rid,kind in selections:
        if rid not in by_id or kind not in KINDS:continue
        row=by_id[rid]
        if row.excluded:continue
        court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
        rit=value(work,row,'rit').strip()
        key=(court,normalize(rit),kind)
        group=groups.setdefault(key,[])
        existing={r.id for r in group}
        for candidate in case_rows.get((court,normalize(rit)),[]):
            if candidate.id not in existing:group.append(candidate);existing.add(candidate.id)
    projects=[];errors=[]
    with TemporaryDirectory() as directory:
        for (court,_,kind),rows in groups.items():
            rit=value(work,rows[0],'rit')
            template=Path(template_dir)/court/(kind+'.docx')
            if not template.is_file():
                errors.append(rit+': falta matriz '+court+'/'+kind)
                continue
            vals=grouped_values(work,rows);missing=[]
            for key in template_variables(template):
                if not vals.get(key):vals[key]='[COMPLETAR '+key.replace('_',' ')+']';missing.append(key)
            project=Project(uuid4().hex,court,rit,kind,[r.id for r in rows],str(template),vals,'','',missing)
            doc=render_project(project,Path(directory)/(project.key+'.docx'))
            project.text='\n'.join(p.text for p in paragraphs(doc));project.original_text=project.text
            projects.append(project)
    return projects,errors


def generate_projects(work,projects,destination):
    """Una resolución comienza en página nueva. No corta textos extensos."""
    if not projects:raise ValueError('No hay proyectos disponibles para generar.')
    with TemporaryDirectory() as directory:
        docs=[]
        for project in projects:
            target=Path(directory)/(project.key+'.docx')
            doc=render_project(project,target)
            if project.text!=project.original_text:
                existing=list(paragraphs(doc));lines=project.text.split('\n')
                for i,p in enumerate(existing):replace_paragraph(p,lines[i] if i<len(lines) else '')
                for line in lines[len(existing):]:doc.add_paragraph(line)
            docs.append(doc)
        composer=Composer(docs[0])
        for doc in docs[1:]:
            composer.doc.add_page_break()
            composer.append(doc)
        write_new_file(Path(destination),composer.save)
    for project in projects:
        work.receipts[uuid4().hex]={'kind':'word','record_ids':project.record_ids,'record_id':project.record_ids[0],'path':str(destination),'type':project.kind}
    if getattr(work,'storage_directory',None):work.save(work.storage_directory)
    return str(destination)
