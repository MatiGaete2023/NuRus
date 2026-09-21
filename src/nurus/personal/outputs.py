"""Salidas desde el trabajo compartido. No modifica RUS ni envía correos."""
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from string import Formatter
from uuid import uuid4
import json
import re

from nurus.domain.models import Product, ProductKind, ProductStatus, Template
from nurus.adapters.outlook import save_draft, DraftSaveUncertain
from nurus.rus.columns import normalize
from nurus.rus.rules import tribunal
from nurus.services.file_output import write_new_file
from .config import CC, emails


@dataclass
class Draft:
    subject: str
    body: str
    to: str = ''
    cc: str = CC
    attachments: list = field(default_factory=list)
    required: bool = False
    key: str = ''
    program: str = ''
    court: str = ''
    due: str = ''
    kind: str = ''
    record_ids: list = field(default_factory=list)


_MONTHS = ('', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')
_UNITS = ('cero','uno','dos','tres','cuatro','cinco','seis','siete','ocho','nueve')
_SPECIAL = {10:'diez',11:'once',12:'doce',13:'trece',14:'catorce',15:'quince',16:'dieciséis',17:'diecisiete',18:'dieciocho',19:'diecinueve',20:'veinte',21:'veintiuno',22:'veintidós',23:'veintitrés',24:'veinticuatro',25:'veinticinco',26:'veintiséis',27:'veintisiete',28:'veintiocho',29:'veintinueve'}
_TENS = {30:'treinta',40:'cuarenta',50:'cincuenta',60:'sesenta',70:'setenta',80:'ochenta',90:'noventa'}
_HUNDREDS = {100:'cien',200:'doscientos',300:'trescientos',400:'cuatrocientos',500:'quinientos',600:'seiscientos',700:'setecientos',800:'ochocientos',900:'novecientos'}


def number_in_words(number):
    """Cardinal español suficiente para días y años de proyectos judiciales."""
    number=int(number)
    if number<0 or number>9999:raise ValueError('Número fuera del rango admitido para fecha.')
    if number<10:return _UNITS[number]
    if number in _SPECIAL:return _SPECIAL[number]
    if number<100:
        tens=(number//10)*10;unit=number%10
        return _TENS[tens]+(' y '+_UNITS[unit] if unit else '')
    if number<1000:
        hundreds=(number//100)*100;rest=number%100
        head='ciento' if hundreds==100 and rest else _HUNDREDS[hundreds]
        return head+(' '+number_in_words(rest) if rest else '')
    thousands=number//1000;rest=number%1000
    head='mil' if thousands==1 else number_in_words(thousands)+' mil'
    return head+(' '+number_in_words(rest) if rest else '')


def date_in_words(value):
    """Fecha completamente en palabras: «catorce de septiembre de dos mil veintiséis»."""
    return f"{number_in_words(value.day)} de {_MONTHS[value.month]} de {number_in_words(value.year)}"


def value(work,row,key):
    """Texto operativo de una celda sin perder ceros ni exponer horas de Excel."""
    raw=row.values.get(work.mapping.get(key,''),'')
    if raw is None:return ''
    if isinstance(raw,(datetime,date)):return raw.strftime('%d/%m/%Y')
    return str(raw)


def resolve_contact(cfg,name):
    key=normalize(name)
    aliases={normalize(k):normalize(v) for k,v in cfg['aliases'].items()}
    key=aliases.get(key,key)
    found={v for k,v in cfg['contactos'].items() if normalize(k)==key and v.strip()}
    return next(iter(found)) if len(found)==1 else ''


def import_contacts(cfg,path):
    import pandas as pd
    frame=pd.read_excel(path,dtype=str).fillna('')
    columns={normalize(c):c for c in frame.columns}
    name=next((columns[k] for k in ('nombre','programa','nombre programa') if k in columns),None)
    address=next((columns[k] for k in ('mail','correo','email','correo electronico') if k in columns),None)
    if not name or not address:raise ValueError('El catastro necesita columnas Nombre/Programa y Mail/Correo.')
    candidates=dict(cfg['contactos']);conflicts=[]
    for record in frame.to_dict('records'):
        n=str(record[name]).strip();m='; '.join(emails(record[address]))
        if not n:continue
        previous={v for k,v in candidates.items() if normalize(k)==normalize(n)}
        if previous and previous!={m}:conflicts.append(n);continue
        candidates[n]=m
    return candidates,conflicts


def _table(work,rows,path,kind=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='Nómina'
    if kind=='programa_espera':
        headers=['RIT','TRIBUNAL','NOMBRE','DERIVACION','T ESPERA']
        keys=('rit','tribunal','nombre','programa','espera')
    elif kind in {'programa_vencido','programa_por_vencer'}:
        headers=['RIT','TRIBUNAL','NOMBRE','DERIVACION','F. VENCIMIENTO']
        keys=('rit','tribunal','nombre','programa','vencimiento')
    else:
        headers=['RIT','TRIBUNAL','RUT','NOMBRE','PROGRAMA','VENCIMIENTO / ESPERA','OBSERVACION']
        keys=None
    sheet.append(headers)
    for row in rows:
        if keys:
            sheet.append([value(work,row,key) for key in keys])
        else:
            sheet.append([value(work,row,k) for k in ('rit','tribunal','rut','nombre','programa')]+[value(work,row,'vencimiento') or value(work,row,'espera'),str(row.review.get('OBSERVACION',row.observation))])
        for cell in sheet[sheet.max_row]:cell.data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True)
    for index,header in enumerate(headers,1):
        sheet.column_dimensions[get_column_letter(index)].width=38 if header=='NOMBRE' else 24
    write_new_file(Path(path),book.save)
    return str(path)


def draft_fingerprint(draft):
    """Identidad estable del borrador realmente revisado.

    Se usa nombre+hash del adjunto, no su ruta temporal. Así regenerar una nómina
    idéntica dentro de otra carpeta UUID no permite crear un duplicado en Outlook,
    mientras que una edición real de destinatarios, texto o bytes sí cambia la clave.
    """
    attachments=[]
    for item in draft.attachments:
        path=Path(item)
        digest=sha256(path.read_bytes()).hexdigest() if path.is_file() else 'MISSING'
        attachments.append((path.name,digest))
    payload={
        'to':'; '.join(emails(draft.to)),
        'cc':'; '.join(emails(CC+';'+draft.cc)),
        'subject':str(draft.subject or ''),
        'body':str(draft.body or ''),
        'attachments':sorted(attachments),
        'required':bool(draft.required),
    }
    return sha256(json.dumps(payload,ensure_ascii=False,sort_keys=True).encode('utf-8')).hexdigest()


def prepare_drafts(work,kind,*,modalities='',period='',confirmed_scope=False,selected=None,directory=None,manual_selection=False,modality_keys=None,recipient_scope='auto'):
    if recipient_scope not in {'auto','programas','tribunales'}:raise ValueError('Destino de correo inválido.')
    work.refresh()
    cfg=work.config
    tpl=cfg['correos']['plantillas'][kind]
    to_program=recipient_scope=='programas' or (recipient_scope=='auto' and kind.startswith('programa_'))
    automatic_kind=kind in {'programa_espera','programa_vencido','programa_por_vencer','medidas'}
    if tpl.get('usa_modalidades') and not modalities.strip():raise ValueError('Indica las modalidades efectivamente comprendidas.')
    from .modalities import selected_row
    groups=defaultdict(list)
    for row in work.rows:
        if row.excluded or (selected is not None and row.id not in selected) or not selected_row(work,row,modality_keys):continue
        # Las comunicaciones personalizadas no tienen una regla del motor.
        if automatic_kind and kind not in row.actions and not (manual_selection or getattr(work,'external_input',False)):continue
        if kind=='proyectos' and not manual_selection and not getattr(work,'external_input',False) and row.id not in {rid for r in work.receipts.values() if r.get('kind')=='word' for rid in r.get('record_ids',[r.get('record_id')])}:continue
        court=tribunal(value(work,row,'tribunal')) or value(work,row,'tribunal')
        program=value(work,row,'programa') if to_program else ''
        groups[(court,normalize(program))].append(row)
    output=[]
    for (court,_),rows in groups.items():
        court_cfg=cfg['correos']['tribunales'].get(court,{'nombre':court,'para':[]})
        program=value(work,rows[0],'programa')
        context={'TRIBUNAL':court_cfg['nombre'],'PROGRAMA':program,'FECHA':date.today().strftime('%d/%m/%Y'),
                 'ALCANCE_MODALIDADES':modalities,'PERIODO':period,'TABLA_REGISTROS':'\n'.join(value(work,r,'rit')+' — '+value(work,r,'nombre') for r in rows)}
        for text in (tpl['asunto'],tpl['cuerpo']):
            for _,key,_,_ in Formatter().parse(text):
                if key and not context.get(key,'').strip():raise ValueError('Completa '+key+' para esta plantilla.')
        to=resolve_contact(cfg,program) if to_program else '; '.join(court_cfg['para'])
        draft=Draft(tpl['asunto'].format_map(context),tpl['cuerpo'].format_map(context),to,
                    '; '.join(emails(CC+';'+cfg['correos'].get('cc_adicional',''))),required=tpl['adjunto']=='obligatorio')
        from nurus.rus.rules import as_date
        dates=[as_date(r.values.get(work.mapping.get('vencimiento',''))) for r in rows]
        dates=[d for d in dates if d]
        draft.program=program if to_program else ''
        draft.court=court_cfg['nombre'];draft.due=min(dates).isoformat() if dates else ''
        draft.kind=kind;draft.record_ids=[r.id for r in rows]
        if cfg.get('firma'):draft.body+='\n\n'+cfg['firma']
        if draft.required:
            folder=Path(directory or Path(work.output).parent);folder.mkdir(parents=True,exist_ok=True)
            by_program=defaultdict(list)
            for row in rows:by_program[value(work,row,'programa')].append(row)
            draft.attachments=[]
            for name,subset in by_program.items():
                target=folder/'adjuntos'/uuid4().hex
                target.mkdir(parents=True,exist_ok=True)
                draft.attachments.append(_table(work,subset,target/(program_filename(name)+'.xlsx'),kind))
        draft.key=draft_fingerprint(draft)
        output.append(draft)
    return output


def prepare_required_drafts(work,*,modalities='',period='',selected=None,directory=None,modality_keys=None,recipient_scope='todos'):
    """Prepara en una sola operación todas las comunicaciones que surgen del trabajo."""
    if recipient_scope not in {'todos','programas','tribunales'}:raise ValueError('Destino de correo inválido.')
    work.refresh()
    cfg=work.config
    templates=cfg['correos']['plantillas']
    mode=str(getattr(work,'mode','')).lower()
    kinds=[]
    if recipient_scope!='programas' and mode in {'espera','cumplimiento','informes'} and mode in templates:kinds.append(mode)
    selected_ids=set(selected) if selected is not None else None
    from .modalities import selected_row
    action_kinds=set()
    for row in work.rows:
        if row.excluded or (selected_ids is not None and row.id not in selected_ids) or not selected_row(work,row,modality_keys):continue
        action_kinds.update(action for action in row.actions if action in templates)
    for kind in ('programa_espera','programa_vencido','programa_por_vencer','medidas'):
        if kind in action_kinds and (recipient_scope!='tribunales' or kind=='medidas'):kinds.append(kind)
    drafts=[];seen=set()
    for kind in kinds:
        for draft in prepare_drafts(work,kind,modalities=modalities,period=period,selected=selected,directory=directory,manual_selection=False,modality_keys=modality_keys,recipient_scope='auto' if recipient_scope=='todos' else recipient_scope):
            if draft.key not in seen:
                drafts.append(draft);seen.add(draft.key)
    return drafts


def program_filename(name):
    result=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',str(name)).strip().rstrip('. ')
    if not result:result='Programa no informado'
    if re.fullmatch(r'(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])',result.split('.')[0]):result='_'+result
    return result[:120].rstrip('. ')


def create_drafts(work,drafts):
    """Un clic guarda el lote. Los resultados parciales se conservan sin repetir Save."""
    result={'created':0,'skipped':0,'errors':[]}
    for draft in drafts:
        try:
            draft.key=draft_fingerprint(draft)
            if draft.key in work.receipts:
                state=work.receipts[draft.key].get('state')
                result['skipped']+=1
                if state!='created':result['errors'].append(draft.subject+': guardado previo incierto; revisa Borradores.')
                continue
            create_draft(work,draft,confirmed=True);result['created']+=1
        except Exception as exc:result['errors'].append(draft.subject+': '+str(exc))
    return result


def create_draft(work,draft,*,confirmed=False):
    if not confirmed:raise ValueError('Revisa destinatarios, texto y adjuntos antes de crear el borrador.')
    draft.key=draft_fingerprint(draft)
    if draft.key in work.receipts:raise ValueError('Este borrador ya se creó o su guardado quedó incierto; revisa Outlook antes de repetir.')
    to='; '.join(emails(draft.to));cc='; '.join(emails(CC+';'+draft.cc))
    if draft.required and not draft.attachments:raise ValueError('Este correo requiere adjunto.')
    attachments=[]
    for file in draft.attachments:
        path=Path(file).resolve()
        if not path.is_file():raise ValueError('Adjunto no disponible: '+str(path))
        attachments.append((str(path),sha256(path.read_bytes()).hexdigest()))
    tpl=Template('personal','Correo revisado',ProductKind.EMAIL,'','',())
    product=Product(ProductKind.EMAIL,tpl,{},recipient=to,cc=cc,attachments=attachments,required_attachment=draft.required,
                    status=ProductStatus.READY,rendered_subject=draft.subject,rendered_body=draft.body)
    work.receipts[draft.key]={'kind':'draft','state':'saving'}
    if getattr(work,'storage_directory',None):work.save(work.storage_directory)
    try:
        receipt=save_draft(product,confirmed=True,account_key=work.config.get('cuenta_outlook') or None,preserve_signature=True)
    except DraftSaveUncertain:
        work.receipts[draft.key]={'kind':'draft','state':'uncertain'}
        if getattr(work,'storage_directory',None):work.save(work.storage_directory)
        raise
    except Exception:
        work.receipts.pop(draft.key,None)
        if getattr(work,'storage_directory',None):work.save(work.storage_directory)
        raise
    work.receipts[draft.key]={'kind':'draft','state':'created','entry_id':receipt.entry_id,'store_id':receipt.store_id}
    if getattr(work,'storage_directory',None):work.save(work.storage_directory)
    return receipt


def fill_docx(template,destination,values):
    """Sustituye tokens entre runs sin reconstruir párrafos ni eliminar estilos."""
    from zipfile import ZipFile, ZIP_DEFLATED
    from lxml import etree as ET
    token=re.compile(r'\{\{([A-Z_]+)\}\}')
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(template) as source:
        content={item.filename:source.read(item.filename) for item in source.infolist()}
    for name,data in list(content.items()):
        if not name.startswith('word/') or not name.endswith('.xml'):continue
        root=ET.fromstring(data,parser=ET.XMLParser(resolve_entities=False,no_network=True))
        changed=False
        for paragraph in root.findall('.//w:p',ns):
            nodes=paragraph.findall('.//w:t',ns)
            full=''.join(n.text or '' for n in nodes)
            for match in reversed(list(token.finditer(full))):
                key=match.group(1)
                replacement=str(values.get(key,'')).strip()
                if not replacement:raise ValueError('Falta la variable Word: '+key)
                pos=0;parts=[]
                for node in nodes:
                    length=len(node.text or '');parts.append((node,pos,pos+length));pos+=length
                affected=[(n,a,b) for n,a,b in parts if a<match.end() and b>match.start()]
                first,a,b=affected[0];last,la,lb=affected[-1]
                prefix=(first.text or '')[:match.start()-a]
                suffix=(last.text or '')[match.end()-la:]
                first.text=prefix+replacement+(suffix if first is last else '')
                first.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                for node,_,_ in affected[1:]:node.text=suffix if node is last else ''
                changed=True
        if changed:content[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    def write(path):
        with ZipFile(path,'w',ZIP_DEFLATED) as out:
            for name,data in content.items():out.writestr(name,data)
    write_new_file(Path(destination),write)
    return str(destination)


def word_values(work,row):
    from nurus.rus.rules import as_date
    durations=[str(v) for k,v in row.values.items() if normalize(k) in ('duracion','plazo','vigencia')]
    vals={key.upper():value(work,row,key) for key in ('rit','rut','nombre','programa')}
    resolution=as_date(row.values.get(work.mapping.get('resolucion','')))
    vals.update(FECHA=date_in_words(date.today()),FECHA_RESOLUCION=date_in_words(resolution) if resolution else '',DURACION=durations[0] if len(durations)==1 else '')
    return vals


def template_variables(path):
    from zipfile import ZipFile
    from lxml import etree as ET
    result=set()
    with ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith('word/') and name.endswith('.xml'):
                root=ET.fromstring(z.read(name),parser=ET.XMLParser(resolve_entities=False,no_network=True))
                for p in root.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                    text=''.join(p.itertext())
                    result.update(re.findall(r'\{\{([A-Z_]+)\}\}',text))
    return result
