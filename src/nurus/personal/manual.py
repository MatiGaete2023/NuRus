"""Correo y Word manuales independientes de cualquier Excel."""
from dataclasses import asdict
from pathlib import Path
from copy import deepcopy
import json
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from .config import atomic_json, emails
from .outputs import Draft, create_draft, fill_docx, template_variables
from .current_tools import confirm_drafts
from nurus.services.file_output import write_new_file


class ManualContext:
    def __init__(self,config,path):
        self.config=deepcopy(config);self.path=Path(path);self.storage_directory=str(self.path.parent)
        self.rows=[];self.mapping={};self.output='';self.receipts={};self.current={'mail':{},'word':{}}
        if self.path.is_file():
            data=json.loads(self.path.read_text(encoding='utf-8'));self.receipts=data.get('receipts',{});self.current.update(data.get('current',{}))
    def save(self,directory=None):
        atomic_json(self.path,{'version':1,'current':self.current,'receipts':self.receipts})


def generate_word(destination,title,body,template='',values=None):
    if template:
        needed=template_variables(template);values=values or {}
        missing=[key for key in needed if not str(values.get(key,'')).strip()]
        if missing:raise ValueError('Completa los campos de la matriz: '+', '.join(sorted(missing)))
        fill_docx(template,destination,values)
    else:
        if not body.strip():raise ValueError('Escribe el contenido del documento.')
        from docx import Document
        doc=Document()
        if title.strip():doc.add_heading(title,0)
        for paragraph in body.split('\n'):doc.add_paragraph(paragraph)
        from .word_quality import require_complete,improve_pagination
        require_complete(doc);improve_pagination(doc)
        write_new_file(Path(destination),doc.save)
    return str(destination)


def open_editor(app,kind='mail'):
    windows=getattr(app,'_manual_windows',{})
    previous=windows.get(kind)
    if previous is not None and previous.winfo_exists():previous.lift();previous.focus_force();return
    app._manual_windows=windows
    context=ManualContext(app.cfg.data,app.cfg.directory/'manual_actual.json');saved=context.current[kind]
    win=tk.Toplevel(app);win.title('Correo manual' if kind=='mail' else 'Word manual · libre o con matriz');win.geometry('820x650');win.transient(app)
    windows[kind]=win
    fields={}
    names=['to','cc','subject'] if kind=='mail' else ['title']
    labels={'to':'Para','cc':'CC','subject':'Asunto','title':'Título'}
    bar=ttk.Frame(win);bar.pack(fill='x',padx=12,pady=10)
    for row,name in enumerate(names):
        fields[name]=tk.StringVar(value=saved.get(name,''));ttk.Label(bar,text=labels[name]).grid(row=row,column=0,sticky='w');ttk.Entry(bar,textvariable=fields[name]).grid(row=row,column=1,sticky='ew',padx=8,pady=3)
    bar.columnconfigure(1,weight=1)
    text=tk.Text(win,wrap='word',undo=True);text.pack(fill='both',expand=True,padx=12);text.insert('1.0',saved.get('body',''))
    attachments=list(saved.get('attachments',[]));template=tk.StringVar(value=saved.get('template',''));values=dict(saved.get('values',{}))
    info=tk.StringVar();ttk.Label(win,textvariable=info,wraplength=780).pack(fill='x',padx=12,pady=6)
    actions=ttk.Frame(win);actions.pack(fill='x',padx=12,pady=10)
    def capture():
        state={name:var.get() for name,var in fields.items()};state.update(body=text.get('1.0','end-1c'),attachments=attachments,template=template.get(),values=values)
        context.current[kind]=state;context.save();return state
    if not hasattr(app,'_manual_captures'):app._manual_captures={}
    app._manual_captures[kind]=(win,capture)
    def guard(fn):
        try:fn()
        except Exception as exc:messagebox.showerror('Producto manual',str(exc),parent=win)
    def close():
        if app.busy:raise ValueError('Espera a que termine la operación antes de cerrar el editor.')
        capture();win.destroy()
    win.protocol('WM_DELETE_WINDOW',lambda:guard(close))
    def favorite(save=False):
        path=app.cfg.directory/'bloques_favoritos.json'
        data=json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}
        if save:
            name=simpledialog.askstring('Guardar bloque favorito','Nombre del bloque de texto:',parent=win)
            if name and name.strip():
                block=text.get('1.0','end-1c');block_variables(block);data[name.strip()]=block;atomic_json(path,data)
        else:
            chooser=tk.Toplevel(win);chooser.title('Insertar bloque favorito');choice=tk.StringVar()
            ttk.Combobox(chooser,textvariable=choice,values=sorted(data),state='readonly').pack(padx=15,pady=15)
            def insert():
                if choice.get() in data:
                    block=data[choice.get()];params={}
                    for key in sorted(block_variables(block)):
                        answer=simpledialog.askstring('Completar bloque',key,parent=chooser)
                        if answer is None:return
                        params[key]=answer
                    text.insert('insert',render_block(block,params));chooser.destroy()
            ttk.Button(chooser,text='Insertar',command=insert).pack(pady=10)
    ttk.Button(actions,text='Guardar texto favorito',command=lambda:guard(lambda:favorite(True))).pack(side='left')
    ttk.Button(actions,text='Insertar favorito',command=lambda:guard(favorite)).pack(side='left',padx=4)
    if kind=='mail':
        def attach():
            chosen=filedialog.askopenfilenames(parent=win)
            attachments.extend(p for p in chosen if p not in attachments);info.set('\n'.join(attachments))
        def save_mail():
            state=capture();draft=Draft(state['subject'],state['body'],to=state['to'],cc=state['cc'],attachments=list(attachments))
            if not emails(draft.to):raise ValueError('Indica al menos un destinatario.')
            if confirm_drafts(win,[draft]):
                app._run('Guardando correo manual en Outlook…',lambda:create_draft(context,draft,confirmed=True),
                         lambda result:info.set('Borrador guardado en Outlook: '+result.entry_id) if win.winfo_exists() else None)
        ttk.Button(actions,text='Adjuntar…',command=attach).pack(side='left',padx=4)
        ttk.Button(actions,text='Quitar adjuntos',command=lambda:(attachments.clear(),info.set(''))).pack(side='left')
        ttk.Button(actions,text='Revisar y guardar en Outlook',command=lambda:guard(save_mail)).pack(side='right')
    else:
        def matrix():
            path=filedialog.askopenfilename(parent=win,initialdir=app.template_dir,filetypes=[('Matriz Word','*.docx')])
            if not path:return
            needed=sorted(template_variables(path));candidate={}
            for key in needed:
                answer=simpledialog.askstring('Datos de la matriz',key,initialvalue=values.get(key,''),parent=win)
                if answer is None:return
                candidate[key]=answer
            template.set(path);values.clear();values.update(candidate);info.set('Matriz: '+Path(path).name+' · '+str(len(needed))+' campos. El texto libre se utiliza sólo sin matriz.');capture()
        def save_word():
            state=capture();path=filedialog.asksaveasfilename(parent=win,defaultextension='.docx',initialfile='Documento_manual.docx')
            if path:generate_word(path,state['title'],state['body'],template.get(),values);info.set('Word generado: '+path)
        ttk.Button(actions,text='Elegir matriz y datos…',command=lambda:guard(matrix)).pack(side='left',padx=4)
        ttk.Button(actions,text='Usar texto libre',command=lambda:(template.set(''),values.clear(),info.set('Texto libre'))).pack(side='left')
        ttk.Button(actions,text='Generar Word',command=lambda:guard(save_word)).pack(side='right')
    def autosave():
        if not win.winfo_exists():return
        try:
            if not app.busy:capture()
        except OSError as exc:info.set('No se pudo guardar la recuperación: '+str(exc))
        win.after(15000,autosave)
    win.after(15000,autosave)


BLOCK_FIELDS={'NOMBRE','RIT','TRIBUNAL','PROGRAMA','FECHA'}


def block_variables(text):
    from .template_validation import variables
    return variables(text,BLOCK_FIELDS)


def render_block(text,values):
    needed=block_variables(text)
    if any(not str(values.get(key,'')).strip() for key in needed):raise ValueError('Completa todas las variables del bloque favorito.')
    return text.format_map(values)
