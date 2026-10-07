"""Informes independientes de varios trabajos; historial y gestiones separados."""
from datetime import date, datetime
from pathlib import Path

from .review_store import identity
from .rus_activity import attach
from .outputs import value
from nurus.rus.rules import as_date, as_int
from nurus.services.file_output import write_new_file


def _book(sheets, destination):
    from openpyxl import Workbook
    from openpyxl.styles import Border, Side, PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();book.remove(book.active)
    edge=Side(style='thin',color='000000')
    try:
        for name,headers,rows in sheets:
            sheet=book.create_sheet(name);sheet.append(headers)
            for row in rows:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value,str):cell.data_type='s'
            for cells in sheet:
                for cell in cells:
                    cell.border=Border(left=edge,right=edge,top=edge,bottom=edge)
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
            for cell in sheet[1]:
                cell.fill=PatternFill('solid',fgColor='1F538D');cell.font=Font(color='FFFFFF',bold=True)
            for number,title in enumerate(headers,1):
                sheet.column_dimensions[get_column_letter(number)].width=60 if title in ('Texto','Detalle','Trámite') else 25
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        write_new_file(destination,book.save)
    finally:book.close()
    return str(Path(destination))


def period(start, end):
    if (not isinstance(start,date) or isinstance(start,datetime) or
            not isinstance(end,date) or isinstance(end,datetime) or start>end):
        raise ValueError('El período del informe no es válido.')


def export_activity(works,destination,start,end):
    """Una fila por firma e ingreso dentro del período; cobertura siempre explícita."""
    period(start,end)
    found={};scopes={}
    for work in works:
        attach(work)
        for code,scope in getattr(work,'activity_sources',{}).get('cobertura',{}).items():
            source_start=as_date(scope.get('desde'));source_end=as_date(scope.get('hasta'))
            overlaps=bool(source_start and source_end and source_start<=end and source_end>=start)
            fully_covers=bool(source_start and source_end and source_start<=start and source_end>=end)
            original_state=str(scope.get('estado','NO_CONSULTADA'))
            if original_state in ('COMPLETA','VACIA_COMPROBADA') and not fully_covers:
                state='PARCIAL' if overlaps else 'NO_CONSULTADA'
                reason=('La consulta de origen no cubre todo el período solicitado.' if overlaps else
                        'El período solicitado queda fuera de la consulta de origen.')
            else:
                state=original_state
                reason='' if fully_covers else 'La cobertura de origen no acredita todo el período solicitado.'
            clipped_start=max(source_start,start) if overlaps else ''
            clipped_end=min(source_end,end) if overlaps else ''
            scopes[(work.source_hash,code)]=[
                Path(work.path).name,code,start.isoformat(),end.isoformat(),state,
                str(scope.get('criterio_temporal','')),source_start.isoformat() if source_start else '',
                source_end.isoformat() if source_end else '',clipped_start.isoformat() if clipped_start else '',
                clipped_end.isoformat() if clipped_end else '',reason]
        for row in work.rows:
            key,kind=identity(work,row)
            for signature in work.signed_activity.get(row.id,{}).get('firmas',[]):
                signed=as_date(signature.get('firma'))
                if signed is None or not start<=signed<=end:continue
                reviewed=work.activity_reviewed.get(row.id,{}).get(signature['huella'],'')
                item=[value(work,row,'tribunal'),value(work,row,'rit'),value(work,row,'nombre'),
                      value(work,row,'programa'),signature['firma'],signature['hora'],signature['tramite'],
                      'Revisada para este ingreso' if reviewed else 'Por revisar',reviewed,
                      signature['identidad'],kind,Path(work.path).name]
                entry=(key,signature['huella'])
                previous=found.get(entry)
                # Una sesión anterior no puede borrar una revisión posterior comprobada.
                if previous is None or reviewed>previous[8]:found[entry]=item
    rows=list(found.values())
    return _book([
        ('Resumen',('Categoría','Cantidad'),[
            ['Período solicitado desde',start.isoformat()],['Período solicitado hasta',end.isoformat()],
            ['Firmas por ingreso detectadas',len(rows)],['Revisadas para el ingreso',sum(bool(r[8]) for r in rows)],
            ['Por revisar',sum(not r[8] for r in rows)],['Fuentes de cobertura',len(scopes)]]),
        ('Firmas por ingreso',('Tribunal','RIT','Persona','Centro','Fecha firma','Hora firma','Trámite',
             'Estado de revisión','Fecha de revisión local','Identidad de firma','Alcance de vínculo','Fuente'),rows),
        ('Cobertura',('Fuente','Tribunal código','Desde pedido','Hasta pedido','Estado','Criterio temporal',
            'Desde fuente','Hasta fuente','Desde consultado en alcance','Hasta consultado en alcance','Detalle'),list(scopes.values()))],destination)


def export_management(works,destination,start,end):
    """Solo recibos de guardado releído cuentan como observaciones nuevas."""
    period(start,end);historical=[];operations={};existing={};products={};seen_sources=set()
    for work in works:
        source_key=(work.source_hash,work.sheet)
        if source_key not in seen_sources:
            seen_sources.add(source_key)
            for row in work.rows:
                review=row.review or {};day=as_date(review.get('FECHA_OBS'))
                if not day or not start<=day<=end or not review.get('OBSERVACION'):continue
                historical.append([value(work,row,'tribunal'),value(work,row,'rit'),value(work,row,'nombre'),
                    value(work,row,'programa'),day.isoformat(),as_int(review.get('TT')),as_int(review.get('CC')),
                    review['OBSERVACION'],'Constancia del Excel; no acredita una gestión nueva',Path(work.path).name])
        for key,item in work.receipts.items():
            if item.get('kind')=='rus_observation':
                if item.get('state') not in ('COMPROBADA','PENDIENTE_EXCEL') or item.get('verified') is not True:continue
                operation=item.get('operation_id');registered=item.get('registered_at')
                if not operation or not registered or not item.get('remote_entry_id'):continue
                try:day=datetime.fromisoformat(registered).date()
                except (TypeError,ValueError):continue
                if not start<=day<=end:continue
                from .bitacoras import cc_for_type
                charge=cc_for_type(item.get('type',''))
                if item.get('cc')!=charge:raise ValueError('El recibo contiene una carga incompatible con el tipo de observación.')
                data=[operation,item.get('tribunal',''),item.get('rit',''),item.get('ingreso_id',''),
                      registered,item.get('author',''),item.get('type',''),item.get('cc'),item.get('text',''),item['remote_entry_id']]
                category=operations if item.get('new_registration') is True else existing
                if operation in category and category[operation]!=data:
                    raise ValueError('Los trabajos contienen recibos contradictorios para una misma operación.')
                if operation in (existing if category is operations else operations):
                    raise ValueError('Un recibo no puede ser simultáneamente nuevo y anterior.')
                category[operation]=data
            else:
                try:day=datetime.fromisoformat(item.get('created_at','')).date()
                except (TypeError,ValueError):continue
                if start<=day<=end:
                    product=[item.get('kind',''),item.get('state',''),item.get('created_at',''),item.get('path',''),
                        'Producto local; no acredita registro en RUS']
                    if key in products and products[key]!=product:raise ValueError('Hay recibos de producto contradictorios.')
                    products[key]=product
    new=list(operations.values())
    return _book([
        ('Resumen',('Categoría','Cantidad'),[
            ['Desde',start.isoformat()],['Hasta',end.isoformat()],['Observaciones nuevas comprobadas',len(new)],
            ['Causas con gestiones nuevas',len({(r[1],r[2]) for r in new})],
            ['Ingresos con gestiones nuevas',len({(r[1],r[3]) for r in new})],
            ['Nuevas con carga',sum(r[7]==1 for r in new)],['Nuevas sin carga',sum(r[7]==0 for r in new)],
            ['Constancias históricas del Excel',len(historical)],['Productos locales',len(products)],
            ['Entradas RUS anteriores o sin atribución de gestión nueva',len(existing)]]),
        ('Nuevas comprobadas',('Operación','Tribunal','RIT','Ingreso RUS','Fecha efectiva','Autor','Tipo','CC','Texto','Entrada RUS'),new),
        ('Entradas ya existentes',('Operación','Tribunal','RIT','Ingreso RUS','Fecha efectiva','Autor','Tipo','CC','Texto','Entrada RUS'),list(existing.values())),
        ('Historial del Excel',('Tribunal','RIT','Persona','Centro','Fecha','TT','CC','Texto','Interpretación','Fuente'),historical),
        ('Productos',('Tipo','Estado','Fecha','Archivo','Interpretación'),list(products.values()))],destination)


def main():
    import argparse
    from .work import Work
    parser=argparse.ArgumentParser(description='Informes CSMP sin conexión a RUS.')
    parser.add_argument('tipo',choices=['firmas','gestion'])
    parser.add_argument('--trabajo',action='append',required=True,help='Carpeta de una sesión CSMP guardada.')
    parser.add_argument('--salida',required=True)
    parser.add_argument('--desde',type=date.fromisoformat,required=True)
    parser.add_argument('--hasta',type=date.fromisoformat,required=True)
    args=parser.parse_args()
    try:period(args.desde,args.hasta)
    except ValueError as exc:parser.error(str(exc))
    works=[Work.load(Path(path)) for path in args.trabajo]
    if args.tipo=='firmas':export_activity(works,args.salida,args.desde,args.hasta)
    else:export_management(works,args.salida,args.desde,args.hasta)


if __name__=='__main__':main()
