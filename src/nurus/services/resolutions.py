"""Proyectos editables desde constancia; no constituyen resolución emitida ni firmada."""
from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import re

from nurus.domain.models import Product, ProductKind, ProductStatus
from nurus.rus.columns import normalize
from nurus.rus.rules import format_date, tribunal
from nurus.services.products import _published_template
from nurus.services.rendering import prepare


def prepare_resolution(db, batch_id, record_id, template_id, *, confirmed_review=False, extra=None):
    if not confirmed_review:
        raise ValueError("Verifica la carpeta judicial y confirma la procedencia del proyecto.")
    snapshot = db.get_snapshot(batch_id)
    if not snapshot["batch"].get("review_import_hash"):
        raise ValueError("El proyecto requiere una constancia importada de la revisión en RUS.")
    rows = [
        row for row in snapshot["records"]
        if row["record_id"] == record_id and row["decision"] == "approved" and row.get("rus_recorded")
    ]
    if len(rows) != 1:
        raise ValueError("Selecciona una fila revisada de la constancia.")
    template = _published_template(db, template_id)
    if template.kind is not ProductKind.RESOLUTION or not template_id.startswith("historica-"):
        raise ValueError("Selecciona una matriz histórica identificada.")
    mapping = json.loads(snapshot["batch"]["column_mapping"])
    values = json.loads(rows[0]["values_json"])
    court = tribunal(values.get(mapping.get("tribunal", ""), ""))
    if not court or not template_id.startswith("historica-" + court.lower() + "-"):
        raise ValueError("La matriz no corresponde al tribunal; no se sustituye silenciosamente.")
    context = {
        key.upper(): str(values.get(mapping.get(key, ""), "") or "")
        for key in ("rit", "rut", "nombre", "programa")
    }
    context["FECHA"] = format_date(date.fromisoformat(rows[0]["review_date"]))
    context["FECHA_RESOLUCION"] = str(values.get(mapping.get("resolucion", ""), "") or "")
    durations = [value for key, value in values.items() if normalize(key) in {"duracion", "plazo", "vigencia"}]
    if len(durations) > 1:
        raise ValueError("Duración ambigua: revisa la fuente antes de preparar el proyecto.")
    context["DURACION"] = str(durations[0] or "") if durations else ""
    for key, value in (extra or {}).items():
        if key not in {"NOMENCLATURA", "FECHA", "FECHA_RESOLUCION", "DURACION"}:
            raise ValueError("Campo adicional de resolución no permitido.")
        context[key] = str(value)
    return prepare(Product(
        kind=template.kind,
        template=template,
        context=context,
        batch_id=batch_id,
        source_snapshot_hash=snapshot["batch"]["snapshot_hash"],
        record_ids=(record_id,),
    ))


def export_resolution(db, product, destination):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    if product.kind is not ProductKind.RESOLUTION or product.status is not ProductStatus.APPROVED:
        raise ValueError("Aprueba y guarda el proyecto antes de producir el Word.")
    with db.connect() as conn:
        row = conn.execute(
            """SELECT p.body,p.source_snapshot_hash,b.status AS batch_status,
                      b.snapshot_hash AS current_snapshot_hash
               FROM products p JOIN batches b ON b.id=p.batch_id
               WHERE p.id=? AND p.status='approved'""",
            (product.id,),
        ).fetchone()
    if not row or row[0] != product.rendered_body or row[1] != product.source_snapshot_hash:
        raise ValueError("El texto no coincide con el producto aprobado.")
    if row[2] != "approved" or row[1] != row[3]:
        raise ValueError(
            "La constancia cambió desde la aprobación del proyecto. Prepara una nueva versión."
        )
    db.get_snapshot(product.batch_id, product.source_snapshot_hash)
    path = Path(destination).resolve()
    if path.suffix.lower() != ".docx" or path.exists():
        raise ValueError("Elige un archivo .docx nuevo; no se sobrescriben documentos.")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1.1)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for line in product.rendered_body.split("\n"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for segment in re.split(r"(\*\*.+?\*\*|\*[^*]+\*)", line):
            bold = segment.startswith("**") and segment.endswith("**")
            italic = not bold and segment.startswith("*") and segment.endswith("*")
            run = paragraph.add_run(segment[2:-2] if bold else segment[1:-1] if italic else segment)
            run.bold, run.italic = bold, italic
    document.core_properties.identifier = product.id
    document.core_properties.comments = "Proyecto no firmado. Snapshot: " + product.source_snapshot_hash
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as output:
        document.save(output)
    digest = sha256(path.read_bytes()).hexdigest()
    db.record_artifact(product.batch_id, product.source_snapshot_hash, str(path), digest, "PROYECTO_DOCX")
    return path
