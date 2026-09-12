"""Estadísticas reconciliables con la constancia, sin reinterpretar códigos administrativos."""
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from nurus.services.file_output import excel_text, write_new_file


@dataclass(frozen=True)
class ReviewStatistics:
    start: date
    end: date
    total_records: int
    reviewed_in_period: int
    excluded: int
    outside_period: int
    missing_or_invalid_date: int
    observations_changed: int
    tt_values: tuple[tuple[str, int], ...]
    workload_values: tuple[tuple[str, int], ...]
    resolution_values: tuple[tuple[str, int], ...]
    unconfirmed: int = 0


def _distribution(values):
    counts = Counter(str(value or "").strip() or "(vacío)" for value in values)
    return tuple(sorted(counts.items(), key=lambda item: (-item[1], item[0])))


def build_review_statistics(db, batch_id: str, start: date, end: date) -> ReviewStatistics:
    if end < start:
        raise ValueError("La fecha final debe ser igual o posterior a la inicial.")
    snapshot = db.get_snapshot(batch_id)
    if not snapshot["batch"].get("review_import_hash"):
        raise ValueError("Las estadísticas requieren una constancia importada.")
    included, outside, missing, excluded, unconfirmed = [], 0, 0, 0, 0
    for row in snapshot["records"]:
        if row["decision"] == "excluded":
            excluded += 1
            continue
        if row["decision"] != "approved" or not row.get("rus_recorded"):
            unconfirmed += 1
            continue
        try:
            reviewed = date.fromisoformat(str(row.get("review_date", "")))
        except ValueError:
            missing += 1
            continue
        if start <= reviewed <= end:
            included.append(row)
        else:
            outside += 1
    return ReviewStatistics(
        start=start,
        end=end,
        total_records=len(snapshot["records"]),
        reviewed_in_period=len(included),
        excluded=excluded,
        outside_period=outside,
        missing_or_invalid_date=missing,
        unconfirmed=unconfirmed,
        observations_changed=sum(
            row["edited_observation"] != row["original_observation"] for row in included
        ),
        tt_values=_distribution(row.get("tt_value", "") for row in included),
        workload_values=_distribution(row.get("workload_value", "") for row in included),
        resolution_values=_distribution(row.get("resolution_value", "") for row in included),
    )


def export_review_statistics(stats: ReviewStatistics, destination):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = Path(destination).resolve()
    if path.suffix.lower() != ".xlsx" or path.exists():
        raise ValueError("Selecciona un archivo .xlsx nuevo.")
    book = Workbook()
    summary = book.active
    summary.title = "Resumen"
    for row in (
        ("Desde", stats.start.isoformat()), ("Hasta", stats.end.isoformat()),
        ("Registros en constancia", stats.total_records),
        ("Revisados en período", stats.reviewed_in_period),
        ("Excluidos", stats.excluded), ("Fuera de período", stats.outside_period),
        ("Fecha ausente o inválida", stats.missing_or_invalid_date),
        ("Sin confirmación de registro en RUS", stats.unconfirmed),
        ("Observaciones modificadas", stats.observations_changed),
    ):
        summary.append(row)
    for name, title, values in (
        ("TT", "Valor TT", stats.tt_values),
        ("CC", "Valor CC/carga", stats.workload_values),
        ("RES", "Valor RES", stats.resolution_values),
    ):
        sheet = book.create_sheet(name)
        sheet.append([title, "Cantidad"])
        for value, count in values:
            sheet.append([excel_text(value), count])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 36
    try:
        write_new_file(path, book.save)
    finally:
        book.close()
    return path
