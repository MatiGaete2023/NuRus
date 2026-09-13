"""Agrupación explícita de comunicaciones desde una constancia congelada."""
from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from nurus.domain.models import Product
from nurus.rus.columns import normalize
from nurus.rus.rules import format_date, is_dce
from nurus.services.contacts import resolve_contact_exact
from nurus.services.products import _published_template
from nurus.services.rendering import prepare
from nurus.services.file_output import excel_text, write_new_file


def prepare_communications(
    db,
    batch_id,
    template_id,
    *,
    modalities: str,
    period: str = "",
    selected_record_ids: tuple[str, ...] | list[str] | None = None,
):
    policy = db.get_communication_policy(template_id)
    snapshot = db.get_snapshot(batch_id)
    batch = snapshot["batch"]
    if not batch.get("review_import_hash"):
        raise ValueError("Los correos requieren una constancia importada de la revisión registrada en RUS.")
    if policy["mode"] and policy["mode"] != batch["mode"]:
        raise ValueError("El tipo de correo no corresponde a la modalidad analizada.")
    if policy["filter"] == "manual":
        raise ValueError("Este correo requiere preparación particular y selección manual de adjuntos.")
    selected = set(selected_record_ids or ())
    if selected_record_ids is not None and not selected:
        return []
    mapping = json.loads(batch["column_mapping"])
    groups = defaultdict(list)
    known = {row["record_id"] for row in snapshot["records"]}
    if selected - known:
        raise ValueError("La selección contiene registros ajenos a la constancia.")
    for row in snapshot["records"]:
        if row["decision"] != "approved" or not row.get("rus_recorded"):
            continue
        rules = set(json.loads(row["rule_ids_json"]))
        chosen = policy["filter"]
        eligible = (
            chosen == "all"
            or (chosen == "medidas" and bool({"C-04", "C-05"}.intersection(rules)))
            or (chosen.startswith(("E-", "I-")) and chosen in rules)
        )
        if not eligible or (selected and row["record_id"] not in selected):
            continue
        values = json.loads(row["values_json"])
        court = str(values.get(mapping.get("tribunal", ""), "")).strip()
        program = str(values.get(mapping.get("programa", ""), "")).strip()
        if not court or (policy["group"] == "programa" and not program):
            raise ValueError("Falta tribunal o programa para agrupar sin ambigüedad.")
        key = (normalize(court), normalize(program) if policy["group"] == "programa" else "")
        groups[key].append((row, values, court, program))

    template = _published_template(db, template_id)
    products = []
    for group in groups.values():
        _, _, court, program = group[0]
        contact = resolve_contact_exact(db, program if policy["group"] == "programa" else court)
        lines = ["RIT | Tribunal | RUT | Nombre | Programa | Fecha vencimiento / Días espera"]
        for row, values, _, _ in group:
            parts = [
                str(values.get(mapping.get(key, ""), "") or "")
                for key in ("rit", "tribunal", "rut", "nombre", "programa")
            ]
            parts.append(str(values.get(mapping.get("vencimiento", mapping.get("espera", "")), "") or ""))
            lines.append(" | ".join(parts))
        context = {
            "TRIBUNAL": court,
            "PROGRAMA": program,
            "ALCANCE_MODALIDADES": modalities,
            "FECHA": format_date(date.fromisoformat(batch["as_of"])),
            "PERIODO": period,
            "TABLA_REGISTROS": "\n".join(lines),
            "ETIQUETA_INFORMES": "informes diagnósticos" if is_dce(program) else "informes de avance",
            "NURUS_POLITICA": json.dumps(policy, ensure_ascii=False, sort_keys=True),
        }
        products.append(prepare(Product(
            kind=template.kind,
            template=template,
            context=context,
            recipient=contact["email"] if contact else "",
            cc=contact["cc"] if contact else "",
            batch_id=batch_id,
            source_snapshot_hash=batch["snapshot_hash"],
            record_ids=tuple(row["record_id"] for row, *_ in group),
            required_attachment=bool(policy["required_attachment"]),
        )))
    return products


def attach_snapshot_table(db, product, directory):
    """Genera solo el subconjunto confirmado; nunca adjunta el libro completo."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    snapshot = db.get_snapshot(product.batch_id, product.source_snapshot_hash)
    selected = set(product.record_ids)
    records = [
        row for row in snapshot["records"]
        if row["record_id"] in selected and row["decision"] == "approved" and row.get("rus_recorded")
    ]
    if not records or len(records) != len(selected):
        raise ValueError("Selección de adjunto vacía o ajena a la constancia.")
    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"Nomina_{product.id}.xlsx"
    if path.exists():
        raise ValueError("El adjunto ya existe. No se sobrescribe.")
    book = Workbook()
    sheet = book.active
    sheet.title = "Nomina"
    mapping = json.loads(snapshot["batch"]["column_mapping"])
    keys = ("rit", "tribunal", "rut", "nombre", "programa", "vencimiento", "espera")
    sheet.append(["RIT", "TRIBUNAL", "RUT", "NOMBRE", "PROGRAMA", "FECHA_VENCIMIENTO", "DIAS_ESPERA"])
    for row in records:
        values = json.loads(row["values_json"])
        sheet.append([excel_text(values.get(mapping.get(key, ""), "")) for key in keys])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column, width in {"A":18,"B":32,"C":16,"D":40,"E":50,"F":24,"G":16}.items():
        sheet.column_dimensions[column].width = width
    try:
        write_new_file(path, book.save)
    finally:
        book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    db.record_artifact(product.batch_id, product.source_snapshot_hash, str(path), digest, "ADJUNTO_NOMINA")
    product.attachments.append((str(path), digest))
    return path
