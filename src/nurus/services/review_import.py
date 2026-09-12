from __future__ import annotations

import hashlib
import math
from dataclasses import dataclass, field
from io import BytesIO
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from nurus.rus.columns import normalize
from nurus.storage.database import Database


class ReviewImportError(ValueError):
    pass


@dataclass(frozen=True)
class ReviewImportPreview:
    path: Path
    sha256: str
    updates: tuple[dict[str, object], ...]
    reviewed_count: int
    excluded_count: int
    changed_observations: int
    errors: tuple[str, ...]
    warnings: tuple[str, ...]
    source_bytes: bytes = field(repr=False, compare=False, default=b"")

    @property
    def valid(self) -> bool:
        return not self.errors


MANAGED_HEADERS = (
    "NURUS_ID_REGISTRO",
    "OBSERVACION",
    "FECHA_OBS",
    "TT",
    "CC",
    "RES",
    "NURUS_ESTADO_REVISION",
)


def _blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and not value.strip()


def _text(value: object) -> str:
    return "" if _blank(value) else str(value).strip()


def _review_date(value: object) -> str:
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = _text(value)
        if not text:
            raise ValueError("falta FECHA_OBS")
        if text.startswith("=") or text.startswith("#"):
            raise ValueError("FECHA_OBS contiene una fórmula o error de Excel")
        parsed = None
        for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                parsed = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                pass
        if parsed is None:
            raise ValueError(f"FECHA_OBS no es una fecha válida: {text!r}")
    if parsed > date.today():
        raise ValueError("FECHA_OBS no puede ser futura")
    return parsed.isoformat()


def _positions(headers: Iterable[object]) -> tuple[dict[str, int], list[str]]:
    found: dict[str, list[int]] = {}
    for index, value in enumerate(headers):
        key = normalize(value)
        if key:
            found.setdefault(key, []).append(index)
    errors: list[str] = []
    positions: dict[str, int] = {}
    for title in MANAGED_HEADERS:
        candidates = found.get(normalize(title), [])
        if not candidates:
            errors.append(f"Falta la columna administrada {title}.")
        elif len(candidates) > 1:
            errors.append(f"La columna {title} está repetida.")
        else:
            positions[title] = candidates[0]
    return positions, errors


def _xlsx_rows(path: Path, sheet_name: str, header_row: int) -> tuple[list[object], list[list[object]], dict[str, str]]:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    try:
        if sheet_name not in book.sheetnames:
            raise ReviewImportError(f"El archivo no contiene la hoja procesada {sheet_name!r}.")
        sheet = book[sheet_name]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=header_row + 1)]
        trace_values: dict[str, str] = {}
        trace_names = [name for name in book.sheetnames if name.startswith("NURUS_TRAZABILIDAD")]
        if trace_names:
            trace = book[trace_names[-1]]
            for row in trace.iter_rows(min_row=2, max_row=7, values_only=True):
                if row and row[0] is not None:
                    trace_values[str(row[0])] = _text(row[1] if len(row) > 1 else "")
        return headers, rows, trace_values
    finally:
        book.close()


def _xls_rows(path: Path, sheet_name: str, header_row: int) -> tuple[list[object], list[list[object]], dict[str, str]]:
    try:
        import pandas as pd

        trace_values: dict[str, str] = {}
        with pd.ExcelFile(path, engine="xlrd") as excel:
            frame = pd.read_excel(excel, sheet_name=sheet_name, header=None, dtype=object)
            trace_names = [name for name in excel.sheet_names if name.startswith("NURUS_TRAZABILIDAD")]
            if trace_names:
                trace = pd.read_excel(excel, sheet_name=trace_names[-1], header=None, dtype=object)
                for row in trace.iloc[1:7].itertuples(index=False, name=None):
                    if row and not _blank(row[0]):
                        trace_values[_text(row[0])] = _text(row[1] if len(row) > 1 else "")
    except ImportError as exc:
        raise ReviewImportError("Para leer .xls instala la dependencia opcional excel-legacy.") from exc
    except ValueError as exc:
        raise ReviewImportError(f"No se pudo leer la hoja {sheet_name!r}: {exc}") from exc
    index = header_row - 1
    if index < 0 or index >= len(frame.index):
        raise ReviewImportError("La fila de encabezados no existe en el Excel revisado.")
    return list(frame.iloc[index]), frame.iloc[index + 1 :].values.tolist(), trace_values


def preview_reviewed_workbook(db: Database, batch_id: str, path: str | Path) -> ReviewImportPreview:
    source = Path(path)
    if not source.is_file():
        raise ReviewImportError("No se encontró el Excel revisado.")
    suffix = source.suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls"}:
        raise ReviewImportError("La constancia debe ser un archivo .xlsx, .xlsm o .xls.")
    content = source.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    batch = db.get_batch(batch_id)
    if batch is None:
        raise KeyError("Lote no encontrado.")
    if suffix == ".xls":
        headers, rows, trace = _xls_rows(BytesIO(content), batch["primary_sheet"], int(batch["header_row"]))
    else:
        headers, rows, trace = _xlsx_rows(BytesIO(content), batch["primary_sheet"], int(batch["header_row"]))

    positions, errors = _positions(headers)
    if trace.get("ESTADO_DOCUMENTO") != "PROPUESTA_NO_REVISADA":
        errors.append("El archivo no es una propuesta NuRus identificable.")
    if trace.get("EVALUACION_SHA256") != batch["evaluation_hash"]:
        errors.append("El archivo corresponde a otra evaluación NuRus.")
    if trace.get("FUENTE_SHA256") != batch["source_hash"]:
        errors.append("El archivo no corresponde al libro de origen de este lote.")

    current = {row["record_id"]: row for row in db.list_review_records(batch_id)}
    seen: set[str] = set()
    updates: list[dict[str, object]] = []
    warnings: list[str] = []
    changed = reviewed = excluded = 0
    if not errors:
        for excel_row, values in enumerate(rows, start=int(batch["header_row"]) + 1):
            record_id = _text(values[positions["NURUS_ID_REGISTRO"]] if positions["NURUS_ID_REGISTRO"] < len(values) else "")
            if not record_id:
                managed_values = [
                    _text(values[positions[title]])
                    for title in MANAGED_HEADERS[1:]
                    if positions[title] < len(values)
                ]
                if any(managed_values):
                    errors.append(
                        f"Fila {excel_row}: contiene datos de revisión pero no NURUS_ID_REGISTRO."
                    )
                continue
            if record_id in seen:
                errors.append(f"Fila {excel_row}: NURUS_ID_REGISTRO repetido.")
                continue
            seen.add(record_id)
            row = current.get(record_id)
            if row is None:
                errors.append(f"Fila {excel_row}: el identificador no pertenece a este lote.")
                continue
            state = _text(values[positions["NURUS_ESTADO_REVISION"]]).upper()
            observation = _text(values[positions["OBSERVACION"]])
            if observation.startswith("=") or observation.startswith("#"):
                errors.append(f"Fila {excel_row}: OBSERVACION contiene una fórmula o error de Excel.")
            if row["decision"] == "excluded":
                excluded += 1
                review_date = ""
                if state != "EXCLUIDO":
                    errors.append(f"Fila {excel_row}: una fila excluida no puede cambiarse desde la constancia.")
            else:
                reviewed += 1
                if not observation:
                    errors.append(f"Fila {excel_row}: falta OBSERVACION.")
                try:
                    review_date = _review_date(values[positions["FECHA_OBS"]])
                except ValueError as exc:
                    errors.append(f"Fila {excel_row}: {exc}.")
                    review_date = ""
            if observation != row["original_observation"]:
                changed += 1
            updates.append({
                "record_id": record_id,
                "observation": observation,
                "review_date": review_date,
                "tt": _text(values[positions["TT"]]),
                "workload": _text(values[positions["CC"]]),
                "resolution": _text(values[positions["RES"]]),
            })

    missing = sorted(set(current) - seen)
    if missing and not any("Falta la columna" in item for item in errors):
        missing_reviewable = sum(current[item]["decision"] != "excluded" for item in missing)
        if missing_reviewable:
            warnings.append(
                f"Faltan {missing_reviewable} registros revisables; permanecerán pendientes."
            )
        missing_excluded = len(missing) - missing_reviewable
        if missing_excluded:
            warnings.append(
                f"Faltan {missing_excluded} registros excluidos; su estado no será modificado."
            )
    return ReviewImportPreview(
        path=source,
        sha256=digest,
        updates=tuple(updates),
        reviewed_count=reviewed,
        excluded_count=excluded,
        changed_observations=changed,
        errors=tuple(errors),
        warnings=tuple(warnings),
        source_bytes=content,
    )


def import_reviewed_workbook(
    db: Database,
    batch_id: str,
    path: str | Path,
    *,
    responsible: str,
    confirmed_in_rus: bool,
    expected_sha256: str | None = None,
) -> ReviewImportPreview:
    if not confirmed_in_rus:
        raise ReviewImportError("Debes confirmar que la revisión y sus observaciones quedaron registradas en RUS.")
    preview = preview_reviewed_workbook(db, batch_id, path)
    if expected_sha256 is not None and preview.sha256 != expected_sha256:
        raise ReviewImportError("El Excel cambió durante la confirmación. Guarda tus cambios y vuelve a validar.")
    if not preview.valid:
        raise ReviewImportError("\n".join(preview.errors[:12]))
    db.apply_review_import(
        batch_id,
        source_name=preview.path.name,
        source_bytes=preview.source_bytes,
        responsible=responsible,
        updates=[dict(item) for item in preview.updates],
    )
    return preview
