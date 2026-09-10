from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import tempfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import ZipFile

from nurus.storage.database import Database


class ExportError(ValueError):
    pass


@dataclass(frozen=True)
class ExportResult:
    path: Path
    sha256: str
    row_count: int


def _safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _is_blank_excel_value(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _snapshot_rows(db: Database, batch_id: str) -> tuple[object, list[dict[str, object]]]:
    try:
        snapshot = db.get_snapshot(batch_id)
    except ValueError as exc:
        raise ExportError(str(exc)) from exc
    batch = snapshot["batch"]
    rows: list[dict[str, object]] = []
    for item in snapshot["records"]:
        if item["decision"] != "approved":
            continue
        values = json.loads(item["values_json"])
        values["NURUS_OBSERVACION"] = item["edited_observation"]
        values["NURUS_HOJA_ORIGEN"] = item["source_sheet"]
        values["NURUS_FILA_ORIGEN"] = item["source_row"]
        values["NURUS_HASH_ORIGEN"] = item["source_hash"]
        rows.append(values)
    if not rows:
        raise ExportError("El lote no contiene registros aprobados exportables.")
    return batch, rows


def _headers(rows: list[dict[str, object]]) -> list[str]:
    result: list[str] = []
    for row in rows:
        for key in row:
            if key not in result:
                result.append(key)
    return result


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_destination(
    batch: dict[str, object], target: Path, *, allow_overwrite: bool = False
) -> None:
    source_path = str(batch.get("source_path", "") or "")
    if not source_path:
        raise ExportError("El lote no identifica la ruta de origen; vuelve a importarlo antes de exportar.")
    source = Path(source_path).expanduser().resolve()
    try:
        is_source = target == source or (
            target.exists() and source.exists() and target.samefile(source)
        )
    except OSError as exc:
        raise ExportError(f"No se pudo comprobar la identidad del destino: {exc}") from exc
    if is_source:
        raise ExportError("El destino corresponde al archivo de origen; elige otro archivo.")
    if target.exists() and not allow_overwrite:
        raise ExportError("El archivo de destino ya existe; confirma un nombre distinto o sobrescritura.")


def export_review_snapshot(
    db: Database,
    batch_id: str,
    destination: str | Path,
    *,
    overwrite: bool = False,
) -> ExportResult:
    """Exporta una tabla de revisión; no conserva la estructura del libro original."""
    target = Path(destination).expanduser().resolve()
    if target.suffix.lower() not in {".xlsx", ".csv"}:
        raise ExportError("La exportación debe ser .xlsx o .csv.")
    batch, rows = _snapshot_rows(db, batch_id)
    _validate_destination(batch, target, allow_overwrite=overwrite)
    target.parent.mkdir(parents=True, exist_ok=True)

    headers = _headers(rows)
    suffix = target.suffix.lower()
    temp_path: Path | None = None
    try:
        if suffix == ".xlsx":
            from openpyxl import Workbook

            handle = tempfile.NamedTemporaryFile(
                prefix=target.stem + ".", suffix=".tmp.xlsx", dir=target.parent, delete=False
            )
            handle.close()
            temp_path = Path(handle.name)
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Revision"
            sheet.append(headers)
            for row in rows:
                sheet.append([_safe_cell(row.get(header, "")) for header in headers])
            trace = workbook.create_sheet("Trazabilidad")
            trace.append(["CAMPO", "VALOR"])
            for key, value in [
                ("batch_id", batch["id"]),
                ("snapshot_hash", batch["snapshot_hash"]),
                ("source_name", batch["source_name"]),
                ("source_hash", batch["source_hash"]),
                ("mode", batch["mode"]),
                ("as_of", batch["as_of"]),
                ("engine_version", batch["engine_version"]),
                ("catalog_hash", batch["catalog_hash"]),
            ]:
                trace.append([key, value])
            trace.sheet_state = "hidden"
            workbook.save(temp_path)
            workbook.close()
        else:
            handle = tempfile.NamedTemporaryFile(
                mode="w", encoding="utf-8-sig", newline="", prefix=target.stem + ".",
                suffix=".tmp.csv", dir=target.parent, delete=False
            )
            temp_path = Path(handle.name)
            with handle:
                writer = csv.writer(handle, delimiter=";", quoting=csv.QUOTE_MINIMAL)
                writer.writerow(headers)
                for row in rows:
                    writer.writerow([_safe_cell(row.get(header, "")) for header in headers])
        os.replace(temp_path, target)
        temp_path = None
    except OSError as exc:
        raise ExportError(f"No se pudo escribir la exportación: {exc}") from exc
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)

    return ExportResult(target, _hash(target), len(rows))


def _normalized_header(value: object) -> str:
    return " ".join(str(value or "").strip().upper().replace("_", " ").split())


def _review_records(snapshot: dict) -> list[dict]:
    records = list(snapshot["records"])
    if not records:
        raise ExportError("El snapshot no contiene filas de revisión.")
    primary = snapshot["batch"]["primary_sheet"]
    for record in records:
        if record["source_sheet"] != primary:
            raise ExportError("La procedencia del snapshot no corresponde a la hoja principal.")
        if int(record["source_row"]) < 1:
            raise ExportError("El snapshot contiene una fila de origen inválida.")
    return records


def _trace_rows(snapshot: dict) -> list[tuple[object, ...]]:
    batch = snapshot["batch"]
    stage = str(batch.get("export_stage", "reviewed"))
    rows: list[tuple[object, ...]] = [
        ("CAMPO", "VALOR"),
        ("ESTADO_DOCUMENTO", "PROPUESTA_NO_REVISADA" if stage == "proposal" else "CONSTANCIA_REVISADA"),
        ("EVALUACION_SHA256", batch.get("evaluation_hash", "")),
        ("SNAPSHOT", batch.get("snapshot_hash", "")),
        ("FUENTE_SHA256", batch["source_hash"]),
        ("MODO", batch["mode"]),
        ("HOJA_PROCESADA", batch["primary_sheet"]),
        ("EXCEPCIONES", json.dumps(snapshot.get("exceptions", []), ensure_ascii=False, sort_keys=True)),
        (),
        ("HOJA", "FILA", "ESTADO", "MOTIVO", "OBSERVACION", "REGLAS", "HASH_ORIGEN"),
    ]
    for record in snapshot["records"]:
        rows.append((
            record["source_sheet"], record["source_row"], record["decision"],
            record["edit_reason"], record["edited_observation"],
            record["rule_ids_json"], record["source_hash"],
        ))
    return rows


def _trace_name(existing: set[str]) -> str:
    base = "NURUS_TRAZABILIDAD"
    result = base
    index = 2
    while result in existing:
        result = f"{base}_{index}"
        index += 1
    return result


def _column_plan(headers: list[object], titles: tuple[str, ...]) -> dict[str, int]:
    positions: dict[str, list[int]] = {}
    for position, value in enumerate(headers, start=1):
        key = _normalized_header(value)
        if key:
            positions.setdefault(key, []).append(position)
    for title in titles:
        name = _normalized_header(title)
        if len(positions.get(name, [])) > 1:
            raise ExportError(f"Hay más de una columna {name!r}; no se puede anotar sin ambigüedad.")
    last = len(headers)
    result: dict[str, int] = {}
    for title in titles:
        name = _normalized_header(title)
        position = positions.get(name, [last + 1])[0]
        result[title] = position
        last = max(last, position)
    return result


def _annotation_values(record: dict, stage: str) -> tuple[str, str]:
    if record["decision"] == "excluded":
        state = "EXCLUIDO"
    elif stage == "proposal" and record["evaluation_status"] == "blocked":
        state = "REQUIERE_REVISION"
    elif stage == "proposal":
        state = "PROPUESTA"
    else:
        state = "REVISADO"
    return str(_safe_cell(record["edited_observation"])), state


def _portable_preserved(content: bytes, target: Path, snapshot: dict) -> None:
    from copy import copy

    from openpyxl import load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    suffix = target.suffix.lower()
    book = load_workbook(
        BytesIO(content), data_only=False, keep_links=True, keep_vba=suffix == ".xlsm"
    )
    try:
        batch = snapshot["batch"]
        if batch["primary_sheet"] not in book.sheetnames:
            raise ExportError("La hoja procesada no existe en los bytes conservados.")
        sheet = book[batch["primary_sheet"]]
        stage = str(batch.get("export_stage", "reviewed"))
        observation_title = "NURUS_PROPUESTA" if stage == "proposal" else "NURUS_OBSERVACION_FINAL"
        header_row = int(batch["header_row"])
        if header_row < 1 or header_row > sheet.max_row:
            raise ExportError("La fila de encabezado del snapshot no es válida en el libro conservado.")
        headers = [sheet.cell(header_row, col).value for col in range(1, sheet.max_column + 1)]
        titles = (
            ("NURUS_ID_REGISTRO", "NURUS_PROPUESTA", "OBSERVACION", "FECHA_OBS", "TT", "CC", "RES", "NURUS_ESTADO_REVISION")
            if stage == "proposal"
            else ("NURUS_ID_REGISTRO", observation_title, "NURUS_ESTADO_REVISION")
        )
        columns = _column_plan(headers, titles)
        observation_column = columns[observation_title]
        state_column = columns["NURUS_ESTADO_REVISION"]
        for title, column in columns.items():
            cell = sheet.cell(header_row, column)
            if cell.value is None:
                source = sheet.cell(header_row, max(1, column - 1))
                cell._style = copy(source._style)
                cell.number_format = source.number_format
                cell.alignment = copy(source.alignment)
                cell.protection = copy(source.protection)
                cell.value = title

        records = _review_records(snapshot)
        for record in records:
            row = int(record["source_row"])
            if row <= header_row or row > sheet.max_row:
                raise ExportError(f"La fila {row} no existe en la hoja procesada.")
            observation, state = _annotation_values(record, stage)
            sheet.cell(row, columns["NURUS_ID_REGISTRO"]).value = _safe_cell(record["record_id"])
            sheet.cell(row, observation_column).value = observation
            sheet.cell(row, state_column).value = state
            if stage == "proposal" and _is_blank_excel_value(sheet.cell(row, columns["OBSERVACION"]).value):
                sheet.cell(row, columns["OBSERVACION"]).value = observation

        first_data_row = header_row + 1
        last_column = max(sheet.max_column, state_column)
        status_letter = get_column_letter(state_column)
        highlight = FormulaRule(
            formula=["$" + status_letter + str(first_data_row) + '="EXCLUIDO"'],
            fill=PatternFill(fill_type="solid", fgColor="FFF2CC"),
        )
        sheet.conditional_formatting.add(
            f"A{first_data_row}:{get_column_letter(last_column)}{sheet.max_row}", highlight
        )

        trace = book.create_sheet(_trace_name(set(book.sheetnames)))
        for values in _trace_rows(snapshot):
            trace.append([_safe_cell(value) for value in values])
        trace.sheet_state = "hidden"
        book.save(target)
    finally:
        book.close()


def _native_preserved(content: bytes, target: Path, snapshot: dict) -> None:
    import platform

    if platform.system() != "Windows":
        raise ExportError(
            "La exportación fiel requiere Windows, Excel de escritorio y pywin32. "
            "Para una salida portable de fidelidad reducida, confírmala expresamente."
        )
    if target.suffix.lower() in {".xlsx", ".xlsm"}:
        with ZipFile(BytesIO(content)) as package:
            names = {name.lower() for name in package.namelist()}
            if "xl/connections.xml" in names or any("macrosheets/" in name for name in names):
                raise ExportError(
                    "El libro contiene conexiones externas o macros XLM y requiere revisión especializada."
                )
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ExportError(
            "Falta pywin32 para usar Excel de escritorio. Instala la dependencia excel-native."
        ) from exc

    target.write_bytes(content)
    pythoncom.CoInitialize()
    app = book = placeholder = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        app.EnableEvents = False
        app.AskToUpdateLinks = False
        app.AutomationSecurity = 3
        placeholder = app.Workbooks.Add()
        app.Calculation = -4135
        app.CalculateBeforeSave = False
        book = app.Workbooks.Open(str(target), UpdateLinks=0, ReadOnly=False)
        placeholder.Close(SaveChanges=False)
        placeholder = None

        batch = snapshot["batch"]
        stage = str(batch.get("export_stage", "reviewed"))
        observation_title = "NURUS_PROPUESTA" if stage == "proposal" else "NURUS_OBSERVACION_FINAL"
        sheet = book.Worksheets(batch["primary_sheet"])
        header_row = int(batch["header_row"])
        used_last = sheet.UsedRange.Column + sheet.UsedRange.Columns.Count - 1
        headers = [sheet.Cells(header_row, column).Value2 for column in range(1, used_last + 1)]
        titles = (
            ("NURUS_ID_REGISTRO", "NURUS_PROPUESTA", "OBSERVACION", "FECHA_OBS", "TT", "CC", "RES", "NURUS_ESTADO_REVISION")
            if stage == "proposal"
            else ("NURUS_ID_REGISTRO", observation_title, "NURUS_ESTADO_REVISION")
        )
        columns = _column_plan(headers, titles)
        observation_column = columns[observation_title]
        state_column = columns["NURUS_ESTADO_REVISION"]
        for title, column in columns.items():
            if sheet.Cells(header_row, column).Value2 in (None, ""):
                sheet.Cells(header_row, column).Value2 = title

        records = _review_records(snapshot)
        last_row = sheet.UsedRange.Row + sheet.UsedRange.Rows.Count - 1
        for record in records:
            row = int(record["source_row"])
            if row <= header_row or row > last_row:
                raise ExportError(f"La fila {row} no existe en la hoja procesada.")
            observation, state = _annotation_values(record, stage)
            sheet.Cells(row, columns["NURUS_ID_REGISTRO"]).NumberFormat = "@"
            sheet.Cells(row, columns["NURUS_ID_REGISTRO"]).Value2 = str(record["record_id"])
            sheet.Cells(row, observation_column).NumberFormat = "@"
            sheet.Cells(row, observation_column).Value2 = observation
            sheet.Cells(row, state_column).NumberFormat = "@"
            sheet.Cells(row, state_column).Value2 = state
            if stage == "proposal" and sheet.Cells(row, columns["OBSERVACION"]).Value2 in (None, ""):
                sheet.Cells(row, columns["OBSERVACION"]).NumberFormat = "@"
                sheet.Cells(row, columns["OBSERVACION"]).Value2 = observation

        last_column = max(used_last, state_column)
        status_column_letter = _excel_column_name(state_column)
        for record in records:
            if record["decision"] != "excluded":
                continue
            row = int(record["source_row"])
            region = sheet.Range(sheet.Cells(row, 1), sheet.Cells(row, last_column))
            rule = region.FormatConditions.Add(
                Type=2, Formula1="=$" + status_column_letter + str(row) + '="EXCLUIDO"'
            )
            rule.Interior.Color = 204 + 242 * 256 + 255 * 65536
            rule.SetFirstPriority()
            rule.StopIfTrue = False

        trace = book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count))
        trace.Name = _trace_name({book.Worksheets(index).Name for index in range(1, book.Worksheets.Count + 1)})
        trace.Visible = 0
        for row_number, values in enumerate(_trace_rows(snapshot), start=1):
            for column, value in enumerate(values, start=1):
                trace.Cells(row_number, column).NumberFormat = "@"
                trace.Cells(row_number, column).Value2 = str(value or "")
        book.Save()
    finally:
        try:
            if book is not None:
                book.Close(SaveChanges=False)
        finally:
            try:
                if placeholder is not None:
                    placeholder.Close(SaveChanges=False)
                if app is not None:
                    app.Quit()
            finally:
                pythoncom.CoUninitialize()


def _excel_column_name(column: int) -> str:
    result = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _export_preserved_payload(
    content: bytes,
    snapshot: dict,
    destination: str | Path,
    *,
    backend: Literal["native", "portable"],
    allow_reduced_fidelity: bool,
) -> ExportResult:
    batch = snapshot["batch"]
    suffix = Path(str(batch["source_name"])).suffix.lower()
    target = Path(destination).expanduser().resolve()
    if suffix not in {".xls", ".xlsx", ".xlsm"}:
        raise ExportError("El origen no es un formato Excel admitido para exportación.")
    if target.suffix.lower() != suffix:
        raise ExportError(f"La salida debe conservar la extensión original {suffix}.")
    if backend not in {"native", "portable"}:
        raise ExportError("Backend de exportación no válido.")
    if backend == "portable" and (suffix == ".xls" or not allow_reduced_fidelity):
        raise ExportError(
            "La salida portable requiere confirmación de fidelidad reducida y no admite .xls."
        )
    _validate_destination(batch, target, allow_overwrite=False)
    target.parent.mkdir(parents=True, exist_ok=True)

    handle = tempfile.NamedTemporaryFile(
        prefix=target.stem + ".", suffix=target.suffix, dir=target.parent, delete=False
    )
    handle.close()
    temporary = Path(handle.name)
    target_created = False
    try:
        if backend == "native":
            _native_preserved(content, temporary, snapshot)
        else:
            _portable_preserved(content, temporary, snapshot)
        with target.open("xb") as output, temporary.open("rb") as source:
            target_created = True
            shutil.copyfileobj(source, output)
    except OSError as exc:
        if target_created:
            target.unlink(missing_ok=True)
        raise ExportError(f"No se pudo escribir la exportación: {exc}") from exc
    finally:
        temporary.unlink(missing_ok=True)

    return ExportResult(target, _hash(target), len(snapshot["records"]))


def export_proposal_workbook(
    db: Database,
    batch_id: str,
    destination: str | Path,
    *,
    backend: Literal["native", "portable"] = "native",
    allow_reduced_fidelity: bool = False,
) -> ExportResult:
    """Exporta la propuesta del motor antes de la revisión humana."""
    try:
        proposal = db.get_proposal_payload(batch_id)
        content = db.get_original_workbook(batch_id)
    except (KeyError, ValueError) as exc:
        raise ExportError(str(exc)) from exc
    return _export_preserved_payload(
        content,
        proposal,
        destination,
        backend=backend,
        allow_reduced_fidelity=allow_reduced_fidelity,
    )


def export_preserved_workbook(
    db: Database,
    batch_id: str,
    destination: str | Path,
    *,
    backend: Literal["native", "portable"] = "native",
    allow_reduced_fidelity: bool = False,
) -> ExportResult:
    """Crea una salida nueva desde los bytes congelados y el snapshot vigente."""
    try:
        snapshot = db.get_snapshot(batch_id)
        content = db.get_original_workbook(batch_id, snapshot_hash=snapshot["batch"]["snapshot_hash"])
    except ValueError as exc:
        raise ExportError(str(exc)) from exc

    snapshot["batch"]["export_stage"] = "reviewed"
    return _export_preserved_payload(
        content,
        snapshot,
        destination,
        backend=backend,
        allow_reduced_fidelity=allow_reduced_fidelity,
    )
