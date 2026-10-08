"""Regresiones de resultados: integridad, procedencia, tipos y comparación."""
from collections import Counter
from datetime import date
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import Workbook,load_workbook
import motor as m
from resultados import load_lot,export_book,export_csv,export_zip,export_ics,compare_lots,headers_and_rows,publish
from preferencias import Preferences,date_preset,pending_plan


def fixture(folder,mode='Espera',names=('Persona ficticia','Persona ficticia'),private=False,complete=True,program='AFT FICTICIO'):
    folder.mkdir(parents=True,exist_ok=True)
    s={'screen':'seguimiento','tribunal':'113','tribunal_nombre':'Jgdo. L. y G. de Laja','modality':'2','tab':mode,
       'report':'1' if mode=='Informes' else None,'state':None,'month':None,'year':None,'fecha_descarga':'2026-10-01T09:00:00'}
    headers=['RIT','NOMBRE','RUT','NOMBRE CENTRO']
    extra={'Espera':(['T ESPERA'],[40]),'Cumplimiento':(['DIAS DE CUMPLIMIENTO','DIAS PARA EGRESAR','FEC.EGRESO PROYECTADO'],[50,90,'30/12/2026']),
           'Informes':(['FECHA VENCIMIENTO'],['10/10/2026'])}
    more,values=extra[mode];headers+=more
    book=Workbook();sheet=book.active;sheet.append(headers)
    for name in names:
        sheet.append(['X-1-2026',name,'00111111-1',program,*values])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    out=io.BytesIO();book.save(out);book.close();content=out.getvalue()
    path=folder/'pagina.xls';path.write_bytes(content)
    item={'archivo':path.name,'pagina':1,'registros':len(names),'sha256':hashlib.sha256(content).hexdigest(),'bytes':len(content),'estado':'OK'}
    manifest={'estado':'VALIDADA','consulta':'Consulta ficticia','seleccion':s,'archivos':[item],'paginas_esperadas':1,'paginas_verificadas':1}
    (folder/'verificacion 001.json').write_text(json.dumps(manifest),encoding='utf-8')
    query={'consulta':'Consulta ficticia','estado':'VALIDADA','seleccion':s,'verificacion':'verificacion 001.json'}
    summary={'estado':'VALIDADA' if complete else 'INCOMPLETA','version_contrato':1,'consultas_esperadas':1,'consultas_completadas':1,
             'consultas':[query],'plan':[s],'otros_filtros':'formulario' if private else 'todos los registros','filtro_fecha':None,'fecha_hora':s['fecha_descarga']}
    (folder/'resumen.json').write_text(json.dumps(summary),encoding='utf-8')
    return folder


class ResultTests(unittest.TestCase):
    def setUp(self):self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def test_consolidation_keeps_duplicates_and_text(self):
        lot=load_lot(fixture(self.root/'lote',program='=1+2'))
        before=(lot.folder/'pagina.xls').read_bytes();path=export_book([lot],csmp=True)
        book=load_workbook(path);sheet=book['ESPERA'];headers=[c.value for c in sheet[1]]
        self.assertEqual(sheet.max_row,3);self.assertEqual(sheet['D2'].data_type,'s');self.assertEqual(sheet['D2'].value,'=1+2')
        self.assertEqual(sheet.cell(2,headers.index('RUT')+1).value,'00111111-1')
        self.assertIn('SITFA_ARCHIVO',headers);self.assertIn('TRIBUNAL',headers)
        self.assertEqual((lot.folder/'pagina.xls').read_bytes(),before);book.close()
    def test_modified_or_missing_original_rejected(self):
        folder=fixture(self.root/'lot');(folder/'pagina.xls').write_bytes(b'modificado')
        with self.assertRaises(m.PocError):load_lot(folder)
    def test_path_escape_and_symlink_rejected(self):
        folder=fixture(self.root/'lot');file=folder/'verificacion 001.json';data=json.loads(file.read_text());data['archivos'][0]['archivo']='../otro.xls';file.write_text(json.dumps(data))
        with self.assertRaises(m.PocError):load_lot(folder)
    def test_comparison_preserves_multiplicity(self):
        old=load_lot(fixture(self.root/'old',names=('Persona ficticia',)))
        new=load_lot(fixture(self.root/'new',names=('Persona ficticia','Persona ficticia')))
        book=load_workbook(compare_lots(old,new));self.assertEqual(book['Nuevas']['D2'].value,1);book.close()
    def test_different_or_private_scope_rejected(self):
        old=load_lot(fixture(self.root/'old'));new=load_lot(fixture(self.root/'new',private=True))
        with self.assertRaises(m.PocError):compare_lots(old,new)
        new=load_lot(fixture(self.root/'other',mode='Informes'))
        with self.assertRaises(m.PocError):compare_lots(old,new)
    def test_partial_requires_explicit_choice(self):
        folder=fixture(self.root/'partial',complete=False)
        with self.assertRaises(m.PocError):load_lot(folder)
        lot=load_lot(folder,partial=True);self.assertIn('INCOMPLETO',export_book([lot]).name)
        with self.assertRaises(m.PocError):export_book([lot],csmp=True)
    def test_csv_ics_and_zip(self):
        lot=load_lot(fixture(self.root/'lot',mode='Informes',program='=1+2'))
        csv=export_csv(lot)[0];self.assertIn("'=1+2",csv.read_text(encoding='utf-8-sig'))
        text=export_ics(lot).read_text();self.assertIn('DTSTART;VALUE=DATE:20261010',text);self.assertNotIn('111111',text);self.assertNotIn('=1+2',text)
        with zipfile.ZipFile(export_zip(lot)) as archive:
            self.assertIn('pagina.xls',archive.namelist());self.assertIn('INDICE.txt',archive.namelist())
    def test_csmp_cross_and_multiple_schemas(self):
        a=load_lot(fixture(self.root/'cumpl',mode='Cumplimiento'))
        b=load_lot(fixture(self.root/'informes',mode='Informes'))
        book=load_workbook(export_book([a,b],csmp=True));self.assertIn('CUMPLIMIENTO',book.sheetnames);self.assertIn('INFORMES',book.sheetnames);book.close()
    def test_publish_failure_does_not_leave_product(self):
        def fail(path):path.write_bytes(b'incompleto');raise OSError('fallo')
        with self.assertRaises(OSError):publish(self.root,'Producto','.xlsx',fail)
        self.assertEqual(list(self.root.iterdir()),[])
    def test_existing_output_is_never_deleted_on_collision(self):
        from datetime import datetime
        with patch('resultados.m.now',return_value=datetime(2026,10,1)):
            first=publish(self.root,'Producto','.txt',lambda p:p.write_text('original'))
            with self.assertRaises(FileExistsError):publish(self.root,'Producto','.txt',lambda p:p.write_text('nuevo'))
            self.assertEqual(first.read_text(),'original')
    def test_dates_favorites_and_pending(self):
        self.assertEqual(date_preset('Mes anterior',date(2024,3,12)),(date(2024,2,1),date(2024,2,29)))
        folder=fixture(self.root/'lot');self.assertEqual(pending_plan(folder),[])
        prefs=Preferences(self.root/'prefs.json');from lotes import Batch
        prefs.favorite('Favorito',Batch(('113',),('2',),('Espera',)))
        self.assertIsNone(json.loads(prefs.path.read_text())['favoritos']['Favorito']['dates'])
    def test_unknown_footer_is_not_silently_dropped(self):
        info=m.ExcelInfo('fixture',Counter({('X','Persona'):1}),[['RIT','NOMBRE'],['X','Persona'],['Texto desconocido','']])
        with self.assertRaises(m.PocError):headers_and_rows(info)
    def test_calendar_trailing_empty_columns_are_ignored(self):
        headers=['FECHA VENCIMIENTO','RIT','TRIBUNAL','RUT MENOR','NOMBRE MENOR','TIPO PROGRAMA','NOMBRE CENTRO']
        rows=[['10/10/2026','X-1-2026','Tribunal ficticio','00111111-1','Persona ficticia','FAE','Centro ficticio'],
              ['11/10/2026','X-2-2026','Tribunal ficticio','00222222-2','Otra persona','FAE','Centro ficticio']]
        grid=[headers+['',None,' '],*[r+['',None,''] for r in rows]]
        info=m.ExcelInfo('fixture',m.records_in_grid(grid),grid)
        actual,parsed=headers_and_rows(info)
        self.assertEqual(actual,tuple(headers))
        self.assertEqual([r for _,r in parsed],[tuple(r) for r in rows])
    def test_trailing_header_cannot_hide_nonempty_data(self):
        grid=[['RIT','NOMBRE',''],['X-1-2026','Persona ficticia','dato sin cabecera']]
        with self.assertRaisesRegex(m.PocError,'fuera de la cabecera'):
            headers_and_rows(m.ExcelInfo('fixture',m.records_in_grid(grid),grid))
    def test_empty_interior_and_duplicate_headers_still_rejected(self):
        for header in (['RIT','','NOMBRE'],['RIT','NOMBRE','NOMBRE']):
            with self.subTest(header=header),self.assertRaises(m.PocError):
                headers_and_rows(m.ExcelInfo('fixture',Counter(),[header]))
    def test_xlsx_1904_dates_are_preserved(self):
        from openpyxl.utils.datetime import CALENDAR_MAC_1904
        book=Workbook();book.epoch=CALENDAR_MAC_1904;book.active.append(['RIT','NOMBRE','FECHA VENCIMIENTO'])
        book.active.append(['X-1-2026','Persona ficticia',date(2026,10,10)])
        out=io.BytesIO();book.save(out);book.close();info=m.read_excel(out.getvalue())
        self.assertEqual(info.grid[1][2].date(),date(2026,10,10))
    def test_xml_sparse_rows_and_dates_are_preserved(self):
        xml=b'''<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet><Table>
        <Row><Cell><Data>RIT</Data></Cell><Cell><Data>NOMBRE</Data></Cell><Cell><Data>FECHA VENCIMIENTO</Data></Cell></Row>
        <Row ss:Index="4"><Cell><Data>X-1-2026</Data></Cell><Cell><Data>Persona ficticia</Data></Cell><Cell><Data ss:Type="DateTime">2026-10-10T00:00:00.000</Data></Cell></Row>
        </Table></Worksheet></Workbook>'''
        header,rows=headers_and_rows(m.read_excel(xml));self.assertEqual(rows[0][0],4);self.assertEqual(rows[0][1][2],date(2026,10,10))


if __name__=='__main__':unittest.main()
