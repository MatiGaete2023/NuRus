"""Shared presentation for newly generated reports, never for preserved sources."""
from datetime import date, datetime
from math import ceil
from openpyxl.styles import Border, Side, PatternFill, Font, Alignment
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='195F7B')
HEADER_FONT = Font(name='Calibri', size=11, color='FFFFFF', bold=True)
BODY_FONT = Font(name='Calibri', size=11)
WRAP = Alignment(vertical='top', wrap_text=True)
EDGE = Side(style='thin', color='FF000000')
BORDER = Border(left=EDGE,right=EDGE,top=EDGE,bottom=EDGE)


def column_width(title):
    title = str(title).upper()
    if any(x in title for x in ('OBSERVACION', 'OBSERVACIÓN', 'DETALLE', 'TEXTO', 'TRÁMITE', 'TRAMITE')): return 65
    if 'NOMBRE' in title or 'PERSONA' in title: return 35
    if 'PROGRAMA' in title or 'DERIVACION' in title: return 30
    if 'TRIBUNAL' in title: return 28
    if 'RIT' == title: return 17
    if 'RUT' in title: return 16
    if any(x in title for x in ('FECHA', 'F.', 'VENCIMIENTO', 'EGRESO')): return 19
    if 'ESPERA' in title or 'DÍAS' in title or 'DIAS' in title: return 14
    return 24


def configure_sheet(sheet, headers):
    sheet.freeze_panes = 'A2'
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = '9'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.45, bottom=.45, header=.2, footer=.2)
    sheet.oddFooter.center.text = 'Página &P de &N'
    sheet.row_dimensions[1].height = 32
    for index, title in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = column_width(title)


def format_cell(cell, header=False):
    cell.alignment = WRAP
    cell.border = BORDER
    cell.font = HEADER_FONT if header else BODY_FONT
    if header:
        cell.fill = HEADER_FILL
    elif isinstance(cell.value, str):
        cell.data_type = 's'
        cell.number_format = '@'
    elif isinstance(cell.value, (date, datetime)):
        cell.number_format = 'dd/mm/yyyy'


def format_sheet(sheet, headers):
    configure_sheet(sheet, headers)
    for cells in sheet:
        for cell in cells:
            format_cell(cell, cell.row == 1)
        if cells[0].row > 1:
            lines = max((sum(max(1, ceil(len(part) / max(8, column_width(headers[c.column-1]) - 3))) for part in str(c.value or '').split('\n')) for c in cells), default=1)
            # Leave very long text to Excel auto-fit instead of forcing a clipped cap.
            if lines <= 12:
                sheet.row_dimensions[cells[0].row].height = max(22, 15 * lines + 5)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_options.horizontalCentered = True
