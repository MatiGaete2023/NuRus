import tempfile,json,hashlib,unittest
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook,Workbook
from icalendar import Calendar
from test_resultados import fixture
from resultados import load_lot,export_book,export_zip,export_ics


class ProductDesignTests(unittest.TestCase):
    def setUp(self):self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def test_printable_excel_and_manifest(self):
        lot=load_lot(fixture(self.root/'lot'))
        book=load_workbook(export_book([lot]));s=book['ESPERA']
        self.assertEqual(s.print_title_rows,'$1:$1');self.assertEqual(s.page_setup.fitToWidth,1)
        self.assertGreater(s.column_dimensions['B'].width,s.column_dimensions['A'].width)
        self.assertTrue(s['A1'].font.bold);book.close()
        with ZipFile(export_zip(lot)) as z:
            manifest=json.loads(z.read('MANIFIESTO.json'))
            for item in manifest['archivos']:self.assertEqual(hashlib.sha256(z.read(item['archivo'])).hexdigest(),item['sha256'])
    def test_calendar_roundtrip_and_distinct_ambiguous_rows(self):
        lot=load_lot(fixture(self.root/'lot',mode='Informes'))
        cal=Calendar.from_ical(export_ics(lot).read_bytes());events=cal.walk('VEVENT')
        self.assertEqual(len(events),2);self.assertNotEqual(str(events[0]['UID']),str(events[1]['UID']))
        self.assertIn('X-1-2026',str(events[0]['SUMMARY']))
        self.assertNotIn('Persona',str(events[0]['SUMMARY']))
    def test_unambiguous_calendar_uid_survives_date_edit(self):
        folder=fixture(self.root/'lot',mode='Informes',names=('Persona',))
        lot=load_lot(folder);first=Calendar.from_ical(export_ics(lot).read_bytes()).walk('VEVENT')[0]
        # Authoritative source and manifest are updated together, as a fresh verified input.
        book=load_workbook(__import__('io').BytesIO((folder/'pagina.xls').read_bytes()))
        s=book.active;s.cell(2,5,'12/10/2026');out=__import__('io').BytesIO();book.save(out);book.close();content=out.getvalue();(folder/'pagina.xls').write_bytes(content)
        path=folder/'verificacion 001.json';data=json.loads(path.read_text());item=data['archivos'][0];item.update(sha256=hashlib.sha256(content).hexdigest(),bytes=len(content));path.write_text(json.dumps(data))
        second=Calendar.from_ical(export_ics(load_lot(folder)).read_bytes()).walk('VEVENT')[0]
        self.assertEqual(str(first['UID']),str(second['UID']))
        self.assertNotEqual(first.decoded('DTSTART'),second.decoded('DTSTART'))

if __name__=='__main__':unittest.main()
