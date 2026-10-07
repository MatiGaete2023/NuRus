"""Motor local: valida respuestas y Excel recibidos desde la extension de Chrome.

No lee HAR al ejecutar ni accede a cookies, contrasenas o perfiles de navegador.
"""
from __future__ import annotations

import argparse
import ast
from html import unescape
from collections import Counter
from dataclasses import dataclass,field
from datetime import datetime
import hashlib
import io
import json
import os
from pathlib import Path
import queue
import re
import sys
import threading
import zipfile
import unicodedata
from urllib.parse import parse_qsl, urlencode, urlsplit
from zoneinfo import ZoneInfo

from lxml import etree, html

ORIGIN = "https://familia.pjud.cl"
ENDPOINT = ORIGIN + "/SITFAWEB/InformesDAction.do"
START = ORIGIN + "/SITFAWEB/jsp/Login/LoginB4.jsp"
PAGE_FIELD = "NUM_PaginaCumplimiento"
TOTAL_FIELD = "NUM_TotalCumplimiento"
TARGET = {"COD_Lengueta": "tdCumplimiento", "COD_Tribunal_sel": "119", "TIP_Consulta": "2"}


@dataclass(frozen=True)
class QueryProfile:
    key: str
    label: str
    filename_prefix: str
    tribunal: str
    modality: str
    excel_template: str
    tab: str = "tdCumplimiento"
    page_field: str = "NUM_PaginaCumplimiento"
    total_field: str = "NUM_TotalCumplimiento"
    table_id: str = "tablaCumplimiento"
    report_type: str | None = None
    screen: str = "seguimiento"
    action: str = "Buscar medida"
    state: str | None = None
    month: str | None = None
    year: str | None = None
    start: str | None = None
    end: str | None = None

    @property
    def endpoint(self):
        return ORIGIN+'/SITFAWEB/MaoDAction.do' if self.screen=='carga' else ENDPOINT

    @property
    def form_name(self):return 'MaoPpalForm' if self.screen=='carga' else 'InformesPpalForm'

    @property
    def target(self):
        if self.screen=='carga':
            return {'COD_TribunalDist':self.tribunal,'FEC_Desde':self.start,'FEC_Hasta':self.end}
        if self.screen == "litigantes":
            return {"COD_Tribunal_sel": self.tribunal, "TIP_Consulta": "2", "COD_EstBusqueda": self.state}
        if self.screen in ("calendario_informes", "calendario_medidas"):
            return {"COD_Tribunal_sel": self.tribunal, "TIP_Consulta": "1", "COD_Mes_sel": self.month,
                    "COD_Anio_Sel": self.year}
        result = {"COD_Lengueta": self.tab, "COD_Tribunal_sel": self.tribunal, "TIP_Consulta": self.modality}
        if self.report_type is not None:
            result["TIP_Informe"] = self.report_type
        return result



ORIGINAL = QueryProfile("mulchen-cumplimiento-ambulatorio", "Mulchen / Cumplimiento / Ambulatorio",
                        "Mulchen_Cumplimiento_Ambulatorio", "119", "2", "amblistcump")
LAJA_FAE = QueryProfile("laja-cumplimiento-fae", "Laja / Cumplimiento / FAE-FAS",
                       "Laja_Cumplimiento_FAE-FAS", "113", "3", "faefaslistcump")
LAJA_ESPERA = QueryProfile("laja-espera-ambulatorio", "Laja / Espera / Ambulatorio",
                          "Laja_Espera_Ambulatorio", "113", "2", "amblistesp",
                          "tdEspera", "NUM_PaginaEspera", "NUM_TotalEspera", "tablaEspera")
PROFILES = {p.key: p for p in (ORIGINAL, LAJA_FAE, LAJA_ESPERA)}
ALLOWED_FIELDS = frozenset("""COD_Centro GLS_TribunalOrigen FLG_Consulta COD_Lengueta
COD_Tribunal_sel TIP_Consulta COD_CentroResidencial COD_PlazoIntervencion TIP_Informe
COD_TiempoEspera TIP_Causa ROL_Causa ERA_Causa RUT_Consulta RUT_DvConsulta
COD_TipLitigante CHK_Consulta FEC_Inicio FEC_Fin irAccion NUM_PaginaEspera
NUM_TotalEspera NUM_PaginaCumplimiento NUM_TotalCumplimiento NUM_PaginaInforme
NUM_TotalInforme NUM_PaginaEgreso NUM_TotalEgreso""".split())
PAGINATION = frozenset(x for x in ALLOWED_FIELDS if x.startswith(("NUM_Pagina", "NUM_Total")))
COMMON_FIELDS = frozenset("COD_Centro GLS_TribunalOrigen FLG_Consulta COD_Tribunal_sel TIP_Consulta TIP_Causa ROL_Causa ERA_Causa RUT_Consulta RUT_DvConsulta CHK_Consulta FEC_Inicio FEC_Fin irAccion".split())
SCREEN_FIELDS = {
    'carga':frozenset('GLS_Tribunal COD_Rus COD_TribunalDist FEC_Desde FEC_Hasta irAccion'.split()),
    "seguimiento": ALLOWED_FIELDS,
    "litigantes": COMMON_FIELDS | frozenset("COD_Medida COD_EstBusqueda NUM_PaginaInforme NUM_Total".split()),
    "calendario_informes": frozenset("FLG_Busqueda COD_TribunalActual GLS_TribunalActual busquedaPorTribunal TIP_Consulta COD_Tribunal_sel FEC_Inicio COD_Mes_sel COD_Anio_Sel CHK_Contrae irAccion".split()),
    "calendario_medidas": frozenset("FLG_Busqueda COD_TribunalActual GLS_TribunalActual busquedaPorTribunal TIP_Consulta COD_Tribunal_sel FEC_Inicio COD_Mes_sel COD_Anio_Sel CHK_Contrae irAccion".split()),
}
PAGINATION = PAGINATION | frozenset({"NUM_Total"})


class PocError(Exception):
    """Solo mensajes controlados: nunca valores del servidor, URL XLS o excepciones crudas."""


def now():
    return datetime.now(ZoneInfo("America/Santiago"))


def normalized(value):
    text = ("" if value is None else str(value)).replace("\xa0", " ")
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return " ".join(text.upper().split())



def record_key(rit, name):
    r = normalized(rit).replace("–", "-").replace("—", "-")
    match = re.search(r"\b[A-Z]+\s*-\s*\d+\s*-\s*\d{4}\b", r)
    if not match or not normalized(name):
        raise PocError("No se reconoce la estructura RIT/Nombre del listado; descarga detenida.")
    return re.sub(r"\s+", "", match.group(0)), normalized(name)


def root_html(data):
    try:
        return html.fromstring(data, parser=html.HTMLParser(no_network=True))
    except (etree.ParserError, ValueError):
        raise PocError("Respuesta HTML vacia o ilegible.") from None


def has_login(root):
    return bool(root.xpath('//input[translate(@type,"PASSWORD","password")="password"]'))


def fields_from_form(form):
    result = {}
    for el in form.xpath('.//input[@name] | .//select[@name] | .//textarea[@name]'):
        name = el.get("name")
        if el.tag=="input" and el.get("type","").lower() in ("button","reset","image"):
            continue
        if el.get("type", "").lower() == "radio":
            group = form.xpath('.//input[@name=$name]', name=name)
            if any(e.get("type", "").lower() != "radio" for e in group):
                raise PocError("Formulario ambiguo: tipos de campo repetidos.")
            checked = [e for e in group if e.get("checked") is not None]
            if len(checked) > 1:
                raise PocError("Formulario ambiguo: opciones de consulta repetidas.")
            result[name] = (checked or group)[0].get("value", "")
            continue
        if name in result:
            raise PocError("Formulario ambiguo: campos repetidos.")
        if el.tag == "select":
            opts = el.xpath('./option[@selected]') or el.xpath('./option')[:1]
            result[name] = opts[0].get("value", "") if opts else ""
        elif el.tag == "textarea":
            result[name] = el.text_content()
        else:
            result[name] = el.get("value", "")
    return result


def response_fields(form,profile):
    fields=fields_from_form(form)
    if profile.screen=='seguimiento' and profile.report_type is not None:
        from filtros_respuesta import selected_report
        selected=selected_report(form.getroottree(),profile.tab)
        if selected is not None:
            if 'TIP_Informe' not in fields:
                raise PocError('Falta el selector de vencimiento de la respuesta.')
            fields['TIP_Informe']=selected
    return fields



@dataclass(frozen=True)
class RecordColumns:
    case: tuple[int, ...]
    name: int


def record_columns(row):
    """Encabezados observados en XLS reales; no confunde Nombre Centro con el NNA."""
    values = [normalized(v) for v in row]
    if "RIT" in values:
        case_labels = ("RIT",)
    elif all(label in values for label in ("TIPO CAUSA", "ROL CAUSA", "ERA CAUSA")):
        case_labels = ("TIPO CAUSA", "ROL CAUSA", "ERA CAUSA")
    else:
        return None
    names = [i for i, value in enumerate(values) if value in ("NOMBRE", "NOMBRE MENOR")]
    if not names:
        return None
    if len(names) != 1 or any(values.count(label) != 1 for label in case_labels):
        raise PocError("El listado contiene columnas de causa o nombre ambiguas.")
    return RecordColumns(tuple(values.index(label) for label in case_labels), names[0])


def record_from_row(row, columns):
    if len(row) <= max(*columns.case, columns.name):
        return None
    parts = [normalized(int(row[i]) if isinstance(row[i], float) and row[i].is_integer() else row[i])
             for i in columns.case]
    if len(parts) == 3:
        if not (re.fullmatch(r"[A-Z]", parts[0]) and re.fullmatch(r"\d+", parts[1])
                and re.fullmatch(r"\d{4}", parts[2])):
            return None
        rit = "-".join(parts)
    else:
        rit = parts[0]
    if not re.search(r"\b[A-Z]+\s*-\s*\d+\s*-\s*\d{4}\b", rit):
        return None
    return record_key(rit, row[columns.name])


def export_url(root, profile=None):
    """Lee solo ShowExcel; nunca ejecuta JavaScript ni inventa un nombre remoto."""
    scripts = "\n".join(root.xpath('//script/text()'))
    start = re.search(r'function\s+ShowExcel\s*\(\s*\)\s*\{', scripts)
    if not start:
        raise PocError("SITFA no entrego un enlace Excel para esta consulta.")
    block = scripts[start.end():]
    block = re.split(r'\bfunction\s+\w+\s*\(', block, maxsplit=1)[0]
    match = re.search(r'window\.open\s*\(\s*[\x22\x27]([^\x22\x27]+)[\x22\x27]', block)
    if not match:
        raise PocError("No se encontro la exportacion XLS auditada.")
    url = match.group(1)
    validate_export_url(url, profile)
    return url


def validate_export_url(url, profile=None):
    u = urlsplit(url)
    if (u.scheme != "https" or u.netloc != "familia.pjud.cl" or u.query or u.fragment
            or not re.fullmatch(r"/sitfa/reportes/[A-Za-z0-9_-]+\.xls", u.path)):
        raise PocError("URL XLS fuera del alcance permitido.")
    if profile is not None and not re.fullmatch(r"/sitfa/reportes/[A-Za-z0-9_-]+_" + re.escape(profile.excel_template) + r"\.xls", u.path):
        raise PocError("El enlace Excel no coincide con la consulta actual.")


def result_records(form, profile):
    if profile.screen=='carga':
        from carga import result_records as load_records
        return load_records(form,profile)
    if profile.screen.startswith("calendario_"):
        scripts = "\n".join(form.getroottree().xpath('//script/text()'))
        values = []
        for name in ("Ritxdias", "Nombresxdias"):
            literals = re.findall(name + r'\[dia\]\s*=\s*("(?:[^"\\]|\\.)*"|\x27(?:[^\x27\\]|\\.)*\x27)\s*;', scripts)
            try:
                items = [unescape(ast.literal_eval(s.replace(r'\/', '/'))) for s in literals]
            except (ValueError, SyntaxError):
                raise PocError("No se pueden comprobar los registros del calendario.") from None
            values.append(items)
        if len(values[0]) != len(values[1]):
            raise PocError("El calendario contiene registros incompletos.")
        tables = form.getroottree().xpath('//table[@id="TablaCalendarioSemana"] | //table[@id="TablaCalendarioDia"]')
        if len(tables) != 2 or any(len(t.xpath('./tr[td] | ./tbody/tr[td]')) != len(values[0]) for t in tables):
            raise PocError("SITFA cambio la estructura del calendario.")
        return Counter(record_key(rit, name) for rit, name in zip(*values))
    if profile.screen == "litigantes":
        tables = [t for t in form.xpath('.//table') if any(record_columns(row) is not None for row in table_grid(t))]
    else:
        tables = form.xpath('.//table[@id=$id]', id=profile.table_id)
    if len(tables) != 1:
        raise PocError("No se encontro un unico listado de la consulta seleccionada.")
    table = tables[0]
    records = records_in_grid(table_grid(table))
    rows = table.xpath('./tr[td] | ./tbody/tr[td]')
    # SITFA también utiliza TD para las cabeceras y mensajes de lista vacía.
    # Solo se excluyen cabeceras reconocidas y mensajes exactos, nunca filas desconocidas.
    empty_messages = {"SIN REGISTROS", "NO HAY REGISTROS", "SIN RESULTADOS",
                      "NO SE ENCONTRARON REGISTROS", "NO EXISTEN REGISTROS"}
    data_rows = []
    for row in rows:
        cells = row.xpath('./th | ./td')
        values = [cell.text_content() for cell in cells]
        if record_columns(values) is not None:
            continue
        if len(cells) == 1 and normalized(cells[0].text_content()).rstrip('.') in empty_messages:
            continue
        data_rows.append(row)
    if sum(records.values()) != len(data_rows):
        raise PocError("No se pudieron identificar todos los registros de la consulta.")
    return records



def records_in_grid(rows):
    header = None
    result = Counter()
    for row in rows:
        columns = record_columns(row)
        if columns is not None:
            header = columns
            continue
        if header is None:
            continue
        key = record_from_row(row, header)
        if key is not None:
            result[key] += 1
    return result



def table_grid(table):
    result=[];cells=0
    for row in table.xpath('./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr'):
        values=[cell.text_content() for cell in row.xpath('./th | ./td')];cells+=len(values)
        if len(result)>=200_000 or len(values)>200 or cells>2_000_000:raise PocError('La tabla supera los límites de filas/celdas.')
        result.append(values)
    return result


def query_pairs(raw, profile=ORIGINAL):
    pairs = parse_qsl(raw or "", keep_blank_values=True, max_num_fields=100)
    if len({k for k, _ in pairs}) != len(pairs):
        raise PocError("Consulta con parametros repetidos; descarga detenida.")
    p = dict(pairs)
    allowed = SCREEN_FIELDS.get(profile.screen)
    if not p or allowed is None or set(p) - allowed:
        raise PocError("La consulta tiene parametros no auditados; requiere revision local.")
    expected_action = {"seguimiento":"Buscar medida", "litigantes":"Consulta Informe",
                       "calendario_informes":"Buscar Inf.", "calendario_medidas":"Buscar", 'carga':'Aud.Carga Func.-'}.get(profile.screen)
    if profile.action != expected_action or p.get("irAccion") != expected_action or any(p.get(k) != v for k, v in profile.target.items()):
        raise PocError("La solicitud no corresponde a la consulta de lectura seleccionada.")
    if profile.screen=='carga':
        from carga import parse_day
        a,b=parse_day(profile.start),parse_day(profile.end)
        if a is None or b is None or not 0<=(b-a).days<30:raise PocError('Carga admite hasta 30 días inclusivos por consulta.')
    return pairs



@dataclass
class SearchPage:
    current: int
    total: int
    excel_url: str
    records: Counter
    pagination: dict
    bindings: list = field(default_factory=list)


def search_page(data, pairs, expected_page=None, expected_total=None, profile=ORIGINAL):
    p = dict(pairs)
    query_pairs(urlencode(pairs), profile)
    root = root_html(data)
    if has_login(root):
        raise PocError("Sesion expirada: se recibio una pagina de acceso.")
    forms = root.xpath('//form[@name=$name]',name=profile.form_name)
    if len(forms) != 1:
        raise PocError("Falta el formulario de Seguimiento; sesion o respuesta inesperada.")
    form = forms[0]
    action = form.get("action", "")
    if action not in (urlsplit(profile.endpoint).path, profile.endpoint) or form.get("method", "").lower() != "post":
        raise PocError("El destino del formulario cambio; descarga detenida.")
    fields = response_fields(form,profile)
    if any(fields.get(k) != p[k] for k in profile.target if k != "COD_Lengueta" and not (profile.screen=="litigantes" and k=="TIP_Consulta")):
        raise PocError("La respuesta pertenece a otro tribunal o modalidad.")
    # COD_Lengueta se llena en Envio(); el HTML auditado lo entrega vacio.
    if profile.screen == "seguimiento" and fields.get("COD_Lengueta") not in ("", p["COD_Lengueta"]):
        raise PocError("La respuesta pertenece a otra pestaña.")
    if any(re.search(r"csrf|token|nonce", k, re.I) for k in fields):
        raise PocError("Aparecio un campo dinamico no auditado; requiere revision.")
    try:
        if profile.screen.startswith("calendario_") or profile.screen=='carga':
            current = total = 1
        else:
            current, total = int(fields[profile.page_field]), int(fields[profile.total_field])
    except (KeyError, ValueError):
        raise PocError("No se puede determinar la paginacion de la pestaña seleccionada.") from None
    if not 1 <= current <= total <= 500:
        raise PocError("Paginacion fuera de rango; descarga detenida.")
    if expected_page is not None and current != expected_page:
        raise PocError("El servidor devolvio una pagina distinta de la solicitada.")
    if expected_total is not None and total != expected_total:
        raise PocError("Cambio el total de paginas durante la consulta; repetir en una ejecucion nueva.")
    url = export_url(root, profile)
    records = result_records(form, profile)
    if not records:
        raise PocError("No se pudieron identificar todos los registros de la pagina.")
    bindings=[]
    if profile.screen=='seguimiento':
        from vinculos import extract
        tables=form.xpath('.//table[@id=$id]',id=profile.table_id)
        if len(tables)==1:bindings=extract(tables[0],profile.tribunal)
    return SearchPage(current, total, url, records, {k: fields.get(k, "") for k in PAGINATION},bindings)



@dataclass
class ExcelInfo:
    format: str
    records: Counter
    grid: list | None = None



def read_excel(data,record_reader=None):
    if not data:
        raise PocError("El XLS esta vacio.")
    if len(data)>80*1024*1024:raise PocError('El Excel supera 80 MB.')
    grids = []
    kind = ""
    try:
        if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            import xlrd
            wb = xlrd.open_workbook(file_contents=data, logfile=io.StringIO())
            try:
                if sum(s.nrows*s.ncols for s in wb.sheets())>2_000_000 or any(s.ncols>200 for s in wb.sheets()):
                    raise PocError('El Excel BIFF excede los límites de celdas.')
                grids=[]
                for sheet in wb.sheets():
                    grid=[]
                    for i in range(sheet.nrows):
                        values=sheet.row_values(i)
                        for j,cell_type in enumerate(sheet.row_types(i)):
                            if cell_type==xlrd.XL_CELL_DATE:values[j]=xlrd.xldate_as_datetime(values[j],wb.datemode)
                        grid.append(values)
                    grids.append(grid)
            finally:
                wb.release_resources()
            kind = "XLS-BIFF"
        elif data.startswith(b"PK\x03\x04"):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries=archive.infolist()
                if len(entries)>2000 or sum(item.file_size for item in entries)>160*1024*1024:
                    raise PocError('El Excel comprimido excede los límites de lectura.')
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                grids=[];cells=0
                for sheet in wb.worksheets:
                    if sheet.max_row and sheet.max_row>200_000 or sheet.max_column and sheet.max_column>200:
                        raise PocError('Las dimensiones del Excel exceden los límites de lectura.')
                    grid=[]
                    for row in sheet.iter_rows(values_only=True):
                        cells+=len(row)
                        if cells>2_000_000:raise PocError('El Excel excede dos millones de celdas.')
                        grid.append(row)
                    grids.append(grid)
            finally:
                wb.close()
            kind = "OOXML-con-extension-xls"
        elif b"urn:schemas-microsoft-com:office:spreadsheet" in data[:4096]:
            root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))
            ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
            for table in root.xpath('//ss:Table', namespaces=ns):
                grid = []
                for row in table.xpath('./ss:Row', namespaces=ns):
                    row_index=row.get('{'+ns['ss']+'}Index')
                    if row_index:
                        target=int(row_index)
                        if not len(grid)<target<=200_000:raise PocError('Índice de fila XML inválido.')
                        grid.extend([[] for _ in range(target-1-len(grid))])
                    vals = []
                    for cell in row.xpath('./ss:Cell', namespaces=ns):
                        index = cell.get('{'+ns['ss']+'}Index')
                        if index:
                            if not 1<=int(index)<=200:raise PocError('El índice XML excede los límites de columnas.')
                            vals.extend([""] * (int(index)-1-len(vals)))
                        value="".join(cell.xpath('./ss:Data//text()', namespaces=ns))
                        data_cell=cell.find('ss:Data',namespaces=ns)
                        if data_cell is not None and data_cell.get('{'+ns['ss']+'}Type')=='DateTime':
                            try:value=datetime.fromisoformat(value.rstrip('Z')).date()
                            except ValueError:pass
                        vals.append(value)
                        if len(vals)>200:raise PocError('La fila XML excede los límites de columnas.')
                    grid.append(vals)
                    if len(grid)>200_000:raise PocError('La tabla XML excede los límites de filas.')
                grids.append(grid)
            kind = "XML-Spreadsheet-2003"
        else:
            root = root_html(data)
            if has_login(root):
                raise PocError("Sesion expirada: la descarga contiene un formulario de acceso.")
            if root.xpath('//form | //script'):
                raise PocError("La descarga es una pagina web activa, no un Excel verificable.")
            grids = [table_grid(t) for t in root.xpath('//table')]
            kind = "HTML-tabular-con-extension-xls"
        candidates = [((record_reader or records_in_grid)(grid),grid) for grid in grids]
        candidates = [(r,g) for r,g in candidates if r]
        if len(candidates) != 1:
            raise PocError("Se recibio el Excel, pero no se reconoce un unico listado de causas y nombres. No es una consulta sin resultados.")
        return ExcelInfo(kind, candidates[0][0], candidates[0][1])
    except PocError:
        raise
    except Exception:
        raise PocError("No se pudo abrir/parsear el archivo Excel; requiere revision local.") from None



def invariant(pairs):
    return [(k, v) for k, v in pairs if k not in PAGINATION]


def excel_record_rows(info):
    """Compara las filas completas: un RIT/nombre puede tener varias medidas."""
    rows = Counter();columns = None;headers = ()
    for row in info.grid or []:
        found = record_columns(row)
        if found is not None:
            columns = found;headers = tuple(normalized(value) for value in row)
        elif columns is not None and record_from_row(row, columns) is not None:
            rows[(headers,tuple(normalized(value) for value in row))] += 1
    return rows


def verified_download(data, expected, prior_bytes, prior_records,profile=ORIGINAL):
    if profile.screen=='carga':
        from carga import records_in_grid as record_reader
    else:record_reader=None
    info = read_excel(data,record_reader)
    if info.records != expected.records:
        raise PocError("Los registros XLS no coinciden con la pagina HTML (puede ser un XLS anterior o de toda la consulta).")
    if profile.screen.startswith('calendario_'):
        verify_month(info,profile)
    digest = hashlib.sha256(data).hexdigest()
    if digest in prior_bytes or excel_record_rows(info) in prior_records:
        raise PocError("Dos paginas contienen el mismo archivo o los mismos registros; descarga detenida.")
    return info, digest


def verify_month(info,profile):
    """Una ruta temporal con las mismas personas no acredita el mes consultado."""
    from datetime import date
    aliases={'FECHA VENCIMIENTO','FEC.VENCIMIENTO','FEC. VENCIMIENTO',
             'F. VENCIMIENTO','F.VENCIMIENTO','F.VENC','F. VENC','VENCIMIENTO'}
    if profile.screen=='calendario_medidas':
        aliases=aliases|{'FEC.EGRESO PROYECTADO','FEC. EGRESO PROYECTADO',
                 'FEC EGRESO PROYECTADO','FECHA EGRESO PROYECTADO'}
    columns=None;position=None;checked=0
    for number,row in enumerate(info.grid,1):
        found=record_columns(row)
        if found is not None:
            candidates=[i for i,value in enumerate(row) if normalized(value) in aliases]
            if len(candidates)!=1:raise PocError('El calendario XLS no permite comprobar sus fechas.')
            columns=found;position=candidates[0]
        elif columns is not None and record_from_row(row,columns) is not None:
            value=row[position] if position<len(row) else None
            if isinstance(value,datetime):day=value.date()
            elif isinstance(value,date):day=value
            else:
                day=None
                for pattern in ('%d/%m/%Y','%Y-%m-%d','%d-%m-%Y'):
                    try:day=datetime.strptime(str(value).strip(),pattern).date();break
                    except ValueError:pass
            if not day:
                raise PocError(f'La fecha del calendario XLS en la fila {number} está vacía o no se reconoce; no se pudo comprobar el período.')
            if (str(day.month),str(day.year))!=(profile.month,profile.year):
                raise PocError(f'La fecha {day:%d/%m/%Y} del calendario XLS (fila {number}) no corresponde al mes {profile.month}/{profile.year} consultado; puede ser una descarga anterior.')
            checked+=1
    if checked!=sum(info.records.values()):raise PocError('No se pudieron comprobar todas las fechas del calendario XLS.')


def save_exclusive(path, data):
    # No sobrescribe descargas anteriores. El archivo se conserva solo si fue escrito completo.
    created = False
    try:
        with path.open("xb") as f:
            created = True
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
    except FileExistsError:
        raise PocError("Ya existe el archivo de destino; usar una ejecucion nueva.") from None
    except OSError:
        if created and path.exists():
            path.unlink(missing_ok=True)
        raise PocError("Fallo la escritura local del archivo.") from None


def write_manifest(folder, records, total, state, profile=ORIGINAL, manifest_name="verificacion.json"):
    payload = {"fecha_hora": now().isoformat(), "caso": profile.key, "consulta": profile.label,
               "estado": state, "paginas_esperadas": total, "paginas_verificadas": len(records),
               "archivos": records}
    temp = folder / (manifest_name+".tmp")
    temp.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    temp.replace(folder / manifest_name)


def run_sequence(transport, initial_pairs, initial_page, folder, timeout_ms=60000, profile=ORIGINAL,
                 progress=None, cancel=None, folder_ready=False, filename_factory=None, manifest_name="verificacion.json"):
    """POST(n) -> XLS(n) -> validar -> fsync -> POST(n+1). Sin reintentos silenciosos."""
    total = initial_page.total
    original = invariant(initial_pairs)
    pairs = list(initial_pairs)
    pagination = initial_page.pagination
    results, digests, record_sets = [], set(), []
    if not folder_ready: folder.mkdir(parents=True, exist_ok=False)
    def manifest(state):write_manifest(folder,results,total,state,profile,manifest_name)
    manifest("EN_CURSO")
    try:
        for number in range(1, total+1):
            if cancel is not None and cancel.is_set():
                raise PocError("Descarga cancelada. Las paginas ya verificadas se conservan.")
            pairs = [(k, str(number) if k == profile.page_field else pagination.get(k, v)) for k, v in pairs]
            if invariant(pairs) != original:
                raise PocError("Cambio la consulta entre paginas; descarga detenida.")
            query_pairs(urlencode(pairs), profile)
            if number == 1:
                # El llamador ya hizo y validó el POST inicial. SITFA reutiliza
                # el XLS temporal: descargarlo antes de pedir otra página.
                current = initial_page
                if current.current != 1 or not current.records:
                    raise PocError("La consulta inicial no corresponde a la primera pagina.")
                validate_export_url(current.excel_url, profile)
            else:
                body = transport.search(pairs, timeout_ms)
                current = search_page(body, pairs, number, total, profile)
            pagination = current.pagination
            # La URL original se usa solo en memoria. No se escribe en manifiestos o logs.
            content = transport.download(current.excel_url, timeout_ms)
            info, digest = verified_download(content, current, digests, record_sets,profile)
            # El nombre reproduce lo solicitado sin identificador de usuario del servidor.
            filename = (filename_factory(number,total) if filename_factory else
                        f"{now():%Y-%m-%d}_{profile.filename_prefix}_P{number:02}.xls")
            if Path(filename).name != filename or not filename.endswith('.xls'):
                raise PocError("Nombre de archivo fuera de alcance.")
            path = folder / filename
            save_exclusive(path, content)
            # Verifica el archivo que efectivamente quedo en disco antes del siguiente POST.
            saved = path.read_bytes()
            if saved != content or verified_download(saved,current,set(),[],profile)[0].records != current.records:
                raise PocError("El archivo guardado no coincide con la descarga validada.")
            digests.add(digest)
            record_sets.append(excel_record_rows(info))
            results.append({"pagina": number, "total": total, "archivo": filename,
                            "bytes": len(saved), "sha256": digest, "formato": info.format,
                            "registros": sum(info.records.values()), "coincide_con_html": True, "estado": "OK"})
            if current.bindings:results[-1]['vinculos_ingreso']=current.bindings
            manifest("EN_CURSO")
            if progress is not None:
                progress(results[-1], folder)
        if len(results) != total or any(not (folder/item['archivo']).is_file() for item in results):
            raise PocError("La cantidad de XLS no coincide con el total de paginas.")
        manifest("VALIDADA")
        return results
    except BaseException:
        manifest("INCOMPLETA")
        raise
