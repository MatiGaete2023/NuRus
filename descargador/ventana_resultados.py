"""Resumen e historial: acciones secundarias fuera del formulario de descarga."""
import os
from pathlib import Path
import queue
import subprocess
import threading
import tkinter as tk
from tkinter import filedialog,messagebox,ttk

import motor as m
from resultados import load_lot,export_book,export_csv,export_ics,export_zip,compare_lots,inside,read_json,tables,compatibility


def open_local(path):
    path=Path(path).resolve()
    if not path.exists():raise m.PocError('El archivo fue movido o ya no existe.')
    if os.name=='nt':os.startfile(str(path))
    else:subprocess.Popen(['xdg-open',str(path)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)


class ResultsWindow(tk.Toplevel):
    def __init__(self,app):
        super().__init__(app.root);self.app=app;self.title('Resultados y herramientas · rama experimental')
        self.geometry('980x700');self.minsize(800,500);self.busy=False;self.events=queue.Queue();self.folders=[];self.lot=None
        app.result_windows.append(self)
        self.partial=tk.BooleanVar(value=False);self.status=tk.StringVar(value='Selecciona un lote del historial o abre su carpeta.')
        self.history=ttk.Treeview(self,columns=('fecha','estado','consultas'),show='headings',height=5)
        for col,label in [('fecha','Lote'),('estado','Estado'),('consultas','Consultas completadas')]:self.history.heading(col,text=label)
        self.history.pack(fill='x',padx=12,pady=12);self.history.bind('<<TreeviewSelect>>',self.select_history)
        controls=ttk.Frame(self);controls.pack(fill='x',padx=12)
        ttk.Button(controls,text='Abrir lote…',command=self.choose).pack(side='left')
        ttk.Button(controls,text='Actualizar historial',command=self.reload).pack(side='left',padx=8)
        ttk.Checkbutton(controls,text='Permitir resultados INCOMPLETOS',variable=self.partial,command=self.load).pack(side='left')
        self.detail=ttk.Treeview(self,columns=('consulta','estado','filas','aptitud'),show='headings',height=8)
        for col,label in [('consulta','Consulta'),('estado','Resultado'),('filas','Registros'),('aptitud','Compatibilidad CSMP')]:self.detail.heading(col,text=label)
        self.detail.column('consulta',width=350);self.detail.pack(fill='both',expand=True,padx=12,pady=12)
        self.detail.bind('<Double-1>',lambda _:self.open_source())
        row=ttk.Frame(self);row.pack(fill='x',padx=12)
        for text,action in [('Consolidar XLSX',lambda:self.run(lambda:export_book([self.lot]))),
            ('Exportar para CSMP',self.csmp),('Añadir informes al cruce…',self.cross),
            ('Comparar con otro lote…',self.compare)]:ttk.Button(row,text=text,command=action).pack(side='left',padx=(0,6))
        row=ttk.Frame(self);row.pack(fill='x',padx=12,pady=8)
        for text,action in [('CSV',lambda:self.run(lambda:export_csv(self.lot))),('Calendario ICS',lambda:self.run(lambda:export_ics(self.lot))),
            ('ZIP para entrega',lambda:self.run(lambda:export_zip(self.lot))),('Abrir carpeta',lambda:self.safe(lambda:open_local(self.lot.folder))),
            ('Preparar pendientes',self.pending)]:ttk.Button(row,text=text,command=action).pack(side='left',padx=(0,6))
        ttk.Label(self,textvariable=self.status,wraplength=940).pack(fill='x',padx=12,pady=12)
        self.protocol('WM_DELETE_WINDOW',self.close);self.reload();self.after(100,self.poll)
    def close(self):
        if self.busy:self.status.set('Espera a que termine la escritura local.');return
        self.destroy()
    def safe(self,action):
        try:
            if self.busy:raise m.PocError('Hay una operación de resultados en curso.')
            if not self.lot:raise m.PocError('Selecciona un lote primero.')
            action()
        except (m.PocError,OSError,ValueError) as exc:self.status.set(str(exc))
    def reload(self):
        if self.busy:return
        self.history.delete(*self.history.get_children());self.folders=[]
        base=Path(self.app.dest.get())
        for folder in sorted(base.glob('*'),reverse=True)[:300]:
            if not folder.is_dir() or not (folder/'resumen.json').is_file():continue
            try:data=read_json(folder/'resumen.json')
            except (m.PocError,OSError):continue
            iid=str(len(self.folders));self.folders.append(folder)
            self.history.insert('','end',iid=iid,values=(folder.name,data.get('estado'),data.get('consultas_completadas')))
        if self.app.last_folder:
            self.folder=self.app.last_folder;self.load()
    def select_history(self,event=None):
        selected=self.history.selection()
        if selected and not self.busy:self.folder=self.folders[int(selected[0])];self.load()
    def choose(self):
        if self.busy:return
        folder=filedialog.askdirectory(parent=self,title='Carpeta del lote')
        if folder:self.folder=Path(folder);self.load()
    def load(self):
        if self.busy or not getattr(self,'folder',None):return
        self.lot=None
        # Capturar variables Tk en el hilo de interfaz.
        partial=self.partial.get();folder=self.folder
        def worker_action():
            lot=load_lot(folder,partial);aptitudes={}
            for query,selection,item,headers,rows in tables(lot):
                check=compatibility(headers,selection,rows)
                aptitudes.setdefault(query['consulta'],[]).append(check['estado']+(' · faltan '+', '.join(check['faltan']) if check['faltan'] else ''))
            return lot,aptitudes
        self.start(worker_action,'load')
    def start(self,action,kind='product'):
        if self.busy:return
        self.busy=True;self.status.set('Comprobando archivos y procesando resultados localmente…')
        def work():
            try:self.events.put((kind,action(),None))
            except Exception as exc:self.events.put((kind,None,str(exc) if isinstance(exc,(m.PocError,OSError,ValueError)) else 'E_RESULTADOS: no se pudo generar el producto; los originales se conservan.'))
        threading.Thread(target=work,daemon=True).start()
    def run(self,action):self.safe(lambda:self.start(action))
    def poll(self):
        try:
            while True:
                kind,result,error=self.events.get_nowait();self.busy=False
                if error:self.status.set(error);continue
                if kind=='load':
                    self.lot,aptitudes=result;self.detail.delete(*self.detail.get_children())
                    for i,(q,manifest) in enumerate(self.lot.queries):
                        self.detail.insert('','end',iid=str(i),values=(q['consulta'],manifest['estado'],sum(v['registros'] for v in manifest.get('archivos',[])),
                            '; '.join(dict.fromkeys(aptitudes.get(q['consulta'],['Sin registros'])))))
                    self.status.set(('INCOMPLETO · ' if self.lot.partial else '')+self.lot.folder.name+' · Doble clic en una consulta para abrir su primer archivo verificado.')
                else:
                    self.status.set('Creado: '+('; '.join(str(p) for p in result) if isinstance(result,list) else str(result)))
                    if kind=='csmp':self.launch_csmp(result)
        except queue.Empty:pass
        self.after(100,self.poll)
    def open_source(self):
        def action():
            selected=self.detail.selection()
            if not selected:return
            _,manifest=self.lot.queries[int(selected[0])];items=manifest.get('archivos',[]) or [manifest.get('evidencia',{})]
            if items and items[0].get('archivo'):open_local(inside(self.lot.folder,items[0]['archivo']))
        self.safe(action)
    def csmp(self):self.safe(lambda:self.start(lambda:export_book([self.lot],csmp=True),'csmp'))
    def cross(self):
        def action():
            folder=filedialog.askdirectory(parent=self,title='Lote de informes compatible para el cruce')
            if folder:
                current=self.lot
                def export():
                    other=load_lot(folder)
                    modes={compatibility(h,s,rows)['modo'] for lot in (current,other) for _,s,_,h,rows in tables(lot)}
                    if modes!={'CUMPLIMIENTO','INFORMES'}:raise m.PocError('El cruce requiere un lote de Cumplimiento y otro de Informes, sin otras pantallas.')
                    return export_book([current,other],csmp=True)
                self.start(export,'csmp')
        self.safe(action)
    def compare(self):
        def action():
            folder=filedialog.askdirectory(parent=self,title='Lote anterior del mismo alcance')
            if folder:
                current=self.lot;self.start(lambda:compare_lots(load_lot(folder),current))
        self.safe(action)
    def pending(self):self.safe(lambda:self.app.prepare_pending(self.lot.folder))
    def launch_csmp(self,path):
        executable=self.app.prefs.data.get('csmp','')
        if not executable:
            if not messagebox.askyesno('Libro preparado','¿Quieres abrirlo en la rama experimental de CSMP? Selecciona csmp-assistant.exe o pythonw.exe de su entorno.',parent=self):return
            executable=filedialog.askopenfilename(parent=self,title='Ejecutable de CSMP o Python de su entorno',filetypes=[('Ejecutable','*.exe')])
            if not executable:return
            self.app.prefs.data['csmp']=executable;self.app.prefs.save()
        try:
            args=[executable]
            if Path(executable).stem.lower() in ('python','pythonw'):args+=['-m','nurus.personal.app']
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True)
            try:modes=[name for name in ('ESPERA','CUMPLIMIENTO','INFORMES') if name in book.sheetnames]
            finally:book.close()
            extra=['--modo',modes[0],'--hoja',modes[0]] if len(modes)==1 else []
            subprocess.Popen([*args,'--archivo',str(path),*extra],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        except OSError:self.status.set('Libro creado. No se pudo abrir CSMP: revisa su ejecutable configurado.')
