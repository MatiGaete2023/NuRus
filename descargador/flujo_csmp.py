"""Descarga conjunta y recuperación por fase. Comparte los contratos del descargador."""
from dataclasses import asdict
from datetime import date,timedelta
from pathlib import Path
import hashlib,json

import motor as m
from lotes import Batch,BatchRunner
from resultados import load_lot,export_book,publish,mapping,tables,append_text,check_file
from perfiles import Selection


def months(today):
    following=(today.replace(day=28)+timedelta(days=4)).replace(day=1)
    return ((str(today.month),str(today.year)),(str(following.month),str(following.year)))


def plan(tribunals,modalities,mode,today,days=60):
    if mode not in ('Espera','Cumplimiento'):raise m.PocError('El flujo conjunto admite Espera o Cumplimiento.')
    if not isinstance(days,int) or not 60<=days<=730:raise m.PocError('Selecciona entre 60 y 730 días para firmas.')
    if not tribunals or not modalities:raise m.PocError('Selecciona tribunales y modalidades.')
    result=[('principal',Batch(tuple(tribunals),tuple(modalities),(mode,),report='0'))]
    for month,year in months(today):
        result.append((f'informes-{year}-{month}',Batch(tuple(tribunals),(),(),screen='calendario_informes',month=month,year=year)))
    # Carga no forma parte del flujo entregable: su navegación en RUS no quedó
    # validada y bloqueaba el retorno incluso con principal e informes completos.
    # `days` se conserva únicamente para recuperar las selecciones anteriores.
    return result


def write_state(path,data):
    temp=path.with_suffix('.tmp');temp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8');temp.replace(path)


def source_folder(root,name):
    if not isinstance(name,str) or not name:raise m.PocError('Falta la carpeta de fuentes de la fase.')
    target=(root/name).resolve()
    if target==root.resolve() or not target.is_relative_to(root.resolve()):raise m.PocError('La fuente de recuperación está fuera del flujo.')
    return target


def selection_key(selection):return tuple((k,selection.get(k)) for k in Selection.__dataclass_fields__)


def merge_phase(lots,expected,destination):
    """Reconstruir una fase con consultas comprobadas; conservar sus originales."""
    folder=Path(destination)/('Reunido '+m.now().strftime('%Y%m%d_%H%M%S_%f'));folder.mkdir(parents=True)
    expected={selection_key(asdict(s)) for s in expected};seen=set();queries=[];planned=[]
    for lot in lots:
        for query,manifest in lot.queries:
            if query['estado'] not in ('VALIDADA','SIN_RESULTADOS'):continue
            selection=query.get('seleccion',{});key=selection_key(selection)
            if key not in expected or key in seen:raise m.PocError('Consulta duplicada o fuera del alcance al recuperar la fase.')
            seen.add(key);index=len(queries)+1;manifest=dict(manifest);items=[]
            for item in manifest.get('archivos',[]):
                if item.get('estado')!='OK':raise m.PocError('Una consulta completa contiene un archivo pendiente.')
                source=check_file(lot.folder,item);name=f'{index:03}-'+source.name
                m.save_exclusive(folder/name,source.read_bytes());items.append({**item,'archivo':name})
            manifest['archivos']=items
            if manifest.get('evidencia'):
                item=manifest['evidencia'];source=check_file(lot.folder,item);name=f'{index:03}-'+source.name
                m.save_exclusive(folder/name,source.read_bytes());manifest['evidencia']={**item,'archivo':name}
            name=f'verificacion {index:03}.json';write_state(folder/name,manifest)
            queries.append({**query,'verificacion':name});planned.append(selection)
    if seen!=expected:raise m.PocError('La recuperación no contiene todas las consultas de la fase.')
    write_state(folder/'resumen.json',{'version_contrato':1,'estado':'VALIDADA','consultas':queries,'plan':planned,
        'consultas_esperadas':len(expected),'consultas_completadas':len(queries),'otros_filtros':'todos los registros','filtro_fecha':None,'fecha_hora':m.now().isoformat()})
    load_lot(folder);return folder


def assemble(lots,signature_rows,coverage_by_court,destination):
    """Una entrada CSMP; la hoja auxiliar conserva el orden mensual/original."""
    base=export_book(lots,destination,csmp=True)
    from openpyxl import load_workbook
    book=load_workbook(base)
    # El calendario masivo carece de RUT en algunos contratos. Completarlo solo
    # con una identidad sin RUT unívoca de la descarga principal, nunca por RIT solo.
    primary=next((s for s in book if s.title in ('ESPERA','CUMPLIMIENTO')),None)
    auxiliary=book['INFORMES'] if 'INFORMES' in book.sheetnames else None
    if primary and auxiliary:
        pcols=mapping([c.value for c in primary[1]]);acols=mapping([c.value for c in auxiliary[1]])
        keys=('rit','nombre','tribunal','programa')
        if all(k in pcols and k in acols for k in keys) and 'rut' in pcols and 'rut' not in acols:
            from collections import defaultdict
            index=defaultdict(set)
            for row in primary.iter_rows(min_row=2,values_only=True):
                key=tuple(m.normalized(row[pcols[k]]) for k in keys)
                if all(key):index[key].add(str(row[pcols['rut']] or '').strip())
            column=auxiliary.max_column+1;auxiliary.cell(1,column,'RUT')
            auxiliary.cell(1,column+1,'SITFA_ORIGEN_RUT')
            for cells in auxiliary.iter_rows(min_row=2):
                key=tuple(m.normalized(cells[acols[k]].value) for k in keys);matches=index.get(key,set())
                if len(matches)==1 and '' not in matches:
                    auxiliary.cell(cells[0].row,column,next(iter(matches)))
                    auxiliary.cell(cells[0].row,column+1,'Identidad única RIT/tribunal/persona/centro de hoja principal')
    details=book.create_sheet('Resoluciones firmadas')
    labels=('tribunal_codigo','tribunal','rit','tramite','firma','hora','documento','identidad','huella')
    details.append(labels)
    for item in signature_rows:
        details.append([item.get(k,'') for k in labels])
        for cell in details[details.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    status=book.create_sheet('SITFA_COBERTURA');status.append(('tribunal_codigo','desde','hasta','estado','criterio_temporal','sin_coincidencias'))
    for code,item in coverage_by_court.items():status.append([code,*[item[k] for k in ('desde','hasta','estado','criterio_temporal','sin_coincidencias')]])
    target=publish(destination,'Flujo CSMP','.xlsx',book.save);book.close()
    return target


class JointRunner:
    def __init__(self,bridge,emit,cancel,runner_factory=BatchRunner):
        self.bridge=bridge;self.emit=emit;self.cancel=cancel;self.runner_factory=runner_factory

    def run(self,tribunals,modalities,mode,destination,days=60,*,today=None,resume=None):
        today=today or m.now().date();phases=plan(tribunals,modalities,mode,today,days)
        identity={'tribunales':list(tribunals),'modalidades':list(modalities),'modo':mode,'corte':today.isoformat(),'dias_firmas':days,
                  'consultar_firmas':False}
        if resume:
            root=Path(resume).resolve();data=json.loads((root/'flujo.json').read_text(encoding='utf-8'))
            previous_identity=dict(data.get('seleccion',{}));previous_identity.setdefault('consultar_firmas',False)
            if previous_identity!=identity:raise m.PocError('La recuperación no coincide con la selección y fecha de corte.')
            data['seleccion']=identity
            retired={key:phase for key,phase in data['fases'].items() if key.startswith('carga-')}
            if retired:
                data.setdefault('fases_fuera_del_flujo',{}).update(retired)
                data['fases']={key:phase for key,phase in data['fases'].items() if key not in retired}
        else:
            root=Path(destination)/('Flujo CSMP '+m.now().strftime('%Y-%m-%d %H-%M-%S-%f'));root.mkdir(parents=True)
            data={'version_contrato':1,'seleccion':identity,'estado':'EN_CURSO','fases':{}}
        state_path=root/'flujo.json';write_state(state_path,data)
        try:
            for key,batch in phases:
                if self.cancel.is_set():raise m.PocError('Flujo cancelado; conserva las fases comprobadas.')
                old=data['fases'].get(key)
                if old and old.get('estado')=='VALIDADA':
                    load_lot(source_folder(root,old['carpeta']));continue
                catalog=self.bridge.call('catalog',{'screen':batch.screen})
                previous=[];pending=None
                if old and old.get('carpeta'):
                    previous_names=list(dict.fromkeys([*old.get('carpetas_parciales',[]),old['carpeta']]))
                    previous=[load_lot(source_folder(root,name),partial=True) for name in previous_names]
                    done={selection_key(query['seleccion']) for lot in previous for query,_ in lot.queries if query['estado'] in ('VALIDADA','SIN_RESULTADOS')}
                    pending=[s for s in batch.plan(catalog) if selection_key(asdict(s)) not in done]
                folders=[]
                def phase_emit(name,value):
                    if name=='folder':folders.append(value)
                    elif name!='done':self.emit(name,value)
                runner=self.runner_factory(self.bridge,catalog,phase_emit,self.cancel)
                try:
                    if pending==[]:
                        combined=merge_phase(previous,batch.plan(catalog),root/key)
                        folders.append(str(combined));completed=list(load_lot(combined).queries)
                    else:
                        completed=runner.run(batch,root/key,plan_override=pending)
                    if len(folders)!=1:raise m.PocError('La fase no identifica su carpeta de fuentes.')
                    if previous and pending:
                        combined=merge_phase([*previous,load_lot(folders[0])],batch.plan(catalog),root/key)
                        folders=[str(combined)];completed=list(load_lot(combined).queries)
                    data['fases'][key]={'estado':'VALIDADA','carpeta':str(Path(folders[0]).relative_to(root)),'consultas':len(completed)}
                except BaseException:
                    folder_name=str(Path(folders[0]).relative_to(root)) if folders else old.get('carpeta','') if old else ''
                    prior_names=[*old.get('carpetas_parciales',[]),old['carpeta']] if old and old.get('carpeta') else []
                    data['fases'][key]={'estado':'INCOMPLETA','carpeta':folder_name,'carpetas_parciales':list(dict.fromkeys(prior_names))};raise
                finally:write_state(state_path,data)
            inputs=[]
            start=today-timedelta(days=days-1)
            for key,batch in phases:
                lot=load_lot(source_folder(root,data['fases'][key]['carpeta']))
                inputs.append(lot)
            covers={code:{'desde':start.isoformat(),'hasta':today.isoformat(),'estado':'NO_CONSULTADA',
                         'criterio_temporal':'consulta_firmas_retirada_del_flujo',
                         'sin_coincidencias':'Resoluciones firmadas no consultadas; no acredita ausencia de movimientos.'}
                    for code in tribunals}
            final=assemble(inputs,[],covers,root)
            data.update(estado='PREPARADO_SIN_CONSULTA_FIRMAS',archivo=final.name,sha256=hashlib.sha256(final.read_bytes()).hexdigest(),cobertura=covers)
            write_state(state_path,data);self.emit('folder',str(root));self.emit('joint_done',{'folder':str(root),'archivo':str(final),'estado':data['estado'],'modo':mode.upper()})
            return final
        except BaseException:
            data['estado']='INCOMPLETO';write_state(state_path,data);self.emit('folder',str(root));raise
