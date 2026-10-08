"""Sonda offline de interoperabilidad; solo datos ficticios, sin Office ni red.

Ejecutar con el Python de pruebas de NuRus:
  python pruebas/compatibilidad_csmp.py /ruta/al/repositorio/NuRus
No es un exportador ni forma parte de la aplicación de uso diario.
"""
import argparse
from datetime import date
from hashlib import sha256
from html import escape
import io
import json
from pathlib import Path
import sys
import tempfile


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('nurus_repo',type=Path)
    args=parser.parse_args()
    sys.path.insert(0,str(args.nurus_repo.resolve()/'src'))
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
    import motor
    from openpyxl import Workbook
    from nurus.personal.config import defaults
    from nurus.personal.work import Work
    from nurus.rus.reader import WorkbookReadError

    base=['RIT','TRIBUNAL','NOMBRE','RUT','NOMBRE CENTRO']
    person=['X-1-2026','Jgdo. L. y G. de Laja','Persona ficticia','11111111-1','AFT FICTICIO']
    tables={
        'ESPERA':[base+['T ESPERA'],person+[40]],
        'CUMPLIMIENTO':[base+['DIAS DE CUMPLIMIENTO','DIAS PARA EGRESAR','FEC.EGRESO PROYECTADO'],
                        person+[50,90,'30/12/2026']],
        'INFORMES':[base+['FECHA VENCIMIENTO'],person+['10/10/2026']],
    }
    report={'datos':'ficticios','consultas_sitfa':0,'office_ejecutado':False,
            'adaptador_productivo_implementado':False,'casos':[]}

    def book_bytes(rows,title):
        book=Workbook();book.active.title=title
        for row in rows:
            book.active.append(row)
            for cell in book.active[book.active.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        out=io.BytesIO();book.save(out);book.close();return out.getvalue()

    with tempfile.TemporaryDirectory(prefix='csmp_compatibilidad_') as tmp:
        folder=Path(tmp)
        for mode,rows in tables.items():
            html=('<html><meta charset="utf-8"><table>'+''.join('<tr>'+''.join(
                '<td>'+escape(str(value))+'</td>' for value in row)+'</tr>' for row in rows)+'</table></html>').encode()
            xml=('<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet><Table>'+
                 ''.join('<Row>'+''.join('<Cell><Data>'+escape(str(value))+'</Data></Cell>' for value in row)+
                         '</Row>' for row in rows)+'</Table></Worksheet></Workbook>').encode()
            sources={'HTML':html,'XML_2003':xml,'OOXML':book_bytes(rows,mode)}
            for kind,content in sources.items():
                source=folder/(mode+'_'+kind+'.xls');source.write_bytes(content)
                info=motor.read_excel(content)
                assert sum(info.records.values())==1
                try:
                    Work(defaults()).analyze(source,mode,as_of=date(2026,10,1))
                except WorkbookReadError:
                    direct=False
                else:
                    direct=True
                # NuRus experimental reconoce OOXML por contenido; el antiguo lo rechazaba.
                if kind!='OOXML':assert direct is False,(mode,kind,'Revisar contrato de HTML/XML')
                target=source.with_suffix('.xlsx')
                target.write_bytes(book_bytes(info.grid,mode))
                work=Work(defaults()).analyze(target,mode,as_of=date(2026,10,1))
                assert len(work.rows)==1 and work.rows[0].observation
                assert sha256(source.read_bytes()).digest()==sha256(content).digest()
                if mode=='ESPERA':assert 'programa_espera' in work.rows[0].actions
                if mode=='INFORMES':assert 'programa_por_vencer' in work.rows[0].actions
                if mode=='CUMPLIMIENTO':assert work.cross_missing
                report['casos'].append({'modo':mode,'formato_real':kind,'lectura_download':True,
                    'xls_directo_csmp':direct,'xlsx_convertido_csmp':True,'filas':len(work.rows),
                    'original_intacto':True,'cruce_faltante':work.cross_missing})

        # Identificación suficiente para Download no equivale a campos para CSMP.
        partial=folder/'identidad_minima.xlsx'
        partial.write_bytes(book_bytes([['RIT','NOMBRE'],['X-1-2026','Persona ficticia']],'ESPERA'))
        assert motor.read_excel(partial.read_bytes()).records
        try:limited=Work(defaults()).analyze(partial,'ESPERA',as_of=date(2026,10,1))
        except (KeyError,WorkbookReadError):
            report['identidad_minima_csmp']={'rechazada':True}
        else:
            assert not all(k in limited.mapping for k in ('tribunal','programa','espera'))
            report['identidad_minima_csmp']={'rechazada':False,'campos_funcionales_incompletos':True,
                'advertencias':bool(limited.warnings or any(row.warnings for row in limited.rows)),
                'acciones':sum(len(row.actions) for row in limited.rows)}

        cross=Workbook();cross.remove(cross.active)
        for mode in ('CUMPLIMIENTO','INFORMES'):
            sheet=cross.create_sheet(mode)
            for row in tables[mode]:sheet.append(row)
        path=folder/'cruce.xlsx';cross.save(path);cross.close()
        work=Work(defaults()).analyze(path,'CUMPLIMIENTO',as_of=date(2026,10,1))
        assert not work.cross_missing
        assert 'CUMPLIMIENTO.C10_HOJA2' in work.rows[0].rules
        report['cruce_c10_con_identidad_completa_y_fecha_futura']=True

        # Cargar externa conserva revisiones; no sustituye PROCESAR.
        external=Work.external(folder/'ESPERA_HTML.xlsx',defaults(),'ESPERA')
        assert len(external.rows)==1 and not external.rows[0].rules and not external.rows[0].actions
        report['cargar_externa_no_ejecuta_reglas']=True
    report['resultado']='OK'
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
