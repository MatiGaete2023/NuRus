"""Recuperación secuencial: reutiliza consultas completas sin repetirlas."""
from dataclasses import asdict
from datetime import date
from pathlib import Path
import hashlib,io,json,tempfile,threading,unittest
from openpyxl import Workbook,load_workbook
import motor as m
from flujo_csmp import JointRunner,merge_phase
from perfiles import Selection
from resultados import load_lot
from test_flujo_csmp import row
from carga import HEADERS


CAT={'tribunales':{'113':'Jgdo. L. y G. de Laja','119':'Jgdo. L. y G. de Mulchén'},'modalidades':{'2':'Ambulatorio'},'pestanas':['Cumplimiento'],
    'informes':{'1':'Informes por vencer'},'medidas':{'0':'Todas'},'meses':{'10':'Octubre','11':'Noviembre'},'anios':{'2026':'2026'}}


def make_lot(folder,selections,complete=True,planned=None):
    folder.mkdir(parents=True);queries=[];selection_meta=[]
    for index,s in enumerate(selections,1):
        name=CAT['tribunales'][s.tribunal];meta={**asdict(s),'tribunal_nombre':name,'fecha_descarga':'2026-10-02T09:00:00'};selection_meta.append(meta)
        book=Workbook();sheet=book.active
        if s.screen=='carga':
            sheet.append(HEADERS)
            values=row();values[1]=name;sheet.append(values)
        else:
            header=['RIT','NOMBRE','RUT','NOMBRE CENTRO','TRIBUNAL']
            values=['X-1-2026','Persona ficticia','11111111-1','PRM FICTICIO',name]
            if s.screen=='calendario_informes':header+=['F. VENCIMIENTO'];values+=[f'10/{int(s.month):02}/2026']
            else:header+=['DIAS DE CUMPLIMIENTO','DIAS PARA EGRESAR','FEC.EGRESO PROYECTADO'];values+=[50,90,'30/12/2026']
            sheet.append(header);sheet.append(values)
        out=io.BytesIO();book.save(out);book.close();content=out.getvalue();filename=f'pagina-{index}.xls';(folder/filename).write_bytes(content)
        item={'archivo':filename,'pagina':1,'registros':1,'sha256':hashlib.sha256(content).hexdigest(),'bytes':len(content),'estado':'OK'}
        manifest={'estado':'VALIDADA','consulta':'Consulta ficticia','seleccion':meta,'archivos':[item],'paginas_esperadas':1,'paginas_verificadas':1}
        manifest_name=f'verificacion {index:03}.json';(folder/manifest_name).write_text(json.dumps(manifest),encoding='utf-8')
        queries.append({'consulta':'Consulta ficticia','estado':'VALIDADA','seleccion':meta,'verificacion':manifest_name})
    summary={'estado':'VALIDADA' if complete else 'INCOMPLETA','version_contrato':1,'consultas':queries,'plan':selection_meta if planned is None else [asdict(s) for s in planned],
        'consultas_esperadas':len(selections) if planned is None else len(planned),'consultas_completadas':len(queries),'otros_filtros':'todos los registros','filtro_fecha':None}
    (folder/'resumen.json').write_text(json.dumps(summary),encoding='utf-8');return folder


class FakeBridge:
    def __init__(self):self.screens=[]
    def call(self,command,payload):
        assert command=='catalog';self.screens.append(payload['screen']);return {**CAT,'pantalla':payload['screen']}


class RecoveryTests(unittest.TestCase):
    def test_csmp_requested_mode_is_enforced_before_download(self):
        from descargador import Application
        from types import SimpleNamespace
        app=SimpleNamespace(csmp_reply=Path('respuesta.json'),csmp_mode='ESPERA')
        Application.check_csmp_mode(app,'Espera')
        with self.assertRaisesRegex(m.PocError,'CSMP solicitó ESPERA'):
            Application.check_csmp_mode(app,'Cumplimiento')
        app.csmp_reply=None
        Application.check_csmp_mode(app,'Cumplimiento')
    def test_legacy_flow_stopped_before_carga_finishes_without_any_browser_query(self):
        from flujo_csmp import plan
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);folder=root/'Flujo anterior';folder.mkdir()
            state={'version_contrato':1,'estado':'INCOMPLETO','seleccion':{
                'tribunales':['113'],'modalidades':['2'],'modo':'Cumplimiento',
                'corte':'2026-10-02','dias_firmas':60},'fases':{}}
            for key,batch in plan(('113',),('2',),'Cumplimiento',date(2026,10,2)):
                cat={**CAT,'pantalla':batch.screen}
                make_lot(folder/key,batch.plan(cat))
                state['fases'][key]={'estado':'VALIDADA','carpeta':key}
            (folder/'flujo.json').write_text(json.dumps(state),encoding='utf-8')
            class NoQueries:
                def call(self,*args):raise AssertionError('No debe repetir ni abrir Carga')
            events=[]
            result=JointRunner(NoQueries(),lambda *event:events.append(event),threading.Event()).run(
                ('113',),('2',),'Cumplimiento',root,today=date(2026,10,2),resume=folder)
            book=load_workbook(result)
            self.assertEqual(book['CUMPLIMIENTO'].max_row,2)
            self.assertEqual(book['INFORMES'].max_row,3)
            self.assertEqual(book['Resoluciones firmadas'].max_row,1)
            self.assertEqual(book['SITFA_COBERTURA'].cell(2,4).value,'NO_CONSULTADA')
            book.close()
            self.assertEqual([name for name,_ in events].count('joint_done'),1)
            self.assertEqual(events[-1][1]['modo'],'CUMPLIMIENTO')
            final=json.loads((folder/'flujo.json').read_text(encoding='utf-8'))
            self.assertFalse(final['seleccion']['consultar_firmas'])
            self.assertEqual(final['estado'],'PREPARADO_SIN_CONSULTA_FIRMAS')

    def test_merge_preserves_complete_queries_and_refuses_modified_files(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);a=Selection('113','Cumplimiento','2');b=Selection('119','Cumplimiento','2')
            old=load_lot(make_lot(root/'old',[a],complete=False,planned=[a,b]),partial=True)
            new=load_lot(make_lot(root/'new',[b]));merged=load_lot(merge_phase([old,new],[a,b],root/'merged'))
            self.assertEqual(len(merged.queries),2);self.assertEqual(len(list(merged.sources())),2)
            (old.folder/'pagina-1.xls').write_bytes(b'alterado')
            with self.assertRaises(m.PocError):merge_phase([old,new],[a,b],root/'bad')

    def test_full_flow_and_resume_request_only_missing_tribunal(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);calls=[];fail=[True]
            class Runner:
                def __init__(self,bridge,catalog,emit,cancel):self.emit=emit;self.catalog=catalog
                def run(self,batch,destination,plan_override=None):
                    selections=batch.plan(self.catalog) if plan_override is None else plan_override
                    calls.append((batch.screen,[s.tribunal for s in selections]))
                    folder=Path(destination)/str(len(calls));self.emit('folder',str(folder))
                    if batch.screen=='seguimiento' and fail[0]:
                        fail[0]=False;make_lot(folder,selections[:1],complete=False,planned=selections);raise m.PocError('Corte ficticio después de la primera consulta')
                    make_lot(folder,selections);return load_lot(folder).summary['consultas']
            bridge=FakeBridge();runner=JointRunner(bridge,lambda *args:None,threading.Event(),runner_factory=Runner)
            with self.assertRaises(m.PocError):runner.run(('113','119'),('2',),'Cumplimiento',root,today=date(2026,10,2))
            folder=next(root.glob('Flujo CSMP *'));final=runner.run(('113','119'),('2',),'Cumplimiento',root,today=date(2026,10,2),resume=folder)
            self.assertEqual(calls[1],('seguimiento',['119']))
            self.assertEqual([s for s,_ in calls[2:]],['calendario_informes','calendario_informes'])
            book=load_workbook(final);self.assertEqual(book['CUMPLIMIENTO'].max_row,3);self.assertEqual(book['INFORMES'].max_row,5)
            self.assertEqual(book['Resoluciones firmadas'].max_row,1);book.close()
            state=json.loads((folder/'flujo.json').read_text(encoding='utf-8'));self.assertEqual(state['estado'],'PREPARADO_SIN_CONSULTA_FIRMAS')
            self.assertTrue(all(v['estado']=='NO_CONSULTADA' for v in state['cobertura'].values()))
