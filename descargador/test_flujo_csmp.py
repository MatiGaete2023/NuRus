"""Cruces y períodos con registros ficticios; no ejecuta consultas reales."""
from collections import Counter
from datetime import date
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import tempfile,unittest
from urllib.parse import urlencode
from openpyxl import Workbook
import motor as m
from carga import HEADERS,date_blocks,record_rows,signatures,coverage,records_in_grid
from flujo_csmp import months,plan,assemble
from resultados import load_lot
from lotes import Batch,Dates
from perfiles import Selection,selection_profile,response_profile


def row(rit='X-1-2026',signature='02/10/2026',invalid='--'):
    values=['']*23;values[0]=rit;values[1]='Tribunal ficticio';values[4]='Trámite ficticio';values[5]=invalid;values[6]='SI'
    values[21]=signature;values[22]='09:30';return values


class JointTests(unittest.TestCase):
    def test_assembled_input_preserves_month_order_sources_and_primary_rows(self):
        from test_resultados import fixture
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            principal=load_lot(fixture(root/'principal',mode='Cumplimiento',names=('Persona ficticia',)))
            current=fixture(root/'actual',mode='Informes',names=('Persona ficticia',));following=fixture(root/'siguiente',mode='Informes',names=('Persona ficticia',))
            book=load_workbook(BytesIO((following/'pagina.xls').read_bytes()))
            book.active.cell(2,5,'10/11/2026');out=BytesIO();book.save(out);book.close();content=out.getvalue();(following/'pagina.xls').write_bytes(content)
            import hashlib,json
            meta=following/'verificacion 001.json';data=json.loads(meta.read_text());data['archivos'][0].update(sha256=hashlib.sha256(content).hexdigest(),bytes=len(content));meta.write_text(json.dumps(data))
            scopes={'113':coverage(date(2026,8,4),date(2026,10,2),[{'estado':'VALIDADA'}])}
            result=assemble([principal,load_lot(current),load_lot(following)],[],scopes,root/'final')
            book=load_workbook(result)
            self.assertEqual(book['CUMPLIMIENTO'].max_row,2)
            self.assertEqual(book['INFORMES'].max_row,3)
            headers=[c.value for c in book['INFORMES'][1]];due=headers.index('FECHA VENCIMIENTO')+1
            self.assertEqual([book['INFORMES'].cell(i,due).value for i in (2,3)],['10/10/2026','10/11/2026'])
            self.assertEqual(book['SITFA_COBERTURA'].cell(2,4).value,'PARCIAL')
            self.assertEqual(book['SITFA_ARCHIVOS'].max_row,4)
            book.close()

    def test_inclusive_periods_and_year_change(self):
        blocks=date_blocks(date(2026,8,4),date(2026,10,2));self.assertEqual(blocks,[(date(2026,8,4),date(2026,9,2)),(date(2026,9,3),date(2026,10,2))])
        self.assertEqual(months(date(2026,12,31)),(('12','2026'),('1','2027')))
        phases=plan(('777','888'),('1','2'),'Cumplimiento',date(2026,10,2))
        self.assertIsNone(phases[0][1].dates);self.assertFalse(phases[0][1].keep_filters)
        self.assertEqual(len(phases),3)
        self.assertNotIn('carga',[batch.screen for _,batch in phases])
        self.assertEqual([s.month for _,s in phases if s.screen=='calendario_informes'],['10','11'])
        with self.assertRaises(m.PocError):plan(('777',),('1',),'Espera',date(2026,10,2),59)

    def test_signed_only_and_fixed_title_is_ignored(self):
        rows=[['Juzgado de Familia Concepción'],list(HEADERS),row(),row(signature='--'),row(invalid='X'),row(signature='01/01/2026')]
        found=signatures(rows,'777','Tribunal ficticio',date(2026,8,4),date(2026,10,2))
        self.assertEqual(len(found),1);self.assertEqual(found[0]['tribunal_codigo'],'777')
        self.assertEqual(len(signatures([list(HEADERS),row(),row()],'777','Tribunal ficticio',date(2026,8,4),date(2026,10,2))),2)
        with self.assertRaises(m.PocError):signatures([list(HEADERS),row()],'888','Otro tribunal',date(2026,8,4),date(2026,10,2))
        with self.assertRaises(m.PocError):record_rows([list(HEADERS),row()+['dato extra']])

    def test_no_rows_never_prove_full_history_by_default(self):
        scope=coverage(date(2026,8,4),date(2026,10,2),[{'estado':'SIN_RESULTADOS'}])
        self.assertEqual(scope['estado'],'PARCIAL');self.assertIn('informes consultados',scope['sin_coincidencias'])
        self.assertEqual(coverage(date(2026,8,4),date(2026,10,2),[{'estado':'INCOMPLETA'}],complete_history=True)['estado'],'PARCIAL')

    def test_carga_specific_request_response_and_full_cells_match(self):
        cat={'pantalla':'carga','tribunales':{'777':'Tribunal ficticio'}}
        selection=Selection('777','','',screen='carga',start='03/09/2026',end='02/10/2026')
        profile=selection_profile(selection,cat,'antr');pairs=list({**profile.target,'irAccion':profile.action}.items())
        fields=''.join(f'<input name="{k}" value="{v}">' for k,v in pairs)
        listing='<table><tr>'+''.join(f'<td>{v}</td>' for v in row())+'</tr></table>'
        source=f'<form name="MaoPpalForm" action="/SITFAWEB/MaoDAction.do" method="post">{fields}{listing}</form><script>function ShowExcel(){{window.open("https://familia.pjud.cl/sitfa/reportes/demo_antr.xls");}}</script>'
        actual,page=response_profile(source,pairs,selection,cat)
        self.assertEqual(actual.screen,'carga');self.assertEqual(page.total,1)
        book=Workbook();book.active.append(list(HEADERS));book.active.append(row());out=BytesIO();book.save(out);book.close()
        self.assertEqual(m.verified_download(out.getvalue(),page,set(),[],profile)[0].records,page.records)
        changed=Workbook();changed.active.append(list(HEADERS));changed.active.append(row());changed.active.cell(2,22,'01/10/2026');out2=BytesIO();changed.save(out2);changed.close()
        with self.assertRaises(m.PocError):m.verified_download(out2.getvalue(),page,set(),[],profile)
        for replacement in ({'irAccion':'Grabar'},{'TIP_Cargo':'46'},{'COD_TribunalDist':'888'}):
            with self.assertRaises(m.PocError):m.query_pairs(urlencode({**dict(pairs),**replacement}),profile)
        with self.assertRaises(m.PocError):selection_profile(replace(selection,start='02/09/2026'),cat)
