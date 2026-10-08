"""Productos locales derivados de archivos verificados. Sin Chrome ni Office."""
from collections import Counter,defaultdict
from dataclasses import dataclass
from datetime import date,datetime
from datetime import timezone
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
import zipfile

import motor as m
from lotes import safe_name

MAX_ROWS=200_000
MAX_JSON=4*1024*1024
MODES={'Espera':'ESPERA','Cumplimiento':'CUMPLIMIENTO','Informes':'INFORMES'}
MODALITIES={'1':'RES','2':'AMB','3':'FAE','4':'DCE'}
ALIASES={
 'rit':('RIT',),'nombre':('NOMBRE','NOMBRE MENOR','NOMBRE COMPLETO'),
 'rut':('RUT','RUT MENOR','RUT NNA'), 'tribunal':('TRIBUNAL',),
 'programa':('DERIVACION','PROGRAMA','NOMBRE CENTRO'),
 'espera':('T ESPERA','T_ESPERA','DIAS_ESPERA','TESPERA','DIAS DE ESPERA'),
 'dias_cumpl':('DIAS DE CUMPLIMIENTO','DIAS CUMPLIMIENTO'),
 'dias_egresar':('DIAS PARA EGRESAR','DIAS EGRESAR'),
 'vencimiento':('FECHA VENCIMIENTO','FEC.VENCIMIENTO','FEC. VENCIMIENTO','F. VENCIMIENTO','F.VENCIMIENTO','F.VENC','F. VENC','VENCIMIENTO'),
 'egreso_proy':('FEC.EGRESO PROYECTADO','FEC. EGRESO PROYECTADO','FEC EGRESO PROYECTADO'),
}
REQUIRED={'ESPERA':('rit','nombre','tribunal','programa','espera'),
          'CUMPLIMIENTO':('rit','nombre','tribunal','programa','dias_cumpl','dias_egresar'),
          'INFORMES':('rit','nombre','tribunal','programa','vencimiento')}
ORIGIN=('SITFA_LOTE','SITFA_CONSULTA','SITFA_ARCHIVO','SITFA_PAGINA','SITFA_FILA','SITFA_SHA256','SITFA_FECHA_DESCARGA')


def read_json(path):
    if path.stat().st_size>MAX_JSON:raise m.PocError('El resumen excede el límite de lectura.')
    try:
        data=json.loads(path.read_text(encoding='utf-8'))
        if not isinstance(data,dict):raise ValueError()
        return data
    except (ValueError,UnicodeError):raise m.PocError('Resumen JSON inválido.') from None


def inside(folder,name):
    folder=Path(folder).resolve()
    if not isinstance(name,str) or not name or Path(name).name!=name or '\\' in name:
        raise m.PocError('El manifiesto contiene una ruta fuera del lote.')
    path=(folder/name).resolve()
    if path.parent!=folder:raise m.PocError('El archivo está fuera del lote.')
    return path


def digest_file(path):
    digest=hashlib.sha256()
    with path.open('rb') as source:
        for block in iter(lambda:source.read(1024*1024),b''):digest.update(block)
    return digest.hexdigest()


def check_file(folder,item):
    path=inside(folder,item.get('archivo'))
    if not path.is_file() or path.stat().st_size>80*1024*1024:
        raise m.PocError('Falta un archivo o supera el límite de lectura.')
    if item.get('bytes') is not None and path.stat().st_size!=item['bytes']:
        raise m.PocError('El tamaño de un archivo verificado cambió.')
    if not item.get('sha256') or digest_file(path)!=item['sha256']:
        raise m.PocError('Un archivo cambió después de la descarga. No se exportó el lote.')
    return path


@dataclass
class Lot:
    folder:Path
    summary:dict
    queries:list
    partial:bool

    def sources(self):
        for query,manifest in self.queries:
            for item in manifest.get('archivos',[]):
                if item.get('estado')=='OK':yield query,manifest,item,check_file(self.folder,item)


def load_lot(folder,partial=False):
    folder=Path(folder).resolve();summary=read_json(folder/'resumen.json')
    incomplete=summary.get('estado')!='VALIDADA'
    if incomplete and not partial:raise m.PocError('El lote es incompleto. Selecciona explícitamente una exportación parcial.')
    queries=[];seen=set()
    for query in summary.get('consultas',[]):
        path=inside(folder,query.get('verificacion'))
        if path.name in seen:raise m.PocError('Verificación repetida en el resumen.')
        seen.add(path.name);manifest=read_json(path)
        if manifest.get('estado') not in ('VALIDADA','SIN_RESULTADOS'):
            raise m.PocError('El estado de una consulta no coincide con el resumen.')
        if query.get('estado')!=manifest['estado']:raise m.PocError('Estados de consulta incoherentes.')
        queries.append((query,manifest))
    # La consulta interrumpida puede no figurar aún en consultas completadas.
    if incomplete and partial:
        for path in sorted(folder.glob('verificacion *.json')):
            if path.name in seen:continue
            manifest=read_json(path)
            if manifest.get('estado')=='INCOMPLETA':
                queries.append(({'consulta':manifest.get('consulta','Consulta incompleta'),
                                 'estado':'INCOMPLETA','seleccion':manifest.get('seleccion',{}),'verificacion':path.name},manifest))
    if not incomplete and len(queries)!=summary.get('consultas_esperadas'):
        raise m.PocError('El lote no contiene todas sus consultas.')
    lot=Lot(folder,summary,queries,incomplete)
    filenames=set()
    for query,manifest,item,path in lot.sources():
        if path.name in filenames:raise m.PocError('El mismo archivo aparece en varias consultas.')
        filenames.add(path.name)
    for query,manifest in queries:
        if manifest.get('estado')=='VALIDADA':
            files=manifest.get('archivos',[])
            if manifest.get('paginas_esperadas') is not None and len(files)!=manifest['paginas_esperadas']:
                raise m.PocError('Faltan páginas en una consulta validada.')
    for query,manifest in queries:
        if manifest.get('evidencia'):check_file(folder,manifest['evidencia'])
    return lot


def headers_and_rows(info):
    header=None;rows=[]
    for number,row in enumerate(info.grid or [],1):
        row=list(row)
        if m.record_columns(row) is not None:
            current=tuple(str(v or '').strip() for v in row)
            # RUS formatea columnas vacías al final del calendario. Los datos
            # fuera de esta cabecera se siguen rechazando al leer cada fila.
            while current and not current[-1]:current=current[:-1]
            if not all(current) or len(set(map(m.normalized,current)))!=len(current):
                raise m.PocError('Cabeceras vacías o ambiguas: requiere revisión local.')
            if header is not None and current!=header:raise m.PocError('La cabecera cambia dentro del Excel.')
            header=current;continue
        if header is None:continue
        if m.record_from_row(row,m.record_columns(header)) is not None:
            if len(row)>len(header) and any(v not in ('',None) for v in row[len(header):]):
                raise m.PocError('Hay datos fuera de la cabecera reconocida.')
            rows.append((number,tuple((row+['']*len(header))[:len(header)])))
        elif any(v not in ('',None) for v in row):
            if m.normalized(' '.join(str(v or '') for v in row)).startswith(('TOTAL','SIN REGISTROS','NO HAY REGISTROS')):continue
            raise m.PocError('Hay filas no reconocidas; no se consolidan silenciosamente.')
    if not header or len(rows)!=sum(info.records.values()):raise m.PocError('No coinciden las filas del Excel reconocido.')
    return header,rows


def mapping(headers):
    result={}
    for key,names in ALIASES.items():
        matches=[i for i,h in enumerate(headers) if m.normalized(h) in set(map(m.normalized,names))]
        if len(matches)>1:raise m.PocError('Columnas ambiguas para '+key+'.')
        if matches:result[key]=matches[0]
    return result


def table_mode(selection):
    if selection.get('screen')=='seguimiento':return MODES.get(selection.get('tab'),'')
    if selection.get('screen')=='calendario_informes':return 'INFORMES'
    return ''


def tables(lot):
    count=0
    for query,manifest,item,path in lot.sources():
        content=path.read_bytes()
        if hashlib.sha256(content).hexdigest()!=item['sha256']:raise m.PocError('El archivo cambió durante la lectura.')
        selection=query.get('seleccion') or manifest.get('seleccion') or {}
        if selection.get('screen')=='carga':
            from carga import HEADERS,records_in_grid,canonical
            import re
            info=m.read_excel(content,records_in_grid);headers=HEADERS
            rows=[(i,tuple(row)) for i,row in enumerate(info.grid,1) if row and re.fullmatch(r'[A-Z]+-\d+-\d{4}',canonical(row[0]))]
        else:
            info=m.read_excel(content);headers,rows=headers_and_rows(info)
        if len(rows)!=item['registros']:raise m.PocError('El recuento no coincide con la verificación.')
        if any(h.startswith('SITFA_') for h in headers):raise m.PocError('La descarga contiene columnas reservadas de procedencia.')
        count+=len(rows)
        if count>MAX_ROWS:raise m.PocError('El lote supera 200.000 filas para exportación local.')
        selection=query.get('seleccion') or manifest.get('seleccion') or {}
        # Un tribunal ausente se añade únicamente con la selección auditada.
        cols=mapping(headers)
        if 'tribunal' not in cols and selection.get('tribunal_nombre'):
            headers=(*headers,'TRIBUNAL');rows=[(n,(*row,selection['tribunal_nombre'])) for n,row in rows]
        # Comprobar tribunal por igualdad normalizada; abreviaturas requieren revisión.
        cols=mapping(headers)
        if 'tribunal' in cols and selection.get('tribunal_nombre'):
            from lotes import tribunal_name
            def tribunal_key(value):
                normalized=m.normalized(tribunal_name(str(value)))
                for label in ('LAJA','MULCHEN','TOME'):
                    if normalized==label or normalized.endswith(' '+label):return label
                return normalized
            for _,row in rows:
                if tribunal_key(row[cols['tribunal']])!=tribunal_key(selection['tribunal_nombre']):
                    raise m.PocError('Un tribunal del Excel no coincide con la consulta; revisa sus etiquetas.')
        if selection.get('screen')=='seguimiento' and selection.get('modality') in MODALITIES:
            if any(m.normalized(h) in ('MODALIDAD','TIPO PROGRAMA') for h in headers):
                # Conservar la columna original; la modalidad auditada queda separada.
                headers=(*headers,'SITFA_MODALIDAD')
            else:headers=(*headers,'MODALIDAD')
            rows=[(n,(*row,MODALITIES[selection['modality']])) for n,row in rows]
        if selection.get('tribunal'):
            headers=(*headers,'SITFA_TRIBUNAL_CODIGO')
            rows=[(n,(*row,selection['tribunal'])) for n,row in rows]
        if selection.get('screen')=='seguimiento':
            from vinculos import match_rows
            headers,rows=match_rows(headers,rows,item.get('vinculos_ingreso',[]),selection.get('tribunal',''))
        yield query,selection,item,headers,rows


def compatibility(headers,selection,rows=()):
    mode=table_mode(selection);cols=mapping(headers)
    if not mode:return {'modo':'','estado':'Solo análisis','faltan':['Pantalla sin modo CSMP']}
    missing=[k for k in REQUIRED[mode] if k not in cols]
    if mode=='INFORMES' and selection.get('screen')=='seguimiento' and selection.get('report')=='3':
        missing.append('Contrato para Informes recibidos (no corresponde a pendientes)')
    blanks=sum(any(row[cols[k]] in ('',None) for k in REQUIRED[mode] if k in cols) for _,row in rows)
    return {'modo':mode,'estado':'Sin datos suficientes' if missing else ('Compatible con limitaciones' if blanks else 'Compatible'),
            'faltan':missing,'filas_incompletas':blanks}


def publish(folder,label,suffix,writer):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    prefix=safe_name(label)+' '+m.now().strftime('%Y%m%d_%H%M%S_%f')
    target=folder/(prefix+suffix)
    fd,temp=tempfile.mkstemp(suffix=suffix,dir=folder);os.close(fd)
    created=False
    try:
        writer(Path(temp))
        with target.open('xb') as output:
            created=True
            with open(temp,'rb') as source:
                shutil.copyfileobj(source,output);output.flush();os.fsync(output.fileno())
    except BaseException:
        if created and target.exists():target.unlink()
        raise
    finally:Path(temp).unlink(missing_ok=True)
    return target


def append_text(sheet,values,header=False):
    from openpyxl.cell import WriteOnlyCell
    cells=[]
    for value in values:
        cell=WriteOnlyCell(sheet,value=value)
        if isinstance(value,str):cell.data_type='s'
        from excel_presentacion import format_cell
        format_cell(cell,header)
        cells.append(cell)
    sheet.append(cells)


def export_book(lots,destination=None,csmp=False):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    lots=list(lots)
    if not lots:raise m.PocError('Selecciona al menos un lote.')
    if csmp:
        if any(l.partial for l in lots):raise m.PocError('Para CSMP se requieren lotes completos.')
        schemas=defaultdict(set);total=0
        for lot in lots:
            for q,s,item,h,rows in tables(lot):
                report=compatibility(h,s,rows)
                if not report['modo'] or report['faltan']:raise m.PocError('No apto para CSMP: '+q['consulta']+' · '+', '.join(report['faltan']))
                schemas[report['modo']].add(h);total+=len(rows)
        if any(len(v)>1 for v in schemas.values()):raise m.PocError('CSMP requiere una estructura compatible por modo. Separa los esquemas diferentes.')
        if total==0:raise m.PocError('El lote no contiene registros para CSMP; consulta sus PDF y resumen.')
        # El libro puede unir Cumplimiento e Informes; otros cruces deben ser conscientes.
        if len(lots)>1:
            if any(l.summary.get('otros_filtros')=='formulario' for l in lots):raise m.PocError('No se cruzan filtros privados no comparables.')
            dates={json.dumps(l.summary.get('filtro_fecha'),sort_keys=True) for l in lots}
            courts=[{s.get('tribunal') for s in l.summary.get('plan',[])} for l in lots]
            if len(dates)>1 or not courts[0] or any(c!=courts[0] for c in courts):raise m.PocError('Los lotes de cruce no comparten tribunales y filtros de fecha.')
    from excel_presentacion import configure_sheet
    book=Workbook(write_only=True);summary=book.create_sheet('Resumen')
    configure_sheet(summary,['Resultado','Lote','Consulta','Hoja','Filas','Compatibilidad CSMP','Columnas faltantes']);append_text(summary,['Resultado','Lote','Consulta','Hoja','Filas','Compatibilidad CSMP','Columnas faltantes'],header=True)
    origin=book.create_sheet('SITFA_ARCHIVOS')
    configure_sheet(origin,['Carpeta origen','Archivo origen','SHA256']);append_text(origin,['Carpeta origen','Archivo origen','SHA256'],header=True)
    groups={};names=Counter();reports=[];source_list=[];source_rows=0
    used_names=set(book.sheetnames)
    for lot in lots:
        for query,manifest in lot.queries:
            if manifest.get('estado')=='SIN_RESULTADOS':
                append_text(summary,['SIN_RESULTADOS',lot.folder.name,query['consulta'],'PDF',0,'Sin registros',''])
            elif manifest.get('estado')=='INCOMPLETA' and not manifest.get('archivos'):
                append_text(summary,['INCOMPLETA',lot.folder.name,query['consulta'],'',0,'Consulta pendiente',''])
        for query,selection,item,headers,rows in tables(lot):
            report=compatibility(headers,selection,rows);reports.append(report)
            if csmp and (not report['modo'] or report['faltan']):
                raise m.PocError('No apto para CSMP: '+query['consulta']+' · '+', '.join(report['faltan']))
            family=report['modo'] or selection.get('screen') or 'Datos'
            key=(family,headers)
            if key not in groups:
                names[family]+=1;name=(family if names[family]==1 else family+'_'+str(names[family]))[:31]
                if name in ('Resumen','SITFA_ARCHIVOS','Vencimientos'):name='Datos_'+name
                while name in used_names:name=(name[:27]+'_'+str(len(used_names)))[:31]
                used_names.add(name)
                sheet=book.create_sheet(name);sheet.freeze_panes='A2'
                configure_sheet(sheet,(*headers,*ORIGIN))
                append_text(sheet,(*headers,*ORIGIN),header=True);groups[key]=[sheet,1]
            sheet,total=groups[key]
            for rownum,row in rows:
                append_text(sheet,(*row,lot.folder.name,query['consulta'],item['archivo'],item['pagina'],rownum,item['sha256'],
                                    selection.get('fecha_descarga',lot.summary.get('fecha_hora',''))))
            groups[key][1]+=len(rows);source_rows+=len(rows)
            if source_rows>MAX_ROWS:raise m.PocError('La exportación supera 200.000 filas.')
            append_text(summary,['INCOMPLETO' if lot.partial else 'VERIFICADO',lot.folder.name,query['consulta'],sheet.title,len(rows),
                                 report['estado'],', '.join(report['faltan'])])
            append_text(origin,[str(lot.folder),item['archivo'],item['sha256']])
            source_list.append({'lote':str(lot.folder),'archivo':item['archivo'],'sha256':item['sha256'],'filas':len(rows)})
    if csmp:
        for family in names:
            if names[family]>1:raise m.PocError('Hay varios esquemas para '+family+'. Exporta cada esquema por separado antes de CSMP.')
    due_sheet=book.create_sheet('Vencimientos')
    configure_sheet(due_sheet,['Vencimiento','Tribunal','Lote','Archivo','Fila','Estado de fecha']);append_text(due_sheet,['Vencimiento','Tribunal','Lote','Archivo','Fila','Estado de fecha'],header=True)
    # Tabla derivada ordenada; no recalcula plazos ni modifica las fuentes.
    due_rows=[]
    for lot in lots:
        for _,selection,item,headers,rows in tables(lot):
            cols=mapping(headers);field='vencimiento' if 'vencimiento' in cols else 'egreso_proy' if 'egreso_proy' in cols else None
            if not field:continue
            for n,row in rows:
                raw=row[cols[field]];day=parse_day(raw)
                due_rows.append((day or date.max,str(raw or ''),selection.get('tribunal_nombre',''),lot.folder.name,item['archivo'],n))
    for day,raw,tribunal,lote,archivo,n in sorted(due_rows):
        append_text(due_sheet,[day.isoformat() if day!=date.max else raw,tribunal,lote,archivo,n,'Fecha válida' if day!=date.max else 'No interpretable'])
    for (_,headers),(sheet,total) in groups.items():sheet.auto_filter.ref=f'A1:{get_column_letter(len(headers)+len(ORIGIN))}{total}'
    label=('Para CSMP' if csmp else 'Consolidado')+(' INCOMPLETO' if any(l.partial for l in lots) else '')
    target=publish(destination or lots[0].folder,label,'.xlsx',book.save);book.close()
    metadata={'version_contrato':1,'archivo':target.name,'sha256':digest_file(target),'filas':source_rows,
              'estado':'INCOMPLETO' if any(l.partial for l in lots) else 'VALIDADA','fuentes':source_list,
              'compatibilidad':reports,'hojas':list(names),'fecha_exportacion':m.now().isoformat()}
    (target.with_suffix('.json')).write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    return target


def parse_day(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    for fmt in ('%d/%m/%Y','%Y-%m-%d','%d-%m-%Y'):
        try:return datetime.strptime(str(value).strip(),fmt).date()
        except ValueError:pass
    return None


def export_csv(lot):
    outputs=[];schemas={}
    for query,selection,item,headers,rows in tables(lot):
        key=(table_mode(selection) or selection.get('screen') or 'Datos',headers)
        schemas.setdefault(key,len(schemas)+1)
    for (family,headers),number in schemas.items():
        def write(path,expected=(family,headers)):
            with path.open('w',encoding='utf-8-sig',newline='') as output:
                writer=csv.writer(output,delimiter=';');writer.writerow((*headers,*ORIGIN))
                for query,selection,item,current,rows in tables(lot):
                    if (table_mode(selection) or selection.get('screen') or 'Datos',current)!=expected:continue
                    for n,row in rows:
                        # Excel abre CSV sin tipos: neutralizar texto que puede ejecutar fórmulas.
                        values=(*row,lot.folder.name,query['consulta'],item['archivo'],item['pagina'],n,item['sha256'],selection.get('fecha_descarga',''))
                        writer.writerow(["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in values])
        outputs.append(publish(lot.folder,'Tabla '+family+' '+str(number),'.csv',write))
    return outputs


def export_ics(lot):
    events=[];occurrences=Counter()
    def escape(value):return str(value).replace('\\','\\\\').replace('\n','\\n').replace(';','\\;').replace(',','\\,').replace('\r','')
    table_data=list(tables(lot));identities=Counter()
    def identity(selection,headers,row,field):
        cols=mapping(headers)
        required=('rit','rut','programa')
        if any(key not in cols or not str(row[cols[key]] or '').strip() for key in required):return None
        return json.dumps([selection.get('tribunal'),selection.get('screen'),field,*[str(row[cols[key]]).strip() for key in required]],ensure_ascii=False)
    for _,selection,item,headers,rows in table_data:
        cols=mapping(headers)
        for field in ('vencimiento','egreso_proy'):
            if field in cols:
                for n,row in rows:
                    key=identity(selection,headers,row,field)
                    if key and parse_day(row[cols[field]]):identities[key]+=1
    for _,selection,item,headers,rows in table_data:
        cols=mapping(headers)
        for field in ('vencimiento','egreso_proy'):
            if field not in cols:continue
            for n,row in rows:
                day=parse_day(row[cols[field]])
                if day is None:continue
                stable=identity(selection,headers,row,field)
                key=stable if stable and identities[stable]==1 else json.dumps([selection.get('tribunal'),field,[str(v) for v in row]],ensure_ascii=False)
                occurrences[key]+=1;uid=hashlib.sha256((key+':'+str(occurrences[key])).encode()).hexdigest()+'@sitfa-local'
                rit=str(row[cols['rit']]).strip() if 'rit' in cols else ''
                label='Vencimiento de informe' if field=='vencimiento' else 'Egreso proyectado'
                meaning=('Verificar vencimiento en el original; no acredita entrega del informe.' if field=='vencimiento' else 'Fecha proyectada de la fuente; no acredita un egreso efectivo.')
                if field=='vencimiento' and selection.get('screen')=='calendario_medidas':
                    label='Vencimiento de medida';meaning='Fecha de vencimiento de la medida en la fuente; no acredita un egreso efectivo.'
                events+=['BEGIN:VEVENT' ,'UID:'+uid,'DTSTAMP:'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ'),
                         'DTSTART;VALUE=DATE:'+day.strftime('%Y%m%d'),
                         'SUMMARY:'+escape(label+' SITFA · '+rit+' · '+selection.get('tribunal_nombre','')),
                         'DESCRIPTION:'+escape(f'Fuente: {item["archivo"]}, fila {n}. Columna: {headers[cols[field]]}. '
                             +meaning+' Cobertura del lote: '+('INCOMPLETO' if lot.partial else 'VALIDADO')+'.'),'END:VEVENT']
    if not events:raise m.PocError('No hay fechas verificables para exportar al calendario.')
    lines=['BEGIN:VCALENDAR','VERSION:2.0','PRODID:-//SITFA local//Vencimientos//ES','CALSCALE:GREGORIAN',*events,'END:VCALENDAR']
    def write(path):
        # Plegado a <=75 octetos sin cortar secuencias UTF-8 (RFC 5545).
        folded=[]
        for line in lines:
            current=''
            for char in line:
                if len((current+char).encode())>73:folded.append(current);current=' '+char
                else:current+=char
            folded.append(current)
        content=('\r\n'.join(folded)+'\r\n').encode()
        from icalendar import Calendar
        parsed=Calendar.from_ical(content)
        if len(parsed.walk('VEVENT'))!=events.count('BEGIN:VEVENT'):raise m.PocError('El calendario no superó la comprobación de eventos.')
        path.write_bytes(content)
    return publish(lot.folder,'Vencimientos','.ics',write)


def export_zip(lot):
    paths={lot.folder/'resumen.json'}
    for query,manifest in lot.queries:
        if query.get('verificacion'):paths.add(inside(lot.folder,query['verificacion']))
        for item in manifest.get('archivos',[]):
            if item.get('estado')=='OK':paths.add(check_file(lot.folder,item))
        if manifest.get('evidencia'):paths.add(check_file(lot.folder,manifest['evidencia']))
    def write(path):
        entries=[{'archivo':p.name,'bytes':p.stat().st_size,'sha256':digest_file(p)} for p in sorted(paths)]
        expected_hashes={item['archivo']:item['sha256'] for item in entries}
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('INDICE.txt',('INCOMPLETO' if lot.partial else 'VERIFICADO')+'\nArchivos del lote actual. Consultar MANIFIESTO.json para comprobar integridad.\n'+'\n'.join(sorted(p.name for p in paths)))
            archive.writestr('MANIFIESTO.json',json.dumps({'version':1,'estado':'INCOMPLETO' if lot.partial else 'VERIFICADO','archivos':entries},ensure_ascii=False,indent=2))
            for source in sorted(paths):
                expected=expected_hashes[source.name];archive.write(source,source.name)
                h=hashlib.sha256()
                with archive.open(source.name) as f:
                    for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
                if h.hexdigest()!=expected:raise m.PocError('Un archivo cambió durante la entrega.')
    return publish(lot.folder,'Entrega'+(' INCOMPLETO' if lot.partial else ''),'.zip',write)


def comparable_scope(lot):
    if lot.partial or lot.summary.get('otros_filtros')=='formulario':raise m.PocError('No se comparan lotes parciales ni filtros privados no verificables.')
    plan=lot.summary.get('plan')
    if not plan:raise m.PocError('Este lote antiguo no tiene alcance estructurado para comparar.')
    clean=lambda s:{k:v for k,v in s.items() if k not in ('fecha_descarga','tribunal_nombre','modalidad_nombre')}
    return json.dumps([sorted((json.dumps(clean(s),sort_keys=True) for s in plan)),lot.summary.get('filtro_fecha')],sort_keys=True)


def compare_lots(old,new):
    from openpyxl import Workbook
    if comparable_scope(old)!=comparable_scope(new):raise m.PocError('Los lotes tienen alcances diferentes.')
    def inventory(lot):
        data=Counter()
        for _,selection,_,headers,rows in tables(lot):
            scope=json.dumps({k:v for k,v in selection.items() if k not in ('fecha_descarga','tribunal_nombre','modalidad_nombre')},sort_keys=True)
            for _,row in rows:data[(scope,headers,tuple('' if v is None else str(v) for v in row))]+=1
        return data
    a,b=inventory(old),inventory(new)
    schemas=lambda inv:{(scope,headers) for scope,headers,_ in inv}
    if a and b and schemas(a)!=schemas(b):raise m.PocError('Las estructuras de columnas no coinciden.')
    book=Workbook(write_only=True)
    for label,delta in [('Nuevas',b-a),('Ausentes',a-b)]:
        sheet=book.create_sheet(label);append_text(sheet,['Consulta','Columnas','Valores de fila','Cantidad'])
        for (scope,headers,row),n in delta.items():append_text(sheet,[scope,json.dumps(headers,ensure_ascii=False),json.dumps(row,ensure_ascii=False),n])
    sheet=book.create_sheet('Resumen')
    for row in [('Lote anterior',str(old.folder)),('Lote nuevo',str(new.folder)),('Nuevas',sum((b-a).values())),('Ausentes',sum((a-b).values())),
                ('Interpretación','Ausente no implica egreso ni eliminación; se comparan filas completas y multiplicidad.')]:append_text(sheet,row)
    target=publish(new.folder,'Comparacion','.xlsx',book.save);book.close();return target
